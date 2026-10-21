"""통합 판매현황(첫 시트 '통합판매현황') 오프라인 검증 — 로그인·쿠팡 호출 없음(결정적). SSOT=designs/SALES_SUMMARY.md.

목표(소유자 2026-10-15·2026-10-21): 계정별·상품별 판매수량·판매금액(윙/로켓그로스 구분)을 결과 파일 첫 시트에 매일
다시 그린다. 위 = 대시보드(①요약·②월별·③주간 8주·④일별 31일 차트·⑤계정 순위·⑥상품 상위 20 열지도), 아래 = 일별 상세표
(최근 31일·최신이 왼쪽). 집계 기준 = 관리 상품(결과 사업자 시트 블록 VID). 일별 기록은 2026-04-01부터 계정당 하루 30일씩
최근부터 거꾸로 소급. '지난달 같은 기간' = 이번 달 1~N일 vs 지난달 1~N일(일별).

    python tools/verify_sales_summary_offline.py
"""
from __future__ import annotations

import sys
import tempfile
from datetime import date, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
try:
    sys.stdout.reconfigure(encoding="utf-8")
except AttributeError:
    pass

from coupang_analytics import config  # noqa: E402
from coupang_analytics import sales_dashboard as SD  # noqa: E402
from coupang_analytics import sales_history as SH  # noqa: E402
from coupang_analytics import sales_summary as SS  # noqa: E402
from coupang_analytics import workbook_sales as WS  # noqa: E402
from coupang_analytics import workbook_sales_dash as WD  # noqa: E402
from coupang_analytics.report import OptionMetric  # noqa: E402
from coupang_analytics.workbook import OutputWorkbook  # noqa: E402

Y = date(2026, 10, 14)   # 전일(수요일) — 이번 주 월요일 = 10-12


def ok(msg):
    print(f"  ✔ {msg}")


def _m(vid, name, reg, qty, gmv, opt="기본"):
    return OptionMetric(vid, name, opt, "i", 0, qty, 0, registration_type=reg, gmv=gmv, orders=qty)


def t1_plan():
    print("[T1] 기간 조회 계획 — 처음 = 1월~이번 달 + 최근 8주, 마감된 기간은 다시 안 함, 이번 달·주는 매일 다시")
    plan = SS.plan_periods({}, Y)
    keys = [k for k, _f, _t in plan]
    months = [k for k in keys if k.startswith("M:")]
    weeks = [k for k in keys if k.startswith("W:")]
    assert months == [f"M:2026-{m:02d}" for m in range(1, 11)], months
    assert weeks[0] == "W:2026-10-12" and weeks[-1] == "W:2026-08-24" and len(weeks) == 8, weeks
    rng = {k: (f, t) for k, f, t in plan}
    assert rng["M:2026-10"] == ("2026-10-01", "2026-10-14") and rng["M:2026-02"] == ("2026-02-01", "2026-02-28")
    assert rng["W:2026-10-12"] == ("2026-10-12", "2026-10-14") and rng["W:2026-10-05"] == ("2026-10-05", "2026-10-11")
    done = {k: {"조회기간": [f, t]} for k, f, t in plan}               # 전부 받음(이번 달·주는 10-14 까지)
    again = SS.plan_periods(done, date(2026, 10, 15))
    assert [k for k, _f, _t in again] == ["M:2026-10", "W:2026-10-12"], again   # 마감된 것은 다시 안 함
    assert SS.plan_periods(done, Y) == [], "같은 날 다시 돌면 조회 없음"
    ok("1~10월 10개 + 8주 · 2월 말일 · 이번 주 = 월~어제 · 다음 날엔 이번 달·주만 다시")


def t2_record_period():
    print("[T2] 기간 기록 저장 — 쿠팡 합계 그대로·옵션별·키별 교체·불러오기")
    with tempfile.TemporaryDirectory() as d:
        SH.record_period(d, "a1", "M:2026-10", "2026-10-01", "2026-10-14",
                         {"v1": _m("v1", "캠핑테이블", "RFM", 3, 30000)})
        SH.record_period(d, "a1", "M:2026-10", "2026-10-01", "2026-10-15",
                         {"v1": _m("v1", "캠핑테이블", "RFM", 4, 40000)})
        data = SH.load_periods(d)
        rec = data["a1"]["M:2026-10"]
        assert rec["조회기간"] == ["2026-10-01", "2026-10-15"] and rec["합계"]["판매금액"] == 40000, rec
        assert rec["옵션"]["v1"]["구분"] == "로켓그로스" and rec["옵션"]["v1"]["판매수량"] == 4
        assert not list(Path(d, SH.PERIOD_DIRNAME).glob("*.tmp")), "임시 파일 남지 않음"
        SH.record(d, "a1", "2026-10-14", "2026-10-14", {"v1": _m("v1", "캠핑테이블", "RFM", 1, 9000)})
        assert SH.load_daily(d)["a1"]["2026-10-14"]["합계"]["판매수량"] == 1
    ok("같은 키는 교체·조회기간 기록·구분 윙/로켓그로스·일별도 같은 방식으로 불러옴")


def _fixture(d):
    SH.record(d, "a1", "2026-10-14", "2026-10-14", {
        "v1": _m("v1", "캠핑테이블", "RFM", 2, 20000, "블랙"), "v2": _m("v2", "캠핑테이블", "RFM", 1, 10000, "화이트"),
        "v3": _m("v3", "허리보호대", "NORMAL", 1, 5000), "v4": _m("v4", "안팔린상품", "RFM", 0, 0)})
    SH.record_period(d, "a1", "W:2026-10-12", "2026-10-12", "2026-10-14", {"v1": _m("v1", "캠핑테이블", "RFM", 5, 50000)})
    SH.record_period(d, "a1", "M:2026-10", "2026-10-01", "2026-10-14", {
        "v1": _m("v1", "캠핑테이블", "RFM", 9, 90000), "v3": _m("v3", "허리보호대", "NORMAL", 2, 10000)})
    SH.record_period(d, "a1", "M:2026-09", "2026-09-01", "2026-09-30", {})          # 판매 0 인 달(미수집 아님)


def t3_build():
    print("[T3] 상세표 — 일별 31칸·최신이 왼쪽·계정→상품(옵션 합산)·윙/로켓그로스 줄 나눔·판매 0 상품 생략·미수집 ≠ 0")
    with tempfile.TemporaryDirectory() as d:
        _fixture(d)
        s = SS.build(SH.load_daily(d), {"a1": "에벤에셀 (a1)", "a2": "니코에이블 (a2)"}, Y)
    assert len(s.sections) == 1, "주간·월별 구역은 상세표에서 제거(대시보드가 대신)"
    day = s.sections[0]
    keys = [c.key for c in day.cols]
    assert len(keys) == 31 and keys[0] == "D:2026-10-14" and keys[-1] == "D:2026-09-14", (keys[0], keys[-1])
    a1 = next(b for b in day.accounts if b.aid == "a1")
    rows = {(r.product, r.kind): r.values["D:2026-10-14"] for r in a1.rows}
    assert rows == {("캠핑테이블", "로켓그로스"): (3, 30000), ("허리보호대", "윙"): (1, 5000)}, rows
    assert a1.total["D:2026-10-14"] == (4, 35000) and a1.total["D:2026-10-13"] is None, "기록 없는 날 = 미수집"
    a2 = next(b for b in day.accounts if b.aid == "a2")
    assert a2.total["D:2026-10-14"] is None and a2.rows == [], "데이터 없는 계정 = 미수집(None)"
    assert day.total["D:2026-10-14"] == (4, 35000) and day.missing["D:2026-10-14"] == 1
    ok("31칸(10.14 → 09.14) · 옵션 합산 3개·30,000원 · 윙 줄 따로 · a2 미수집 · 기록 없는 날 미수집")


def _wb_with_accounts(path):
    wb = OutputWorkbook.empty()
    for biz, aid, vids in (("에벤에셀", "a1", ["v1", "v2", "v3", "v4"]), ("니코에이블", "a2", ["v9"])):
        wb.ensure_account(biz)
        wb.set_account_id(biz, aid)
        wb.ensure_product_block(biz, f"{biz} 상품", config.KIND_CONTRACT, ["kw"], registered=f"{biz} 상품")
        wb.set_product_vids(biz, f"{biz} 상품", vids)          # 관리 상품 = 블록 VID(집계 기준)
        wb.ensure_date(biz, "2026-10-14")
    wb.apply_style()
    wb.save(path)
    return wb


def t4_sheet():
    print("[T4] 결과 파일 첫 시트 — 통합판매현황 맨 앞·계정 목록 두 번째·사업자 시트로 오인 안 됨·다시 그려도 1개")
    with tempfile.TemporaryDirectory() as d:
        _fixture(d)
        path = Path(d) / "m.xlsx"
        wb = _wb_with_accounts(path)
        before = wb.account_sheets()
        from coupang_analytics import login_today as LT
        LT.record_failure("a2", LT.PW_WRONG, out_dir=d)                  # 오늘 a2 로그인 비밀번호 오류(D-068)
        msg = WS.refresh_sales_sheet(wb, d, today=date(2026, 10, 15))
        wb.apply_style()
        WS.refresh_sales_sheet(wb, d, today=date(2026, 10, 15))           # 두 번 그려도 시트 1개
        wb.apply_style()
        wb.save(path)
        wb2 = OutputWorkbook.load(path)
        names = wb2.wb.sheetnames
        assert names[0] == WS.SALES_SHEET and names[1] == "계정 목록", names
        assert names.count(WS.SALES_SHEET) == 1
        assert wb2.account_sheets() == before, (wb2.account_sheets(), before)
        assert wb2.wb.active.title == WS.SALES_SHEET, "파일을 열면 통합판매현황"
        cells = [[c for c in r] for r in wb2.wb[WS.SALES_SHEET].iter_rows(values_only=True)]
        flat = [v for r in cells for v in r if v is not None]
        assert any(str(v).startswith("① 요약") for v in flat) and any(str(v).startswith("상세 — 일별") for v in flat)
        assert not any(str(v).startswith(("2. 주간", "3. 월별")) for v in flat), "옛 주간·월별 구역 없음"
        tot = next(r for r in cells if r[0] == "전체 합계")
        assert tot[3] == 4 and tot[4] == 35000, tot
        assert any(r[0] == "니코에이블 (a2)" and r[3] == "미수집" for r in cells), "미수집 계정 표시"
        subs = {r[0]: r[1] for r in cells if r[1] and str(r[1]).startswith("소계")}
        assert subs == {"니코에이블 (a2)": "소계 — ⚠ 로그인 비밀번호 오류", "에벤에셀 (a1)": "소계"}, subs   # 소유자 2026-10-16
        assert any(r[1] == "캠핑테이블" and r[2] == "로켓그로스" and r[4] == 30000 for r in cells)
        kpi = next(r for r in cells if str(r[0]).startswith("전일 "))
        assert (kpi[1], kpi[2]) == (35000, 4), f"대시보드 전일 = 상세 전체 합계: {kpi}"
        assert "기록" in msg or "갱신" in msg, msg
    ok("첫 시트·계정 목록 두 번째·활성 시트·사업자 시트 목록 불변·전체 합계 4개 35,000원·미수집 표시·미수집 사유(로그인 비밀번호 오류)")


def t5_gsheet():
    print("[T5] 결과 구글시트 — 통합판매현황을 만들고 맨 앞(index 0)으로·사업자 미러링 대상 아님")
    from coupang_analytics import gsheet_stats as GS
    with tempfile.TemporaryDirectory() as d:
        _fixture(d)
        wb = _wb_with_accounts(Path(d) / "m.xlsx")
        WS.refresh_sales_sheet(wb, d, today=date(2026, 10, 15))
        wb.apply_style()

    class _Client:
        def __init__(self):
            self.reqs, self.made = [], []

        def ensure_sheets(self, titles):
            self.made += titles
            return {t: i + 10 for i, t in enumerate(titles)}

        def ensure_sheet(self, title):
            self.made.append(title)
            return 99

        def chart_ids(self, title):
            return [501, 502] if title == WS.SALES_SHEET else []

        def batch_update(self, reqs):
            self.reqs += reqs
    c = _Client()
    GS.push_statistics(c, wb)
    assert WS.SALES_SHEET not in c.made, "사업자 시트로 미러링 안 함"
    WS.push_sales_sheet(c, wb)
    moves = [r for r in c.reqs if "updateSheetProperties" in r
             and r["updateSheetProperties"]["properties"].get("index") == 0]
    assert moves and moves[-1]["updateSheetProperties"]["properties"]["sheetId"] == 99, moves
    assert any("updateCells" in r for r in c.reqs)
    bad = _merges_across_freeze(c.reqs)   # 라이브 2026-10-15 400 "You can't merge frozen and non-frozen columns"
    assert not bad, f"틀고정 경계를 가로지르는 병합(구글시트 400): {bad[:3]}"
    kinds = [next(iter(r)) for r in c.reqs]
    dels = [r["deleteEmbeddedObject"]["objectId"] for r in c.reqs if "deleteEmbeddedObject" in r]
    first_cells = kinds.index("updateCells", len(kinds) - kinds[::-1].index("deleteEmbeddedObject"))
    assert dels == [501, 502], "옛 차트를 지움"
    charts = [r["addChart"]["chart"] for r in c.reqs if "addChart" in r]
    assert len(charts) == 3 and kinds.index("addChart") > first_cells, "값을 쓴 뒤 차트 3개 다시 추가"
    assert sorted(ch["spec"]["basicChart"]["chartType"] for ch in charts) == ["COLUMN", "COLUMN", "LINE"]
    assert all(ch["spec"]["hiddenDimensionStrategy"] == "SHOW_ALL" for ch in charts), "숨긴 열 자료도 그림"
    hid = [r["updateDimensionProperties"] for r in c.reqs if "updateDimensionProperties" in r
           and r["updateDimensionProperties"]["properties"].get("hiddenByUser")]
    cols = {ch["spec"]["basicChart"]["domains"][0]["domain"]["sourceRange"]["sources"][0]["startColumnIndex"]
            for ch in charts}
    assert hid and all(any(h["range"]["startIndex"] <= ci < h["range"]["endIndex"] for h in hid) for ci in cols), \
        "차트 자료 = 숨긴 열"
    assert kinds[-1] == "updateSheetProperties", "index 0 이동이 마지막"
    ok("옛 차트 삭제 → 값·서식 교체 → 숨긴 열·차트 3개 추가 → index 0 · 사업자 미러링 아님 · 틀고정 넘는 병합 없음")


def _merges_across_freeze(reqs: list) -> list:
    """구글시트는 고정/비고정 행·열에 걸친 병합을 거부(400) — 같은 시트의 틀고정 값으로 병합 범위 검사."""
    frozen = {}
    for r in reqs:
        g = r.get("updateSheetProperties", {}).get("properties", {})
        if "gridProperties" in g:
            gp = g["gridProperties"]
            frozen[g["sheetId"]] = (gp.get("frozenRowCount", 0), gp.get("frozenColumnCount", 0))
    bad = []
    for r in reqs:
        m = r.get("mergeCells", {}).get("range")
        if not m:
            continue
        fr, fc = frozen.get(m["sheetId"], (0, 0))
        if m["startColumnIndex"] < fc < m["endColumnIndex"] or m["startRowIndex"] < fr < m["endRowIndex"]:
            bad.append(m)
    return bad


def t6_collect():
    print("[T6] ① 판매수집 기간 조회 — 계획한 기간만 조회·저장 · 조회 실패면 그 계정 나머지 중단(로그)")
    from coupang_analytics import collector
    from coupang_analytics import pipeline_sales as PS
    calls, logs = [], []

    def fake(page, f, t, log=None):
        calls.append((f, t))
        if len(calls) == 3:
            raise collector.SalesFetchError("status=500")
        return {"v1": _m("v1", "캠핑테이블", "RFM", 1, 1000)}
    orig = (collector.fetch_sales_details, PS.time.sleep)
    try:
        collector.fetch_sales_details, PS.time.sleep = fake, (lambda s: None)
        with tempfile.TemporaryDirectory() as d:
            PS._record_sales_periods("PAGE", "a1", d, logs.append, yesterday=Y)
            got = SH.load_periods(d).get("a1", {})
    finally:
        collector.fetch_sales_details, PS.time.sleep = orig
    assert len(calls) == 3 and len(got) == 2, (calls, got)
    assert any("중단" in m for m in logs), logs
    logs.clear()

    def broken(page, f, t, log=None):          # 핀 A 가 잡은 결함: 판매분석 오류가 아닌 브라우저 오류도 수집을 막으면 안 됨
        raise RuntimeError("Target page, context or browser has been closed")
    try:
        collector.fetch_sales_details = broken
        with tempfile.TemporaryDirectory() as d:
            PS._record_sales_periods("PAGE", "a1", d, logs.append, yesterday=Y)   # 예외가 밖으로 나오지 않음
    finally:
        collector.fetch_sales_details = orig[0]
    assert any("RuntimeError" in m for m in logs), logs
    ok("기간 2개 저장 · 3번째 실패에서 그 계정 중단 · 브라우저 오류도 ① 수집을 막지 않음(로그)")


def t7_backfill_plan():
    print("[T7] 일별 소급 계획 — 2026-04-01까지 최근부터 거꾸로·있는 날 빼고·계정당 하루 최대 30일(오늘 받은 것 포함)")
    y = date(2026, 10, 20)
    have = {"2026-10-20": {"수집": "2026-10-21T18:05:00"}, "2026-10-18": {"수집": "2026-10-19T18:05:00"}}
    plan = SS.plan_backfill(have, y, today=date(2026, 10, 21))
    assert plan[:3] == ["2026-10-19", "2026-10-17", "2026-10-16"] and len(plan) == 29, plan[:5]   # 오늘 1건 = 29
    assert plan == sorted(plan, reverse=True) and "2026-10-18" not in plan
    assert SS.plan_backfill({}, date(2026, 4, 3)) == ["2026-04-03", "2026-04-02", "2026-04-01"], "4/1 전은 안 함"
    full = {f"x{i}": {"수집": "2026-10-21T19:00:00"} for i in range(30)}
    assert SS.plan_backfill(full, y, today=date(2026, 10, 21)) == [], "오늘 30일 받았으면 내일"
    days = (y - SS.BACKFILL_FROM).days + 1
    assert days == 203 and -(-days // SS.BACKFILL_PER_DAY) == 7, days
    ok("10.19부터 거꾸로 29일(오늘 1건) · 있는 날 건너뜀 · 4/1 끝 · 하루 30일 상한 · 4/1~10/20 = 203일 → 약 7일")


def t8_backfill_collect():
    print("[T8] ① 일별 소급 — 같은 세션 하루씩 조회·저장 · 조회 사이 쉼 · 실패면 그 계정 소급 중단(비치명·로그)")
    from coupang_analytics import collector
    from coupang_analytics import pipeline_sales as PS
    calls, sleeps, logs = [], [], []

    def fake(page, f, t, log=None):
        calls.append((f, t))
        if len(calls) == 4:
            raise collector.SalesFetchError("status=500")
        return {"v1": _m("v1", "캠핑테이블", "RFM", 1, 1000)} if f.endswith("7") else {}
    orig = (collector.fetch_sales_details, PS.time.sleep)
    try:
        collector.fetch_sales_details, PS.time.sleep = fake, sleeps.append
        with tempfile.TemporaryDirectory() as d:
            SH.record(d, "a1", "2026-10-19", "2026-10-19", {})
            PS._record_sales_backfill("PAGE", "a1", d, logs.append, yesterday=date(2026, 10, 20))
            got = SH.load_daily(d)["a1"]
    finally:
        collector.fetch_sales_details, PS.time.sleep = orig
    assert [f for f, _t in calls] == ["2026-10-20", "2026-10-18", "2026-10-17", "2026-10-16"], calls
    assert all(f == t for f, t in calls), "하루 단위 조회"
    assert sorted(got) == ["2026-10-17", "2026-10-18", "2026-10-19", "2026-10-20"], sorted(got)
    assert got["2026-10-18"]["합계"]["판매금액"] == 0 and got["2026-10-17"]["합계"]["판매금액"] == 1000, \
        "판매 0 인 날도 기록(미수집과 구분)"
    assert len(sleeps) == 3 and all(1.0 <= x <= 2.5 for x in sleeps), sleeps
    assert any("중단" in m for m in logs) and any("일별 소급 3/" in m for m in logs), logs
    ok("어제부터 거꾸로·있는 날(10.19) 건너뜀 · 판매 0 날 기록 · 1~2.5초 쉼 · 4번째 실패에서 중단(3일 저장)")


def t9_scope():
    print("[T9] 관리 상품 기준 — 결과 사업자 시트 블록 VID(판매중지 포함)만 · 결과 파일에 없는 계정 제외")
    with tempfile.TemporaryDirectory() as d:
        wb = _wb_with_accounts(Path(d) / "m.xlsx")
        wb.set_product_vids("에벤에셀", "에벤에셀 상품", ["v1"])
        wb.ensure_product_block("에벤에셀", "옛 상품", config.KIND_CONTRACT, ["kw"], registered="옛 상품")
        wb.set_product_vids("에벤에셀", "옛 상품", ["v3"])
        wb.set_discontinued("에벤에셀", "옛 상품", True)
        managed = WS.managed_vids(wb)
    assert managed["a1"] == {"v1", "v3"} and managed["a2"] == {"v9"}, managed
    store = {"a1": {"2026-10-14": {"옵션": {"v1": {"판매금액": 1}, "v2": {"판매금액": 9}, "v3": {"판매금액": 2}}}},
             "zz": {"2026-10-14": {"옵션": {"v9": {}}}}}
    got = SS.scope(store, managed)
    assert set(got) == {"a1"} and set(got["a1"]["2026-10-14"]["옵션"]) == {"v1", "v3"}, got
    assert set(SS.scope(store, None)) == {"a1", "zz"}, "None = 거르지 않음"
    ok("a1 = v1 + 판매중지 v3 · 관리 안 하는 v2 제외 · 결과 파일에 없는 계정 zz 제외")


def _dash_fixture(d):
    """a1 = 9월 하루 1,000원·10월 하루 1,500원(관리 v1) + 관리 안 하는 vx 999원 · a2 = 9/1~10/15 하루 2,000원
    (10/16~ 미수집) · a1 1월 기간 합계 50,000원."""
    day = date(2026, 9, 1)
    while day <= date(2026, 10, 20):
        amt = 1000 if day.month == 9 else 1500
        SH.record(d, "a1", day.isoformat(), day.isoformat(),
                  {"v1": _m("v1", "캠핑테이블", "RFM", 1, amt), "vx": _m("vx", "남의상품", "NORMAL", 1, 999)})
        if day <= date(2026, 10, 15):
            SH.record(d, "a2", day.isoformat(), day.isoformat(), {"v2": _m("v2", "기저귀백팩", "NORMAL", 2, 2000)})
        day += timedelta(days=1)
    SH.record_period(d, "a1", "M:2026-01", "2026-01-01", "2026-01-31", {"v1": _m("v1", "캠핑테이블", "RFM", 5, 50000)})


def t10_dashboard():
    print("[T10] 대시보드 값 — 전일·지난주(그 전주 대비)·이번 달 누계(지난달 같은 기간 대비)·월별·주간·일별·계정 순위·상품")
    y = date(2026, 10, 20)   # 화요일
    assert SD.prev_month_same(y) == (date(2026, 9, 1), date(2026, 9, 20))
    assert SD.prev_month_same(date(2026, 3, 31)) == (date(2026, 2, 1), date(2026, 2, 28)), "지난달이 짧으면 말일"
    assert SD.last_full_week(y) == date(2026, 10, 12) and SD.last_full_week(date(2026, 10, 18)) == date(2026, 10, 12)
    labels = {"a1": "에벤에셀 (a1)", "a2": "니코에이블 (a2)"}
    with tempfile.TemporaryDirectory() as d:
        _dash_fixture(d)
        managed = {"a1": {"v1"}, "a2": {"v2"}}
        daily, per = SS.scope(SH.load_daily(d), managed), SS.scope(SH.load_periods(d), managed)
    b = SD.build(daily, per, labels, y)
    k_day, k_week, k_month = b.kpis
    assert (k_day.amount, k_day.missing, b.missing_today) == (1500, 1, 1), k_day
    assert k_week.span == "10.12~10.18" and k_week.amount == 10500 and k_week.missing == 1, k_week
    assert k_week.change == 0.0, "그 전주 10.05~10.11 도 a1 10,500 (a2 는 이번 주 미수집이라 비교에서 뺌)"
    assert k_month.amount == 30000 and abs(k_month.change - 0.5) < 1e-9, k_month       # 30,000 vs 9/1~20 20,000
    assert "09.01~09.20" in k_month.basis
    months = {p.label: p for p in b.months}
    assert list(months)[0] == "1월" and list(months)[-1] == "10월" and len(months) == 10
    assert months["1월"].amount == 50000 and months["1월"].missing == 1, "1월 = 쿠팡 기간 합계(일별 없음)"
    assert months["9월"].amount == 30 * 1000 + 30 * 2000 and months["9월"].missing == 0
    assert months["10월"].partial and months["10월"].amount == 30000 and months["10월"].missing == 1
    assert months["5월"].amount == 0 and months["5월"].missing == 2, "자료 없는 달 = 미수집"
    assert [p.label for p in b.weeks][-1] == "10.12" and len(b.weeks) == 8
    assert len(b.days) == 31 and b.days[-1].label == "10.20" and b.days[0].label == "09.20", "일별 = 오래된 → 최신"
    assert b.days[-1].amount == 1500, "관리 안 하는 vx 999원 빠짐"
    r1, r2 = b.accounts
    assert (r1.aid, r1.amount, r1.prev, r1.share) == ("a1", 30000, 20000, 1.0) and abs(r1.change - 0.5) < 1e-9
    assert r1.spark == "··▆▆▆▇██", r1.spark      # 8/24·8/31 주 = 일별 빠짐(9/1 시작) → 미수집
    assert r2.aid == "a2" and r2.amount is None and r2.change is None, "미수집 계정 = 맨 아래"
    assert [(p.account, p.product) for p in b.products] == [("니코에이블 (a2)", "기저귀백팩"), ("에벤에셀 (a1)", "캠핑테이블")]
    assert b.products[0].weeks[-1] is None and b.products[0].weeks[2] == 14000, b.products[0].weeks
    ok("전일 1,500 · 지난주 10,500(0%) · 이번 달 30,000(+50%) · 1월 기간합계 · 10월 진행 중 · 순위 a1·a2 미수집 · ··▆▆▆▇██")


def t11_charts_persist():
    print("[T11] 엑셀 차트 3개 — 숨긴 열 자료·다시 열어도(③ 등이 마스터를 다시 열고 저장) 차트 유지·중복 없음")
    from openpyxl.utils import get_column_letter
    with tempfile.TemporaryDirectory() as d:
        _dash_fixture(d)
        path = Path(d) / "m.xlsx"
        wb = _wb_with_accounts(path)
        wb.set_product_vids("에벤에셀", "에벤에셀 상품", ["v1"])
        wb.set_product_vids("니코에이블", "니코에이블 상품", ["v2"])
        WS.refresh_sales_sheet(wb, d, today=date(2026, 10, 21))
        ws = wb.wb[WS.SALES_SHEET]
        assert len(ws._charts) == 3
        specs = WD.read_specs(ws)
        data_cols = {c for sp in specs for c in [sp["cat"][2]] + [s[2] for s in sp["series"]]}
        assert all(ws.column_dimensions[get_column_letter(c)].hidden for c in data_cols), "차트 자료 = 숨긴 열"
        assert min(data_cols) > 3 + 2 * 31, "상세표 오른쪽"
        wb.save(path)
        wb2 = OutputWorkbook.load(path)
        assert len(wb2.wb[WS.SALES_SHEET]._charts) == 3, "다시 열면 설정으로 다시 붙임"
        wb2.save(path)
        wb3 = OutputWorkbook.load(path)
        assert len(wb3.wb[WS.SALES_SHEET]._charts) == 3, "열고 저장을 거듭해도 3개"
        cells = list(wb3.wb[WS.SALES_SHEET].iter_rows(values_only=True))
    head = next(r for r in cells if r[0] == "계정" and r[3] == "10.20")
    assert head[3:6] == ("10.20", None, "10.19") and head[3 + 2 * 30] == "09.20", "상세 = 최신이 왼쪽 31일"
    tot = next(r for r in cells if r[0] == "전체 합계")
    kpi = next(r for r in cells if str(r[0]).startswith("전일 "))
    assert tot[4] == kpi[1] == 1500, (tot[4], kpi[1])
    rank = [r[0] for r in cells if isinstance(r[1], str) and r[1] and set(r[1]) <= set(SD.SPARK + SD.SPARK_MISSING)]
    assert rank == ["에벤에셀 (a1)", "니코에이블 (a2)"], rank
    assert any(v == "▲ 50.0%" for r in cells for v in r), "증감 표시"
    fill, color = WD.heat_style(90, 100)
    assert fill.fgColor.rgb.endswith("2F75B5") and color == "FFFFFF" and WD.heat_style(None, 100) == (None, "9AA7B6")
    ok("차트 3개 · 자료 = 상세표 오른쪽 숨긴 열 · 다시 열고 저장 2회에도 3개 · 상세 10.20 → 09.20 · 전일 = 상세 합계")


def main() -> int:
    t1_plan()
    t2_record_period()
    t3_build()
    t4_sheet()
    t5_gsheet()
    t6_collect()
    t7_backfill_plan()
    t8_backfill_collect()
    t9_scope()
    t10_dashboard()
    t11_charts_persist()
    print("통합 판매현황 오프라인 검증 통과")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
