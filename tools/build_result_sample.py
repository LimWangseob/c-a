"""결과/정산 파일(파일④) 샘플 생성 — 앱이 매일 산출하는 통계+정산 파일의 구조 예시.

실데이터가 아니라 **포맷 설명용 샘플**이다(통계=상품×키워드 노출순위·판매 시계열, 정산=계약대비 정산현황).
실제 파일은 파이프라인(workbook.py)이 매 실행 생성·누적한다. 이 도구는 틀만 보여준다.

    python tools/build_result_sample.py [출력.xlsx]
"""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

_DL = Path.home() / "Downloads"
DEF_OUT = _DL / "결과정산_샘플.xlsx"

_HDR = PatternFill("solid", fgColor="2E75B6")
_IDF = PatternFill("solid", fgColor="548235")
_SUB = PatternFill("solid", fgColor="D9E1F2")
_UP = PatternFill("solid", fgColor="C6EFCE")
_DN = PatternFill("solid", fgColor="FFC7CE")
_WHITE = Font(color="FFFFFF", bold=True)
_RIGHT = Alignment(horizontal="right")
_CTR = Alignment(horizontal="center")
_BORDER = Border(*[Side(style="thin", color="BFBFBF")] * 4)
_NUM = "#,##0"


def _hrow(ws, r, headers, fill=_HDR):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(r, c, h)
        cell.fill = fill
        cell.font = _WHITE
        cell.alignment = _CTR
        cell.border = _BORDER


def _autofit(ws):
    for c in range(1, (ws.max_column or 0) + 1):
        ml = max((len(str(ws.cell(r, c).value or "")) for r in range(1, (ws.max_row or 0) + 1)), default=6)
        ws.column_dimensions[get_column_letter(c)].width = min(44, max(9, int(ml * 1.7) + 2))


def build(path: Path) -> None:
    wb = openpyxl.Workbook()

    # ── 안내 ──
    gd = wb.active
    gd.title = "안내"
    gd.sheet_properties.tabColor = "BFBFBF"
    notes = [
        "결과/정산 파일(파일④) — 앱이 매일 자동 생성/누적 (※ 이 파일은 포맷 설명용 '샘플', 실데이터 아님)",
        "",
        "[계정목록] 전 계정/상품 한눈 요약 미러 — 사업자·상품·계정·상태·체험단효과. 상품명=쿠팡 노출상품 링크.",
        "[통계_OOO]  사업자별 시트 — 상품 블록마다 키워드 노출순위 + 판매/방문/재고 지표를 '날짜 가로'로 누적.",
        "[정산현황]  계약 대비 정산 샘플 — 지급일·매출·원가·수익·배분(갑/을)·달성률. 계약 조건은 운영대장 '계약'에서 연동.",
        "",
        "원본 입력=운영대장(사람) · 물류=재고현황 · 이력=셀독등록원장. 결과는 읽기 전용(앱 산출).",
    ]
    for i, t in enumerate(notes, 1):
        gd.cell(i, 1, t)
    gd.column_dimensions["A"].width = 110

    # ── 계정목록 ──
    al = wb.create_sheet("계정목록")
    al.sheet_properties.tabColor = "4472C4"
    heads = ["대표자", "사업자명", "계정ID", "상품(물류)명", "체험단판매가", "체험단리뷰", "체험단택배", "상태", "체험단효과"]
    _hrow(al, 1, heads)
    for c, h in enumerate(heads, 1):
        if h in ("대표자", "사업자명", "계정ID"):
            al.cell(1, c).fill = _IDF
    rows = [
        ["송창호", "(주)웰빙곳간", "wellbing1107", "웰빙곳간 활력 볶은 맥문동 환 프리미엄 30포", "", "", "", "판매중", "판매 +38% · 순위 32→18 ↑"],
        ["조진형", "플랜잇", "plan_it", "구강세정기 B31", "계약자", "계약자", "토탈셀러", "판매중", "판매 +12% · 순위 45→40 ↑"],
        ["안재영", "휴라엘", "sg0141n", "신형타프 R008 (블랙 Free)", "", "", "", "판매중지", ""],
    ]
    band = ["FFF2CC", "DDEBF7", "E2EFDA"]
    for i, row in enumerate(rows):
        r = i + 2
        for c, v in enumerate(row, 1):
            cell = al.cell(r, c, v)
            cell.fill = PatternFill("solid", fgColor=band[i % 3])
            cell.border = _BORDER
            if c == 4 and row[7] != "판매중지":
                cell.hyperlink = "https://www.coupang.com/vp/products/0?vendorItemId=0000000000"
                cell.font = Font(color="0563C1", underline="single")
            if c == 9 and "↑" in str(v):
                cell.fill = _UP
    al.freeze_panes = "A2"
    _autofit(al)

    # ── 통계_웰빙곳간 ──
    st = wb.create_sheet("통계_웰빙곳간")
    st.sheet_properties.tabColor = "70AD47"
    dates = ["09.27", "09.28", "09.29", "09.30"]
    _hrow(st, 1, ["구분", "키워드/지표", "", "", "", ""] + dates)
    st.cell(2, 1, "웰빙곳간 활력 볶은 맥문동 환 프리미엄 30포").font = Font(bold=True)
    st.cell(2, 1).fill = PatternFill("solid", fgColor="FCE4D6")
    kw = [("노출순위", "맥문동환", [32, 25, 20, 18]), ("노출순위", "맥문동 환", [40, 38, 33, 30]),
          ("노출순위", "활력환", [55, 50, 48, 45])]
    r = 3
    for _k, name, vals in kw:
        st.cell(r, 2, name)
        for j, v in enumerate(vals):
            cell = st.cell(r, 7 + j, v)
            cell.alignment = _RIGHT
            cell.number_format = _NUM
        r += 1
    metrics = [("재고현황", [1200, 1150, 1090, 1020]), ("판매가", [19900, 19900, 19900, 18900]),
               ("판매상태", ["판매중", "판매중", "판매중", "판매중"]), ("노출수", [3200, 3550, 4100, 4800]),
               ("판매건수", [12, 15, 19, 24]), ("방문자수", [210, 245, 300, 360])]
    for name, vals in metrics:
        cell0 = st.cell(r, 2, name)
        cell0.fill = _SUB
        for j, v in enumerate(vals):
            cell = st.cell(r, 7 + j, v)
            if isinstance(v, int):
                cell.alignment = _RIGHT
                cell.number_format = _NUM
        r += 1
    st.freeze_panes = "G2"
    _autofit(st)

    # ── 정산현황 ──
    se = wb.create_sheet("정산현황")
    se.sheet_properties.tabColor = "C00000"
    heads = ["사업자명", "상품(물류)명", "귀속월", "지급일", "매출", "원가", "판매수수료", "광고비",
             "배송/입출고", "수익(순액)", "배분기준", "갑(위탁자)", "을(수탁자)", "목표수익률", "달성률"]
    _hrow(se, 1, heads)
    data = [
        ["효성", "파미젠 와사비잎 추출물 plus max 120정", "2026-09", "2026-10-15",
         8400000, 3000000, 1008000, 900000, 420000, 3072000, "판매수익 갑70/을30", 2150400, 921600, "30%", "36.6%"],
        ["(주)웰빙곳간", "웰빙곳간 활력 볶은 맥문동 환 프리미엄 30포", "2026-09", "2026-10-15",
         4300000, 1500000, 516000, 350000, 180000, 1754000, "순이익 갑70/을30", 1227800, 526200, "-", "-"],
    ]
    for i, row in enumerate(data):
        r = i + 2
        for c, v in enumerate(row, 1):
            cell = se.cell(r, c, v)
            cell.border = _BORDER
            if isinstance(v, int):
                cell.alignment = _RIGHT
                cell.number_format = _NUM
            if c == 15 and isinstance(v, str) and v.endswith("%") and v != "-":
                cell.fill = _UP if float(v[:-1]) >= 30 else _DN
    se.freeze_panes = "A2"
    _autofit(se)

    wb.save(path)


def main() -> int:
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else DEF_OUT
    build(out)
    print(f"[결과/정산 샘플] 생성: {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
