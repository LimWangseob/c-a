"""재고현황 파일(파일②) 템플릿 생성 — 물류팀이 관리하는 '제고(재고)' 입력 파일의 포맷.

4파일 모델(SSOT=designs/OPERATION_DATA_MODEL.md §2·§5) 중 파일②.
- 물류팀 편집·비번 격리. 운영대장이 물류명/바코드로 상품코드 매핑 + 앱이 현재고 미러(읽기).
- 구조는 migrate_to_operation_ledger.read_stock_rows 가 **열 위치로** 읽는다(아래 순서 고정):
  Sheet1 · 1행=제목('재고현황 YY.MM.DD 기준') · 2행=헤더 · 3행~데이터.
  열: 1창고 · 2구분 · 3바코드 · 4물류명 · 5현재고 · 6셀독 · 7당근 · 8자사.

    python tools/build_inventory_template.py [출력.xlsx]
"""
from __future__ import annotations

import sys
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

_DL = Path.home() / "Downloads"
DEF_OUT = _DL / "재고현황_템플릿.xlsx"

HEADERS = ["창고", "구분", "바코드", "물류명", "현재고", "셀독", "당근", "자사"]
# (포맷 예시 행 — 실데이터 아님. 물류팀이 실제 재고로 채운다.)
SAMPLE = [
    ["검단", "건기식", "8800000000001", "샘플 상품명 A 60캡슐", 1200, "셀독", "당근", ""],
    ["김포1", "공산품", "8800000000002", "샘플 상품명 B R010", 540, "셀독", "", ""],
    ["검단", "건기식", "8800000000003", "샘플 상품명 C 120정", 0, "셀독", "당근", "자사"],
]

_HDR = PatternFill("solid", fgColor="2E75B6")
_TITLE = PatternFill("solid", fgColor="D9E1F2")
_WHITE = Font(color="FFFFFF", bold=True)
_CTR = Alignment(horizontal="center", vertical="center")
_RIGHT = Alignment(horizontal="right")
_THIN = Border(*(Side(style="thin", color="BFBFBF"),) * 4)


def _autofit(ws, start_row: int) -> None:
    for c in range(1, len(HEADERS) + 1):
        lens = [len(str(ws.cell(r, c).value or "").split("\n")[0])
                for r in range(start_row, (ws.max_row or 0) + 1)]
        ml = max(lens, default=6)
        cap = 12 if HEADERS[c - 1] in ("현재고",) else 40
        ws.column_dimensions[get_column_letter(c)].width = min(cap, max(8, int(ml * 1.6) + 2))


def build(path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"                       # migrate 가 wb["Sheet1"] 로 읽음(고정)

    ws.cell(1, 1, f"재고현황 {date.today():%y.%m.%d} 기준").font = Font(bold=True, size=12)
    ws.cell(1, 1).fill = _TITLE

    for c, h in enumerate(HEADERS, 1):        # 2행 헤더
        cell = ws.cell(2, c, h)
        cell.fill = _HDR
        cell.font = _WHITE
        cell.alignment = _CTR
        cell.border = _THIN

    for i, row in enumerate(SAMPLE, 3):       # 3행~ 데이터(예시)
        for c, v in enumerate(row, 1):
            cell = ws.cell(i, c, v)
            cell.border = _THIN
            if HEADERS[c - 1] == "현재고" and isinstance(v, int):
                cell.alignment = _RIGHT
                cell.number_format = "#,##0"

    ws.freeze_panes = "A3"
    _autofit(ws, 2)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def main() -> int:
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else DEF_OUT
    p = build(out)
    print(f"[재고현황 템플릿(파일②)] 생성: {p}")
    print("  열: " + " · ".join(HEADERS) + "  (Sheet1 · 1행 제목 · 2행 헤더 · 3행~데이터)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
