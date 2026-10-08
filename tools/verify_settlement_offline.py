"""정산 계산 모듈 오프라인 검증 — designs/coupang_golden_cases.json 100% + 규칙 핀. SSOT=designs/SETTLEMENT_MODULE.md §6.

로그인·공휴일 API·정산 API 호출 없음(결정적). 공휴일은 골든 `_meta.holidays_used` 주입.
실패 시 AssertionError → exit 1. run_checks 게이트 편입 대상.

    python tools/verify_settlement_offline.py
"""
from __future__ import annotations

import json
import sys
import tempfile
from datetime import date, timedelta
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
try:
    sys.stdout.reconfigure(encoding="utf-8")
except AttributeError:
    pass

from coupang_analytics import holiday_kr as HK  # noqa: E402
from coupang_analytics import payout as P  # noqa: E402
from coupang_analytics import settlement_amount as SA  # noqa: E402
from coupang_analytics import settlement_parse as SP  # noqa: E402

GOLDEN = json.loads((ROOT / "designs" / "coupang_golden_cases.json").read_text(encoding="utf-8"))
HOL = {date.fromisoformat(d) for d in GOLDEN["_meta"]["holidays_used"]}
D = date.fromisoformat


def ok(msg):
    print(f"  ✔ {msg}")


def expect(exc, fn, what):
    try:
        fn()
    except exc:
        return
    raise AssertionError(f"{what} — {exc.__name__} 안 남")


# ── 골든 ──────────────────────────────────────────────────────────
def g1_payout_dates():
    print("[G1] 지급일 골든 12케이스")
    cases = GOLDEN["payout_date"]
    for c in cases:
        week = c.get("confirm_week") or c.get("sales_week")
        got = P.payout_date(c["policy"], week_end=D(week[1]) if week else None,
                            revenue_month=c.get("revenue_month"), holidays=HOL)
        assert got == D(c["expected"]), (c, got)
    assert len(cases) == 12, len(cases)
    ok(f"{len(cases)}/{len(cases)} 일치(마감일+15/20영업일·월 경계 분할=같은 주 일요일·MP FINAL 익익월 1일 무보정)")


def g2_amounts():
    print("[G2] 금액 골든 2케이스")
    by = {c["policy"]: c for c in GOLDEN["amount_formula"]}
    m = by["MP_REVENUE_REPORT"]
    r = SA.mp_revenue(sales=m["A_sales"], seller_coupon=m["B_seller_coupon"], fee=m["C_fee"],
                      myshop_fee=m["D_myshop_fee"], deduction=m["F_deduction"])
    assert (r.settlement_target, r.final_payout) == (m["expected_E"], m["expected_final"]), r
    h = by["RG_HOME_PROFIT"]
    p = SA.rg_home_profit(sales=h["sales"], seller_coupon=h["seller_coupon"], fee=h["fee"],
                          fulfillment=h["fulfillment"], ads=h["ads"], review_event=h["review_event"],
                          inventory_compensation=h["inventory_compensation"])
    assert (p.cost, p.profit, p.margin_pct) == (h["expected_cost"], h["expected_profit"], h["expected_margin_pct"]), p
    assert len(by) == 2
    ok("윙 E·최종 / RG 비용·이익·마진% 일치")


def g3_invariants():
    print("[G3] 판매분석 불변식 골든 4행(음수 행 통과)")
    rows = GOLDEN["insights_row_invariants"]
    for c in rows:
        row = SimpleNamespace(**{k: c[k] for k in ("sales", "orders", "units", "views", "total_sales",
                                                   "total_units", "cancel_amount", "cancel_units")},
                              cvr_pct=Decimal(str(c["cvr_pct"])))
        assert SP.check_invariants(row) == [], (c, SP.check_invariants(row))
    assert len(rows) == 4 and len(GOLDEN["invariant_rules"]) == 3
    bad = SimpleNamespace(sales=100, total_sales=90, cancel_amount=0, units=2, total_units=2, cancel_units=0,
                          orders=1, views=3, cvr_pct=Decimal("33.00"))
    assert len(SP.check_invariants(bad)) == 2                          # 매출식 + 전환율(33.33) 위반
    ok("골든 4행 위반 0(판매자배송 -35000 포함)·3식 적용·위반 행은 잡음")


# ── 규칙 핀 ───────────────────────────────────────────────────────
def p1_holidays():
    print("[P1] 영업일·대체공휴일·공휴일 소스 실패")
    assert HK.add_business_days(D("2026-08-09"), 1, HOL) == D("2026-08-10")
    assert HK.next_business_day(D("2026-09-26"), HOL) == D("2026-09-28")        # 추석 토 → 월
    assert HK.next_business_day(D("2026-09-28"), HOL) == D("2026-09-28")        # 영업일 자신
    expect(ValueError, lambda: HK.add_business_days(D("2026-08-09"), 0, HOL), "0영업일")
    subs = HK.substitute_holidays({D("2026-08-15"): HK.KIND_NATIONAL,           # 광복절 토 → 8/17
                                   D("2026-09-24"): HK.KIND_LUNAR, D("2026-09-25"): HK.KIND_LUNAR,
                                   D("2026-09-26"): HK.KIND_LUNAR})             # 추석 목~토 → 대체 없음
    assert subs == {D("2026-08-17")}, subs
    s25 = HK.substitute_holidays({D("2025-05-05"): (HK.KIND_NATIONAL, HK.KIND_NATIONAL),   # 어린이날+부처님오신날
                                  D("2025-10-03"): HK.KIND_NATIONAL, D("2025-10-05"): HK.KIND_LUNAR,
                                  D("2025-10-06"): HK.KIND_LUNAR, D("2025-10-07"): HK.KIND_LUNAR})
    assert s25 == {D("2025-05-06"), D("2025-10-08")}, s25                     # 겹침→5/6·추석 일요일→10/8
    assert HK.substitute_holidays({D("2027-02-06"): HK.KIND_LUNAR, D("2027-02-07"): HK.KIND_LUNAR,
                                   D("2027-02-08"): HK.KIND_LUNAR}) == {D("2027-02-09")}   # 설 토·일 → 일요일분 1일
    with tempfile.TemporaryDirectory() as tmp:
        expect(HK.HolidaySourceError, lambda: HK.holidays(2026, cache_dir=tmp), "소스·캐시 없음")
        expect(HK.HolidaySourceError, lambda: HK.holidays(2026, cache_dir=tmp, fetch=lambda y: []), "빈 결과")
        expect(HK.HolidaySourceError, lambda: HK.holidays(2026, cache_dir=tmp, fetch=lambda y: 1 / 0), "조회 예외")
        expect(HK.HolidaySourceError, lambda: HK.holidays(2026, cache_dir=tmp, fetch=lambda y: ["2025-01-01"]),
               "다른 해")
        calls = []
        got = HK.holidays(2026, cache_dir=tmp, fetch=lambda y: calls.append(y) or ["2026-08-17", "2026-09-24"])
        again = HK.holidays(2026, cache_dir=tmp, fetch=lambda y: calls.append(y) or [], extra=["2026-06-03"])
        assert got == {D("2026-08-17"), D("2026-09-24")} and calls == [2026]           # 두 번째는 캐시
        assert again == got | {D("2026-06-03")}                                           # 임시공휴일 수동 추가
    ok("영업일 가산·다음 영업일·대체공휴일(광복절 토·추석 토 무대체·겹침·추석 일요일·설 토일)·실패=예외·캐시·수동 추가")


def p2_payout_rules():
    print("[P2] 지급 규칙 — 골든 밖 정책·인자 검증·행 단위 반올림")
    assert P.week_sunday(D("2026-08-31")) == D("2026-09-06") and P.week_sunday(D("2026-09-06")) == D("2026-09-06")
    assert P.payout_date(P.PAYOUT_MP_MONTHLY, revenue_month="2026-08", holidays=HOL) == \
        HK.add_business_days(D("2026-08-31"), 15, HOL)
    assert P.payout_date(P.PAYOUT_RG_MONTHLY, revenue_month="2026-08", holidays=HOL) == \
        HK.add_business_days(D("2026-08-31"), 20, HOL)
    assert P.payout_date(P.PAYOUT_RG_WEEKLY_FINAL, revenue_month="2025-01", holidays=HOL) == D("2025-03-03")  # 토→월
    assert P.payout_date(P.PAYOUT_RG_WEEKLY_FINAL, revenue_month="2024-11",
                         holidays={D("2025-01-01")}) == D("2025-01-02")          # 익익월 1일 공휴일 → 다음 영업일
    expect(ValueError, lambda: P.payout_date("PAYOUT_X", week_end=D("2026-08-09"), holidays=HOL), "미지 정책")
    expect(ValueError, lambda: P.payout_date(P.PAYOUT_MP_WEEKLY_1ST, holidays=HOL), "week_end 누락")
    expect(ValueError, lambda: P.payout_date(P.PAYOUT_MP_MONTHLY, holidays=HOL), "revenue_month 누락")
    expect(ValueError, lambda: P.payout_date(P.PAYOUT_MP_MONTHLY, revenue_month="2026/08", holidays=HOL), "월 형식")
    # 실측 규칙(2026-10-06): 윙 70%=전체×70% · 로켓그로스 70%=줄별 70% 사사오입 합 · 30%=전체−70%
    assert P.payout_amount(P.PAYOUT_MP_WEEKLY_1ST, [15, 25]) == 28             # 40×0.7 (줄별이면 29)
    assert P.payout_amount(P.PAYOUT_MP_WEEKLY_FINAL, [15, 25]) == 40 - (11 + 18)  # 윙 최종액=주별 합계마다 70% 뺀 나머지
    assert P.payout_amount(P.PAYOUT_MP_WEEKLY_FINAL, [49721, 16364, 11191, 337770]) == 124513   # 실측 8월 최종액 화면값
    assert P.payout_amount(P.PAYOUT_RG_WEEKLY_1ST, [15, 25]) == 11 + 18        # 10.5→11·17.5→18(줄별)
    assert P.payout_amount(P.PAYOUT_RG_WEEKLY_FINAL, [15, 25]) == 40 - 29      # 나머지(줄별 30% 합 5+8=13 아님)
    assert P.payout_amount(P.PAYOUT_RG_MONTHLY, [15, 25]) == 40
    assert P.payout_amount(P.PAYOUT_MP_WEEKLY_1ST, [337770]) == 236439          # 실측 윙 8/24~30 주 화면값
    expect(TypeError, lambda: P.payout_amount(P.PAYOUT_MP_WEEKLY_1ST, [1.5]), "정수 아닌 금액")
    e = P.estimate(P.PAYOUT_RG_WEEKLY_1ST, week_end=D("2026-08-30"), holidays=HOL, row_amounts=[100])
    assert (e.date, e.amount, e.is_estimate) == (D("2026-09-29"), 70, True)
    ok("월정산·RG 최종액(첫 영업일)·인자 검증·70/30 행 단위 사사오입·예측 표기")


def p3_amount_rules():
    print("[P3] 금액 규칙 — RG 실지급·매출 0 마진·정수 강제")
    assert SA.rg_payout(settlement_target=1000, milkrun=100, ads=200, cfs=50) == 650
    z = SA.rg_home_profit(sales=0, seller_coupon=0, fee=0, fulfillment=0, ads=10, review_event=0)
    assert (z.profit, z.margin_pct) == (-10, None)                              # 매출 0 = 마진 정의 불가(0% 아님)
    c = SA.rg_home_profit(sales=1000, seller_coupon=0, fee=0, fulfillment=0, ads=0, review_event=0,
                          inventory_compensation=50)
    assert (c.cost, c.profit, c.margin_pct) == (-50, 1050, 105.0)              # 재고손실보상 = 이익 가산
    h = SA.rg_home_profit(sales=400, seller_coupon=0, fee=0, fulfillment=0, ads=351, review_event=0)
    assert h.margin_pct == 12.3, h                                              # 12.25 → 사사오입 12.3(은행가 12.2 아님)
    expect(TypeError, lambda: SA.mp_revenue(sales="1000", seller_coupon=0, fee=0), "문자열 금액")
    expect(TypeError, lambda: SA.rg_payout(settlement_target=1.0, milkrun=0, ads=0, cfs=0), "소수 금액")
    ok("RG 실지급·매출 0 → 마진 None·재고손실보상 가산·정수 아닌 금액 거부")


HDR = list(SP.REQUIRED_HEADERS)


def _raw(**over):
    base = {"옵션 ID": "94632081597", "옵션명": "1개", "상품명": "예시", "등록상품ID": "1234567890",
            "카테고리": "건강", "판매방식": "로켓그로스", "매출(원)": "285,900", "주문": "12", "판매량": "13",
            "방문자": "105", "조회": "159", "장바구니": "20", "구매전환율": "7.55%",
            "아이템위너 비율(%)": "1.10000000000000008881", "총 매출(원)": "285900", "총 판매수": "13",
            "총 취소 금액(원)": "0", "총 취소된 상품수": "0", "즉시 취소된 상품수": "0"}
    base.update(over)
    return [base[h] for h in HDR]


def p4_parse():
    print("[P4] 판매분석 파싱 — 헤더 이름·전 셀 문자열·ID 문자열·실패=오류·메타 강제·불변식 경고")
    meta = {"period_start": "2026-09-01", "period_end": 46293, "account_id": "ex01"}      # 시리얼 46293=2026-09-28
    shuffled = list(reversed(HDR))
    rows = [shuffled, list(reversed(_raw())),
            list(reversed(_raw(**{"옵션 ID": "95907346795", "판매방식": "판매자배송", "매출(원)": "-35000",
                                  "주문": "-1", "판매량": "-1", "조회": "13", "구매전환율": "-7.69%",
                                  "총 매출(원)": "0", "총 판매수": "0", "총 취소 금액(원)": "-35000",
                                  "총 취소된 상품수": "-1"}))),
            [""] * len(HDR),
            list(reversed(_raw(**{"옵션 ID": "95222903297", "매출(원)": "100"})))]              # 매출식 위반
    logs: list = []
    f = SP.parse_insights_rows(rows, **meta, on_log=logs.append)
    assert f.meta == SP.InsightsMeta(D("2026-09-01"), D("2026-09-28"), "ex01")
    assert [r.option_id for r in f.rows] == ["94632081597", "95907346795", "95222903297"]    # 빈 행 건너뜀·열 순서 무관
    r0 = f.rows[0]
    assert (r0.sales, r0.cvr_pct, r0.item_winner_pct, r0.product_id) == (285900, Decimal("7.55"), Decimal("1.10"),
                                                                         "1234567890")
    assert f.rows[1].sales == -35000                                               # 음수 정상·필터링 없음
    assert len(f.warnings) == 1 and "95222903297" in f.warnings[0] and any("불변식" in m for m in logs)
    expect(SP.SettlementParseError, lambda: SP.parse_insights_rows(rows, period_start="2026-09-01", period_end="",
                                                                   account_id="ex01"), "메타 없음")
    expect(SP.SettlementParseError, lambda: SP.parse_insights_rows(rows, period_start="2026-09-30",
                                                                   period_end="2026-09-01", account_id="a"), "기간 역전")
    expect(SP.SettlementParseError, lambda: SP.parse_insights_rows([HDR[:-1], _raw()[:-1]], **meta), "헤더 누락")
    for bad in ({"매출(원)": "12a"}, {"주문": ""}, {"판매량": "1.5"}, {"구매전환율": "없음"}, {"옵션 ID": "9.46E+10"}):
        expect(SP.SettlementParseError, lambda b=bad: SP.parse_insights_rows([HDR, _raw(**b)], **meta), f"값 {bad}")
    try:
        SP.parse_insights_rows([HDR[1:], _raw()[1:]], **meta)
    except SP.SettlementParseError as exc:
        assert "옵션 ID" in str(exc)                                                # 누락 목록 명시
    assert SP.parse_date("2026.9.1") == D("2026-09-01") and SP.parse_date(45658) == D("2025-01-01")
    ok("열 순서 무관·빈 행·음수 행·ID/비율 문자열·메타 강제(시리얼 날짜)·헤더/값 실패 오류·불변식 위반=경고")


class _Resp:
    """requests.Response 대역(녹화 응답)."""
    def __init__(self, payload=None, *, status=200, text=None):
        self.status_code = status
        self._payload = payload
        self.text = text if text is not None else json.dumps(payload, ensure_ascii=False)

    def json(self):
        if self._payload is None:
            raise ValueError("not json")
        return self._payload


def _kasi(items, total=None, code="00"):
    body_items = "" if items == [] else {"item": items[0] if len(items) == 1 else items}
    return {"response": {"header": {"resultCode": code, "resultMsg": "NORMAL SERVICE."},
                         "body": {"items": body_items, "numOfRows": 100, "pageNo": 1,
                                  "totalCount": len(items) if total is None else total}}}


def p5_holiday_source():
    print("[P5] 공휴일 API(특일정보) — 녹화 응답 파싱·오류 코드·holiday_kr 연결(실 API 미호출)")
    from coupang_analytics import holiday_source as HS
    y2026 = [{"dateKind": "01", "dateName": "광복절", "isHoliday": "Y", "locdate": 20260815, "seq": 1},
             {"dateKind": "01", "dateName": "대체공휴일", "isHoliday": "Y", "locdate": 20260817, "seq": 1},
             {"dateKind": "01", "dateName": "추석", "isHoliday": "Y", "locdate": 20260924, "seq": 1},
             {"dateKind": "01", "dateName": "추석", "isHoliday": "Y", "locdate": 20260925, "seq": 1},
             {"dateKind": "01", "dateName": "추석", "isHoliday": "Y", "locdate": 20260926, "seq": 1},
             {"dateKind": "01", "dateName": "제헌절", "isHoliday": "N", "locdate": 20260717, "seq": 1}]
    sent = []

    def get(resp):
        def http_get(url, params=None, timeout=None):
            sent.append((url, dict(params)))
            return resp
        return http_get
    days = HS.make_fetch("KEY", http_get=get(_Resp(_kasi(y2026))))(2026)
    assert days == {D("2026-08-15"), D("2026-08-17"), D("2026-09-24"), D("2026-09-25"), D("2026-09-26")}, days
    url, params = sent[-1]
    assert url == HS.ENDPOINT and params["solYear"] == "2026" and params["_type"] == "json" and params["ServiceKey"] == "KEY"
    assert HS.make_fetch("K", http_get=get(_Resp(_kasi(y2026[:1]))))(2026) == {D("2026-08-15")}   # 1건 = dict 응답
    assert HS.parse_rest_days(_kasi([]), 2026) == set()                                          # 0건 = items ""
    xml = ("<OpenAPI_ServiceResponse><cmmMsgHeader><returnAuthMsg>SERVICE_KEY_IS_NOT_REGISTERED_ERROR"
           "</returnAuthMsg><returnReasonCode>30</returnReasonCode></cmmMsgHeader></OpenAPI_ServiceResponse>")
    bad = [("인증오류 XML", _Resp(None, text=xml)), ("HTTP 500", _Resp(None, status=500, text="err")),
           ("서비스 오류코드", _Resp(_kasi(y2026, code="22"))), ("건수 불일치", _Resp(_kasi(y2026, total=120))),
           ("JSON 아님", _Resp(None, text="oops")),
           ("다른 해", _Resp(_kasi([{"isHoliday": "Y", "locdate": 20250101}])))]
    for what, resp in bad:
        expect(HS.HolidayApiError, lambda r=resp: HS.make_fetch("K", http_get=get(r))(2026), what)
    for resp in (_Resp(None, text=xml), _Resp(None, status=500, text="err"), _Resp(None, text="oops")):
        try:
            HS.make_fetch("SECRET-9f3a", http_get=get(resp))(2026)
        except HS.HolidayApiError as exc:
            assert "SECRET-9f3a" not in str(exc), exc                            # 오류 문구에 키 노출 금지
    try:
        HS.make_fetch("K", http_get=get(_Resp(None, text=xml)))(2026)
    except HS.HolidayApiError as exc:
        msg = str(exc)                                                           # 원인 문구·코드를 풀어서 보여줌
        assert "인증" in msg and "SERVICE_KEY_IS_NOT_REGISTERED_ERROR" in msg and "코드 30" in msg, msg
    expect(HS.HolidayApiError, lambda: HS.make_fetch("  "), "키 없음")
    expect(HS.HolidayApiError, lambda: HS.fetch_from_store(SimpleNamespace(get_password=lambda k: None)), "저장 키 없음")
    with tempfile.TemporaryDirectory() as tmp:                                     # holiday_kr 연결: 실패=HolidaySourceError
        expect(HK.HolidaySourceError,
               lambda: HK.holidays(2026, cache_dir=tmp, fetch=HS.make_fetch("K", http_get=get(_Resp(None, text=xml)))),
               "API 오류 → 지급일 계산 중단")
        hol = HK.holidays(2026, cache_dir=tmp, fetch=HS.make_fetch("K", http_get=get(_Resp(_kasi(y2026)))))
        assert P.payout_date(P.PAYOUT_MP_WEEKLY_1ST, week_end=D("2026-08-09"), holidays=hol) == D("2026-08-31")
    ok("대체·연휴 포함·isHoliday=N 제외·1건 dict·0건·인증 XML/HTTP/코드/건수/형식/다른 해 오류·키 없음·홀리데이 연결")


def p6_diag_no_values():
    print("[P6] 정산 주소 실측 도구 — 구조만 기록·값(금액·이름·주문번호·질의값) 비저장")
    import importlib.util
    spec = importlib.util.spec_from_file_location("diag_se", ROOT / "tools" / "diag_settlement_endpoints.py")
    DG = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(DG)
    payload = {"data": [{"orderId": "31234567890", "productName": "비밀상품명", "settlementAmount": 3764453,
                         "detail": {"fee": 414947, "items": [{"vendorItemId": "94632081597"}]}}], "total": 7}
    e = DG.entry("POST", "https://wing.coupang.com/tenants/settlement/v2/search?vendorId=A00123&from=2026-09-01",
                 200, "fetch", "application/json;charset=UTF-8",
                 json.dumps({"vendorId": "A00123", "period": {"from": "2026-09-01"}}), payload)
    dumped = json.dumps(e, ensure_ascii=False)
    for secret in ("31234567890", "비밀상품명", "3764453", "414947", "94632081597", "A00123", "2026-09-01"):
        assert secret not in dumped, (secret, dumped)
    assert e["query"] == ["from", "vendorId"] and e["path"] == "/tenants/settlement/v2/search"
    assert e["response"]["data"]["item"]["detail"]["fee"] == "int" and e["response"]["data"]["__list__"] == 1
    assert e["request"] == {"vendorId": "str", "period": {"from": "str"}}
    assert DG.body_keys("a=1&b=secret") == {"__form__": ["a", "b"]} and DG.body_keys("rawtext") == {"__text_len__": 7}
    keyed = DG.shape({"94632081597": {"qty": 3}, "95222903297": {"qty": 1}, "2026-09-01": 5, "20260902": 6,
                      "status": "DONE"})
    assert keyed == {"<숫자키>": {"qty": "int"}, "<날짜키>": "int", "status": "str"}, keyed   # ID·날짜 키 비노출
    other = DG.entry("GET", "https://wing.coupang.com/tenants/common/menu", 200, "xhr", "text/html", None, None)
    lines = DG.summarize([other, e, e])
    assert len(lines) == 2 and lines[0].startswith("★") and "settlement" in lines[0]    # 정산 주소 먼저·중복 1줄
    ok("값 7종 비노출·질의 이름만·요청/응답 구조·폼/텍스트 본문·정산 주소 우선 요약·중복 제거")


# ── 정산 파일·요청 기록·집계 (실측 머리글 그대로의 가상 데이터 — 실파일은 개인정보라 저장소에 넣지 않음) ──
WING_HDR = ["주문번호", "과세유형", "상품 ID", "상품명", "옵션 ID", "옵션명", "판매가", "판매수량", "환불수량", "판매액",
            "판매자 할인쿠폰(A+B)", "판매자 할인쿠폰(A.즉시할인)", "판매자 할인쿠폰(B.다운로드)", "판매수수료",
            "서비스이용율(%,VAT별도)", "서비스이용료소급분", "마이샵수수료할인", "정산금액", "구매자명", "결제완료일",
            "배송완료일", "구매확정일", "취소완료일", "구매확정(출고)유형", "정산예정일"]
RG_HDR = ["정산유형", "정산주기(종료일)", "세금계산서 발행월", "발생일(결제완료일)", "매출인식일", "주문ID", "거래유형",
          "카테고리ID", "카테고리명", "과세유형", "등록상품 ID", "옵션ID", "SKU ID", "등록상품명", "옵션명", "판매가(A)",
          "판매수량(B)", "판매액(A*B)", "쿠팡지원할인(C)", "매출금액(A*B-C)", "즉시할인쿠폰(D)", "다운로드쿠폰(E)",
          "판매자할인쿠폰(D+E)", "정산대상액", "판매수수료율(%,VAT별도)", "할인적용 판매수수료율(%,VAT별도)", "판매수수료",
          "판매수수료 VAT"]


def _wing_row(order, pid, name, oid, price, qty, coupon, fee, settle, buyer="홍길동", conf="2026-08-27",
              due="2026-09-18", sales=None):
    v = dict.fromkeys(WING_HDR, "")
    v.update({"주문번호": order, "상품 ID": pid, "상품명": name, "옵션 ID": oid, "옵션명": name and "옵션",
              "판매가": float(price), "판매수량": float(qty), "환불수량": 0.0,
              "판매액": float(price * qty if sales is None else sales),
              "판매자 할인쿠폰(A+B)": float(coupon), "판매자 할인쿠폰(A.즉시할인)": float(coupon),
              "판매자 할인쿠폰(B.다운로드)": 0.0, "판매수수료": float(fee), "서비스이용율(%,VAT별도)": "10.8",
              "마이샵수수료할인": 0.0, "정산금액": float(settle), "구매자명": buyer if name else "",
              "구매확정일": conf if name else "", "정산예정일": due})
    return [v[h] for h in WING_HDR]


def _wing_grid():
    return [WING_HDR,
            _wing_row("1001", "77", "보냉백", "95467404277", 12700, 1, 0, 1372, 11328),
            _wing_row("1001", "", "", "<기본배송료>", 0, 0, 0, 0, 0),
            _wing_row("1001", "", "", "<추가배송료>", 0, 0, 0, 0, 0),
            _wing_row("1002", "88", "압축포장기", "95425233537", 37900, 1, 20000, 1933, 15967),
            _wing_row("1003", "", "", "<기본배송료>", 0, 0, 0, 0, 6000, sales=6000),     # 실측 모양: 판매가·수량 0·금액만
            _wing_row("1004", "77", "보냉백", "95467404277", 12700, 2, 0, 2744, 99999)]   # 검산 위반 = 경고


def _rg_row(order, kind, pid, oid, name, price, qty, settle_target, fee, vat, rec="2026-08-05", cycle="2026-08-09"):
    v = dict.fromkeys(RG_HDR, "")
    v.update({"정산유형": "주정산", "정산주기(종료일)": cycle, "세금계산서 발행월": "2026-08", "매출인식일": rec,
              "주문ID": order, "거래유형": kind, "등록상품 ID": pid, "옵션ID": oid, "SKU ID": "5" + oid[-4:],
              "등록상품명": name, "옵션명": name, "판매가(A)": float(price), "판매수량(B)": float(qty),
              "판매액(A*B)": float(price * qty), "쿠팡지원할인(C)": 0.0, "매출금액(A*B-C)": float(price * qty),
              "즉시할인쿠폰(D)": 0.0, "다운로드쿠폰(E)": 0.0, "판매자할인쿠폰(D+E)": 0.0, "정산대상액": float(settle_target),
              "판매수수료율(%,VAT별도)": 10.8, "할인적용 판매수수료율(%,VAT별도)": 10.8, "판매수수료": float(fee),
              "판매수수료 VAT": float(vat)})
    return [v[h] for h in RG_HDR]


def _rg_grid(cycle="2026-08-09"):
    return [[None] * len(RG_HDR), RG_HDR,
            _rg_row("2001", "주문 정산", "77", "95222903297", "보냉백", 9310, 1, 8205, 1005, 100, cycle=cycle),
            _rg_row("2002", "주문 정산취소", "77", "95222903297", "보냉백", 9310, -1, -8205, -1005, -100, cycle=cycle)]


def p7_settle_files():
    print("[P7] 정산 파일 읽기 — 실측 머리글·0원 배송비 줄 제외·검산 경고·기간 대조·구매자명 삭제·파일 이름 규칙")
    from coupang_analytics import settlement_files as SF
    w = SF.parse_wing_detail(_wing_grid(), account="계정A", period_start=D("2026-08-24"), period_end=D("2026-08-30"))
    assert (len(w.rows), w.dropped_zero) == (4, 2), (len(w.rows), w.dropped_zero)
    assert [r.option_id for r in w.rows][2] == "<기본배송료>" and w.rows[2].settle == 6000          # 금액 있는 배송비 줄
    assert len(w.warnings) == 1 and "정산금액" in w.warnings[0], w.warnings                          # 1004 검산 위반
    assert not any(hasattr(r, "buyer") for r in w.rows) and "홍길동" not in str([vars(r) for r in w.rows])
    assert w.rows[0].option_id == "95467404277" and w.rows[1].coupon == 20000
    cg = _rg_grid()                                                                   # 쿠팡지원할인(C) 줄(실측 모양)
    cg[2][RG_HDR.index("쿠팡지원할인(C)")], cg[2][RG_HDR.index("매출금액(A*B-C)")] = 2000.0, 7310.0
    assert SF.parse_rg_fee(cg, account="A", period_start=D("2026-08-03"), period_end=D("2026-08-09")).warnings == []
    cg[2][RG_HDR.index("정산대상액")] = 6205.0                                         # 매출금액 기준 값 = 어긋남 경고
    assert len(SF.parse_rg_fee(cg, account="A", period_start=D("2026-08-03"), period_end=D("2026-08-09")).warnings) == 1
    end = D("2025-12-21")
    summ = [[None] * 5, ["쿠팡풀필먼트서비스(CFS) 보관비 정산 내역"], ["정산주기(종료일)", "보관비 합계", "세액", "최종비용"],
            ["2025-12-21", 182986.0, 18301.0, 201287.0], [], [], ["정산유형", "정산주기(종료일)", "옵션ID"]]
    comp = [[None] * 3, ["발생일", "정산주기(종료일)", "주문ID", "옵션ID", "보상 금액"],
            [None, None, None, None, None], ["2025-12-01", "2025-12-21", "1", "9", 4775.0],
            ["2025-12-02", "2025-12-21", "2", "8", 1700.0]]
    assert SF.rg_cost_totals({"보관비": summ, "재고 손실 보상": comp}, end) == {"보관비": 201287, "재고 손실 보상": 6475}
    expect(SF.SettlementParseError, lambda: SF.rg_cost_totals({"보관비": summ}, D("2025-12-28")), "다른 주 요약")
    comp_bad = [*comp, ["2025-12-03", "2025-12-28", "3", "7", 100.0]]
    expect(SF.SettlementParseError, lambda: SF.rg_cost_totals({"재고 손실 보상": comp_bad}, end), "다른 주 보상 줄")
    expect(SF.SettlementParseError, lambda: SF.rg_cost_totals({"x": [["a"], ["b"]]}, end), "모르는 모양")
    notice = [[], ["고객 반품된 상품을 재판매하지 않고 자동으로 반출받는 경우 …"], [], [],
              ["세금계산서 발행월", "발생일", "반출ID", "옵션ID", "반출 완료 수량(A)", "반출비 청구 제외 수량(B)",
               "최종 반출 수량(A-B) (쿠팡 귀책 제외)"], ["2025-12", "2025-12-02", "9", "8", 3.0, 1.0, 2.0]]
    assert SF.rg_cost_totals({"보관비": summ, "자동반출(고객반품) - 쿠팡귀책": notice}, end) == {"보관비": 201287}
    expect(SF.SettlementParseError, lambda: SF.rg_cost_totals({"안내": notice[:4]}, end), "수량 머리글 없는 안내문")
    SF.assert_no_pii({"보관비": summ})
    expect(SF.SettlementParseError, lambda: SF.assert_no_pii({"x": [["주문ID", "구매자명"]]}), "개인정보 칸")
    ci = WING_HDR.index("판매자 할인쿠폰(A+B)")
    old = _wing_grid()
    for row in old[2:4]:
        row[ci] = row[WING_HDR.index("마이샵수수료할인")] = None                       # 옛 파일 모양: 배송비 줄 금액칸 빈칸
    assert SF.parse_wing_detail(old, account="A", period_start=D("2026-08-24"), period_end=D("2026-08-30")).dropped_zero == 2
    old[1][ci] = None                                                                 # 상품 줄 빈칸은 여전히 오류
    expect(SF.SettlementParseError, lambda: SF.parse_wing_detail(old, account="A", period_start=D("2026-08-24"),
                                                                 period_end=D("2026-08-30")), "상품 줄 빈 쿠폰")
    g = SF.parse_rg_fee(_rg_grid(), account="계정A", period_start=D("2026-08-03"), period_end=D("2026-08-09"))
    assert (len(g.rows), g.settle_total, g.rows[0].fee, g.warnings) == (2, 0, 1105, []), g
    expect(SF.SettlementParseError, lambda: SF.parse_rg_fee(_rg_grid("2026-08-16"), account="계정A",
                                                             period_start=D("2026-08-03"), period_end=D("2026-08-09")),
           "다른 기간 파일")
    bad = _wing_grid()
    bad[1][WING_HDR.index("판매액")] = "12,7OO"
    expect(SF.SettlementParseError, lambda: SF.parse_wing_detail(bad, account="a", period_start=D("2026-08-24"),
                                                                  period_end=D("2026-08-30")), "숫자 변환 실패")
    expect(SF.SettlementParseError, lambda: SF.parse_wing_detail([WING_HDR[:-3]], account="a",
                                                                  period_start=D("2026-08-24"),
                                                                  period_end=D("2026-08-30")), "머리글 누락")
    n = SF.settle_file_name("셀독 계정_A", D("2026-09-18"), SF.CH_WING, "주정산", "주문상세", D("2026-08-24"), D("2026-08-30"))
    assert n == "셀독-계정-A_20260918_윙_주정산_주문상세_20260824-20260830.xlsx", n
    meta = SF.parse_file_name(n)
    assert (meta["account"], meta["settle_date"], meta["channel"], meta["period_end"]) == (
        "셀독-계정-A", D("2026-09-18"), "윙", D("2026-08-30"))
    expect(ValueError, lambda: SF.parse_file_name("MSF_PAYMENT_REVENUE_DETAIL-2026-10-06.xlsx"), "규칙 밖 이름")
    import openpyxl
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "raw.xlsx"
        wb = openpyxl.Workbook()
        for row in _wing_grid():
            wb.active.append(row)
        wb.save(src)
        dst = Path(tmp) / n
        assert SF.scrub_pii(src, dst) == 1
        assert "구매자명" not in SF.read_grid(dst)[0] and "홍길동" not in str(SF.read_grid(dst))
        f = SF.load_settle_file(dst)
        assert (f.account, f.kind, f.report, len(f.rows)) == ("셀독-계정-A", "주정산", "주문상세", 4)
        empty = Path(tmp) / "empty.xlsx"
        openpyxl.Workbook().save(empty)
        expect(SF.SettlementParseError, lambda: SF.scrub_pii(empty, Path(tmp) / "x.xlsx"), "머리글 없는 파일 저장")
    ok("윙 0원 배송비 2줄 제외·금액 배송비 유지·검산 경고·구매자명 미보관/삭제·로켓그로스 취소·기간 대조·이름 왕복")


def p8_settle_jobs():
    print("[P8] 요청 기록 — 같은 기간 한 번·비용 리포트는 금액 있는 것만·하루 상한·미래 제외·재요청·목록 대조")
    from datetime import datetime
    from coupang_analytics import settlement_jobs as SJ
    ev = [SJ.SettleEvent("A", "로켓그로스", "주정산", D("2026-09-07"), D("2026-08-03"), D("2026-08-09")),
          SJ.SettleEvent("A", "로켓그로스", "주정산", D("2026-10-01"), D("2026-08-03"), D("2026-08-09"), ("보관비",)),
          SJ.SettleEvent("A", "윙", "주정산", D("2026-09-18"), D("2026-08-24"), D("2026-08-30")),
          SJ.SettleEvent("A", "윙", "최종액", D("2026-10-01"), D("2026-08-03"), D("2026-08-30")),
          SJ.SettleEvent("A", "윙", "주정산", D("2026-10-30"), D("2026-10-05"), D("2026-10-11")),   # 미래
          SJ.SettleEvent("B", "윙", "주정산", D("2026-09-18"), D("2026-08-24"), D("2026-08-30"))]
    jobs: list = []
    now = datetime(2026, 10, 6, 19, 0)
    pick = SJ.plan_requests(ev, jobs, today=D("2026-10-06"), now=now, per_account_cap=3)
    keys = sorted((j.account, j.channel, j.kind, j.report) for j in jobs)
    assert keys == [("A", "로켓그로스", "주정산", "보관비"), ("A", "로켓그로스", "주정산", "판매수수료"),
                    ("A", "윙", "최종액", "주문상세"),
                    ("B", "윙", "주정산", "주문상세")], keys                       # 70/30 1번·미래 제외·A 윙 주정산=월별로 대체
    assert sum(j.account == "A" for j in pick) == 3 and sum(j.account == "B" for j in pick) == 1
    assert pick[0].settle_date == "2026-09-07"                                     # 오래된 것부터
    for k, j in enumerate(pick):
        SJ.mark_requested(j, now, f"req-{k}")
    assert SJ.plan_requests(ev, jobs, today=D("2026-10-07"), now=datetime(2026, 10, 7, 19), per_account_cap=3) == []
    later = SJ.plan_requests(ev, jobs, today=D("2026-10-10"), now=datetime(2026, 10, 10, 19), per_account_cap=9)
    assert len(later) == 4                                                         # 3일 지나도 미완료 → 재요청 대상
    capped = SJ.plan_requests(ev, [], today=D("2026-10-06"), now=now, per_account_cap=2)
    assert sum(j.account == "A" for j in capped) == 2                               # 계정당 하루 상한
    # 월별 파일이 나오기 전엔 주정산을 계획 → 나온 뒤엔 아직 요청 안 한 주정산은 '생략(월별로 대체)'
    evc = [SJ.SettleEvent("C", "윙", "주정산", D("2026-09-18"), D("2026-08-24"), D("2026-08-30")),
           SJ.SettleEvent("C", "윙", "주정산", D("2026-09-29"), D("2026-08-31"), D("2026-08-31")),   # 월 경계 분할 행
           SJ.SettleEvent("C", "윙", "최종액", D("2026-10-01"), D("2026-08-03"), D("2026-08-30"))]
    jc: list = []
    early = SJ.plan_requests(evc, jc, today=D("2026-09-20"), now=datetime(2026, 9, 20, 19), per_account_cap=5)
    assert [(j.kind, j.period_start) for j in early] == [("주정산", "2026-08-24")]         # 최종액은 아직 미도래
    after = SJ.plan_requests(evc, jc, today=D("2026-10-06"), now=now, per_account_cap=5)
    assert jc[0].status == SJ.ST_SKIPPED and "월별" in jc[0].note
    assert sorted((j.kind, j.period_start) for j in after) == [("주정산", "2026-08-31"), ("최종액", "2026-08-03")]
    done = SJ.Job("C", "윙", "주정산", "주문상세", "2026-08-24", "2026-08-30", "2026-09-18", status=SJ.ST_DONE)
    SJ.plan_requests(evc, [done], today=D("2026-10-06"), now=now, per_account_cap=5)
    assert done.status == SJ.ST_DONE                                               # 이미 받은 건 그대로
    j0 = next(j for j in jobs if j.kind == "최종액")
    j0.attempts = SJ.MAX_ATTEMPTS
    SJ.plan_requests(ev, jobs, today=D("2026-10-10"), now=datetime(2026, 10, 10, 19), per_account_cap=9)
    assert j0.status == SJ.ST_FAILED                                               # 무한 재시도 금지
    assert SJ.parse_condition("구매확정일:2026-08-24 - 2026-08-30") == (D("2026-08-24"), D("2026-08-30"))
    expect(ValueError, lambda: SJ.parse_condition("구매확정일:"), "기간 없는 조건")
    wing_b = next(j for j in jobs if j.account == "B")
    rg_fee = next(j for j in jobs if j.report == "판매수수료")
    t_b, t_rg = datetime.fromisoformat(wing_b.requested_at), datetime.fromisoformat(rg_fee.requested_at)
    rows = [SJ.DownloadRow("윙", t_b, "WAIT", period_start=D("2026-08-24"), period_end=D("2026-08-30")),
            SJ.DownloadRow("윙", t_b + timedelta(minutes=1), "FINISHED", period_start=D("2026-08-24"),
                           period_end=D("2026-08-30")),
            SJ.DownloadRow("로켓그로스", t_rg + timedelta(seconds=4), "COMPLETED", report="판매수수료",
                           req_id=rg_fee.req_id)]
    pairs, notes = SJ.match_downloads(rows, [wing_b, rg_fee])
    assert [(j.account, j.report, r.status) for j, r in pairs] == [("B", "주문상세", "FINISHED"),
                                                                   ("A", "판매수수료", "COMPLETED")], pairs
    pairs, notes = SJ.match_downloads([rows[0]], [wing_b])
    assert pairs == [] and "아직 WAIT" in notes[0]
    # 로켓그로스 = 요청번호로만 짝: 같은 시각·같은 리포트라도 번호가 다르면 남의 줄(시각으로 짐작하지 않음)
    t0 = datetime(2026, 10, 6, 20, 0, 0)
    w1 = SJ.Job("A", "로켓그로스", "주정산", "판매수수료", "2026-08-03", "2026-08-09", "2026-10-01",
                status=SJ.ST_REQUESTED, requested_at=t0.isoformat(), req_id="r1")
    w2 = SJ.Job("A", "로켓그로스", "주정산", "판매수수료", "2026-08-10", "2026-08-16", "2026-10-01",
                status=SJ.ST_REQUESTED, requested_at=t0.isoformat(), req_id="r2")
    lst = [SJ.DownloadRow("로켓그로스", t0, "COMPLETED", report="판매수수료", handle="t2", req_id="r2"),
           SJ.DownloadRow("로켓그로스", t0, "PENDING", report="판매수수료", handle="t1", req_id="r1")]
    pairs, notes = SJ.match_downloads(lst, [w1, w2])
    assert [(j.period_start, r.handle) for j, r in pairs] == [("2026-08-10", "t2")] and "아직 PENDING" in notes[0]
    w3 = SJ.Job("A", "로켓그로스", "주정산", "판매수수료", "2026-08-17", "2026-08-23", "2026-10-01",
                status=SJ.ST_REQUESTED, requested_at=t0.isoformat())
    pairs, notes = SJ.match_downloads(lst, [w3])
    assert pairs == [] and "요청 번호" in notes[0]                                   # 번호 없으면 받지 않음
    mb = [SJ.Job("A", "윙", "주정산", "주문상세", "2026-08-24", "2026-08-30", "2026-09-18", status=SJ.ST_DONE)]
    mx = [SJ.Job("A", "윙", "주정산", "주문상세", "2026-08-24", "2026-08-30", "2026-09-18"),            # 같은 키 = 이 PC 우선
          SJ.Job("W", "윙", "최종액", "주문상세", "2026-01-01", "2026-01-31", "2026-03-02", status=SJ.ST_DONE)]
    assert SJ.merge_jobs(mb, mx) == (1, 1) and mb[0].status == SJ.ST_DONE and mb[1].account == "W"
    with tempfile.TemporaryDirectory() as tmp:
        p = Path(tmp) / "기록.json"
        SJ.save_jobs(p, jobs)
        assert [j.key for j in SJ.load_jobs(p)] == [j.key for j in jobs] and not (Path(tmp) / "기록.json.tmp").exists()
    ok("70/30 한 번·보관비만 추가·미래 제외·상한·오래된 순·만료 재요청·3회 실패·월별 나온 달 주정산 생략·윙 기간/RG 요청번호 대조·번호 없음=보류·기록 왕복")


def p9_settle_stats():
    print("[P9] 병합·집계 — 최종액 파일 제외·중복 파일/줄 1회·배송비 묶음·월=매출인식일·계약자 정산금액")
    from coupang_analytics import settlement_files as SF
    from coupang_analytics import settlement_stats as ST
    w = SF.parse_wing_detail(_wing_grid(), account="A", period_start=D("2026-08-24"), period_end=D("2026-08-30"))
    w.kind, w.settle_date = "주정산", D("2026-09-18")
    wfinal = SF.parse_wing_detail([*_wing_grid(), _wing_row("1009", "77", "보냉백", "95467404277", 12700, 1, 0, 1372,
                                                            11328, conf="2026-08-12", due="2026-08-31")],
                                  account="A", period_start=D("2026-08-03"), period_end=D("2026-08-30"))
    wfinal.kind = "최종액"
    g = SF.parse_rg_fee(_rg_grid(), account="A", period_start=D("2026-08-03"), period_end=D("2026-08-09"))
    g.kind, g.report = "주정산", "판매수수료"
    res = ST.aggregate([w, wfinal, g, g])
    by = {(p.month, p.product_id): p for p in res.products}
    assert by[("2026-08", "77")].wing_settle == 11328 + 99999 + 11328 and by[("2026-08", "77")].rg_settle == 0
    assert by[("2026-08", "77")].rg_qty == 0 and by[("2026-08", "77")].product_name == "보냉백"
    assert by[("2026-08", ST.SHIPPING_KEY)].wing_settle == 6000                      # 날짜 없는 배송비 = 기간 끝 달
    assert res.account_totals() == {"A": 11328 + 15967 + 6000 + 99999 + 11328}       # 최종액은 주정산 없는 주(1009)만
    assert any("주정산 파일이 없는 주의 줄 1건" in x for x in res.warnings), res.warnings
    assert any("두 번" in x for x in res.warnings) and any("날짜 없는" in x for x in res.warnings)
    assert not any("겹친 같은 줄" in x for x in res.warnings)                        # 주정산 있는 주는 최종액서 제외
    lines = {(f.kind, f.channel, label): amt for f, label, amt in ST.payout_lines([w, wfinal, g])}
    assert lines[("최종액", "윙", "주정산 70%(정산예정 2026-08-31·월별 파일에서 계산)")] == 7930   # 주정산 없는 주만
    assert not any("2026-09-18·월별" in k[2] for k in lines)                           # 주정산 있는 주는 중복 계산 안 함
    assert lines[("주정산", "윙", "주정산 70%")] == 93306                               # 133,294×0.7
    assert lines[("최종액", "윙", "최종액 30%")] == 144622 - (93306 + 7930)            # 주별(09-18·08-31) 70% 뺀 나머지
    assert lines[("주정산", "로켓그로스", "주정산 70%")] == 0 and ("주정산", "로켓그로스", "2차 30%") in lines
    assert ST.contractor_amount(10_000_000, 3_764_453) == 6_235_547
    expect(ValueError, lambda: ST.contractor_amount(None, 1), "계약금액 미입력")
    expect(TypeError, lambda: ST.contractor_amount(1.5, 1), "정수 아닌 계약금액")
    import openpyxl
    with tempfile.TemporaryDirectory() as tmp:
        out = ST.write_stats(Path(tmp) / "집계.xlsx", res, {"A": 200_000}, files=[w, wfinal, g])
        wb = openpyxl.load_workbook(out)
        assert wb.sheetnames == ["상품별 월별", "계정별 월별", "계약자 정산", "지급액 검산", "경고"]
        assert [c.value for c in wb["계약자 정산"][2]] == ["A", 144622, 200000, 55378]
        assert wb["지급액 검산"].max_row == 1 + 5                 # 윙 주정산·최종액 30%·월별서 계산한 08-31 주 70% + RG 2
        out3 = ST.write_stats(Path(tmp) / "집계3.xlsx", res, amounts=[
            {"계정": "A", "채널": "로켓그로스", "정산일": "2026-01-06", "기간 시작": "2025-12-01", "기간 끝": "2025-12-07",
             "지급비율": 70, "최종지급액": 3479206, "totalPayableAmount": 5658592}],
            costs=[("A", D("2026-01-06"), D("2025-12-01"), D("2025-12-07"), "보관비", "보관비", 205791)])
        wb3 = openpyxl.load_workbook(out3)
        assert [c.value for c in wb3["쿠팡 지급 내역"][1]][-1] == "totalPayableAmount"
        assert [c.value for c in wb3["쿠팡 지급 내역"][2]][6:] == [3479206, 5658592]
        assert wb3["로켓그로스 비용"].cell(2, 7).value == 205791
        out4 = ST.write_stats(Path(tmp) / "집계4.xlsx", res, amounts=[
            {"계정": "A", "채널": "윙", "정산일": "2026-07-01", "기간 시작": "2026-05-01", "기간 끝": "2026-05-31", "지급비율": 30,
             "최종지급액": 2235280, "차감 사유": [{"사유": "판매자서비스이용료", "금액": 55000}]}])
        c = [x.value for x in openpyxl.load_workbook(out4)["쿠팡 지급 내역"][2]]
        assert '"사유": "판매자서비스이용료"' in c[-1], c                                # 목록 = 글자로 펼쳐 저장
        out2 = ST.write_stats(Path(tmp) / "집계2.xlsx", res)
        assert [c.value for c in openpyxl.load_workbook(out2)["계약자 정산"][2]][2:] == [None, None]   # 미입력=빈칸
    ok("최종액=주정산 없는 주만·중복 파일 1회·배송비 묶음·월 귀속·지급액 검산(윙 최종액 30% 주별)·계약자 정산금액·엑셀 5시트")


def p10_wing_api_parse():
    print("[P10] 정산 주소 응답 → 정산 일정·요청 본문·다운로드 목록(실측 응답 모양·값은 가상·라이브 미호출)")
    from datetime import datetime
    from coupang_analytics import settlement_jobs as SJ
    from coupang_analytics import settlement_wing_api as API
    wresp = {"paymentReports": [
        {"payDate": "2026-01-02", "transactionCycleCode": "R", "recognitionFrom": "2025-11-03",
         "recognitionTo": "2025-11-30", "paymentStatus": "DONE", "ratio": 30},
        {"payDate": "2026-01-23", "transactionCycleCode": "W", "recognitionFrom": "2026-01-01",
         "recognitionTo": "2026-01-04", "paymentStatus": "DONE", "ratio": 70}], "totalAmount": 0}
    ev = API.wing_events(wresp, "계정A")
    assert [(e.kind, e.settle_date, e.period_start, e.period_end) for e in ev] == [
        ("최종액", D("2026-01-02"), D("2025-11-03"), D("2025-11-30")),
        ("주정산", D("2026-01-23"), D("2026-01-01"), D("2026-01-04"))], ev
    expect(API.SiteChangedError, lambda: API.wing_events(
        {"paymentReports": [{**wresp["paymentReports"][0], "transactionCycleCode": "X"}]}, "a"), "모르는 주기 코드")
    expect(API.SiteChangedError, lambda: API.wing_events({"reports": []}, "a"), "응답 칸 없음")
    rresp = {"settlementStatusReports": [
        {"settlementDate": "2026-09-06T15:00:00.000Z", "settlementPeriodStartDate": "2026-07-26T15:00:00.000Z",
         "settlementPeriodEndDate": "2026-07-30T15:00:00.000Z", "settlementRatio": 30, "settlementCycle": "WEEKLY",
         "settlementGroupKey": "V1-2026-07-27-2026-07-31", "finalSettlementAmount": 1}]}
    rev = API.rg_events(rresp, "계정A")
    assert [(e.kind, e.settle_date, e.period_start, e.period_end, e.ref) for e in rev] == [
        ("최종액", D("2026-09-07"), D("2026-07-27"), D("2026-07-31"), "V1-2026-07-27-2026-07-31")], rev   # UTC→한국 날짜
    expect(API.SiteChangedError, lambda: API.rg_events({"settlementStatusReports": [
        {**rresp["settlementStatusReports"][0], "settlementCycle": "DAILY"}]}, "a"), "모르는 RG 주기")
    assert API.rg_events_body(D("2026-01-01"), D("2026-01-31")) == {
        "startDate": "2025-12-31T15:00:00.000Z", "endDate": "2026-01-31T15:00:00.000Z", "searchDateType": "PAYMENT"}
    assert API.wing_events_body(D("2026-01-01"), D("2026-01-31"))["fromDate"] == "2026-01-01"
    jobs: list = []
    picked = SJ.plan_requests(ev + rev, jobs, today=D("2026-10-06"), now=datetime(2026, 10, 6, 15), per_account_cap=9)
    rj = next(j for j in picked if j.channel == "로켓그로스")
    assert rj.ref == "V1-2026-07-27-2026-07-31"                                     # 일정의 키가 요청까지 전달
    assert API.rg_request_body(rj, 1791000000000) == {"sellerReportType": "CATEGORY_TR", "requestTime": "1791000000000",
                                                      "settlementGroupKeys": [rj.ref], "locale": "ko"}
    wj = next(j for j in picked if j.channel == "윙" and j.kind == "주정산")
    assert API.wing_request_body(wj) == {"excelType": "MSF_PAYMENT_REVENUE_DETAIL",
                                         "recognitionDateRange": {"start": "2026-01-01", "end": "2026-01-04"},
                                         "searchDateType": "CONFIRM_DATE",
                                         "searchDateRange": {"start": "2026-01-01", "end": "2026-01-04"}}
    expect(API.SiteChangedError, lambda: API.rg_request_body(SJ.Job("a", "로켓그로스", "주정산", "판매수수료", "x", "y", "z"),
                                                              1), "묶음 키 없음")
    expect(API.SiteChangedError, lambda: API.rg_request_body(SJ.Job("a", "로켓그로스", "주정산", "새 비용", "x", "y", "z",
                                                                     ref="k"), 1), "코드 미확인 리포트")
    assert API.wing_request_id({"success": True, "reason": "OK", "data": 134}) == "134"
    expect(API.SiteChangedError, lambda: API.wing_request_id({"success": False, "reason": "LIMIT", "data": None}), "거절")
    assert API.rg_request_id({"requestId": "abc", "duplicateRequest": False, "remainingTime": 0}) == "abc"
    expect(API.RequestThrottled, lambda: API.rg_request_id({"requestId": "", "duplicateRequest": True,
                                                            "remainingTime": 30}), "대기시간 응답")
    # 비용 리포트: 그 주 금액 있는 종류만(입출고+배송 → 1종으로 합침)·0원 제외·미확인 비용은 알림
    det = {"totalStorageFeeDeductionAmount": 201287, "totalFulfillmentFeeDeductionAmount": 843508,
           "totalWarehousingFeeDeductionAmount": 867526, "totalBarcodeLabelingFeeDeductionAmount": 0,
           "totalCfsInventoryCompensationAmount": 119770, "totalAdSalesDeductionAmount": 1003081,
           "totalContainerUnloadingFeeDeductionAmount": 5000}
    assert API.cost_reports(det) == ("보관비", "입출고/배송비", "재고 손실 보상")
    r70 = {**rresp["settlementStatusReports"][0], "settlementRatio": 70, "settlementStatusReportDetail": det}
    ev70 = API.rg_events({"settlementStatusReports": [r70]}, "계정A")[0]
    assert ev70.reports == ("보관비", "입출고/배송비", "재고 손실 보상") and SJ.reports_for(ev70)[0] == "판매수수료"
    assert API.unmapped_costs({"settlementStatusReports": [r70]}) == [
        "V1-2026-07-27-2026-07-31: totalContainerUnloadingFeeDeductionAmount=5000"]
    cj = SJ.Job("a", "로켓그로스", "주정산", "입출고/배송비", "x", "y", "z", ref="k")
    assert API.rg_request_body(cj, 1)["sellerReportType"] == "WAREHOUSING_SHIPPING"
    am = API.amount_rows({"paymentReports": [{**wresp["paymentReports"][1], "finalPaidAmount": 19664,
                                              "bankAccountInfo": {"bank": "가림"}, "isAdditionalPayment": False,
                                              "detail": {"paidAmount": 19664, "isActualPayment": True,
                                                         "payableSummarySeqList": [1, 2],
                                                         "deductionDetail": [{"serviceFeeAmount": 55000, "serviceType": "CLR_3SF",
                                                                              "serviceTypeNameKR": "판매자서비스이용료",
                                                                              "serviceTypeNameEN": "SELLER SERVICE FEE"}]}}]}, "윙")
    assert am == [{"정산일": "2026-01-23", "기간 시작": "2026-01-01", "기간 끝": "2026-01-04", "지급비율": 70,
                   "최종지급액": 19664, "지급상태": "DONE", "실지급": True,
                   "차감 사유": [{"사유": "판매자서비스이용료", "금액": 55000}], "paidAmount": 19664}], am   # 계좌·내부번호 제외
    r70x = {**r70, "settlementStatusReportDetail": {**det, "pastDeductedCfsFeeDetails": [{"paymentRatio": 30, "pastDeductedCfsFeeAmount": 24959.0}],
                                                    "adSalesOffsetYearMonthBreakdown": {"2026-08": 59266.0}, "salesAdjustmentDetails": []}}
    ar = API.amount_rows({"settlementStatusReports": [r70x]}, "로켓그로스")[0]
    assert ar["정산일"] == "2026-09-07" and ar["totalStorageFeeDeductionAmount"] == 201287
    assert ar["pastDeductedCfsFeeDetails"] == [{"paymentRatio": 30, "pastDeductedCfsFeeAmount": 24959.0}]   # 상계 상세 보존
    assert ar["adSalesOffsetYearMonthBreakdown"] == {"2026-08": 59266.0} and "salesAdjustmentDetails" not in ar   # 빈 목록은 생략
    wl = API.wing_list_rows([
        {"id": 1, "excelType": "MSF_PAYMENT_REVENUE_DETAIL", "status": "FINISHED", "startedAt": "2026-10-06 14:49:59",
         "downloadUrl": "https://x/dl?id=1", "jsonItems": '[{"key":"구매확정일", "value":"2026-09-01 - 2026-09-06", "view":true}]'},
        {"id": 2, "excelType": "MSF_VAT_DETAIL", "status": "FINISHED", "startedAt": "2025-02-19 05:07:42",
         "downloadUrl": "https://x/dl?id=2", "jsonItems": '[{"key":"구매확정일", "value":"2025-01-01 - 2025-01-31"}]'}])
    assert [(r.requested_at, r.status, r.period_start, r.period_end, r.handle) for r in wl] == [
        (datetime(2026, 10, 6, 14, 49, 59), "FINISHED", D("2026-09-01"), D("2026-09-06"), "https://x/dl?id=1")]
    rl = API.rg_list_rows([{"requestId": "abc", "requestTime": "1791000000000", "downloadStatus": "COMPLETED",
                            "sellerReportType": "CATEGORY_TR", "recognitionDateFrom": None}])
    assert (rl[0].report, rl[0].status, rl[0].req_id, rl[0].handle) == ("판매수수료", "COMPLETED", "abc", "1791000000000")
    assert API.rg_list_body(datetime(2026, 10, 6, 0, 0), datetime(2026, 10, 6, 0, 0, 1)) == {
        "requestTimeFrom": str(int(datetime(2026, 10, 6).timestamp() * 1000)),
        "requestTimeTo": str(int(datetime(2026, 10, 6).timestamp() * 1000) + 1000)}
    assert API.month_windows(D("2026-01-15"), D("2026-03-02")) == [
        (D("2026-01-15"), D("2026-01-31")), (D("2026-02-01"), D("2026-02-28")), (D("2026-03-01"), D("2026-03-02"))]

    class FakePage:
        def __init__(self, status, text):
            self.r = {"status": status, "ctype": "x", "text": text}

        def evaluate(self, js, arg=None):
            return self.r
    assert API.call(FakePage(200, '{"a":1}'), "GET", "/p") == {"a": 1}
    expect(API.ApiBlocked, lambda: API.call(FakePage(403, "x"), "GET", "/p"), "403=차단")
    expect(API.ApiBlocked, lambda: API.call(FakePage(200, "<html>Access Denied"), "GET", "/p"), "HTML=차단")
    expect(API.SiteChangedError, lambda: API.call(FakePage(500, "{}"), "GET", "/p"), "500=중단")
    for st in (502, 503, 504):                                                      # 실측 2026-10-08: 504 HTML = 서버 일시 오류
        expect(API.ServerBusy, lambda st=st: API.call(FakePage(st, "<html><title>%d Gateway Time-out</title>" % st), "GET", "/p"),
               f"{st}=일시 오류(차단 아님)")
        assert not issubclass(API.ServerBusy, API.ApiBlocked)
    ok("윙/RG 일정(UTC→한국 날짜·묶음키)·요청 본문·요청번호·목록(부가세 메뉴 무시)·달 구간·403/HTML=차단·500=중단")


def p11_runlog():
    print("[P11] 실행 기록 — 처리기록 CSV(실행·누적)·로그인/차단/실패 기록·오류 전체 추적·요약")
    import csv as _csv
    from datetime import datetime
    from coupang_analytics import settlement_runlog as RLG
    with tempfile.TemporaryDirectory() as tmp:
        out: list = []
        lg = RLG.RunLog(tmp, now=datetime(2026, 10, 6, 21, 0, 0), echo=out.append)
        lg.record("A-1", "로그인", RLG.OK, "success: 로그인 완료")
        lg.record("A-1", "요청", RLG.OK, "다운로드 요청함", channel="윙", kind="주정산·주문상세", settle_date="2026-01-12",
                  period="2025-12-29~2026-01-04")
        lg.record("B-2", "차단감지", RLG.BLOCK, "차단 화면 — https://wing.coupang.com/x")
        try:
            raise ValueError("표 머리글 없음")
        except ValueError as exc:
            lg.error("B-2", "조회", exc, channel="윙", settle_date="2026-01-13")
        lg2 = RLG.RunLog(tmp, now=datetime(2026, 10, 7, 21, 0, 0), echo=out.append)
        lg2.record("A-1", "받기", RLG.WAIT, "아직 WAIT", channel="윙")
        rows = list(_csv.DictReader(open(lg.csv, encoding="utf-8-sig")))
        assert [r["결과"] for r in rows] == ["정상", "정상", "차단", "실패"] and rows[1]["정산일"] == "2026-01-12"
        assert (rows[3]["채널"], rows[3]["정산일"]) == ("윙", "2026-01-13")                    # 실패 줄에도 정산일
        total = list(_csv.DictReader(open(lg.total, encoding="utf-8-sig")))
        assert len(total) == 5 and {r["실행ID"] for r in total} == {"261006_210000", "261007_210000"}   # 누적
        tb = lg.errors.read_text(encoding="utf-8")
        assert "Traceback" in tb and "ValueError: 표 머리글 없음" in tb and "B-2 · 조회" in tb
        assert "⛔ [차단감지] B-2" in lg.text.read_text(encoding="utf-8")
        summ = lg.summary()
        assert "[요약] A-1: 정상 2" in summ and "[요약] B-2: 실패 1 · 차단 1" in summ, summ
        assert "[비정상] 조회 실패: 1건" in summ and "[비정상] 차단감지 차단: 1건" in summ
        # heartbeat: 창 없는 자동 실행 모니터링 — 고정 경로(_현재상태.txt)에 **덮어쓰기**(append 아님)·개인정보 없음.
        lg.heartbeat("대기중", "앱 ①판매수집 완료 기다림")
        st = lg.status
        assert st.name == "_현재상태.txt" and st.parent == lg.dir                    # 로그 폴더 안 고정 파일
        txt1 = st.read_text(encoding="utf-8")
        assert "상태: 대기중" in txt1 and "앱 ①판매수집 완료 기다림" in txt1
        assert f"실행ID: {lg.run_id}" in txt1 and f"상세 로그: {lg.text.name}" in txt1
        lg.heartbeat("작동중", "요청 누르는 중", extra=("추가줄",))
        txt2 = st.read_text(encoding="utf-8")
        assert "상태: 작동중" in txt2 and "추가줄" in txt2
        assert "대기중" not in txt2 and txt2.count("상태:") == 1                      # 누적 아님 — 최신 1건만
        assert not st.with_name(st.name + ".tmp").exists()                           # 임시파일은 남기지 않음(os.replace)
        lg2.heartbeat("완료", "명령 run")                                             # 다른 RunLog 도 같은 고정 파일을 공유
        assert lg2.status == st and "상태: 완료" in st.read_text(encoding="utf-8")
    ok("실행·누적 CSV(엑셀용)·정산일 칸·차단/로그인 기록·오류 전체 추적·계정별/단계별 요약·상태 파일 덮어쓰기 모니터링")


def p12_watch_decide():
    print("[P12] 24h 감시 판단 — ①판매수집 중 정지·소급 연속·받을 것 없으면 다음 ①완료까지 대기")
    import os as _os
    from datetime import datetime
    from coupang_analytics import settlement_watch as W
    T = datetime
    # plan_watch: ① 진행 중이면 정지(busy 사유 전달)
    a, r = W.plan_watch(True, "①판매수집 진행 중", T(2026, 10, 7, 19, 0), None)
    assert a == "wait" and "진행 중" in r
    # 소급(last_sales_at=None): ①완료 기록이 있든 없든 항상 실행
    assert W.plan_watch(False, "", None, None) == ("run", "소급 수집(밀린 정산 받는 중)")
    assert W.plan_watch(False, "", T(2026, 10, 7, 19, 30), None)[0] == "run"
    # 정상(소급 끝·last_sales_at 있음): 새 ①완료가 있어야만 1회
    a, r = W.plan_watch(False, "", T(2026, 10, 7, 19, 30), T(2026, 10, 7, 19, 30))
    assert a == "wait" and "없음" in r                                              # 같은 ①완료 = 이미 함
    a, r = W.plan_watch(False, "", T(2026, 10, 8, 19, 30), T(2026, 10, 7, 19, 30))
    assert a == "run" and "새 ①판매수집" in r                                        # 다음날 새 ①완료 = 1회
    assert W.plan_watch(False, "", None, T(2026, 10, 7, 19, 30))[0] == "wait"        # ①완료 기록 없음 = 대기
    # sales_in_progress: _진행중.json 수정 시각(신선=진행 중)
    with tempfile.TemporaryDirectory() as tmp:
        prog = Path(tmp) / "진행중.json"
        assert W.sales_in_progress(prog)[0] is False                                # 파일 없음
        prog.write_text("{}", encoding="utf-8")
        _os.utime(prog, (T(2026, 10, 7, 12, 0, 0).timestamp(),) * 2)
        assert W.sales_in_progress(prog, now=T(2026, 10, 7, 12, 10))[0] is True      # 10분 전 = 진행 중
        assert W.sales_in_progress(prog, now=T(2026, 10, 7, 12, 40))[0] is False     # 40분 전 = 아님(①끝남)
        # read_marker 왕복
        mp = Path(tmp) / "단계.json"
        assert W.read_marker(mp)[0] is None and "없음" in W.read_marker(mp)[1]
        mp.write_text('{"date": "2026-10-07", "stage": "sales", "at": "2026-10-07T19:30:00"}', encoding="utf-8")
        assert W.read_marker(mp)[0] == {"stage": "sales", "at": T(2026, 10, 7, 19, 30)}
        mp.write_text("{깨짐", encoding="utf-8")
        assert W.read_marker(mp)[0] is None and "읽기 실패" in W.read_marker(mp)[1]
    ok("① 진행 중 정지·소급 연속(last_sales_at None)·정상 1회(새 ①완료)·진행 파일 신선도·단계 기록 왕복")


def p13_accounts_file():
    print("[P13] 정산 계정 파일 — 탭/공백 섞임·뒤 공백·3열=비번 줄은 이름 없음(비번 비노출)·중복/칸 없음=오류·계정명 맞춤")
    from coupang_analytics import settlement_accounts as SA
    from coupang_analytics import settlement_files as SF
    text = ("acc1\tpw1!!\t홍길동-주식회사 가나\n"           # 이름에 공백
            "acc2 \tsecret#9\tsecret#9\n"                  # 3열 = 비번(실제 파일에 있는 모양)
            "acc3  pw3\t\n"                                 # 이름 없음·공백 구분
            "plan_it \tpw4* \t 김다라_플랜\n"               # 뒤 공백·밑줄 ID
            "\n\t\n")
    accts, warns = SA.parse_accounts_text(text)
    assert [(a.account_id, a.label, a.password) for a in accts] == [
        ("acc1", "홍길동-주식회사 가나", "pw1!!"), ("acc2", "acc2", "secret#9"), ("acc3", "acc3", "pw3"),
        ("plan_it", "김다라_플랜", "pw4*")]
    assert len(warns) == 2 and "비밀번호와 같아" in warns[0] and not any("secret" in w for w in warns)   # 비번 비노출
    assert "secret" not in repr(accts) and "pw1" not in repr(accts)                                       # repr 에도 없음
    expect(SA.AccountsFileError, lambda: SA.parse_accounts_text("acc1 a 이름\nacc1 b 이름"), "중복 계정ID")
    expect(SA.AccountsFileError, lambda: SA.parse_accounts_text("acc1\n"), "비번 칸 없음")
    expect(SA.AccountsFileError, lambda: SA.parse_accounts_text("\n\n"), "빈 파일")
    with tempfile.TemporaryDirectory() as tmp:
        expect(SA.AccountsFileError, lambda: SA.read_accounts_file(Path(tmp) / "없음.txt"), "파일 없음")
        (Path(tmp) / "a.txt").write_bytes("acc1\tpw\t이름".encode("cp949"))
        expect(SA.AccountsFileError, lambda: SA.read_accounts_file(Path(tmp) / "a.txt"), "UTF-8 아님")
        (Path(tmp) / "b.txt").write_text("﻿acc1\tpw\t이름", encoding="utf-8")
        assert SA.read_accounts_file(Path(tmp) / "b.txt")[0][0].account_id == "acc1"                  # BOM 허용
    assert SA.display_name(accts[0]) == "홍길동-주식회사 가나-acc1"
    olds = ["(주)가나-acc1", "홍길동-주식회사 가나-acc1", "옛이름-plan_it", "남의-xacc3", "모름-zzz"]
    assert SA.renames(olds, accts) == {"(주)가나-acc1": "홍길동-주식회사 가나-acc1", "옛이름-plan_it": "김다라_플랜-plan_it"}
    tok = lambda x: SF._token(x, "계정명")
    assert SA.renames(["(주)가나-acc1", "옛이름-plan-it"], accts, tok=tok) == {
        "(주)가나-acc1": "홍길동-주식회사-가나-acc1", "옛이름-plan-it": "김다라-플랜-plan-it"}
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    import settlement_download as TD
    from coupang_analytics import settlement_jobs as SJ
    here = [SJ.Job("홍길동-주식회사 가나-acc1", "윙", "주정산", "주문상세", "2026-01-05", "2026-01-11", "2026-01-30",
                   status=SJ.ST_DONE, file="홍길동-주식회사-가나-acc1_20260130_윙_주정산_주문상세_20260105-20260111.xlsx")]
    came = [SJ.Job("(주)가나-acc1", "윙", "주정산", "주문상세", "2026-01-05", "2026-01-11", "2026-01-30",
                   file="(주)가나-acc1_20260130_윙_주정산_주문상세_20260105-20260111.xlsx"),
            SJ.Job("(주)가나-acc1", "윙", "최종액", "주문상세", "2026-01-01", "2026-01-31", "2026-03-02",
                   file="비용/(주)가나-acc1_20260302_윙_최종액_주문상세_20260101-20260131.xlsx")]
    TD._rename_jobs(came, accts)
    assert came[1].file == "비용/홍길동-주식회사-가나-acc1_20260302_윙_최종액_주문상세_20260101-20260131.xlsx"
    assert SJ.merge_jobs(here, came) == (1, 1) and here[0].status == SJ.ST_DONE     # 옛 이름이어도 같은 작업은 한 번
    ok("공백/탭 혼합·3열=비번→이름 없음(경고·repr 비노출)·중복/칸 없음/빈 파일/UTF-8 아님=오류·BOM·계정ID 끝 기준 이름 맞춤"
       "·다른 PC 기록 합치기(옛 이름 맞춘 뒤 중복 없음)")


def p14_status_reader():
    print("[P14] 상태 화면 리더 — heartbeat(_현재상태.txt)↔UI 리더 계약·멈춤 의심 판정·오늘 집계")
    import os as _os
    import re as _re
    from datetime import datetime as _dtm
    sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "ui"))
    import settlement_status_panel_qt as UI               # PySide6 는 이 프로젝트 필수 의존(앱 UI)
    from coupang_analytics import settlement_runlog as RLG
    fixed = _dtm(2026, 10, 7, 12, 55, 0)
    with tempfile.TemporaryDirectory() as tmp:
        lg = RLG.RunLog(tmp, now=fixed)                             # RunLog.dir = <tmp>/로그
        lg.heartbeat("작동중", "①판매수집 완료 — 재개 · 멈춤 10-07 17:40")
        lg.record("A-1", "요청", RLG.OK, "다운로드 요청함", channel="윙", settle_date="2026-01-12")
        lg.record("B-2", "요청", RLG.BLOCK, "차단 화면", channel="윙", settle_date="2026-01-12")
        # heartbeat 의 '갱신'은 실제 now(생존신호) → 시간 판정만 고정(형식·라벨은 heartbeat 실제 출력 그대로 검증).
        lg.status.write_text(_re.sub(r"갱신: .*", f"갱신: {fixed:%Y-%m-%d %H:%M:%S}",
                                     lg.status.read_text(encoding="utf-8")), encoding="utf-8")
        _os.utime(lg.text, (fixed.timestamp(),) * 2)               # 로그 mtime 도 고정
        info = UI.read_settlement_status(lg.dir, now=_dtm(2026, 10, 7, 12, 57, 0))
        assert info["exists"] and info["상태"] == "작동중" and not info["stale"], info   # 2분 전 = 정상
        assert info["내용"].startswith("①판매수집 완료") and info["log_name"] == lg.text.name
        cnt = UI.read_settlement_status(lg.dir, now=_dtm.now())["today_counts"]    # 누적 CSV 는 실제 '시각'(오늘)
        assert cnt.get(RLG.OK) == 1 and cnt.get(RLG.BLOCK) == 1, cnt              # 오늘 쓴 2줄 집계
        late = UI.read_settlement_status(lg.dir, now=_dtm(2026, 10, 7, 13, 20, 0))
        assert late["stale"] is True, late                                       # 25분 무변화 = 멈춤 의심
        empty = Path(tmp) / "빈폴더"
        assert UI.read_settlement_status(empty)["exists"] is False               # 로그 없음 = 아직 실행 안 됨
        assert UI._counts_text({}) == "오늘 받은 기록 없음" and "분 전" in UI._age_text(180)
    ok("heartbeat 덮어쓰기 파싱·작동중/대기중 멈춤 의심(STALE_SEC)·오늘 결과 집계·로그 없음/빈값 표기")


def p15_browser_closed_transient():
    print("[P15] 브라우저 닫힘(reap)=일시적 — '차단'·연속 실패로 안 셈·다음 바퀴 재시도")
    sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "tools"))
    import settlement_download as SD
    # 브라우저/탭 닫힘 오류는 True(일시적), 그 외 오류는 False(진짜 실패)
    assert SD._is_browser_closed(Exception("Target page, context or browser has been closed"))
    assert SD._is_browser_closed(type("TargetClosedError", (Exception,), {})("x"))
    assert SD._is_browser_closed(Exception("Connection closed while reading from the driver"))
    assert not SD._is_browser_closed(Exception("Access Denied"))
    assert not SD._is_browser_closed(ValueError("표 머리글 없음"))
    ok("TargetClosedError·'has been closed'·'Connection closed'=일시적(True)·그 외(차단/파싱 오류)=False")


def p16_cost_by_product():
    print("[P16] 로켓그로스 비용 상품별 — 상세 줄→상품·월(VAT 별도 정확·세액 비율 배분 합 일치)·보상=옵션→상품·박스 여러 상품=묶음")
    from coupang_analytics import settlement_files as SF
    from coupang_analytics import settlement_stats as ST
    end = D("2025-12-21")
    head = [[], ["쿠팡풀필먼트서비스(CFS) 배송비 정산 내역"], ["정산주기(종료일)", "배송비 합계", "세액", "최종비용"],
            ["2025-12-21", 1000.0, 101.0, 1101.0], [], [],
            ["정산유형", "정산주기(종료일)", "매출인식일", "등록상품 ID", "옵션ID", "등록상품명", "판매수량",
             "쿠팡풀필먼트서비스(CFS) 배송비 (VAT 별도)", None, None, None, None],
            [None, None, None, None, None, None, None, "발생비용(A)", "할인가(B)", "할인적용가(A-B)", "추가비용", "최종비용"]]
    ship = head + [["주정산", "2025-12-21", "2025-12-15", "77", "701", "보냉백", 1, 900, 0, 900, 0, 700.0],
                   ["주정산", "2025-12-21", "2025-12-16", "77", "701", "보냉백", 1, 900, 0, 900, 0, -100.0],   # 음수 조정 줄
                   ["주정산", "2025-12-21", "2026-01-02", "88", "801", "압축팩", 1, 400, 0, 400, 0, 400.0]]
    lines = SF.rg_cost_lines({"배송비": ship}, "A", end)
    got = [(c.product_id, c.month, c.ex_vat, c.vat) for c in lines]
    assert got == [("77", "2025-12", 600, 61), ("88", "2026-01", 400, 40)], got   # '최종비용' 칸·월=매출인식일·세액 합 101
    bad = [r[:] for r in ship]
    bad[8][11] = 701.0                                                              # 상세 합 ≠ 요약 합계
    expect(SF.SettlementParseError, lambda: SF.rg_cost_lines({"배송비": bad}, "A", end), "상세 합 불일치")
    ret = [[], ["반출 배송"], ["정산주기(종료일)", "합계", "세액", "최종비용"], ["2025-12-21", 3000.0, 300.0, 3300.0], [], [],
           ["정산유형", "정산주기(종료일)", "매출인식일", "대표 등록상품 ID", "대표 등록상품명", "비용 청구 수량", "X", None],
           [None, None, None, None, None, None, "발생비용(A)", "최종비용(A-B-C)"]]
    ret += [["주정산", "2025-12-21", "2025-12-15", "77,77", "보냉백,보냉백", 1, 0, 1000.0],      # 같은 상품 반복 = 그 상품
            ["주정산", "2025-12-21", "2025-12-15", "77,88", "보냉백,압축팩", 1, 0, 1000.0],      # 서로 다른 상품 = 묶음
            ["주정산", "2025-12-21", "2025-12-15", "-", "-", 1, 0, 1000.0]]                    # 미표기
    rl = SF.rg_cost_lines({"반출 배송 서비스비 리포트": ret}, "A", end)
    assert sorted(c.product_id for c in rl) == sorted(["77", SF.MULTI_PRODUCT, SF.NO_PRODUCT])
    assert sum(c.vat for c in rl) == 300
    comp = [[None] * 3, ["발생일", "정산주기(종료일)", "주문ID", "옵션ID", "등록상품명", "보상 금액"],
            [None] * 6, ["2025-12-02", "2025-12-21", "1", "701", "보냉백", 4775.0],
            ["2025-12-09", "2025-12-21", "2", "999", "모르는상품", 1000.0]]
    cl = SF.rg_cost_lines({"재고 손실 보상": comp}, "A", end)
    assert [(c.kind, c.option_id, c.ex_vat, c.vat) for c in cl] == [
        (SF.COMPENSATION, "701", 4775, 0), (SF.COMPENSATION, "999", 1000, 0)]
    assert SF._split(101, [900, -100, 400]) == [76, -9, 34] and sum(SF._split(7, [1, 1, 1])) == 7
    expect(SF.SettlementParseError, lambda: SF._split(5, [1, -1]), "합 0 인데 나눌 값")
    res = ST.aggregate([], lines + rl + cl)
    by = {p.product_id: p for p in res.products if p.month == "2025-12"}
    assert by["77"].rg_cost == 661 + 1100 and by["77"].rg_comp == 4775            # 배송 + 반출배송(같은 상품 반복)
    assert by["77"].settle_after_costs == 0 - 1761 + 4775 and by["77"].costs["배송비"] == 661
    assert "옵션 999" in by and any("옵션ID 로 상품을 못 찾아" in w for w in res.warnings)
    assert any(SF.MULTI_PRODUCT in w for w in res.warnings) and any(SF.NO_PRODUCT in w for w in res.warnings)
    import openpyxl
    with tempfile.TemporaryDirectory() as tmp:
        wb = openpyxl.load_workbook(ST.write_stats(Path(tmp) / "c.xlsx", res))
        ws = wb["로켓그로스 비용 상품별"]
        hdr = [c.value for c in ws[1]]
        assert hdr[:5] == ["계정", "월(매출인식)", "등록상품ID", "상품명", "배송비"] and hdr[-2:] == ["물류비 합계(VAT포함)", "재고 손실 보상"]
        am = [c.value for c in wb["계정별 월별"][2]]
        assert am[5:] == [sum(p.rg_cost for p in by.values()), 5775, -sum(p.rg_cost for p in by.values()) + 5775]
    ok("상세→상품·월(음수 줄 포함)·세액 배분 합 일치·상세 합 불일치=오류·박스 같은 상품/여러 상품/미표기·보상 옵션→상품·"
       "못 찾음=옵션 줄+경고·상품/계정 시트 물류비·보상·반영 후")


def p17_server_busy():
    print("[P17] 서버 일시 오류(502·503·504) — 재시도 후 성공·끝내 실패=대기(차단·연속실패 아님)·전체 중단 안 함")
    import contextlib
    from datetime import datetime
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    import settlement_download as TD
    from coupang_analytics import settlement_runlog as RLG
    from coupang_analytics import settlement_wing_api as API

    class SeqPage:                                                      # 응답 차례대로
        def __init__(self, seq):
            self.seq = list(seq)

        def evaluate(self, js, arg=None):
            st, text = self.seq.pop(0)
            return {"status": st, "ctype": "x", "text": text}

    class B:
        def __init__(self, seq):
            self.page = SeqPage(seq)
    old_retry, old_log = TD.SERVER_BUSY_RETRY, TD.LOG
    with tempfile.TemporaryDirectory() as tmp:
        TD.LOG = RLG.RunLog(tmp, now=datetime(2026, 10, 8, 12), echo=lambda m: None)
        TD.SERVER_BUSY_RETRY = (0, 2)
        try:
            assert TD.api(B([(504, "<html>"), (502, "<html>"), (200, '{"ok":1}')]), "A", "POST", "/p", {}) == {"ok": 1}
            expect(API.ServerBusy, lambda: TD.api(B([(504, "<html>")] * 3), "A", "POST", "/p", {}), "3번 다 504")
            expect(TD.BlockDetected, lambda: TD.api(B([(403, "x")]), "A", "POST", "/p", {}), "403=차단")
            assert [r["결과"] for r in TD.LOG.rows] == [RLG.BLOCK]                 # 504 는 차단 기록 없음

            @contextlib.contextmanager
            def fake_session(a, hidden):
                yield None

            def busy_run(a, b, args, jobs):
                raise API.ServerBusy("/tenants/x → 504 서버 일시 오류")

            class Acc:
                def __init__(self, i):
                    self.account_id, self.label = i, i

            class Args:
                until, until_at, hidden, command = "", None, True, "run"
            old_session, old_name = TD.session, TD.name_of
            TD.session, TD.name_of = fake_session, (lambda a: a.account_id)
            try:
                TD.LOG = RLG.RunLog(tmp, now=datetime(2026, 10, 8, 13), echo=lambda m: None)
                rc = TD._run_accounts([Acc("a1"), Acc("a2"), Acc("a3")], busy_run, Args)
            finally:
                TD.session, TD.name_of = old_session, old_name
            assert rc == 0, rc                                                      # 연속 2계정이어도 전체 중단 안 함
            res = [(r["계정"], r["결과"]) for r in TD.LOG.rows if r["단계"] == "계정"]
            assert res == [("a1", RLG.WAIT), ("a2", RLG.WAIT), ("a3", RLG.WAIT)], res
            assert not any(r["결과"] == RLG.BLOCK for r in TD.LOG.rows)
        finally:
            TD.SERVER_BUSY_RETRY, TD.LOG = old_retry, old_log
    ok("504→502→성공=값 반환·3번 실패=ServerBusy·403=차단 기록·계정 단위 일시 오류=대기·연속 3계정도 전체 중단 안 함")


def main():
    g1_payout_dates()
    g2_amounts()
    g3_invariants()
    p1_holidays()
    p2_payout_rules()
    p3_amount_rules()
    p4_parse()
    p5_holiday_source()
    p6_diag_no_values()
    p7_settle_files()
    p8_settle_jobs()
    p9_settle_stats()
    p10_wing_api_parse()
    p11_runlog()
    p12_watch_decide()
    p13_accounts_file()
    p14_status_reader()
    p15_browser_closed_transient()
    p16_cost_by_product()
    p17_server_busy()
    print("정산 계산 모듈 오프라인 검증 통과(골든 payout 12·amount 2·불변식 4행/3식 100%)")


if __name__ == "__main__":
    main()
