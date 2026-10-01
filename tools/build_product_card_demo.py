"""상품 통합카드(마스터-디테일) 구현가능성 데모 — 소유자 구상 검토용.

구상(2026-10-01 구체화): 왼쪽에 **계정별·상품별 목록**, 각 행 옆칸에 **👉 방향표시 + '통합내용보기' 클릭**,
클릭하면 **오른쪽 옆 영역**에 그 상품의 여러 시트 흩어진 항목을 **정해진 포맷(그룹블록)**으로 통합표시.
일부 수정 → 각 원본 시트 반영.

이 데모가 증명/설명하는 것(자기완결 1파일):
  (A) 선택→오른쪽 패널 통합표시+그룹핑 = **수식(INDEX/MATCH)만으로 가능**.
  (B) '통합내용보기' 한 번 클릭 → 그 행이 패널에 로드 = **드롭다운 선택은 수식으로 즉시**,
      진짜 '클릭→자동 로드'는 Apps Script(onSelect/버튼) 필요. 👉셀은 HYPERLINK로 패널 이동.
  (C) 패널 노랑칸 수정 → 원본 시트 반영(write-back) = 순수 수식 불가 → Apps Script 또는 앱(gsheet_api).

    python tools/build_product_card_demo.py [출력.xlsx]
"""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

DEF_OUT = Path.home() / "Downloads" / "상품카드_데모.xlsx"

_HDR = PatternFill("solid", fgColor="2E75B6")
_BLOCK = PatternFill("solid", fgColor="548235")       # 블록 헤더(초록)
_LBL = PatternFill("solid", fgColor="D9E1F2")          # 라벨칸(연파랑)
_EDIT = PatternFill("solid", fgColor="FFF2CC")         # 수정입력(노랑)
_SEL = PatternFill("solid", fgColor="FCE4D6")          # 선택칸(살구)
_BTN = PatternFill("solid", fgColor="FFE699")          # 통합내용보기 버튼(노랑)
_W = Font(color="FFFFFF", bold=True)
_B = Font(bold=True)
_LINK = Font(color="0563C1", bold=True, underline="single")
_CTR = Alignment("center", "center")
_THIN = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
_NUM = "#,##0"

# ── 미니 원본 시트(실제로는 운영대장/결과파일에 흩어져 있음) ──
PM = ["사업자명", "상품(물류)명", "노출상품명", "판매방식", "재고현황", "그로스재고", "관리상태", "상품코드"]
PM_ROWS = [
    ["효성", "파미젠 와사비잎 추출물 plus max 120정", "파미젠 와사비잎 추출물 plus max 120정 2개월분", "그로스판매", 2151, 418, "판매중", "G2609-001"],
    ["(주)웰빙곳간", "웰빙곳간 활력 볶은 맥문동 환 프리미엄 30포", "활력 맥문동환 30포 선물세트", "둘다", 1618, 1, "판매중", "G2608-014"],
    ["휴라엘", "신형타프 R008", "신형 대형 타프 R008 블랙 Free", "그로스판매", 88, 331, "판매중", "P2609-003"],
]
ACCOUNT_ID = {"효성": "hyosung01", "(주)웰빙곳간": "wellbing1107", "휴라엘": "bws247"}
CH = ["상품(물류)명", "요청일", "키워드", "건수", "상품코드"]
CH_ROWS = [
    ["파미젠 와사비잎 추출물 plus max 120정", "2026-09-20", "와사비추출물, 숙취", 30, "G2609-001"],
    ["신형타프 R008", "2026-09-18", "대형타프, 캠핑타프", 20, "P2609-003"],
]
SA = ["상품(물류)명", "귀속월", "최종지급예정액", "상품코드"]
SA_ROWS = [
    ["파미젠 와사비잎 추출물 plus max 120정", "2026-09", 8400000, "G2609-001"],
    ["신형타프 R008", "2026-09", 3764453, "P2609-003"],
    ["웰빙곳간 활력 볶은 맥문동 환 프리미엄 30포", "2026-09", 4300000, "G2608-014"],
]
CS = ["상품(물류)명", "계약유형", "수익(순액)", "갑배분(상품주)", "을배분(토탈셀러)", "상품코드"]
CS_ROWS = [
    ["파미젠 와사비잎 추출물 plus max 120정", "셀독(OEM)", 3072000, 2150400, 921600, "G2609-001"],
    ["신형타프 R008", "위탁(운영대행)", 980000, 686000, 294000, "P2609-003"],
    ["웰빙곳간 활력 볶은 맥문동 환 프리미엄 30포", "위탁(운영대행)", 1754000, 1227800, 526200, "G2608-014"],
]

# 패널 선택키 셀(오른쪽 패널이 이 셀의 상품명으로 전 시트 취합)
SEL_CELL = "H1"
# 패널 좌표: G=라벨 · H=값(읽기) · I=수정입력
P_LBL, P_VAL, P_EDIT = 7, 8, 9


def _style(cell, *, fill=None, font=None, align=None, border=True, num=False):
    if fill:
        cell.fill = fill
    if font:
        cell.font = font
    if align:
        cell.alignment = align
    if border:
        cell.border = _THIN
    if num:
        cell.alignment = Alignment(horizontal="right")
        cell.number_format = _NUM


def _mini(wb, title, color, heads, rows):
    ws = wb.create_sheet(title)
    ws.sheet_properties.tabColor = color
    for c, h in enumerate(heads, 1):
        _style(ws.cell(1, c, h), fill=_HDR, font=_W, align=_CTR)
    for i, row in enumerate(rows, 2):
        for c, v in enumerate(row, 1):
            _style(ws.cell(i, c, v), num=isinstance(v, int))
    for c in range(1, len(heads) + 1):
        ml = max((len(str(ws.cell(r, c).value or "")) for r in range(1, ws.max_row + 1)), default=8)
        ws.column_dimensions[get_column_letter(c)].width = min(40, max(9, int(ml * 1.5) + 2))
    ws.freeze_panes = "A2"
    return ws


def _idx(sheet, ret_col, key_col):
    """패널 선택상품(SEL_CELL)으로 sheet에서 ret_col 값을 끌어오는 INDEX/MATCH."""
    rc, kc = get_column_letter(ret_col), get_column_letter(key_col)
    return f"=IFERROR(INDEX('{sheet}'!{rc}:{rc},MATCH(${SEL_CELL[0]}${SEL_CELL[1:]},'{sheet}'!{kc}:{kc},0)),\"-\")"


def _panel(card):
    """오른쪽 '통합카드' 패널 — 정해진 포맷(그룹블록). 선택상품으로 전 시트 자동취합."""
    _style(card.cell(1, P_LBL, "선택상품 ▶"), fill=_LBL, font=_B)
    _style(card.cell(1, P_VAL), fill=_SEL, font=_B)                 # = SEL_CELL (H1)
    card.cell(1, P_VAL).value = "파미젠 와사비잎 추출물 plus max 120정"
    dv = DataValidation(type="list", formula1="=관리상품!$B$2:$B$4", allow_blank=False)
    card.add_data_validation(dv)
    dv.add(card.cell(1, P_VAL))
    _style(card.cell(1, P_EDIT, "← 통합내용보기 클릭 시 로드"), font=Font(italic=True, color="808080"))

    _style(card.cell(2, P_LBL, "항목"), fill=_HDR, font=_W, align=_CTR)
    _style(card.cell(2, P_VAL, "내용(읽기·자동)"), fill=_HDR, font=_W, align=_CTR)
    _style(card.cell(2, P_EDIT, "수정입력"), fill=_HDR, font=_W, align=_CTR)

    blocks = [
        ("기본정보", [("사업자명", _idx("관리상품", 1, 2)), ("노출상품명", _idx("관리상품", 3, 2)),
                   ("판매방식", _idx("관리상품", 4, 2)), ("관리상태", _idx("관리상품", 7, 2)),
                   ("상품코드", _idx("관리상품", 8, 2))]),
        ("재고", [("재고현황", _idx("관리상품", 5, 2)), ("그로스재고", _idx("관리상품", 6, 2))]),
        ("체험단", [("키워드", _idx("체험단", 3, 1)), ("건수", _idx("체험단", 4, 1)),
                 ("최근요청일", _idx("체험단", 2, 1))]),
        ("정산①쿠팡", [("최종지급예정액", _idx("매출", 3, 1)), ("귀속월", _idx("매출", 2, 1))]),
        ("정산②계약", [("계약유형", _idx("계약정산", 2, 1)), ("수익(순액)", _idx("계약정산", 3, 1)),
                   ("갑배분(상품주)", _idx("계약정산", 4, 1)), ("을배분(토탈셀러)", _idx("계약정산", 5, 1))]),
    ]
    num_lbl = {"재고현황", "그로스재고", "건수", "최종지급예정액", "수익(순액)", "갑배분(상품주)", "을배분(토탈셀러)"}
    editable = {"노출상품명", "판매방식", "관리상태", "키워드"}
    r = 3
    for bname, items in blocks:
        bh = card.cell(r, P_LBL, bname)
        _style(bh, fill=_BLOCK, font=_W, align=_CTR)
        card.merge_cells(start_row=r, start_column=P_LBL, end_row=r, end_column=P_EDIT)
        r += 1
        for lbl, formula in items:
            _style(card.cell(r, P_LBL, lbl), fill=_LBL, font=_B)
            _style(card.cell(r, P_VAL, formula), num=lbl in num_lbl)
            ecell = card.cell(r, P_EDIT)
            _style(ecell, fill=_EDIT if lbl in editable else None)
            if lbl in editable:
                ecell.value = "(수정입력)"
                ecell.font = Font(italic=True, color="BF9000")
            r += 1
    return r


def build(path: Path) -> Path:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    card = wb.create_sheet("통합목록카드")
    card.sheet_properties.tabColor = "C00000"
    _mini(wb, "관리상품", "FFC000", PM, PM_ROWS)
    _mini(wb, "체험단", "4472C4", CH, CH_ROWS)
    _mini(wb, "매출", "7030A0", SA, SA_ROWS)
    _mini(wb, "계약정산", "ED7D31", CS, CS_ROWS)

    # ── 왼쪽: 계정별·상품별 목록 + 옆칸 👉 통합내용보기 ──
    for col, w in (("A", 16), ("B", 40), ("C", 14), ("D", 10), ("E", 20),
                   ("F", 3), ("G", 16), ("H", 44), ("I", 16)):
        card.column_dimensions[col].width = w
    card.cell(1, 1, "계정·상품 목록 — 옆칸 👉 통합내용보기 클릭 → 오른쪽 패널에 통합표시").font = Font(bold=True, size=12)
    card.merge_cells("A1:E1")
    for c, h in enumerate(["사업자명", "상품(물류)명", "계정ID", "상태", "통합내용보기"], 1):
        _style(card.cell(2, c, h), fill=_HDR, font=_W, align=_CTR)
    for i, row in enumerate(PM_ROWS, 3):
        _style(card.cell(i, 1, row[0]))
        _style(card.cell(i, 2, row[1]))
        _style(card.cell(i, 3, ACCOUNT_ID.get(row[0], "")))
        _style(card.cell(i, 4, row[6]))
        btn = card.cell(i, 5, "👉 통합내용보기")       # 방향표시 + 클릭 버튼
        _style(btn, fill=_BTN, font=_LINK, align=_CTR)
        btn.hyperlink = f"#'통합목록카드'!{SEL_CELL}"    # 클릭→패널로 이동(진짜 자동로드는 스크립트)

    # ── 오른쪽: 통합카드 패널(정해진 포맷) ──
    last = _panel(card)

    # ── 안내 ──
    r = max(last, len(PM_ROWS) + 4) + 1
    note = [
        "■ 왼쪽 목록의 '👉 통합내용보기' = 오른쪽 패널로 이동(HYPERLINK). 패널 선택상품(H1 드롭다운) 바꾸면 전 항목 자동 취합.",
        "■ 읽기+그룹핑(초록 블록·정해진 포맷) = 수식(INDEX/MATCH)만으로 구현.",
        "■ 진짜 '행 클릭 → 그 상품 자동 로드' = Apps Script(버튼/선택 트리거)로 H1에 그 상품명 기록(쉬움).",
        "■ 노랑 '수정입력' → 원본 시트 반영(write-back)은 수식 불가 → (A) Apps Script onEdit 또는 (B) 앱(gsheet_api). 키=상품코드.",
    ]
    for t in note:
        card.cell(r, 1, t)
        card.merge_cells(start_row=r, start_column=1, end_row=r, end_column=9)
        r += 1
    card.freeze_panes = "A3"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def main() -> int:
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else DEF_OUT
    p = build(out)
    print(f"[상품 통합목록카드 데모] 생성: {p}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
