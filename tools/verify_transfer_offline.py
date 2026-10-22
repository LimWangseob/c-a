"""셀독전환자 정산 1단계(로그인 확인) 오프라인 검증 — 로그인·쿠팡 호출 없음(결정적). D-111.

목표(소유자 2026-10-22): 미정산 셀독전환자 ①셀독진행·②계정수령_상품미등록 계정이 정상 로그인되는지부터 확인 →
계정별 결과표 + 비밀번호 오류 등으로 접속 못 한 계정은 별도 리스트(개별 연락용). 비밀번호 빈 칸 = 관리대장 값.
비밀번호 값은 결과·기록·경고 어디에도 남지 않아야 한다.

    python tools/verify_transfer_offline.py
"""
from __future__ import annotations

import json
import sys
import tempfile
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
try:
    sys.stdout.reconfigure(encoding="utf-8")
except AttributeError:
    pass

import openpyxl  # noqa: E402

from coupang_analytics import transfer_accounts as TA  # noqa: E402
from coupang_analytics import transfer_login as TL  # noqa: E402
from coupang_analytics.registry_lock import registry_lock  # noqa: E402

PW = "S3cret!pw"          # 결과 파일에 나오면 안 되는 값
FP = lambda aid, p: f"fp-{aid}-{len(p)}"   # noqa: E731  가짜 지문(실제 credstore 키 안 씀)
HEAD = ["대표자명", "연락처", "사업자명", "계정명", "비밀번호", "상품명", "생산가", "공급가", None, "판매가"]


def ok(msg):
    print(f"  ✔ {msg}")


def _sheet(title, rows):
    return [[title], HEAD] + rows


def _file(d, one, two, extra_sheets=("③계정미수령",)) -> Path:
    wb = openpyxl.Workbook()
    wb.active.title = "전체 리스트"
    for name, rows in (("①셀독진행", one), ("②계정수령_상품미등록", two)):
        ws = wb.create_sheet(name)
        for r in _sheet(name, rows):
            ws.append(r)
    for name in extra_sheets:
        ws = wb.create_sheet(name)
        for r in _sheet(name, [["무관", "", "무관", "nottarget1", PW, ""]]):
            ws.append(r)
    p = Path(d) / "전환자.xlsx"
    wb.save(p)
    return p


class _Store:
    def __init__(self, vals):
        self.vals = vals

    def get_password(self, aid):
        return self.vals.get(aid)


def t1_parse():
    print("[T1] 파일 읽기 — ①·② 시트만·머리줄 이름으로 칸 찾기·빈/형식 오류 계정명 건너뜀·공백 비번 경고(값 없음)·상품 줄바꿈")
    with tempfile.TemporaryDirectory() as d:
        p = _file(d, [["정인조", "010-1", "에이스인", "wjddlswh2", PW, "기저귀 가방\n캠핑테이블"],
                      ["홍길동", "010-2", "", "", PW, ""],                       # 계정명 없음
                      ["김형식", "010-3", "형식오류", "bad id!", PW, ""],          # 아이디 형식 아님
                      [None] * 6],                                                 # 빈 줄
                  [["이몽룡", "010-4", "", "mongryong7", f" {PW} ", "상품"],      # 공백 비번 · 사업자명 없음
                   ["성춘향", "010-5", "춘향상회", "chunhyang", None, ""]])        # 비번 빈 칸
        accts, warns = TA.read_transfer_file(p)
    assert [a.account_id for a in accts] == ["wjddlswh2", "mongryong7", "chunhyang"], [a.account_id for a in accts]
    a0, a1, a2 = accts
    assert (a0.kind, a0.row, a0.label, a0.phone, a0.products) == ("①셀독진행", 3, "에이스인", "010-1", ("기저귀 가방", "캠핑테이블"))
    assert (a1.kind, a1.row, a1.label, a1.password) == ("②계정수령_상품미등록", 3, "이몽룡", PW), "공백 떼고·이름 = 대표자명"
    assert a2.password == "", "빈 칸 = 빈 값"
    assert len(warns) == 3 and all(PW not in w for w in warns), warns
    assert any("계정명 칸이 비어" in w for w in warns) and any("형식이 아니라" in w for w in warns)
    assert any("앞뒤 공백" in w and "mongryong7" in w for w in warns)
    assert PW not in repr(a0), "repr 에 비밀번호 없음"
    assert "nottarget1" not in {a.account_id for a in accts}, "③ 시트는 대상 아님"
    ok("①·② 3계정 · ③ 제외 · 빈/형식 오류 건너뜀(경고) · 공백 비번 떼고 경고 · 경고·repr 에 비밀번호 없음 · 상품 2개")


def t2_file_errors():
    print("[T2] 파일 오류 — 없는 파일·② 시트 없음·머리줄 칸 없음·계정명 중복(시트 사이) = 중단(비밀번호 안 나옴)")
    with tempfile.TemporaryDirectory() as d:
        for bad, want in ((Path(d) / "없음.xlsx", "파일이 없음"),):
            try:
                TA.read_transfer_file(bad)
                raise AssertionError("없는 파일인데 통과")
            except TA.TransferFileError as exc:
                assert want in str(exc)
        wb = openpyxl.Workbook()
        wb.active.title = "①셀독진행"
        p = Path(d) / "one.xlsx"
        wb.save(p)
        try:
            TA.read_transfer_file(p)
            raise AssertionError("② 시트 없는데 통과")
        except TA.TransferFileError as exc:
            assert "모두 있어야" in str(exc)
        try:
            TA.parse_sheet("①x", [["제목"], ["대표자명", "계정명"]])
            raise AssertionError("머리줄 칸 없는데 통과")
        except TA.TransferFileError as exc:
            assert "비밀번호" in str(exc) and "연락처" in str(exc)
        p = _file(d, [["a", "1", "A", "dupid1", PW, ""]], [["b", "2", "B", "dupid1", "다른값", ""]])
        try:
            TA.read_transfer_file(p)
            raise AssertionError("중복인데 통과")
        except TA.TransferFileError as exc:
            assert "dupid1" in str(exc) and "두 번" in str(exc) and PW not in str(exc)
    ok("없는 파일 · ② 없음 · 머리줄 칸 없음 · 시트 사이 계정명 중복 = 중단")


def t3_passwords():
    print("[T3] 비밀번호 후보 — 파일 먼저 → 관리대장(이 PC 저장) · 빈 칸 = 관리대장 · 다르면 최대 2개 · 없으면 빈 목록")
    mk = lambda pw: TA.TransferAccount("acc1", "A", "", "", "", "①", 3, (), pw)   # noqa: E731
    assert TA.login_passwords(mk(PW), _Store({})) == ([PW], "파일")
    assert TA.login_passwords(mk(PW), _Store({"acc1": PW})) == ([PW], "파일"), "같으면 1개"
    pws, src = TA.login_passwords(mk(PW), _Store({"acc1": "ledger"}))
    assert pws == [PW, "ledger"] and "파일 먼저" in src and PW not in src
    assert TA.login_passwords(mk(""), _Store({"acc1": "ledger"})) == (["ledger"], "파일 칸 비어 관리대장 값")
    pws, src = TA.login_passwords(mk(""), _Store({}))
    assert pws == [] and "비밀번호 없음" in src
    ok("파일 · 같은 값 1개 · 다르면 파일→관리대장 2개 · 빈 칸 = 관리대장 · 둘 다 없음 = 빈 목록")


def _accts():
    return [TA.TransferAccount(f"acc{i}", f"사업자{i}", f"대표{i}", f"사업자{i}", f"010-{i}", "①셀독진행", i + 2, (), PW)
            for i in range(1, 6)]


class _Clock:
    def __init__(self, hhmm="10:00"):
        self.t = datetime(2026, 10, 22, *map(int, hhmm.split(":")))

    def __call__(self):
        return self.t


def t4_run_and_report():
    print("[T4] 실행 — 계정마다 결과 즉시 저장 · 비밀번호 없음 = 로그인 안 함 · 쉼은 로그인 사이만 · 엑셀(결과·연락 필요) · 비번 미기록")
    accts = _accts()
    plan = {"acc1": (TL.OK, "로그인 됨"), "acc2": (TL.FAIL, "로그인 비밀번호 오류 — 쿠팡 안내: 비밀번호가 다릅니다"),
            "acc4": (TL.FAIL, "2단계 인증 필요 — 계약자 휴대폰 인증번호 필요"), "acc5": (TL.SKIP, "로그인 미완료")}
    called, slept, logs = [], [], []

    def check(a, pws, hidden, log):
        called.append((a.account_id, tuple(pws), hidden))
        return plan[a.account_id]

    def pw(a):
        return ([], "비밀번호 없음(파일·관리대장 모두)") if a.account_id == "acc3" else ([PW], "파일")
    with tempfile.TemporaryDirectory() as d:
        rc = TL.run(accts, logs.append, fp=FP, out_dir=d, check=check, passwords=pw, busy=lambda: "", sleep=slept.append,
                    now=_Clock(), until="17:40")
        state = TL.load_state(d)
        files = sorted(Path(d).glob("로그인확인_*.xlsx"))
        assert rc == 0 and len(files) == 1, (rc, files)
        wb = openpyxl.load_workbook(files[0])
        res = [list(r) for r in wb["로그인 결과"].iter_rows(values_only=True)]
        need = [list(r) for r in wb[TL.CONTACT].iter_rows(values_only=True)]
        raw = (Path(d) / TL.STATE_FILE).read_text(encoding="utf-8")
        all_text = raw + json.dumps(res, ensure_ascii=False, default=str) + json.dumps(need, ensure_ascii=False, default=str)
    assert [c[0] for c in called] == ["acc1", "acc2", "acc4", "acc5"], "acc3(비번 없음) 은 로그인 안 함"
    assert all(c[1] == (PW,) and c[2] is False for c in called), "보이는 창(사람이 볼 수 있게)·비번 후보 전달"
    assert len(slept) == 3 and all(TL.GAP_SEC[0] <= s <= TL.GAP_SEC[1] for s in slept), f"로그인 사이 쉼 3번: {slept}"
    assert state["acc3"]["결과"] == TL.FAIL and "비밀번호 없음" in state["acc3"]["사유"]
    assert state["acc1"] == {"결과": TL.OK, "사유": "로그인 됨", "비밀번호 출처": "파일", "시도": "2026-10-22 10:00:00",
                             "지문": [FP("acc1", PW)]}
    assert res[0] == list(TL.HEAD) and len(res) == 6 and res[1][3] == "acc1" and res[1][5] == TL.OK
    assert need[0] == list(TL.CONTACT_HEAD)
    assert [r[2] for r in need[1:]] == ["acc2", "acc3", "acc4"], f"연락 필요 = 비번 오류·비번 없음·2단계 인증: {need}"
    assert need[1][3] == "010-2" and "비밀번호 오류" in need[1][5], "연락처·사유"
    assert PW not in all_text, "기록·엑셀에 비밀번호 값 없음"
    assert "연락처" not in raw and "010-" not in raw, "상태 파일엔 연락처 안 남김"
    assert any("성공 1 · 실패(연락 필요) 3 · 차단 0 · 건너뜀 1" in m for m in logs), logs[-1]
    ok("4계정 로그인·acc3 로그인 안 함 · 쉼 3번(45~75초) · 결과 5줄 · 연락 필요 3줄(연락처·사유) · 비번·연락처 기록 안 됨")


def t5_stops():
    print("[T5] 멈춤 — 차단 2계정 연속 = 전체 중단(rc 1) · 다른 실행 중 = 멈춤(rc 2) · 멈춤 시각 지남 = 새 계정 안 함 · 재시도 = 성공 아닌 것만")
    accts = _accts()
    pw = lambda a: ([PW], "파일")   # noqa: E731
    seq = iter([(TL.OK, "로그인 됨"), (TL.BLOCK, "차단"), (TL.BLOCK, "차단"), (TL.OK, "x"), (TL.OK, "x")])
    called = []

    def check(a, pws, hidden, log):
        called.append(a.account_id)
        return next(seq)
    with tempfile.TemporaryDirectory() as d:
        rc = TL.run(accts, lambda m: None, fp=FP, out_dir=d, check=check, passwords=pw, busy=lambda: "", sleep=lambda s: None,
                    now=_Clock())
        assert rc == 1 and called == ["acc1", "acc2", "acc3"], (rc, called)
        called.clear()
        rc = TL.run(accts, lambda m: None, fp=FP, out_dir=d, retry=True, passwords=pw, busy=lambda: "", sleep=lambda s: None,
                    now=_Clock(), check=lambda a, p, h, lg: (called.append(a.account_id), (TL.OK, "로그인 됨"))[1])
        assert rc == 0 and called == ["acc2", "acc3", "acc4", "acc5"], f"재시도 = 성공 아닌 계정만(acc1 제외): {called}"
        called.clear()
        busy = iter(["", "앱 실행(판매수집·순위)이 도는 중"])
        rc = TL.run(accts, lambda m: None, fp=FP, out_dir=d, passwords=pw, busy=lambda: next(busy), sleep=lambda s: None,
                    now=_Clock(), check=lambda a, p, h, lg: (called.append(a.account_id), (TL.OK, "로그인 됨"))[1])
        assert rc == 2 and called == ["acc1"], (rc, called)
        called.clear()
        rc = TL.run(accts, lambda m: None, fp=FP, out_dir=d, passwords=pw, busy=lambda: "", sleep=lambda s: None,
                    now=_Clock("17:40"), check=lambda a, p, h, lg: (called.append(a.account_id), (TL.OK, ""))[1])
        assert rc == 0 and called == [], "17:40 지나면 새 계정 시작 안 함(18:00 앱 실행 보호)"
    assert not TL.in_quiet("", datetime(2026, 10, 22, 18, 0)) and TL.in_quiet("17:40", datetime(2026, 10, 22, 19, 0))
    assert not TL.in_quiet("17:40", datetime(2026, 10, 22, 21, 0)) and not TL.in_quiet("17:40", datetime(2026, 10, 23, 2, 0))
    assert TL.in_quiet("17:40", datetime(2026, 10, 22, 20, 39)) and not TL.in_quiet("17:40", datetime(2026, 10, 22, 17, 39))
    ok("차단 연속 2 → 중단 · 재시도 = 성공 아닌 4계정 · 다른 실행 중 → 멈춤 · 17:40~20:40(운용 PC 실행 시간대) 시작 안 함")


def t6_busy_lock():
    print("[T6] 다른 실행 감지 — 앱 파이프라인·정산 자동 다운로드 잠금을 쥔 실행이 있으면 '도는 중'·없으면 빈 값")
    with tempfile.TemporaryDirectory() as d:
        pipe, watch = Path(d) / "_잠금" / "파이프라인.lock", Path(d) / "정산" / "_잠금" / "_watch.lock"
        assert TL.busy_reason((pipe, watch)) == ""
        with registry_lock(pipe, wait_sec=0):
            assert "앱 실행" in TL.busy_reason((pipe, watch))
        with registry_lock(watch, wait_sec=0):
            assert "정산 자동 다운로드" in TL.busy_reason((pipe, watch))
        assert TL.busy_reason((pipe, watch)) == "", "잠금 풀리면 다시 빈 값(확인만 하고 쥐지 않음)"
    ok("파이프라인 잠금·정산 잠금 감지 · 확인 뒤 잠금 안 쥠")


def t7_check_one_mapping():
    print("[T7] 계정 하나 판정 — 앱 ①과 같은 로그인 함수 결과 → 성공·실패(비번 오류·2단계 인증)·차단·건너뜀 · 브라우저는 항상 닫음")
    from coupang_analytics import browser as BR, login_today as LT, pipeline_sales as PS
    closed = []

    class FakeB:
        code = ("success", "")

        def __init__(self, **kw):
            self.kw = kw

        def __enter__(self):
            return self

        def __exit__(self, *a):
            closed.append(self.kw["offscreen"])

        def classify_login(self):
            return FakeB.code
    orig = (BR.WingBrowser, PS._login_with_candidates, LT.reasons_today)
    a = _accts()[0]
    cases = [(lambda *x, **k: True, ("success", ""), (TL.OK, "로그인 됨")),
             (lambda *x, **k: (_ for _ in ()).throw(PS.LoginBlocked()), ("blocked", ""), TL.BLOCK),
             (lambda *x, **k: (_ for _ in ()).throw(PS.LoginCredentialError("acc1")),
              ("error", "아이디 또는 비밀번호가 다릅니다"), TL.FAIL),
             (lambda *x, **k: False, ("otp", "인증번호"), TL.FAIL),
             (lambda *x, **k: False, ("akamai", ""), TL.BLOCK),
             (lambda *x, **k: False, ("form", "폼"), TL.SKIP)]
    got = []
    try:
        BR.WingBrowser = FakeB
        LT.reasons_today = lambda **k: {"acc1": LT.PW_WRONG}
        for fn, code, want in cases:
            PS._login_with_candidates = fn
            FakeB.code = code
            st, why = TL.check_one(a, [PW], False, lambda m: None)
            got.append((st, why))
            assert (st, why) == want if isinstance(want, tuple) else st == want, (code, st, why)
    finally:
        BR.WingBrowser, PS._login_with_candidates, LT.reasons_today = orig
    assert "비밀번호 오류" in got[2][1] and "비밀번호가 다릅니다" in got[2][1], got[2]
    assert "2단계 인증" in got[3][1] and "비밀번호 맞음" in got[3][1], got[3]
    assert closed == [False] * len(cases), "보이는 창·매번 닫음"
    ok("성공 · Akamai 차단 · 비번 오류(쿠팡 안내 문구) · 2단계 인증 · 챌린지 = 차단 · 미완료 = 건너뜀 · 브라우저 6번 닫음")


def t8_night_rules():
    print("[T8] 야간 이어 하기 — 멈춤 시각 = 시작 뒤 첫 그 시각 · 시작 대기 · 차단/미완료/오류 2연속 중단 · 재시도 대상 · 비번 지문")
    from datetime import timedelta
    d21, d14 = datetime(2026, 10, 22, 21, 0), datetime(2026, 10, 22, 14, 0)
    assert TL.deadline_of("17:40", d21) == datetime(2026, 10, 23, 17, 40), "밤에 시작 = 다음 날 17:40"
    assert TL.deadline_of("17:40", d14) == datetime(2026, 10, 22, 17, 40)
    assert TL.deadline_of("", d14) is None
    slept, logs = [], []
    clock = {"t": datetime(2026, 10, 22, 15, 20)}

    def fake_sleep(s):
        slept.append(s)
        clock["t"] += timedelta(seconds=s)
    TL.wait_until("21:00", logs.append, now=lambda: clock["t"], sleep=fake_sleep)
    assert clock["t"] >= datetime(2026, 10, 22, 21, 0) and any("21:00" in m for m in logs), (clock, logs)
    n = len(slept)
    TL.wait_until("21:00", logs.append, now=lambda: clock["t"], sleep=fake_sleep)
    assert len(slept) == n, "이미 지났으면 바로 시작"
    # 차단·미완료·오류 2연속 = 중단(오늘 15:08 실측: 차단 화면 → '미완료' → 다음 계정 연결 오류)
    accts = _accts()
    pw = lambda a: ([PW], "파일")   # noqa: E731
    seq = iter([(TL.SKIP, "미완료"), (TL.OK, "x"), (TL.BLOCK, "차단"), (TL.SKIP, "오류"), (TL.OK, "x")])
    called = []
    with tempfile.TemporaryDirectory() as d:
        rc = TL.run(accts, lambda m: None, out_dir=d, passwords=pw, busy=lambda: "", sleep=lambda s: None, now=_Clock(),
                    fp=lambda aid, p: f"fp:{aid}:{p}",
                    check=lambda a, p, h, lg: (called.append(a.account_id), next(seq))[1])
        assert rc == 1 and called == ["acc1", "acc2", "acc3", "acc4"], (rc, called)
        st = TL.load_state(d)
        assert st["acc1"]["지문"] == [f"fp:acc1:{PW}"] and PW not in json.dumps(
            {k: {kk: vv for kk, vv in v.items() if kk != "지문"} for k, v in st.items()}, ensure_ascii=False)
    # 재시도 대상: 성공·인증번호 제외 · 비번 오류는 비번이 바뀐 계정만 · 미완료·비번 없음은 다시
    state = {"acc1": {"결과": TL.OK, "사유": "로그인 됨"},
             "acc2": {"결과": TL.FAIL, "사유": "로그인 비밀번호 오류", "지문": [f"fp:acc2:{PW}"]},
             "acc3": {"결과": TL.FAIL, "사유": "로그인 비밀번호 오류", "지문": ["fp:acc3:옛값"]},
             "acc4": {"결과": TL.FAIL, "사유": "비밀번호 맞음 — 2단계 인증 필요"},
             "acc5": {"결과": TL.SKIP, "사유": "로그인 미완료"}}
    got = [a.account_id for a in TL.pick(accts, state, True, pw, lambda aid, p: f"fp:{aid}:{p}")]
    assert got == ["acc3", "acc5"], f"재시도 = 비번 바뀐 acc3·미완료 acc5 (acc5 다음 미확인 없음): {got}"
    accts6 = accts + [TA.TransferAccount("acc6", "x", "", "", "", "①", 9, (), "")]
    got = [a.account_id for a in TL.pick(accts6, state, True, lambda a: ([], "비밀번호 없음") if a.account_id == "acc6"
                                         else pw(a), lambda aid, p: f"fp:{aid}:{p}")]
    assert got == ["acc3", "acc5", "acc6"], got
    ok("밤 시작 → 다음 날 17:40 · 21:00 까지 대기 · 차단+오류 2연속 중단 · 재시도 = 비번 바뀐 계정·미완료·미확인만 · 지문만 저장")


def main() -> int:
    t1_parse()
    t2_file_errors()
    t3_passwords()
    t4_run_and_report()
    t5_stops()
    t6_busy_lock()
    t7_check_one_mapping()
    t8_night_rules()
    print("셀독전환자 로그인 확인 오프라인 검증 통과")
    return 0


if __name__ == "__main__":
    sys.exit(main())
