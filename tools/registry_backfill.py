"""셀독등록원장 소급 구축 — 앱이 실행마다 남긴 관리대장 백업 사본들로 과거 이력을 만든다. SSOT=designs/LEDGER_REGISTRY.md §12.

사본 출처: output 압축 파일들(기본 = 다운로드 폴더 output*.zip) 안의 `백업/관리대장_yymmdd_hhmmss.xlsx` +
로컬 output 폴더의 관리대장 백업. 같은 시각 사본은 1개만. 마지막으로 **현재 관리대장(구글시트)** 을 이어 넣는다.
빈 원장 파일에서만 실행된다(이미 데이터가 있으면 거부).

사용:
  python tools/registry_backfill.py --dry-run          # 미리보기(원장에 쓰지 않음)
  python tools/registry_backfill.py                    # 실제 구축
  python tools/registry_backfill.py --zips "C:/경로/output*.zip" --registry <원장 URL>
"""
from __future__ import annotations

import argparse
import glob
import os
import re
import sys
import tempfile
import zipfile
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "tools"))
try:
    sys.stdout.reconfigure(encoding="utf-8")
except AttributeError:
    pass

import openpyxl  # noqa: E402

from coupang_analytics import gsheet_api, registry_gsheet  # noqa: E402
from coupang_analytics.credstore import CredStore  # noqa: E402
from coupang_analytics.input_list import read_ledger_rows  # noqa: E402
from coupang_analytics.registry import RegistryGuardError, RegistryIntegrityError  # noqa: E402
from registry_sync import _shared_setting  # noqa: E402

_STAMP = re.compile(r"관리대장_(\d{6}_\d{6})\.xlsx$")


def _zip_name(info: zipfile.ZipInfo) -> str:
    if info.flag_bits & 0x800:
        return info.filename
    try:
        return info.filename.encode("cp437").decode("cp949")
    except (UnicodeEncodeError, UnicodeDecodeError):
        return info.filename


def collect_copies(zip_glob: str, out_dir: Path) -> dict:
    """압축·로컬에서 관리대장 사본을 모아 {시각: 파일경로}. 같은 시각은 1개만."""
    found: dict = {}
    for z in sorted(glob.glob(zip_glob), key=os.path.getmtime, reverse=True):
        with zipfile.ZipFile(z) as zf:
            for info in zf.infolist():
                m = _STAMP.search(_zip_name(info))
                if m and m.group(1) not in found:
                    dest = out_dir / f"관리대장_{m.group(1)}.xlsx"
                    dest.write_bytes(zf.read(info))
                    found[m.group(1)] = dest
    for f in glob.glob(str(ROOT / "output" / "**" / "*관리대장_*.xlsx"), recursive=True):
        m = _STAMP.search(f)
        if m and m.group(1) not in found:
            found[m.group(1)] = Path(f)
    return {datetime.strptime(k, "%y%m%d_%H%M%S"): v for k, v in found.items()}


def xlsx_reader(path: Path):
    """사본 xlsx → (시트명, 값격자, 취소선격자). 취소선=거래중지 판정에 필요해 서식 포함 로드."""
    def read():
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb["셀독리스트"] if "셀독리스트" in wb.sheetnames else wb.worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        strike = [[bool(getattr(c.font, "strike", False)) for c in r] for r in ws.iter_rows()]
        return ws.title, rows, strike
    return read


def main() -> int:
    ap = argparse.ArgumentParser(description="관리대장 백업 사본으로 셀독등록원장 소급 구축")
    ap.add_argument("--zips", default=str(Path.home() / "Downloads" / "output*.zip"), help="output 압축 파일 경로 패턴")
    ap.add_argument("--registry", help="셀독등록원장 URL(기본=설정 registry/url)")
    ap.add_argument("--ledger", help="현재 관리대장 URL(기본=설정 gsheet/input_url)")
    ap.add_argument("--dry-run", action="store_true", help="미리보기(원장에 쓰지 않음)")
    args = ap.parse_args()
    registry_url = args.registry or _shared_setting("registry/url")
    ledger_url = args.ledger or _shared_setting("gsheet/input_url")
    if not registry_url or not ledger_url:
        print("[중단] 원장(registry/url) 또는 관리대장(gsheet/input_url) 주소가 없습니다.")
        return 1
    store = CredStore()
    with tempfile.TemporaryDirectory() as tmp:
        copies = collect_copies(args.zips, Path(tmp))
        if not copies:
            print(f"[중단] 관리대장 사본을 찾지 못했습니다: {args.zips}")
            return 1
        print(f"== [소급] 사본 {len(copies)}개: {min(copies):%m/%d %H:%M} ~ {max(copies):%m/%d %H:%M} ==")
        snaps = [(ts, p.name, xlsx_reader(p)) for ts, p in copies.items()]
        snaps.append((datetime.now(), "현재 관리대장(구글시트)", lambda: read_ledger_rows(ledger_url, store=store)))
        try:
            client = gsheet_api.GSheetClient(registry_url, store=store, on_log=print)
            registry_gsheet.run_backfill(client, snaps, log=print, dry_run=args.dry_run)
        except (RegistryGuardError, RegistryIntegrityError) as exc:
            print(f"[중단] {exc}")
            return 2
        except (gsheet_api.GSheetError, ValueError) as exc:
            print(f"[중단] 구글시트/대장 오류 — 원장은 변경되지 않았습니다: {exc}")
            return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
