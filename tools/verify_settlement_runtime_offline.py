"""정산 다운로드·검증 런타임 오프라인 검증 — 주소 응답 해석·실행 기록·감시 판단·계정 파일·상태 화면·차단/일시 오류
구분·검증 자료 수집·검증 시트·지급일 대조·계좌 입금·정산캘린더. verify_settlement_offline.py(계산·파싱 골든)에서 분리.

로그인·쿠팡 주소·공휴일 API 호출 없음(결정적·가짜 응답). 실패 시 AssertionError → exit 1. run_checks 게이트.

    python tools/verify_settlement_runtime_offline.py
"""
from __future__ import annotations

import sys
import tempfile
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
try:
    sys.stdout.reconfigure(encoding="utf-8")
except AttributeError:
    pass

D = date.fromisoformat


def ok(msg):
    print(f"  ✔ {msg}")


def expect(exc, fn, what):
    try:
        fn()
    except exc:
        return
    raise AssertionError(f"{what} — {exc.__name__} 안 남")


def p10_wing_api_parse():
    print("[P10] 정산 주소 응답 → 정산 일정·요청 본문·다운로드 목록(실측 응답 모양·값은 가상·라이브 미호출)")
    from datetime import datetime
    from coupang_analytics import settlement_jobs as SJ
    from coupang_analytics import settlement_wing_api as API
    wresp = {"paymentReports": [
        {"payDate": "2026-01-02", "transactionCycleCode": "R", "recognitionFrom": "2025-11-03",
         "recognitionTo": "2025-11-30", "paymentStatus": "DONE", "ratio": 30},
        {"payDate": "2026-01-23", "transactionCycleCode": "W", "recognitionFrom": "2026-01-01",
         "recognitionTo": "2026-01-04", "paymentStatus": "DONE", "ratio": 70}], "totalAmount": 0}
    ev = API.wing_events(wresp, "계정A")
    assert [(e.kind, e.settle_date, e.period_start, e.period_end) for e in ev] == [
        ("최종액", D("2026-01-02"), D("2025-11-03"), D("2025-11-30")),
        ("주정산", D("2026-01-23"), D("2026-01-01"), D("2026-01-04"))], ev
    expect(API.SiteChangedError, lambda: API.wing_events(
        {"paymentReports": [{**wresp["paymentReports"][0], "transactionCycleCode": "X"}]}, "a"), "모르는 주기 코드")
    expect(API.SiteChangedError, lambda: API.wing_events({"reports": []}, "a"), "응답 칸 없음")
    rresp = {"settlementStatusReports": [
        {"settlementDate": "2026-09-06T15:00:00.000Z", "settlementPeriodStartDate": "2026-07-26T15:00:00.000Z",
         "settlementPeriodEndDate": "2026-07-30T15:00:00.000Z", "settlementRatio": 30, "settlementCycle": "WEEKLY",
         "settlementGroupKey": "V1-2026-07-27-2026-07-31", "finalSettlementAmount": 1}]}
    rev = API.rg_events(rresp, "계정A")
    assert [(e.kind, e.settle_date, e.period_start, e.period_end, e.ref) for e in rev] == [
        ("최종액", D("2026-09-07"), D("2026-07-27"), D("2026-07-31"), "V1-2026-07-27-2026-07-31")], rev   # UTC→한국 날짜
    expect(API.SiteChangedError, lambda: API.rg_events({"settlementStatusReports": [
        {**rresp["settlementStatusReports"][0], "settlementCycle": "DAILY"}]}, "a"), "모르는 RG 주기")
    assert API.rg_events_body(D("2026-01-01"), D("2026-01-31")) == {
        "startDate": "2025-12-31T15:00:00.000Z", "endDate": "2026-01-31T15:00:00.000Z", "searchDateType": "PAYMENT"}
    assert API.wing_events_body(D("2026-01-01"), D("2026-01-31"))["fromDate"] == "2026-01-01"
    jobs: list = []
    picked = SJ.plan_requests(ev + rev, jobs, today=D("2026-10-06"), now=datetime(2026, 10, 6, 15), per_account_cap=9)
    rj = next(j for j in picked if j.channel == "로켓그로스")
    assert rj.ref == "V1-2026-07-27-2026-07-31"                                     # 일정의 키가 요청까지 전달
    assert API.rg_request_body(rj, 1791000000000) == {"sellerReportType": "CATEGORY_TR", "requestTime": "1791000000000",
                                                      "settlementGroupKeys": [rj.ref], "locale": "ko"}
    wj = next(j for j in picked if j.channel == "윙" and j.kind == "주정산")
    assert API.wing_request_body(wj) == {"excelType": "MSF_PAYMENT_REVENUE_DETAIL",
                                         "recognitionDateRange": {"start": "2026-01-01", "end": "2026-01-04"},
                                         "searchDateType": "CONFIRM_DATE",
                                         "searchDateRange": {"start": "2026-01-01", "end": "2026-01-04"}}
    expect(API.SiteChangedError, lambda: API.rg_request_body(SJ.Job("a", "로켓그로스", "주정산", "판매수수료", "x", "y", "z"),
                                                              1), "묶음 키 없음")
    expect(API.SiteChangedError, lambda: API.rg_request_body(SJ.Job("a", "로켓그로스", "주정산", "새 비용", "x", "y", "z",
                                                                     ref="k"), 1), "코드 미확인 리포트")
    assert API.wing_request_id({"success": True, "reason": "OK", "data": 134}) == "134"
    expect(API.SiteChangedError, lambda: API.wing_request_id({"success": False, "reason": "LIMIT", "data": None}), "거절")
    assert API.rg_request_id({"requestId": "abc", "duplicateRequest": False, "remainingTime": 0}) == "abc"
    expect(API.RequestThrottled, lambda: API.rg_request_id({"requestId": "", "duplicateRequest": True,
                                                            "remainingTime": 30}), "대기시간 응답")
    # 비용 리포트: 그 주 금액 있는 종류만(입출고+배송 → 1종으로 합침)·0원 제외·미확인 비용은 알림
    det = {"totalStorageFeeDeductionAmount": 201287, "totalFulfillmentFeeDeductionAmount": 843508,
           "totalWarehousingFeeDeductionAmount": 867526, "totalBarcodeLabelingFeeDeductionAmount": 0,
           "totalCfsInventoryCompensationAmount": 119770, "totalAdSalesDeductionAmount": 1003081,
           "totalContainerUnloadingFeeDeductionAmount": 5000}
    assert API.cost_reports(det) == ("보관비", "입출고/배송비", "재고 손실 보상")
    r70 = {**rresp["settlementStatusReports"][0], "settlementRatio": 70, "settlementStatusReportDetail": det}
    ev70 = API.rg_events({"settlementStatusReports": [r70]}, "계정A")[0]
    assert ev70.reports == ("보관비", "입출고/배송비", "재고 손실 보상") and SJ.reports_for(ev70)[0] == "판매수수료"
    assert API.unmapped_costs({"settlementStatusReports": [r70]}) == [
        "V1-2026-07-27-2026-07-31: totalContainerUnloadingFeeDeductionAmount=5000"]
    cj = SJ.Job("a", "로켓그로스", "주정산", "입출고/배송비", "x", "y", "z", ref="k")
    assert API.rg_request_body(cj, 1)["sellerReportType"] == "WAREHOUSING_SHIPPING"
    am = API.amount_rows({"paymentReports": [{**wresp["paymentReports"][1], "finalPaidAmount": 19664,
                                              "bankAccountInfo": {"bank": "가림"}, "isAdditionalPayment": False,
                                              "detail": {"paidAmount": 19664, "isActualPayment": True,
                                                         "payableSummarySeqList": [1, 2],
                                                         "deductionDetail": [{"serviceFeeAmount": 55000, "serviceType": "CLR_3SF",
                                                                              "serviceTypeNameKR": "판매자서비스이용료",
                                                                              "serviceTypeNameEN": "SELLER SERVICE FEE"}]}}]}, "윙")
    assert am == [{"정산일": "2026-01-23", "기간 시작": "2026-01-01", "기간 끝": "2026-01-04", "지급비율": 70,
                   "최종지급액": 19664, "지급상태": "DONE", "실지급": True,
                   "차감 사유": [{"사유": "판매자서비스이용료", "금액": 55000}], "paidAmount": 19664}], am   # 계좌·내부번호 제외
    r70x = {**r70, "settlementStatusReportDetail": {**det, "pastDeductedCfsFeeDetails": [{"paymentRatio": 30, "pastDeductedCfsFeeAmount": 24959.0}],
                                                    "adSalesOffsetYearMonthBreakdown": {"2026-08": 59266.0}, "salesAdjustmentDetails": []}}
    ar = API.amount_rows({"settlementStatusReports": [r70x]}, "로켓그로스")[0]
    assert ar["정산일"] == "2026-09-07" and ar["totalStorageFeeDeductionAmount"] == 201287
    assert ar["pastDeductedCfsFeeDetails"] == [{"paymentRatio": 30, "pastDeductedCfsFeeAmount": 24959.0}]   # 상계 상세 보존
    assert ar["adSalesOffsetYearMonthBreakdown"] == {"2026-08": 59266.0} and "salesAdjustmentDetails" not in ar   # 빈 목록은 생략
    wl = API.wing_list_rows([
        {"id": 1, "excelType": "MSF_PAYMENT_REVENUE_DETAIL", "status": "FINISHED", "startedAt": "2026-10-06 14:49:59",
         "downloadUrl": "https://x/dl?id=1", "jsonItems": '[{"key":"구매확정일", "value":"2026-09-01 - 2026-09-06", "view":true}]'},
        {"id": 2, "excelType": "MSF_VAT_DETAIL", "status": "FINISHED", "startedAt": "2025-02-19 05:07:42",
         "downloadUrl": "https://x/dl?id=2", "jsonItems": '[{"key":"구매확정일", "value":"2025-01-01 - 2025-01-31"}]'}])
    assert [(r.requested_at, r.status, r.period_start, r.period_end, r.handle) for r in wl] == [
        (datetime(2026, 10, 6, 14, 49, 59), "FINISHED", D("2026-09-01"), D("2026-09-06"), "https://x/dl?id=1")]
    rl = API.rg_list_rows([{"requestId": "abc", "requestTime": "1791000000000", "downloadStatus": "COMPLETED",
                            "sellerReportType": "CATEGORY_TR", "recognitionDateFrom": None}])
    assert (rl[0].report, rl[0].status, rl[0].req_id, rl[0].handle) == ("판매수수료", "COMPLETED", "abc", "1791000000000")
    assert API.rg_list_body(datetime(2026, 10, 6, 0, 0), datetime(2026, 10, 6, 0, 0, 1)) == {
        "requestTimeFrom": str(int(datetime(2026, 10, 6).timestamp() * 1000)),
        "requestTimeTo": str(int(datetime(2026, 10, 6).timestamp() * 1000) + 1000)}
    assert API.month_windows(D("2026-01-15"), D("2026-03-02")) == [
        (D("2026-01-15"), D("2026-01-31")), (D("2026-02-01"), D("2026-02-28")), (D("2026-03-01"), D("2026-03-02"))]

    class FakePage:
        def __init__(self, status, text):
            self.r = {"status": status, "ctype": "x", "text": text}

        def evaluate(self, js, arg=None):
            return self.r
    assert API.call(FakePage(200, '{"a":1}'), "GET", "/p") == {"a": 1}
    expect(API.ApiBlocked, lambda: API.call(FakePage(403, "x"), "GET", "/p"), "403=차단")
    expect(API.ApiBlocked, lambda: API.call(FakePage(200, "<html>Access Denied"), "GET", "/p"), "HTML=차단")
    expect(API.SiteChangedError, lambda: API.call(FakePage(500, "{}"), "GET", "/p"), "500=중단")
    for st in (502, 503, 504):                                                      # 실측 2026-10-08: 504 HTML = 서버 일시 오류
        expect(API.ServerBusy, lambda st=st: API.call(FakePage(st, "<html><title>%d Gateway Time-out</title>" % st), "GET", "/p"),
               f"{st}=일시 오류(차단 아님)")
        assert not issubclass(API.ServerBusy, API.ApiBlocked)
    ok("윙/RG 일정(UTC→한국 날짜·묶음키)·요청 본문·요청번호·목록(부가세 메뉴 무시)·달 구간·403/HTML=차단·500=중단")


def p11_runlog():
    print("[P11] 실행 기록 — 처리기록 CSV(실행·누적)·로그인/차단/실패 기록·오류 전체 추적·요약")
    import csv as _csv
    from datetime import datetime
    from coupang_analytics import settlement_runlog as RLG
    with tempfile.TemporaryDirectory() as tmp:
        out: list = []
        lg = RLG.RunLog(tmp, now=datetime(2026, 10, 6, 21, 0, 0), echo=out.append)
        lg.record("A-1", "로그인", RLG.OK, "success: 로그인 완료")
        lg.record("A-1", "요청", RLG.OK, "다운로드 요청함", channel="윙", kind="주정산·주문상세", settle_date="2026-01-12",
                  period="2025-12-29~2026-01-04")
        lg.record("B-2", "차단감지", RLG.BLOCK, "차단 화면 — https://wing.coupang.com/x")
        try:
            raise ValueError("표 머리글 없음")
        except ValueError as exc:
            lg.error("B-2", "조회", exc, channel="윙", settle_date="2026-01-13")
        lg2 = RLG.RunLog(tmp, now=datetime(2026, 10, 7, 21, 0, 0), echo=out.append)
        lg2.record("A-1", "받기", RLG.WAIT, "아직 WAIT", channel="윙")
        rows = list(_csv.DictReader(open(lg.csv, encoding="utf-8-sig")))
        assert [r["결과"] for r in rows] == ["정상", "정상", "차단", "실패"] and rows[1]["정산일"] == "2026-01-12"
        assert (rows[3]["채널"], rows[3]["정산일"]) == ("윙", "2026-01-13")                    # 실패 줄에도 정산일
        total = list(_csv.DictReader(open(lg.total, encoding="utf-8-sig")))
        assert len(total) == 5 and {r["실행ID"] for r in total} == {"261006_210000", "261007_210000"}   # 누적
        tb = lg.errors.read_text(encoding="utf-8")
        assert "Traceback" in tb and "ValueError: 표 머리글 없음" in tb and "B-2 · 조회" in tb
        assert "⛔ [차단감지] B-2" in lg.text.read_text(encoding="utf-8")
        summ = lg.summary()
        assert "[요약] A-1: 정상 2" in summ and "[요약] B-2: 실패 1 · 차단 1" in summ, summ
        assert "[비정상] 조회 실패: 1건" in summ and "[비정상] 차단감지 차단: 1건" in summ
        # heartbeat: 창 없는 자동 실행 모니터링 — 고정 경로(_현재상태.txt)에 **덮어쓰기**(append 아님)·개인정보 없음.
        lg.heartbeat("대기중", "앱 ①판매수집 완료 기다림")
        st = lg.status
        assert st.name == "_현재상태.txt" and st.parent == lg.dir                    # 로그 폴더 안 고정 파일
        txt1 = st.read_text(encoding="utf-8")
        assert "상태: 대기중" in txt1 and "앱 ①판매수집 완료 기다림" in txt1
        assert f"실행ID: {lg.run_id}" in txt1 and f"상세 로그: {lg.text.name}" in txt1
        lg.heartbeat("작동중", "요청 누르는 중", extra=("추가줄",))
        txt2 = st.read_text(encoding="utf-8")
        assert "상태: 작동중" in txt2 and "추가줄" in txt2
        assert "대기중" not in txt2 and txt2.count("상태:") == 1                      # 누적 아님 — 최신 1건만
        assert not st.with_name(st.name + ".tmp").exists()                           # 임시파일은 남기지 않음(os.replace)
        lg2.heartbeat("완료", "명령 run")                                             # 다른 RunLog 도 같은 고정 파일을 공유
        assert lg2.status == st and "상태: 완료" in st.read_text(encoding="utf-8")
    ok("실행·누적 CSV(엑셀용)·정산일 칸·차단/로그인 기록·오류 전체 추적·계정별/단계별 요약·상태 파일 덮어쓰기 모니터링")


def p12_watch_decide():
    print("[P12] 24h 감시 판단 — ①판매수집 중 정지·소급 연속·받을 것 없으면 다음 ①완료까지 대기")
    import os as _os
    from datetime import datetime
    from coupang_analytics import settlement_watch as W
    T = datetime
    # plan_watch: ① 진행 중이면 정지(busy 사유 전달)
    a, r = W.plan_watch(True, "①판매수집 진행 중", T(2026, 10, 7, 19, 0), None)
    assert a == "wait" and "진행 중" in r
    # 소급(last_sales_at=None): ①완료 기록이 있든 없든 항상 실행
    assert W.plan_watch(False, "", None, None) == ("run", "소급 수집(밀린 정산 받는 중)")
    assert W.plan_watch(False, "", T(2026, 10, 7, 19, 30), None)[0] == "run"
    # 정상(소급 끝·last_sales_at 있음): 새 ①완료가 있어야만 1회
    a, r = W.plan_watch(False, "", T(2026, 10, 7, 19, 30), T(2026, 10, 7, 19, 30))
    assert a == "wait" and "없음" in r                                              # 같은 ①완료 = 이미 함
    a, r = W.plan_watch(False, "", T(2026, 10, 8, 19, 30), T(2026, 10, 7, 19, 30))
    assert a == "run" and "새 ①판매수집" in r                                        # 다음날 새 ①완료 = 1회
    assert W.plan_watch(False, "", None, T(2026, 10, 7, 19, 30))[0] == "wait"        # ①완료 기록 없음 = 대기
    # sales_in_progress: _진행중.json 수정 시각(신선=진행 중)
    with tempfile.TemporaryDirectory() as tmp:
        prog = Path(tmp) / "진행중.json"
        assert W.sales_in_progress(prog)[0] is False                                # 파일 없음
        prog.write_text("{}", encoding="utf-8")
        _os.utime(prog, (T(2026, 10, 7, 12, 0, 0).timestamp(),) * 2)
        assert W.sales_in_progress(prog, now=T(2026, 10, 7, 12, 10))[0] is True      # 10분 전 = 진행 중
        assert W.sales_in_progress(prog, now=T(2026, 10, 7, 12, 40))[0] is False     # 40분 전 = 아님(①끝남)
        # read_marker 왕복
        mp = Path(tmp) / "단계.json"
        assert W.read_marker(mp)[0] is None and "없음" in W.read_marker(mp)[1]
        mp.write_text('{"date": "2026-10-07", "stage": "sales", "at": "2026-10-07T19:30:00"}', encoding="utf-8")
        assert W.read_marker(mp)[0] == {"stage": "sales", "at": T(2026, 10, 7, 19, 30)}
        mp.write_text("{깨짐", encoding="utf-8")
        assert W.read_marker(mp)[0] is None and "읽기 실패" in W.read_marker(mp)[1]
    ok("① 진행 중 정지·소급 연속(last_sales_at None)·정상 1회(새 ①완료)·진행 파일 신선도·단계 기록 왕복")


def p13_accounts_file():
    print("[P13] 정산 계정 파일 — 탭/공백 섞임·뒤 공백·3열=비번 줄은 이름 없음(비번 비노출)·중복/칸 없음=오류·계정명 맞춤")
    from coupang_analytics import settlement_accounts as SA
    from coupang_analytics import settlement_files as SF
    text = ("acc1\tpw1!!\t홍길동-주식회사 가나\n"           # 이름에 공백
            "acc2 \tsecret#9\tsecret#9\n"                  # 3열 = 비번(실제 파일에 있는 모양)
            "acc3  pw3\t\n"                                 # 이름 없음·공백 구분
            "plan_it \tpw4* \t 김다라_플랜\n"               # 뒤 공백·밑줄 ID
            "\n\t\n")
    accts, warns = SA.parse_accounts_text(text)
    assert [(a.account_id, a.label, a.password) for a in accts] == [
        ("acc1", "홍길동-주식회사 가나", "pw1!!"), ("acc2", "acc2", "secret#9"), ("acc3", "acc3", "pw3"),
        ("plan_it", "김다라_플랜", "pw4*")]
    assert len(warns) == 2 and "비밀번호와 같아" in warns[0] and not any("secret" in w for w in warns)   # 비번 비노출
    assert "secret" not in repr(accts) and "pw1" not in repr(accts)                                       # repr 에도 없음
    expect(SA.AccountsFileError, lambda: SA.parse_accounts_text("acc1 a 이름\nacc1 b 이름"), "중복 계정ID")
    expect(SA.AccountsFileError, lambda: SA.parse_accounts_text("acc1\n"), "비번 칸 없음")
    expect(SA.AccountsFileError, lambda: SA.parse_accounts_text("\n\n"), "빈 파일")
    with tempfile.TemporaryDirectory() as tmp:
        expect(SA.AccountsFileError, lambda: SA.read_accounts_file(Path(tmp) / "없음.txt"), "파일 없음")
        (Path(tmp) / "a.txt").write_bytes("acc1\tpw\t이름".encode("cp949"))
        expect(SA.AccountsFileError, lambda: SA.read_accounts_file(Path(tmp) / "a.txt"), "UTF-8 아님")
        (Path(tmp) / "b.txt").write_text("﻿acc1\tpw\t이름", encoding="utf-8")
        assert SA.read_accounts_file(Path(tmp) / "b.txt")[0][0].account_id == "acc1"                  # BOM 허용
    assert SA.display_name(accts[0]) == "홍길동-주식회사 가나-acc1"
    olds = ["(주)가나-acc1", "홍길동-주식회사 가나-acc1", "옛이름-plan_it", "남의-xacc3", "모름-zzz"]
    assert SA.renames(olds, accts) == {"(주)가나-acc1": "홍길동-주식회사 가나-acc1", "옛이름-plan_it": "김다라_플랜-plan_it"}
    tok = lambda x: SF._token(x, "계정명")
    assert SA.renames(["(주)가나-acc1", "옛이름-plan-it"], accts, tok=tok) == {
        "(주)가나-acc1": "홍길동-주식회사-가나-acc1", "옛이름-plan-it": "김다라-플랜-plan-it"}
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    import settlement_download as TD
    from coupang_analytics import settlement_jobs as SJ
    here = [SJ.Job("홍길동-주식회사 가나-acc1", "윙", "주정산", "주문상세", "2026-01-05", "2026-01-11", "2026-01-30",
                   status=SJ.ST_DONE, file="홍길동-주식회사-가나-acc1_20260130_윙_주정산_주문상세_20260105-20260111.xlsx")]
    came = [SJ.Job("(주)가나-acc1", "윙", "주정산", "주문상세", "2026-01-05", "2026-01-11", "2026-01-30",
                   file="(주)가나-acc1_20260130_윙_주정산_주문상세_20260105-20260111.xlsx"),
            SJ.Job("(주)가나-acc1", "윙", "최종액", "주문상세", "2026-01-01", "2026-01-31", "2026-03-02",
                   file="비용/(주)가나-acc1_20260302_윙_최종액_주문상세_20260101-20260131.xlsx")]
    TD._rename_jobs(came, accts)
    assert came[1].file == "비용/홍길동-주식회사-가나-acc1_20260302_윙_최종액_주문상세_20260101-20260131.xlsx"
    assert SJ.merge_jobs(here, came) == (1, 1) and here[0].status == SJ.ST_DONE     # 옛 이름이어도 같은 작업은 한 번
    ok("공백/탭 혼합·3열=비번→이름 없음(경고·repr 비노출)·중복/칸 없음/빈 파일/UTF-8 아님=오류·BOM·계정ID 끝 기준 이름 맞춤"
       "·다른 PC 기록 합치기(옛 이름 맞춘 뒤 중복 없음)")


def p14_status_reader():
    print("[P14] 상태 화면 리더 — heartbeat(_현재상태.txt)↔UI 리더 계약·멈춤 의심 판정·오늘 집계")
    import os as _os
    import re as _re
    from datetime import datetime as _dtm
    sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "ui"))
    import settlement_status_panel_qt as UI               # PySide6 는 이 프로젝트 필수 의존(앱 UI)
    from coupang_analytics import settlement_runlog as RLG
    fixed = _dtm(2026, 10, 7, 12, 55, 0)
    with tempfile.TemporaryDirectory() as tmp:
        lg = RLG.RunLog(tmp, now=fixed)                             # RunLog.dir = <tmp>/로그
        lg.heartbeat("작동중", "①판매수집 완료 — 재개 · 멈춤 10-07 17:40")
        lg.record("A-1", "요청", RLG.OK, "다운로드 요청함", channel="윙", settle_date="2026-01-12")
        lg.record("B-2", "요청", RLG.BLOCK, "차단 화면", channel="윙", settle_date="2026-01-12")
        # heartbeat 의 '갱신'은 실제 now(생존신호) → 시간 판정만 고정(형식·라벨은 heartbeat 실제 출력 그대로 검증).
        lg.status.write_text(_re.sub(r"갱신: .*", f"갱신: {fixed:%Y-%m-%d %H:%M:%S}",
                                     lg.status.read_text(encoding="utf-8")), encoding="utf-8")
        _os.utime(lg.text, (fixed.timestamp(),) * 2)               # 로그 mtime 도 고정
        info = UI.read_settlement_status(lg.dir, now=_dtm(2026, 10, 7, 12, 57, 0))
        assert info["exists"] and info["상태"] == "작동중" and not info["stale"], info   # 2분 전 = 정상
        assert info["내용"].startswith("①판매수집 완료") and info["log_name"] == lg.text.name
        cnt = UI.read_settlement_status(lg.dir, now=_dtm.now())["today_counts"]    # 누적 CSV 는 실제 '시각'(오늘)
        assert cnt.get(RLG.OK) == 1 and cnt.get(RLG.BLOCK) == 1, cnt              # 오늘 쓴 2줄 집계
        late = UI.read_settlement_status(lg.dir, now=_dtm(2026, 10, 7, 13, 20, 0))
        assert late["stale"] is True, late                                       # 25분 무변화 = 멈춤 의심
        empty = Path(tmp) / "빈폴더"
        assert UI.read_settlement_status(empty)["exists"] is False               # 로그 없음 = 아직 실행 안 됨
        assert UI._counts_text({}) == "오늘 받은 기록 없음" and "분 전" in UI._age_text(180)
    ok("heartbeat 덮어쓰기 파싱·작동중/대기중 멈춤 의심(STALE_SEC)·오늘 결과 집계·로그 없음/빈값 표기")


def p15_browser_closed_transient():
    print("[P15] 브라우저 닫힘(reap)=일시적 — '차단'·연속 실패로 안 셈·다음 바퀴 재시도")
    sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "tools"))
    import settlement_download as SD
    # 브라우저/탭 닫힘 오류는 True(일시적), 그 외 오류는 False(진짜 실패)
    assert SD._is_browser_closed(Exception("Target page, context or browser has been closed"))
    assert SD._is_browser_closed(type("TargetClosedError", (Exception,), {})("x"))
    assert SD._is_browser_closed(Exception("Connection closed while reading from the driver"))
    assert not SD._is_browser_closed(Exception("Access Denied"))
    assert not SD._is_browser_closed(ValueError("표 머리글 없음"))
    ok("TargetClosedError·'has been closed'·'Connection closed'=일시적(True)·그 외(차단/파싱 오류)=False")


def p16_cost_by_product():
    print("[P16] 로켓그로스 비용 상품별 — 상세 줄→상품·월(VAT 별도 정확·세액 비율 배분 합 일치)·보상=옵션→상품·박스 여러 상품=묶음")
    from coupang_analytics import settlement_files as SF
    from coupang_analytics import settlement_stats as ST
    end = D("2025-12-21")
    head = [[], ["쿠팡풀필먼트서비스(CFS) 배송비 정산 내역"], ["정산주기(종료일)", "배송비 합계", "세액", "최종비용"],
            ["2025-12-21", 1000.0, 101.0, 1101.0], [], [],
            ["정산유형", "정산주기(종료일)", "매출인식일", "등록상품 ID", "옵션ID", "등록상품명", "판매수량",
             "쿠팡풀필먼트서비스(CFS) 배송비 (VAT 별도)", None, None, None, None],
            [None, None, None, None, None, None, None, "발생비용(A)", "할인가(B)", "할인적용가(A-B)", "추가비용", "최종비용"]]
    ship = head + [["주정산", "2025-12-21", "2025-12-15", "77", "701", "보냉백", 1, 900, 0, 900, 0, 700.0],
                   ["주정산", "2025-12-21", "2025-12-16", "77", "701", "보냉백", 1, 900, 0, 900, 0, -100.0],   # 음수 조정 줄
                   ["주정산", "2025-12-21", "2026-01-02", "88", "801", "압축팩", 1, 400, 0, 400, 0, 400.0]]
    lines = SF.rg_cost_lines({"배송비": ship}, "A", end)
    got = [(c.product_id, c.month, c.ex_vat, c.vat) for c in lines]
    assert got == [("77", "2025-12", 600, 61), ("88", "2026-01", 400, 40)], got   # '최종비용' 칸·월=매출인식일·세액 합 101
    bad = [r[:] for r in ship]
    bad[8][11] = 701.0                                                              # 상세 합 ≠ 요약 합계
    expect(SF.SettlementParseError, lambda: SF.rg_cost_lines({"배송비": bad}, "A", end), "상세 합 불일치")
    ret = [[], ["반출 배송"], ["정산주기(종료일)", "합계", "세액", "최종비용"], ["2025-12-21", 3000.0, 300.0, 3300.0], [], [],
           ["정산유형", "정산주기(종료일)", "매출인식일", "대표 등록상품 ID", "대표 등록상품명", "비용 청구 수량", "X", None],
           [None, None, None, None, None, None, "발생비용(A)", "최종비용(A-B-C)"]]
    ret += [["주정산", "2025-12-21", "2025-12-15", "77,77", "보냉백,보냉백", 1, 0, 1000.0],      # 같은 상품 반복 = 그 상품
            ["주정산", "2025-12-21", "2025-12-15", "77,88", "보냉백,압축팩", 1, 0, 1000.0],      # 서로 다른 상품 = 묶음
            ["주정산", "2025-12-21", "2025-12-15", "-", "-", 1, 0, 1000.0]]                    # 미표기
    rl = SF.rg_cost_lines({"반출 배송 서비스비 리포트": ret}, "A", end)
    assert sorted(c.product_id for c in rl) == sorted(["77", SF.MULTI_PRODUCT, SF.NO_PRODUCT])
    assert sum(c.vat for c in rl) == 300
    comp = [[None] * 3, ["발생일", "정산주기(종료일)", "주문ID", "옵션ID", "등록상품명", "보상 금액"],
            [None] * 6, ["2025-12-02", "2025-12-21", "1", "701", "보냉백", 4775.0],
            ["2025-12-09", "2025-12-21", "2", "999", "모르는상품", 1000.0]]
    cl = SF.rg_cost_lines({"재고 손실 보상": comp}, "A", end)
    assert [(c.kind, c.option_id, c.ex_vat, c.vat) for c in cl] == [
        (SF.COMPENSATION, "701", 4775, 0), (SF.COMPENSATION, "999", 1000, 0)]
    assert SF._split(101, [900, -100, 400]) == [76, -9, 34] and sum(SF._split(7, [1, 1, 1])) == 7
    expect(SF.SettlementParseError, lambda: SF._split(5, [1, -1]), "합 0 인데 나눌 값")
    res = ST.aggregate([], lines + rl + cl)
    by = {p.product_id: p for p in res.products if p.month == "2025-12"}
    assert by["77"].rg_cost == 661 + 1100 and by["77"].rg_comp == 4775            # 배송 + 반출배송(같은 상품 반복)
    assert by["77"].settle_after_costs == 0 - 1761 + 4775 and by["77"].costs["배송비"] == 661
    assert "옵션 999" in by and any("옵션ID 로 상품을 못 찾아" in w for w in res.warnings)
    assert any(SF.MULTI_PRODUCT in w for w in res.warnings) and any(SF.NO_PRODUCT in w for w in res.warnings)
    import openpyxl
    with tempfile.TemporaryDirectory() as tmp:
        wb = openpyxl.load_workbook(ST.write_stats(Path(tmp) / "c.xlsx", res))
        ws = wb["로켓그로스 비용 상품별"]
        hdr = [c.value for c in ws[1]]
        assert hdr[:5] == ["계정", "월(매출인식)", "등록상품ID", "상품명", "배송비"] and hdr[-2:] == ["물류비 합계(VAT포함)", "재고 손실 보상"]
        am = [c.value for c in wb["계정별 월별"][2]]
        assert am[5:] == [sum(p.rg_cost for p in by.values()), 5775, -sum(p.rg_cost for p in by.values()) + 5775]
    ok("상세→상품·월(음수 줄 포함)·세액 배분 합 일치·상세 합 불일치=오류·박스 같은 상품/여러 상품/미표기·보상 옵션→상품·"
       "못 찾음=옵션 줄+경고·상품/계정 시트 물류비·보상·반영 후")


def p17_server_busy():
    print("[P17] 서버 일시 오류(502·503·504) — 재시도 후 성공·끝내 실패=대기(차단·연속실패 아님)·전체 중단 안 함")
    import contextlib
    from datetime import datetime
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    import settlement_download as TD
    from coupang_analytics import settlement_runlog as RLG
    from coupang_analytics import settlement_wing_api as API

    class SeqPage:                                                      # 응답 차례대로
        def __init__(self, seq):
            self.seq = list(seq)

        def evaluate(self, js, arg=None):
            st, text = self.seq.pop(0)
            return {"status": st, "ctype": "x", "text": text}

    class B:
        def __init__(self, seq):
            self.page = SeqPage(seq)
    old_retry, old_log = TD.SERVER_BUSY_RETRY, TD.LOG
    with tempfile.TemporaryDirectory() as tmp:
        TD.LOG = RLG.RunLog(tmp, now=datetime(2026, 10, 8, 12), echo=lambda m: None)
        TD.SERVER_BUSY_RETRY = (0, 2)
        try:
            assert TD.api(B([(504, "<html>"), (502, "<html>"), (200, '{"ok":1}')]), "A", "POST", "/p", {}) == {"ok": 1}
            expect(API.ServerBusy, lambda: TD.api(B([(504, "<html>")] * 3), "A", "POST", "/p", {}), "3번 다 504")
            expect(TD.BlockDetected, lambda: TD.api(B([(403, "x")]), "A", "POST", "/p", {}), "403=차단")
            assert [r["결과"] for r in TD.LOG.rows] == [RLG.BLOCK]                 # 504 는 차단 기록 없음

            @contextlib.contextmanager
            def fake_session(a, hidden):
                yield None

            def busy_run(a, b, args, jobs):
                raise API.ServerBusy("/tenants/x → 504 서버 일시 오류")

            class Acc:
                def __init__(self, i):
                    self.account_id, self.label = i, i

            class Args:
                until, until_at, hidden, command = "", None, True, "run"
            old_session, old_name = TD.session, TD.name_of
            TD.session, TD.name_of = fake_session, (lambda a: a.account_id)
            try:
                TD.LOG = RLG.RunLog(tmp, now=datetime(2026, 10, 8, 13), echo=lambda m: None)
                rc = TD._run_accounts([Acc("a1"), Acc("a2"), Acc("a3")], busy_run, Args)
            finally:
                TD.session, TD.name_of = old_session, old_name
            assert rc == 0, rc                                                      # 연속 2계정이어도 전체 중단 안 함
            res = [(r["계정"], r["결과"]) for r in TD.LOG.rows if r["단계"] == "계정"]
            assert res == [("a1", RLG.WAIT), ("a2", RLG.WAIT), ("a3", RLG.WAIT)], res
            assert not any(r["결과"] == RLG.BLOCK for r in TD.LOG.rows)
        finally:
            TD.SERVER_BUSY_RETRY, TD.LOG = old_retry, old_log
    ok("504→502→성공=값 반환·3번 실패=ServerBusy·403=차단 기록·계정 단위 일시 오류=대기·연속 3계정도 전체 중단 안 함")


def p18_verify_collect():
    print("[P18] 검증 자료 수집 — 월 끝=어제·화면 먼저 열기·월렛 여러 쪽·계좌 칸 제거·한 달 일시 오류여도 계속")
    from datetime import date
    from coupang_analytics import settlement_verify_collect as VC
    from coupang_analytics import settlement_wing_api as API
    ms = VC.months(D("2026-01-01"), D("2026-10-08"))
    assert len(ms) == 10 and ms[0] == ("202601", D("2026-01-01"), D("2026-01-31")) and ms[-1] == ("202610", D("2026-10-01"), D("2026-10-07"))
    assert VC.months(D("2026-10-01"), D("2026-10-01")) == []                          # 어제 이전 달 없음
    log: list = []

    def open_page(url):
        log.append(("화면", url))

    def call(method, path, body):
        log.append((method, path))
        if path == VC.WALLET_HIST:
            return {"walletHistories": [{"amount": 10 + body["pageNumber"], "bankAccountOwner": "예금주"}], "totalPage": 2}
        if path == VC.WALLET_BAL:
            return {"content": {"balanceAmount": 21, "bankAccountNumber": "123"}}
        if path == VC.PURCHASE and body["fromDate"] == "2026-09-01":
            raise API.ServerBusy("504")
        return {"ok": path, "body": body}
    d = VC.collect(call, open_page, D("2026-08-01"), D("2026-10-08"))
    assert d["wallet"] == [{"amount": 10}, {"amount": 11}] and d["wallet_balance"] == 21    # 2쪽·예금주 제거
    assert "예금주" not in str(d) and "123" not in str(d)
    assert d["purchase"]["202609"] is None and d["purchase"]["202608"]["body"]["toDate"] == "2026-08-31"
    assert d["purchase"]["202610"]["body"]["toDate"] == "2026-10-07"                         # 오늘 이후 날짜 안 넣음
    assert len(d["조회 실패"]) == 1 and "매출내역 202609" in d["조회 실패"][0]
    assert d["wing_vat"]["body"] == {"from": 202608, "to": 202610}
    assert d["rg_vat"]["ok"].endswith("?fromYearMonth=2026-08&toYearMonth=2026-10")
    assert d["pending"]["202608"]["body"]["year"] == 2026 and d["pending"]["202608"]["body"]["month"] == 8
    i = log.index(("화면", VC.SALES_URL))
    assert log[i + 1] == ("POST", VC.PURCHASE)                                             # 매출내역 화면 연 뒤 조회
    assert log.index(("화면", VC.PENDING_URL)) < log.index(("POST", VC.PENDING))
    ok("월 끝=어제·빈 달 없음·화면 먼저·월렛 2쪽·예금주/계좌 제거·504 달=조회 실패 기록 후 계속·부가세 요청 모양")


def p19_verify_sheet():
    print("[P19] 정산 검증 — RG계산식·RG입금(보상 제외)·월렛장부·윙판매/지급·부가세·보류/추가지급·판정 규칙")
    from coupang_analytics import settlement_verify as SV
    rg = [{"정산일": "2026-08-03", "기간 시작": "2026-06-29", "기간 끝": "2026-06-30", "지급비율": 70, "최종지급액": 412909,
           "totalPayableAmount": 359852, "totalCfsInventoryCompensationAmount": 53057},
          {"정산일": "2026-08-03", "기간 시작": "2026-06-08", "기간 끝": "2026-06-14", "지급비율": 30, "최종지급액": 158356,
           "totalPayableAmount": 152130, "totalSalesAdjustment": 6226},
          {"정산일": "2026-10-07", "기간 시작": "2026-08-31", "기간 끝": "2026-08-31", "지급비율": 70, "최종지급액": 69394,
           "totalPayableAmount": 112832, "totalAdditionalDeductionAmount": 112832, "totalCfsInventoryCompensationAmount": 69394},
          {"정산일": "2026-10-15", "기간 시작": "2026-08-31", "기간 끝": "2026-08-31", "지급비율": 30, "최종지급액": 7111,
           "totalPayableAmount": 48354, "totalAdditionalDeductionAmount": 16284, "totalFinalCfsFeeDeductionAmount": 24959}]
    f = SV.rg_formula("A", rg)
    assert f[0]["우리(계산)"] == 4 and f[0]["판정"] == SV.OK and len(f) == 1                  # 4줄 모두 H−I−J+K+조정
    bad = SV.rg_formula("A", [{**rg[0], "최종지급액": 1}])
    assert bad[1]["판정"] == SV.DIFF and bad[1]["차이"] == 412908
    v = {"수집일": "2026-10-08", "기간": ["2026-01-01", "2026-10-07"], "조회 실패": [], "wallet_balance": 300000,
         "wallet": [{"walletEventType": "DEPOSIT", "paymentDate": "20260803000729", "amount": 359852 + 158356},
                    {"walletEventType": "WITHDRAWAL", "paymentDate": "20260805", "amount": 518208 - 300000}],
         "purchase": {"202605": {"purchaseReports": [{"revenueAmount": 8602500, "sellerDiscountCoupon": 0, "feeAmount": 968234,
                                                      "apAmount": 7634266, "deductionAmount": 55000, "finalPaidAmount": 7579266}]},
                      "202609": {"purchaseReports": [{"revenueAmount": 2065700, "sellerDiscountCoupon": 1116300, "feeAmount": 94452,
                                                      "apAmount": 854948, "deductionAmount": 87556, "finalPaidAmount": 767392}]},
                      "202610": None},
         "wing_vat": {"paymentMethodReports": [{"yearMonth": "202605", "total": 8602500}]},
         "rg_vat": {"vatResponseAggregatedDtos": [{"yearMonth": "2026-07", "creditCardPaymentAmountAgg": 7106626,
                                                   "cashPaymentAmountAgg": 2231507, "otherPaymentAmountAgg": 10701127,
                                                   "sellerFundedCouponAggAmount": 10134700}]},
         "pending": {"202605": {"totalRecordCount": 0}}, "additional": {"202605": {"reports": [{"x": 1}]}}}
    w = {x["기준"]: x for x in SV.rg_wallet("A", rg, v)}
    assert w["지급일 2026-08-03"]["판정"] == SV.OK and "보상 53,057 제외" in w["지급일 2026-08-03"]["비고"]
    assert w["지급일 2026-10-07"]["판정"] == SV.OK and w["지급일 2026-10-07"]["쿠팡"] == 0      # 보상만 있는 날: 0 = 입금 없음
    assert w["지급일 2026-10-15"]["판정"] == SV.PENDING                                      # 수집일 이후 = 미정산
    assert w["~2026-10-08"]["판정"] == SV.OK                                                 # 입금 − 인출 = 잔액
    comp = [x for x in SV.rg_wallet("A", rg, v) if x["대조"] == "재고손실보상"]
    assert comp[0]["우리(계산)"] == 53057 + 69394 and comp[0]["판정"] == SV.CONFIRMED          # 공식: 대표계좌 별도 입금
    assert "대표 정산 계좌" in comp[0]["비고"]
    ours = {"202605": {"매출": 8602500, "쿠폰": 0, "수수료": 968234, "정산대상": 7634266}}
    ws = SV.wing_sales("A", ours, v)
    assert [x["판정"] for x in ws if x["기준"] == "202605"] == [SV.OK] * 4
    assert all(x["판정"] == SV.PENDING for x in ws if x["기준"] == "202609")                  # 최근 달·우리 0 = 미정산
    assert any(x["기준"] == "202610" and x["판정"] == SV.NODATA for x in ws)                   # 조회 실패 달
    wp = SV.wing_pay("A", [{"정산일": "2026-07-01", "기간 시작": "2026-05-01", "최종지급액": 7579266, "paidAmount": 7634266,
                            "totalDeductionAmount": 55000, "지급상태": "DONE"},
                           {"정산일": "2026-06-01", "기간 시작": "2026-05-04", "최종지급액": 0, "지급상태": "HOLD"}], v)
    assert [x["판정"] for x in wp if x["기준"] == "202605"] == [SV.OK] * 3
    assert any(x["항목"] == "지급상태" and x["판정"] == SV.CHECK and x["비고"] == "HOLD" for x in wp)
    vt = SV.vat("A", ours, {"202607": {"판매액": 20039260, "쿠폰": 10134700}}, v)
    assert [x["판정"] for x in vt] == [SV.OK, SV.OK, SV.OK], vt
    h = {x["대조"]: x["판정"] for x in SV.holds("A", v)}
    assert h == {"보류": SV.OK, "추가지급": SV.CHECK}
    assert SV.build("A", {}, {}, [], [], None)[1]["판정"] == SV.NODATA                       # 검증 자료 없음
    import openpyxl
    from coupang_analytics import settlement_stats as ST
    rows = SV.build("A", ours, {}, [], rg, v)
    with tempfile.TemporaryDirectory() as tmp:
        res = ST.aggregate([])
        wb = openpyxl.load_workbook(ST.write_stats(Path(tmp) / "v.xlsx", res, verify=rows))
        assert [c.value for c in wb["검증"][1]] == SV.HEAD and wb["검증"].max_row == 1 + len(rows)
        assert any("검증 '확인 필요'" in w for w in res.warnings)                        # 보상·추가지급 = 확인 필요
    ok("RG계산식 4줄·틀리면 다름·입금=최종−보상·수집일 이후=미정산·월렛장부·보상=확인 필요·윙판매/지급(최근 달 미정산)·"
       "지급상태 HOLD=확인 필요·부가세 윙/RG·보류0/추가지급1·자료 없음")


def p20_rg_payout_dates():
    print("[P20] 로켓그로스 지급일 — 70%·100%=주마감+20영업일 · 30% 월걸침 주=+25 · 30% 일반=익익월 첫 영업일(실측 629/629)")
    from coupang_analytics import payout as P
    hol = {D(x) for x in ("2026-01-01", "2026-02-16", "2026-02-17", "2026-02-18", "2026-03-02", "2026-05-01", "2026-05-05",
                          "2026-05-25", "2026-06-03", "2026-07-17", "2026-08-17", "2026-09-24", "2026-09-25",
                          "2026-10-05")}   # 테스트 주입(쿠팡 실지급일과 맞춘 영업일 기준)
    cases = [  # (지급비율, 기간 시작, 기간 끝, 실제 지급일) — nicoable·bux1004 등 운용 PC 정산현황 실값
        (70, "2026-07-27", "2026-07-31", "2026-08-31"), (30, "2026-07-27", "2026-07-31", "2026-09-07"),   # 소유자 화면 확인
        (30, "2026-08-01", "2026-08-02", "2026-09-07"), (30, "2026-08-03", "2026-08-09", "2026-10-01"),
        (70, "2026-08-31", "2026-08-31", "2026-10-07"),
        (100, "2026-01-26", "2026-01-31", "2026-03-05"), (100, "2026-05-25", "2026-05-31", "2026-06-29"),
        (30, "2026-06-01", "2026-06-07", "2026-08-03"), (70, "2026-04-13", "2026-04-19", "2026-05-19"),
        (70, "2026-06-15", "2026-06-21", "2026-07-20")]
    for ratio, ps, pe, want in cases:
        got = P.rg_payout_date(ratio, D(ps), D(pe), hol)
        assert got == D(want), (ratio, ps, pe, got, want)
    expect(ValueError, lambda: P.rg_payout_date(50, D("2026-08-03"), D("2026-08-09"), hol), "모르는 지급비율")
    old = P.payout_date(P.PAYOUT_RG_WEEKLY_FINAL, revenue_month="2026-07", holidays=hol)
    assert old != D("2026-09-07")                                                      # 옛 규칙(익익월)은 월걸침 주를 못 맞춤
    ok("70%·100%=+20(5/1·7/17 비영업일 반영)·30% 월걸침=+25(07-27~31·08-01~02→09-07)·30% 일반=익익월·모르는 비율=오류")


def p21_bank_inflows():
    print("[P21] 계좌 입금 — 윙 지급·월렛 입금/인출·재고 손실 보상(대표계좌)·물류비 환급(익월 21일)·지급월별 합계")
    from coupang_analytics import settlement_verify as SV
    wing = [{"정산일": "2026-07-01", "최종지급액": 2235280, "지급상태": "DONE"},
            {"정산일": "2026-10-20", "최종지급액": 999, "지급상태": "SCHEDULED"}]                      # 미래 = 제외
    rg = [{"정산일": "2026-08-03", "기간 시작": "2026-06-29", "최종지급액": 412909, "totalCfsInventoryCompensationAmount": 53057},
          {"정산일": "2026-07-27", "기간 시작": "2026-06-22", "최종지급액": 0, "totalCfsFeeAdjustment": 10125}]
    v = {"수집일": "2026-10-08", "wallet": [
        {"walletEventType": "DEPOSIT", "paymentDate": "20260803000729", "amount": 359852},
        {"walletEventType": "WITHDRAWAL", "paymentDate": "20260805121648", "amount": 359852}]}
    rows = {r["지급월"]: r for r in SV.inflows("A", wing, rg, v, today=D("2026-10-08"))}
    assert rows["2026-07"]["윙 지급"] == 2235280 and rows["2026-07"]["물류비 환급"] == 10125         # 6월분 환급 = 7월 21일
    assert "07-21" in rows["2026-07"]["비고"]
    a = rows["2026-08"]
    assert (a["RG 월렛 입금"], a["RG 월렛 인출"], a["재고 손실 보상"]) == (359852, 359852, 53057)
    assert a["쿠팡 지급 합계"] == 359852 + 53057 and a["대표계좌 입금 합계"] == 359852 + 53057
    assert "2026-10" not in rows                                                         # 지급 전 회차 제외
    none = SV.inflows("A", wing, rg, None, today=D("2026-10-08"))
    assert all(r["RG 월렛 입금"] is None for r in none) and "월렛 자료 없음" in none[0]["비고"]
    import openpyxl
    from coupang_analytics import settlement_stats as ST
    with tempfile.TemporaryDirectory() as tmp:
        wb = openpyxl.load_workbook(ST.write_stats(Path(tmp) / "i.xlsx", ST.aggregate([]),
                                                   inflows=SV.inflows("A", wing, rg, v, today=D("2026-10-08"))))
        assert [c.value for c in wb["계좌 입금"][1]] == SV.INFLOW_HEAD and wb["계좌 입금"].max_row == 3
    ok("윙=지급일 지난 회차·월렛 입금/인출·보상=정산일 대표계좌·환급=매출인식 익월 21일·합계 2종·지급 전 제외·월렛 없음 표시")


def p22_paydates_verify():
    print("[P22] 지급일 대조 — RG 전부·윙 70%·윙 30%/100%는 대상 아님·공휴일 없으면 자료 없음")
    from coupang_analytics import settlement_verify as SV
    hol = {D(x) for x in ("2026-08-17", "2026-09-24", "2026-09-25", "2026-10-05")}
    rg = [{"정산일": "2026-08-31", "기간 시작": "2026-07-27", "기간 끝": "2026-07-31", "지급비율": 70, "최종지급액": 1},
          {"정산일": "2026-09-07", "기간 시작": "2026-07-27", "기간 끝": "2026-07-31", "지급비율": 30, "최종지급액": 1},
          {"정산일": "2026-09-08", "기간 시작": "2026-08-01", "기간 끝": "2026-08-02", "지급비율": 30, "최종지급액": 1}]   # 하루 틀림
    wing = [{"정산일": "2026-09-18", "기간 시작": "2026-08-24", "기간 끝": "2026-08-30", "지급비율": 70, "최종지급액": 1},
            {"정산일": "2026-10-01", "기간 시작": "2026-08-01", "기간 끝": "2026-08-31", "지급비율": 30, "최종지급액": 1},
            {"정산일": "2026-07-20", "기간 시작": "2026-05-17", "기간 끝": "2026-05-17", "지급비율": 100, "최종지급액": 1}]
    rows = SV.paydates("A", wing, rg, hol)
    summ = {r["항목"]: r for r in rows if r["기준"] == "요약"}
    assert summ["로켓그로스"]["우리(계산)"] == 2 and summ["로켓그로스"]["쿠팡"] == 3 and summ["로켓그로스"]["판정"] == SV.DIFF
    assert summ["윙 70%"]["판정"] == SV.OK and summ["윙 70%"]["쿠팡"] == 1
    bad = [r for r in rows if r["판정"] == SV.DIFF and r["기준"] != "요약"]
    assert len(bad) == 1 and "2026-08-01" in bad[0]["항목"] and bad[0]["비고"] == "계산 2026-09-07"
    assert not any("30%" in r["항목"] and r["대조"] == "지급일" and "윙" in r["항목"] for r in rows)     # 윙 30%·100% 대상 아님
    none = SV.paydates("A", wing, rg, None)
    assert len(none) == 1 and none[0]["판정"] == SV.NODATA and "공휴일" in none[0]["비고"]
    import os
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    import settlement_download as TD
    from datetime import datetime
    from coupang_analytics import settlement_runlog as RLG
    old_cwd, old_log = os.getcwd(), TD.LOG
    with tempfile.TemporaryDirectory() as tmp:
        os.chdir(tmp)
        try:
            TD.LOG = RLG.RunLog(tmp, now=datetime(2026, 10, 8), echo=lambda m: None)
            (Path(tmp) / "output").mkdir()
            (Path(tmp) / "output" / "_holidays_2026.json").write_text('["2026-08-17", "2026-10-05"]', encoding="utf-8")
            got = TD.load_holidays([{"정산일": "2026-08-31"}])
            assert got == {D("2026-08-17"), D("2026-10-05")}                                 # 캐시 우선(키 없어도)
            assert TD.load_holidays([{"정산일": "2025-12-01"}]) is None                       # 캐시·키 없는 해 = None
            assert any("지급일 대조 건너뜀" in r for r in TD.LOG.text.read_text(encoding="utf-8").splitlines())
        finally:
            os.chdir(old_cwd)
            TD.LOG = old_log
    ok("RG 3줄 중 2 일치·틀린 줄 하나만 상세·윙 70% 일치·윙 30%/100% 제외·공휴일 없음=자료 없음 1줄")


def p23_calendar():
    print("[P23] 정산캘린더 — 윙 지급일은 캘린더로 대조(30% 포함·공휴일 불필요)·수집은 폼 본문·캘린더 없으면 규칙")
    from coupang_analytics import settlement_verify as SV
    from coupang_analytics import settlement_verify_collect as VC
    cal = [{"title": "9월 최종액 정산", "start": "2026-11-02 00:00:00", "transactionCycleCode": "R",
            "recognitionFrom": "2026-09-01", "recognitionTo": "2026-09-30"},
           {"title": "[주정산]\n08/31 ~ 09/06", "start": "2026-09-29 00:00:00", "transactionCycleCode": "W",
            "recognitionFrom": "2026-08-31", "recognitionTo": "2026-09-06"},
           {"title": "신정", "start": "2025-01-01 00:00:00", "transactionCycleCode": "UK", "recognitionFrom": None,
            "recognitionTo": None}]
    wing = [{"정산일": "2026-09-29", "기간 시작": "2026-08-31", "기간 끝": "2026-08-31", "지급비율": 70, "최종지급액": 1},
            {"정산일": "2026-09-29", "기간 시작": "2026-09-01", "기간 끝": "2026-09-06", "지급비율": 70, "최종지급액": 1},
            {"정산일": "2026-11-01", "기간 시작": "2026-09-01", "기간 끝": "2026-09-30", "지급비율": 30, "최종지급액": 1},   # 하루 틀림
            {"정산일": "2026-07-20", "기간 시작": "2026-05-17", "기간 끝": "2026-05-17", "지급비율": 100, "최종지급액": 1}]
    rows = SV.paydates("A", wing, [], None, calendar=cal)
    summ = {r["항목"]: r for r in rows if r["기준"] == "요약"}
    assert summ["윙(캘린더)"]["우리(계산)"] == 2 and summ["윙(캘린더)"]["쿠팡"] == 3 and summ["윙(캘린더)"]["판정"] == SV.DIFF
    bad = [r for r in rows if r["기준"] != "요약" and r["판정"] == SV.DIFF]
    assert len(bad) == 1 and bad[0]["비고"] == "캘린더 2026-11-02"                          # 휴일 보정된 최종액 날짜
    assert not any("RG" in r["항목"] or r["항목"] == "공휴일" for r in rows)                  # RG 줄 없으면 공휴일 불요
    miss = SV.paydates("A", [{**wing[0], "기간 시작": "2026-01-05", "기간 끝": "2026-01-11"}], [], None, calendar=cal)
    assert any(r["판정"] == SV.NODATA and "캘린더에 없음" in r["비고"] for r in miss)
    log: list = []

    def call(method, path, body):
        log.append((method, path, body))
        return cal if path == VC.CALENDAR else {}
    d = VC.collect(call, lambda url: log.append(("화면", url)), D("2026-01-01"), D("2026-10-08"))
    assert d["calendar"] == cal
    i = log.index(("화면", VC.CALENDAR_URL))
    assert log[i + 1] == ("POST", VC.CALENDAR, "startDate=2025-12-25&endDate=2026-12-17")    # 폼 본문·시작 7일 전~오늘+70일
    ok("윙 캘린더 대조(분할 주·최종액·틀린 날 1)·100% 제외·RG 없으면 공휴일 불요·캘린더에 없는 회차=자료 없음·수집 폼 본문")



def p24_local_store():
    print("[P24] 정산 로컬 기록 — 다른 PC 폴더 합치기·예전 계정명 맞춤·금액 기록 키 병합(이 PC 값 우선)")
    import json
    import os
    from datetime import datetime
    from types import SimpleNamespace
    sys.path.insert(0, str(ROOT / "tools"))
    import settlement_download as TD
    from coupang_analytics import settlement_accounts as SA
    from coupang_analytics import settlement_files as SF
    from coupang_analytics import settlement_jobs as SJ
    from coupang_analytics import settlement_runlog as RLG
    accts, _ = SA.parse_accounts_text("acc1\tpw1\t홍길동-주식회사 가나\n")
    new, old = "홍길동-주식회사 가나-acc1", "(주)가나-acc1"
    tnew, told = SF._token(new, "계정명"), SF._token(old, "계정명")

    def fname(acct, settle, kind="주정산", report="주문상세", ch="윙", ps="2026-01-05", pe="2026-01-11"):
        return SF.settle_file_name(acct, D(settle), ch, kind, report, D(ps), D(pe))

    def amt(settle, paid):
        return {"정산일": settle, "기간 시작": "2026-01-05", "기간 끝": "2026-01-11", "지급비율": 70, "최종지급액": paid}

    def wjson(path, rows):
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(rows, ensure_ascii=False), encoding="utf-8")

    def rjson(path):
        return json.loads(path.read_text(encoding="utf-8"))

    old_cwd, old_log, old_load = os.getcwd(), TD.LOG, TD.load_accounts
    with tempfile.TemporaryDirectory() as tmp:
        os.chdir(tmp)
        try:
            TD.LOG = RLG.RunLog(Path(tmp) / "기록", now=datetime(2026, 10, 8), echo=lambda m: None)
            TD.load_accounts = lambda only: accts
            # (a) _migrate_names: 이 PC 에 남은 예전 이름 파일·요청 기록·금액 기록 → 지금 이름(금액은 키 병합·지금 이름 값 우선)
            f_old, f_new = fname(old, "2026-01-30"), fname(new, "2026-01-30")
            TD.FILES.mkdir(parents=True)
            (TD.FILES / f_old).write_bytes(b"old")
            SJ.save_jobs(TD.JOBS, [SJ.Job(old, "윙", "주정산", "주문상세", "2026-01-05", "2026-01-11", "2026-01-30",
                                          status=SJ.ST_DONE, file=f_old)])
            wjson(TD.AMOUNTS / f"{told}_윙.json", [amt("2026-01-30", 999), amt("2026-02-06", 7)])
            wjson(TD.AMOUNTS / f"{tnew}_윙.json", [amt("2026-01-30", 100)])
            TD._migrate_names(accts)
            assert sorted(p.name for p in TD.FILES.glob("*.xlsx")) == [f_new]
            j = SJ.load_jobs(TD.JOBS)
            assert [(x.account, x.file) for x in j] == [(new, f_new)]
            assert sorted(p.name for p in TD.AMOUNTS.glob("*.json")) == [f"{tnew}_윙.json"]
            assert [(r["정산일"], r["최종지급액"]) for r in rjson(TD.AMOUNTS / f"{tnew}_윙.json")] == [
                ("2026-01-30", 100), ("2026-02-06", 7)]                                   # 키 병합·지금 이름 값 우선·정렬
            TD._migrate_names(accts)                                                       # 멱등
            assert sorted(p.name for p in TD.FILES.glob("*.xlsx")) == [f_new] and len(SJ.load_jobs(TD.JOBS)) == 1
            # (b) cmd_merge: 옮겨 온 폴더(예전 이름) 합치기 — 같은 작업 1회·같은 이름 파일은 원본에 남김·비용 폴더·금액 병합
            src = Path(tmp) / "옮겨온"
            f2_old, c_old = fname(old, "2026-02-06", ps="2026-01-12", pe="2026-01-18"), fname(old, "2026-01-30", "주정산", "보관비", "로켓그로스")
            SJ.save_jobs(src / "_요청기록.json", [
                SJ.Job(old, "윙", "주정산", "주문상세", "2026-01-05", "2026-01-11", "2026-01-30", file=f_old),
                SJ.Job(old, "윙", "주정산", "주문상세", "2026-01-12", "2026-01-18", "2026-02-06", status=SJ.ST_DONE,
                       file=f2_old)])
            (src / "파일" / "비용").mkdir(parents=True)
            (src / "파일" / f_old).write_bytes(b"other")
            (src / "파일" / f2_old).write_bytes(b"f2")
            (src / "파일" / "비용" / c_old).write_bytes(b"c")
            wjson(src / "쿠팡지급내역" / f"{told}_윙.json", [amt("2026-01-30", 555), amt("2026-02-13", 3)])
            assert TD.cmd_merge(SimpleNamespace(src=str(src))) == 0
            j = SJ.load_jobs(TD.JOBS)
            assert [(x.account, x.settle_date, x.status, x.file) for x in j] == [
                (new, "2026-01-30", SJ.ST_DONE, f_new), (new, "2026-02-06", SJ.ST_DONE, fname(new, "2026-02-06", ps="2026-01-12", pe="2026-01-18"))], j
            assert (TD.FILES / f_new).read_bytes() == b"old" and (src / "파일" / f_old).exists()   # 같은 이름 = 그대로 둠
            assert (TD.FILES / fname(new, "2026-02-06", ps="2026-01-12", pe="2026-01-18")).read_bytes() == b"f2" and not (src / "파일" / f2_old).exists()
            assert [p.name for p in TD.COSTS.glob("*.xlsx")] == [fname(new, "2026-01-30", "주정산", "보관비", "로켓그로스")]
            # ⚠현재 동작 고정(2026-10-08): 옮겨 온 금액 기록은 예전 이름 그대로 남음 — 파일·요청 기록은 합치는 중에 이미
            # 지금 이름이 돼 _migrate_names 의 바꿈표가 비기 때문(결함 후보·별도 수정 사이클에서 이 단언을 바꿀 것).
            assert sorted(p.name for p in TD.AMOUNTS.glob("*.json")) == [f"{told}_윙.json", f"{tnew}_윙.json"]
            assert [(r["정산일"], r["최종지급액"]) for r in rjson(TD.AMOUNTS / f"{told}_윙.json")] == [
                ("2026-01-30", 555), ("2026-02-13", 3)]
            assert [(r["정산일"], r["최종지급액"]) for r in rjson(TD.AMOUNTS / f"{tnew}_윙.json")] == [
                ("2026-01-30", 100), ("2026-02-06", 7)]
            rec = [r for r in TD.LOG.rows if r["단계"] == "합치기"]
            assert len(rec) == 1 and rec[0]["결과"] == RLG.OK and "1건 추가·1건 이미 있음" in rec[0]["사유"] \
                and "파일 2개 옮김·1개 같은 이름" in rec[0]["사유"], rec
            # (c) _save_amounts: 정산현황 응답 → 계정·채널 JSON 에 키로 덮어 모음
            resp = {"paymentReports": [{"payDate": "2026-02-13", "recognitionFrom": "2026-01-05",
                                        "recognitionTo": "2026-01-11", "ratio": 70, "finalPaidAmount": 4, "detail": {}}]}
            TD._save_amounts(new, "윙", resp)
            got = rjson(TD.AMOUNTS / f"{tnew}_윙.json")
            assert [(r["정산일"], r["최종지급액"]) for r in got] == [("2026-01-30", 100), ("2026-02-06", 7), ("2026-02-13", 4)]
            assert not list(TD.AMOUNTS.glob("*.tmp"))
        finally:
            os.chdir(old_cwd)
            TD.LOG, TD.load_accounts = old_log, old_load
    ok("예전 이름 파일·요청·금액 → 지금 이름(멱등)·합치기(같은 작업 1회·같은 이름 파일 원본 보존·비용 폴더·금액 키 병합)"
       "·정산현황 금액 덮어 모음")

def main():
    p10_wing_api_parse()
    p11_runlog()
    p12_watch_decide()
    p13_accounts_file()
    p14_status_reader()
    p15_browser_closed_transient()
    p16_cost_by_product()
    p17_server_busy()
    p18_verify_collect()
    p19_verify_sheet()
    p20_rg_payout_dates()
    p21_bank_inflows()
    p22_paydates_verify()
    p23_calendar()
    p24_local_store()
    print("정산 런타임 오프라인 검증 통과")


if __name__ == "__main__":
    main()
