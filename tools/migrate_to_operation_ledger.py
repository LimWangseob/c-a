"""2단계 데이터 전환(초안) — 관리대장(109열)+재고현황 → 운영대장(designs/OPERATION_DATA_MODEL.md).

best-effort 이관 + 상품코드 발번(분류·입고YYMM 날짜우선순위·순번) + **자동연동(수식)·셀잠금**.
- 파생열(사업자명·상품(물류)명·현재고·그로스재고)=INDEX/MATCH 수식(원본 바뀌면 자동 갱신)·잠금(원본만 수정).
- 관리코드(상품코드·계정아이디 등)=숨김 키. 숫자=원단위 정수 콤마 우측. 날짜=시간 제외. 계정별 8색 밴드.
- 결과는 검토용 DRAFT. 재고 자동매칭 실패분은 '매핑검토' 시트.

    python tools/migrate_to_operation_ledger.py [관리대장.xlsx] [재고현황.xlsx] [출력.xlsx]
"""
from __future__ import annotations

import re
import sys
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "tools"))
import build_operation_template as T  # noqa: E402

_DL = Path.home() / "Downloads"
DEF_LEDGER = _DL / "토탈셀러_셀독 관리 대장 (7).xlsx"
DEF_STOCK = _DL / "재고 현황 (1).xlsx"
DEF_OUT = _DL / "운영대장_전환초안.xlsx"

C_MEMO = list(range(3, 21))
C_REP, C_BIZ, C_AID, C_PW, C_DEPOSIT, C_PROMO = 21, 22, 23, 24, 25, 26
C_PNAME, C_COSTOCK, C_GSTOCK, C_URL, C_HIST, C_UNIT, C_PRICE, C_NOTE = 27, 28, 29, 30, 31, 32, 33, 34
C_GREQDATE, C_GREQQTY, C_GWORKQTY, C_GBOX, C_GPAL, C_GDONE, C_GSHIP = 35, 36, 37, 38, 39, 40, 41
_CLASS = {"건기식": "G", "공산품": "P"}
_NUM = {"계약금", "계약단가", "판매가", "계약수량", "계약금액", "현재고", "그로스재고", "요청수량",
        "작업수량", "박스", "파레트", "광고비", "노출", "클릭", "전환", "매출", "건수", "견적서판매가"}
_BAND = ("F8CBAD", "FFE699", "C6E0B4", "BDD7EE", "D9C2E9", "F4B6C2", "B7DEE8", "D9D9D9")
_RIGHT = Alignment(horizontal="right")


def _s(v) -> str:
    return "" if v is None else str(v).strip()


def _norm(name: str) -> str:
    return re.sub(r"\s+", "", name or "").lower()


def _ym(v) -> str:
    if isinstance(v, datetime):
        return f"{v.year % 100:02d}{v.month:02d}"
    m = re.search(r"(?:20)?(\d{2})[.\-/](\d{1,2})", _s(v))
    return f"{int(m.group(1)):02d}{int(m.group(2)):02d}" if m else ""


def _date(v) -> str:
    """날짜값 → 'YYYY-MM-DD'(시간 제외). 실패 시 원문."""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    s = _s(v)
    m = re.search(r"(?:20)?(\d{2})[.\-/](\d{1,2})[.\-/](\d{1,2})", s)
    return f"20{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}" if m else s


def _logdate(label: str) -> tuple[str, str]:
    s = label.replace("관리내용", "").strip()
    m = re.fullmatch(r"(\d{1,2})-(\d{1,2})", s)
    if m:
        return f"2026-{int(m.group(1)):02d}-{int(m.group(2)):02d}", ""
    return "", f"[{s}] " if s else ""


def _split_promo(txt: str) -> tuple[str, str]:
    """'판매가 : 계약자 / 리뷰비 : 계약자' → (판매가주체, 리뷰비주체)."""
    sale = rev = ""
    for part in re.split(r"[/\n]", txt or ""):
        if "판매가" in part:
            sale = part.split(":", 1)[-1].strip() if ":" in part else part.strip()
        elif "리뷰" in part:
            rev = part.split(":", 1)[-1].strip() if ":" in part else part.strip()
    if not sale and not rev and txt:
        sale = txt.strip()
    return sale, rev


def read_stock(path: Path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Sheet1"]
    ym = _ym(ws.cell(1, 1).value) or ""
    idx: dict[str, dict] = {}
    for r in range(3, (ws.max_row or 0) + 1):
        nm = _s(ws.cell(r, 4).value)
        if nm:
            idx.setdefault(_norm(nm), {"구분": _s(ws.cell(r, 2).value), "바코드": _s(ws.cell(r, 3).value),
                                       "물류명": nm})
    return idx, ym


def read_stock_rows(path: Path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Sheet1"]
    m = re.search(r"(?:20)?(\d{2})[.\-/](\d{1,2})[.\-/](\d{1,2})", _s(ws.cell(1, 1).value))
    upd = f"20{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}" if m else ""
    rows = []
    for r in range(3, (ws.max_row or 0) + 1):
        nm = _s(ws.cell(r, 4).value)
        if nm:
            rows.append({"물류명": nm, "창고": _s(ws.cell(r, 1).value), "구분": _s(ws.cell(r, 2).value),
                         "바코드": _s(ws.cell(r, 3).value), "현재고": _s(ws.cell(r, 5).value),
                         "셀독": _s(ws.cell(r, 6).value), "당근": _s(ws.cell(r, 7).value),
                         "자사": _s(ws.cell(r, 8).value), "갱신일": upd})
    return rows


def read_ledger(path: Path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["셀독리스트"]
    memo_hdr = [_s(ws.cell(2, c).value) for c in C_MEMO]
    accts: list[dict] = []
    cur: dict | None = None
    for r in range(3, (ws.max_row or 0) + 1):
        aid = _s(ws.cell(r, C_AID).value)
        if aid:
            sale, rev = _split_promo(_s(ws.cell(r, C_PROMO).value))
            cur = {"계정아이디": aid, "대표자명": _s(ws.cell(r, C_REP).value),
                   "사업자명": _s(ws.cell(r, C_BIZ).value), "비밀번호": _s(ws.cell(r, C_PW).value),
                   "계약금": _s(ws.cell(r, C_DEPOSIT).value), "판매가주체": sale, "리뷰비주체": rev,
                   "담당자": _s(ws.cell(r, 2).value), "products": [], "memos": []}
            accts.append(cur)
            for h, c in zip(memo_hdr, C_MEMO):
                if _s(ws.cell(r, c).value):
                    cur["memos"].append((h, _s(ws.cell(r, c).value)))
        if cur is None:
            continue
        pname = _s(ws.cell(r, C_PNAME).value)
        if pname:
            lines = [ln.strip() for ln in pname.splitlines() if ln.strip()]
            cur["products"].append({
                "상품명": lines[0] if lines else pname,
                "노출상품명": " ".join(lines[1:]).strip("() ") if len(lines) > 1 else "",
                "그로스재고": _s(ws.cell(r, C_GSTOCK).value), "상품url": _s(ws.cell(r, C_URL).value),
                "계약단가": _s(ws.cell(r, C_UNIT).value), "판매가": _s(ws.cell(r, C_PRICE).value),
                "비고": _s(ws.cell(r, C_NOTE).value), "히스토리": _s(ws.cell(r, C_HIST).value),
                "그로스": {"요청일자": _s(ws.cell(r, C_GREQDATE).value), "요청수량": _s(ws.cell(r, C_GREQQTY).value),
                          "작업수량": _s(ws.cell(r, C_GWORKQTY).value), "박스": _s(ws.cell(r, C_GBOX).value),
                          "파레트": _s(ws.cell(r, C_GPAL).value), "완료일자": _s(ws.cell(r, C_GDONE).value),
                          "출고일자": _s(ws.cell(r, C_GSHIP).value)},
            })
    return accts


def read_promo(path: Path):
    wb = openpyxl.load_workbook(path, data_only=True)
    if "체험단" not in wb.sheetnames:
        return []
    ws = wb["체험단"]
    hdr = {_s(ws.cell(3, c).value): c for c in range(1, (ws.max_column or 0) + 1) if _s(ws.cell(3, c).value)}
    out = []
    for r in range(4, (ws.max_row or 0) + 1):
        def g(name, _r=r):
            return _s(ws.cell(_r, hdr[name]).value) if name in hdr else ""
        if not (g("상호") or g("제품명")):
            continue
        out.append({"사업자명": g("상호"), "상품명": g("제품명"), "요청일": g("요청일"), "키워드": g("키워드"),
                    "판매가": g("판매가"), "리뷰/택배": g("리뷰/택배"), "건수": g("건수"),
                    "신규/추가": g("체험단 신규/추가"), "견적서판매가": g("견적서 판매가"),
                    "견적서리뷰/택배": g("견적서 리뷰/택배"), "밑작업": g("밑작업"),
                    "포토/텍스트": g("포토/텍스트"), "진행여부": g("진행 여부"), "비고": g("비고")})
    return out


CONTRACTS = [
    {"사업자명": "봉이네농원", "대표자명": "곽재봉", "상품명": "곶감 상품", "계약서종류": "온라인판매 플랫폼 위탁운영 계약서",
     "계약일자": "2026-09-08", "계약기간시작": "2026-09-08", "계약기간종료": "2027-09-07",
     "수탁자": "㈜디프픽 (이원규)", "수익배분": "갑70/을30", "수익기준": "순이익", "광고비부담": "갑",
     "재고배송부담": "갑(재고·배송·반품/환불)", "정산시점": "플랫폼 정산 후 7일 이내", "귀속기준일": "플랫폼 정산일",
     "특약": "최근 2개월 평균 대비 매출 10%↑ 시 갑 일방해지 불가 · 비밀유지 2년 · 계약종료 시 데이터 이전",
     "계약서위치": "별도PC / Modusign d69fc960-ab52-11f1"},
    {"사업자명": "효성", "대표자명": "박원준", "상품명": "파미젠 와사비잎 추출물 plus max 120정", "계약서종류": "상품공급 OEM 계약서",
     "계약일자": "2026-09-18", "계약기간시작": "2026-09-18", "계약기간종료": "2027-03-17",
     "수탁자": "㈜동방이노션 (차윤태)", "수익배분": "갑70/을30", "수익기준": "판매수익", "계약금": "10000000",
     "계약단가": "5000", "계약수량": "2000", "계약금액": "10000000", "광고비부담": "갑(쿠팡·외부광고 일체, KC인증비=을)",
     "재고배송부담": "을 창고 납품 · 발주 선결제 후 출고", "정산시점": "판매수익 분배 즉시(월 판매 기준)", "귀속기준일": "월 판매 기준",
     "특약": "판매 2개월 후 목표수익률 판단 · 미달 시 대체매칭/추가 무상공급, 최종 갑 선택",
     "계약서위치": "별도PC / Modusign ca7ca480-b275-11f1"},
]


def assign_code(classc: str, ym: str, seq: dict) -> str:
    seq[(classc, ym)] = seq.get((classc, ym), 0) + 1
    return f"{classc}{ym}-{seq[(classc, ym)]:03d}"


def _hmap(ws) -> dict:
    return {_s(ws.cell(1, c).value): c for c in range(1, (ws.max_column or 0) + 1) if _s(ws.cell(1, c).value)}


def _lk(master: str, ret_c: int, key_c: int, keycell: str) -> str:
    rc, kc = get_column_letter(ret_c), get_column_letter(key_c)
    return f"=IFERROR(INDEX('{master}'!{rc}:{rc},MATCH({keycell},'{master}'!{kc}:{kc},0)),\"\")"


def _opts(sheet: str) -> dict:
    """시트 컬럼명 → (group, opts-set). build_operation_template 스키마 재사용."""
    return {name: (g, set(o.split(",")) if o else set()) for name, g, o, _d in T.SHEETS[sheet]}


def _to_int(v):
    try:
        return int(round(float(str(v).replace(",", "").strip())))
    except (ValueError, TypeError):
        return None


def _style_protect(wb) -> None:
    """계정 밴드 + 숫자 원단위 콤마(우측) + 컬럼 폭 자동 + 파생/키 잠금(시트 보호)."""
    for sn in wb.sheetnames:
        if sn in ("대시보드", "정의", "안내", "매핑검토"):
            continue
        ws = wb[sn]
        cols = T.SHEETS.get(sn)
        if not cols:
            continue
        mr = ws.max_row or 0
        hdr = {name: c for c, (name, _g, _o, _d) in enumerate(cols, 1)}
        keycol = hdr.get("계정아이디") or hdr.get("물류명") or hdr.get("사업자명")
        group, prev = -1, None
        for r in range(2, mr + 1):
            if keycol:
                k = _s(ws.cell(r, keycol).value)
                if k and k != prev:
                    group, prev = group + 1, k
            band = _BAND[group % len(_BAND)] if group >= 0 else None
            for c, (name, g, o, _d) in enumerate(cols, 1):
                cell = ws.cell(r, c)
                if band:
                    cell.fill = PatternFill("solid", fgColor=band)
                editable = g in ("id", "content") and "derived" not in o
                cell.protection = Protection(locked=not editable)
                if name in _NUM:
                    cell.alignment = _RIGHT
                    cell.number_format = "#,##0"
                    if not str(cell.value or "").startswith("="):
                        iv = _to_int(cell.value)
                        cell.value = iv if iv is not None else cell.value
        for c in range(1, (ws.max_column or 0) + 1):
            maxlen = max((len(_s(ws.cell(r, c).value)) for r in range(1, mr + 1)), default=4)
            ws.column_dimensions[get_column_letter(c)].width = min(50, max(10, int(maxlen * 1.6) + 2))
        ws.protection.sheet = True          # 파생/키 잠금 발효(원본만 수정)


def _review(wb):
    rv = wb.create_sheet("매핑검토")
    rv.sheet_properties.tabColor = "FF0000"
    for c, h in enumerate(["상품코드", "사업자명", "상품(물류)명", "사유"], 1):
        rv.cell(1, c, h).font = Font(bold=True)
    rv.freeze_panes = "A2"
    return rv


def migrate(ledger: Path, stock: Path, out: Path) -> dict:
    stock_idx, stock_ym = read_stock(stock)
    stock_rows = read_stock_rows(stock)
    accts = read_ledger(ledger)
    promos = read_promo(ledger)
    T.build(out)
    wb = openpyxl.load_workbook(out)
    H = {sn: _hmap(wb[sn]) for sn in ("관리상품", "재고", "그로스입고", "체험단", "광고", "업무일지", "계정", "계약")}
    # 마스터 참조 열(수식용)
    pm, ac, stk = H["관리상품"], H["계정"], H["재고"]
    rv = _review(wb)
    seq: dict = {}
    code_by_biz_name: dict[tuple, str] = {}
    code_by_stockname: dict[str, str] = {}
    stats = dict.fromkeys(("계정", "관리상품", "재고", "그로스입고", "체험단", "광고", "업무일지", "계약", "매핑검토"), 0)
    rows = {sn: 1 for sn in H}                 # 각 시트 마지막 기록 행

    def put(sn, values, formulas=None):
        rows[sn] += 1
        r = rows[sn]
        ws = wb[sn]
        for name, col in H[sn].items():
            if name in values:
                ws.cell(r, col, values[name])
        for name, fn in (formulas or {}).items():
            ws.cell(r, col := H[sn][name], fn(r))
        return r

    for a in accts:
        put("계정", {"대표자명": a["대표자명"], "사업자명": a["사업자명"], "계정아이디": a["계정아이디"],
                    "비밀번호": a["비밀번호"], "위탁상태": "관리중", "판매가주체": a["판매가주체"],
                    "리뷰비주체": a["리뷰비주체"]})
        stats["계정"] += 1
        # 계약(계정 1행·계약금) — 상세 계약조건은 계약서 수령분만(아래 enrich)
        put("계약", {"사업자명": a["사업자명"], "대표자명": a["대표자명"], "계약금": a["계약금"],
                    "계정아이디": a["계정아이디"]})
        stats["계약"] += 1
        for h, txt in a["memos"]:
            ld, pre = _logdate(h)
            r = put("업무일지", {"일자": ld, "작성자": a["담당자"], "내용": pre + txt, "상태": "완료",
                               "계정아이디": a["계정아이디"]})
            wb["업무일지"].cell(r, H["업무일지"]["사업자명"],
                              _lk("계정", ac["사업자명"], ac["계정아이디"], f"$%s{r}" % get_column_letter(H["업무일지"]["계정아이디"])))
            stats["업무일지"] += 1
        for p in a["products"]:
            st = stock_idx.get(_norm(p["상품명"]))
            classc = _CLASS.get(st["구분"], "E") if st else "E"
            g = p["그로스"]
            ym = _ym(g["출고일자"]) or _ym(g["완료일자"]) or stock_ym or "2610"
            code = assign_code(classc, ym, seq)
            code_by_biz_name[(a["사업자명"], p["상품명"])] = code
            pn = st["물류명"] if st else p["상품명"]       # 상품(물류)명 일원화(재고 기준·없으면 내부명)
            if st:
                code_by_stockname[_norm(st["물류명"])] = code
            r = put("관리상품", {"상품(물류)명": pn, "노출상품명": p["노출상품명"],
                               "카테고리": st["구분"] if st else "", "판매방식": "그로스판매" if _to_int(p["그로스재고"]) else "",
                               "계약단가": p["계약단가"], "판매가": p["판매가"], "그로스재고": p["그로스재고"],
                               "관리상태": "판매중", "비고": p["비고"], "계정아이디": a["계정아이디"],
                               "상품코드": code, "바코드": st["바코드"] if st else ""})
            ws = wb["관리상품"]
            if p["상품url"]:                               # 상품(물류)명에 쿠팡 링크
                ws.cell(r, pm["상품(물류)명"]).hyperlink = p["상품url"]
                ws.cell(r, pm["상품(물류)명"]).font = Font(color="0563C1", underline="single")
            ws.cell(r, pm["사업자명"], _lk("계정", ac["사업자명"], ac["계정아이디"], f"$%s{r}" % get_column_letter(pm["계정아이디"])))
            ws.cell(r, pm["현재고"], _lk("재고", stk["현재고"], stk["상품코드"], f"$%s{r}" % get_column_letter(pm["상품코드"])))
            stats["관리상품"] += 1
            if any(g.values()):
                r = put("그로스입고", {"요청일자": _date(g["요청일자"]), "요청수량": g["요청수량"],
                                    "작업수량": g["작업수량"], "박스": g["박스"], "파레트": g["파레트"],
                                    "완료일자": _date(g["완료일자"]), "출고일자": _date(g["출고일자"]), "상품코드": code})
                kc = "$%s{r}".replace("{r}", str(r)) % get_column_letter(H["그로스입고"]["상품코드"])
                gi = wb["그로스입고"]
                gi.cell(r, H["그로스입고"]["사업자명"], _lk("관리상품", pm["사업자명"], pm["상품코드"], kc))
                gi.cell(r, H["그로스입고"]["상품(물류)명"], _lk("관리상품", pm["상품(물류)명"], pm["상품코드"], kc))
                gi.cell(r, H["그로스입고"]["현재고"], _lk("재고", stk["현재고"], stk["상품코드"], kc))
                gi.cell(r, H["그로스입고"]["그로스재고"], _lk("관리상품", pm["그로스재고"], pm["상품코드"], kc))
                stats["그로스입고"] += 1
            if p["히스토리"]:
                r = put("업무일지", {"일자": "", "작성자": a["담당자"], "내용": p["히스토리"], "상태": "완료",
                                   "계정아이디": a["계정아이디"], "상품코드": code})
                uj = wb["업무일지"]
                uj.cell(r, H["업무일지"]["사업자명"], _lk("계정", ac["사업자명"], ac["계정아이디"], f"$%s{r}" % get_column_letter(H["업무일지"]["계정아이디"])))
                uj.cell(r, H["업무일지"]["상품(물류)명"], _lk("관리상품", pm["상품(물류)명"], pm["상품코드"], f"$%s{r}" % get_column_letter(H["업무일지"]["상품코드"])))
                stats["업무일지"] += 1
            if not st:
                rr = rv.max_row + 1
                for cc, v in enumerate((code, a["사업자명"], p["상품명"], "재고 매칭 실패(분류 E·입고월 기본) — 수동 매핑 필요"), 1):
                    rv.cell(rr, cc, v)
                stats["매핑검토"] += 1

    wsk = wb["계약"]                            # 계약서 분석 내용 반영(사업자명 정확일치 보강·없으면 신규행)
    biz_row: dict[str, int] = {}
    for r in range(2, rows["계약"] + 1):
        b = _norm(_s(wsk.cell(r, H["계약"]["사업자명"]).value))
        if b:
            biz_row.setdefault(b, r)
    for ct in CONTRACTS:
        r = biz_row.get(_norm(ct["사업자명"]))
        if r is None:
            r = put("계약", {"사업자명": ct["사업자명"]})
        for name, col in H["계약"].items():       # 계약서는 계약조건의 권위 소스 → 해당 필드 덮어씀
            if name in ct:
                wsk.cell(r, col, ct[name])

    for sr in stock_rows:                       # 재고 시트(물류 미러) + 그로스재고 수식연동
        r = put("재고", {"물류명": sr["물류명"], "창고": sr["창고"], "구분": sr["구분"], "현재고": sr["현재고"],
                        "셀독": sr["셀독"], "당근": sr["당근"], "자사": sr["자사"], "갱신일": sr["갱신일"],
                        "바코드": sr["바코드"], "상품코드": code_by_stockname.get(_norm(sr["물류명"]), "")})
        wb["재고"].cell(r, stk["그로스재고"], _lk("관리상품", pm["그로스재고"], pm["상품코드"], f"$%s{r}" % get_column_letter(stk["상품코드"])))
        stats["재고"] += 1

    for pr in promos:
        code = code_by_biz_name.get((pr["사업자명"], pr["상품명"]), "")
        aid = next((a["계정아이디"] for a in accts if a["사업자명"] == pr["사업자명"]), "")
        r = put("체험단", {"요청일": _date(pr["요청일"]), "키워드": pr["키워드"], "판매가": pr["판매가"],
                         "리뷰/택배": pr["리뷰/택배"], "건수": pr["건수"], "신규/추가": pr["신규/추가"],
                         "견적서판매가": pr["견적서판매가"], "견적서리뷰/택배": pr["견적서리뷰/택배"],
                         "밑작업": pr["밑작업"], "포토/텍스트": pr["포토/텍스트"], "진행여부": pr["진행여부"],
                         "비고": pr["비고"], "계정아이디": aid, "상품코드": code})
        ch = wb["체험단"]
        ch.cell(r, H["체험단"]["사업자명"], _lk("계정", ac["사업자명"], ac["계정아이디"], f"$%s{r}" % get_column_letter(H["체험단"]["계정아이디"])))
        ch.cell(r, H["체험단"]["상품(물류)명"], _lk("관리상품", pm["상품(물류)명"], pm["상품코드"], f"$%s{r}" % get_column_letter(H["체험단"]["상품코드"])))
        stats["체험단"] += 1

    _style_protect(wb)
    wb.save(out)
    return stats


def main() -> int:
    a = sys.argv
    stats = migrate(Path(a[1]) if len(a) > 1 else DEF_LEDGER,
                    Path(a[2]) if len(a) > 2 else DEF_STOCK,
                    Path(a[3]) if len(a) > 3 else DEF_OUT)
    print("[전환 초안]")
    for k, v in stats.items():
        print(f"  {k}: {v}")
    print("  [주의] 초안 - '매핑검토' 확인/수동 보정. 파생열은 수식연동(잠금)/원본만 수정.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
