"""통합 앱 홈 '변동 추이' 오프라인 검증 — 시안(Concept ④⑤⑥·Main 변동 추이) 기대 동작을 가짜 데이터로 확인.

엑셀·로그인·실 API 없음(결정적). 집계(ui/home_trend) + 위젯 순수함수(ui/widgets_qt) + 홈 화면 offscreen 스모크.

    python tools/verify_home_trend_offline.py
"""
from __future__ import annotations

import os
import sys
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT / "src"), str(ROOT)]
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
try:
    sys.stdout.reconfigure(encoding="utf-8")
except AttributeError:
    pass

import openpyxl  # noqa: E402

from coupang_analytics import config  # noqa: E402
from coupang_analytics.workbook import OutputWorkbook  # noqa: E402
from ui import home_trend as H  # noqa: E402
from ui import widgets_qt as W  # noqa: E402


def ok(msg):
    print(f"  ✔ {msg}")


class _FakeWb:
    """collect() 가 읽는 OutputWorkbook 인덱스만 흉내(시트 2개·날짜 라벨 '월.일')."""
    _rank_num = staticmethod(OutputWorkbook._rank_num)

    def __init__(self):
        self.wb = openpyxl.Workbook()
        self.wb.remove(self.wb.active)
        self._metric_row, self._kw_row, self._date_col = {}, {}, {}

    def account_sheets(self):
        return list(self.wb.sheetnames)

    def add(self, biz, metrics: dict, kws: dict, dates: list[str]):
        ws = self.wb.create_sheet(biz)
        self._date_col[biz] = {d: 8 + i for i, d in enumerate(dates)}
        r = 1
        for (prod, m), vals in metrics.items():
            self._metric_row[(biz, prod, m)] = r
            for i, v in enumerate(vals):
                ws.cell(row=r, column=8 + i).value = v
            r += 1
        for (prod, kw), vals in kws.items():
            self._kw_row[(biz, prod, kw)] = r
            for i, v in enumerate(vals):
                ws.cell(row=r, column=8 + i).value = v
            r += 1


def _daily():
    y = date.today().year
    fw = _FakeWb()
    fw.add("가사업자", {("상품1", config.M_SALES): [3, "1,200", None],
                      ("상품1", config.M_VISITORS): [10, None, None],
                      ("상품1", "판매가"): [9900, 9900, 9900]},                 # 집계 대상 아님
           {("상품1", "캠핑"): ["12위", "44위밖", None], ("상품1", "타프"): [70, "차단", None]},
           ["09.01", "09.02", "09.03"])
    fw.add("나사업자", {("상품9", config.M_SALES): [5, None, None]},
           {("상품9", "의자"): ["30위", "7위", None]}, ["09.01", "09.02", "09.03"])
    return H.collect(fw), date(y, 9, 1), date(y, 9, 2), date(y, 9, 3)


def check_collect():
    d, d1, d2, d3 = _daily()
    assert d[d1].sums == {config.M_SALES: 8, config.M_VISITORS: 10}, d[d1].sums
    ok("지표 합계 = 전 사업자·전 상품 합(판매가 등 비대상 지표 제외)")
    assert d[d2].sums == {config.M_SALES: 1200.0}, d[d2].sums
    ok("쉼표 숫자('1,200') 인식 · 빈 칸은 합에서 제외")
    assert sorted(d[d1].ranks) == [12, 30, 70] and d[d2].ranks == [7], (d[d1].ranks, d[d2].ranks)
    ok("순위 = 정확 순위만('44위밖'·'차단'·공란 제외)")
    assert not d[d3].sums and not d[d3].ranks and H.last_day(d) == d2
    ok("값 없는 날짜 칸(미실행)은 마지막 데이터일에서 제외")


def check_values():
    d, d1, d2, d3 = _daily()
    v, sub = H.value_of(d[d1], H.RANK)
    assert H.PAGE1 == config.RANK_ORGANIC_PER_PAGE and v == 2 and sub == "평균 37위", (v, sub)
    ok(f"노출순위 막대 = 1페이지({H.PAGE1}위) 안 키워드 수 · 보조 = 평균 순위(반올림)")
    assert H.value_of(d[d3], H.RANK) == (None, "") and H.value_of(d[d3], "판매현황") == (None, "")
    ok("빈 날 = None(0 아님·공란)")
    assert H.value_of(d[d1], "검색량") == (None, "") and H.METRICS["검색량"] == config.M_VIEWS
    ok("검색량 탭 = 노출량 지표")


def check_series_period():
    d, d1, d2, d3 = _daily()
    s = H.series(d, "판매현황", d1, d3 + (d3 - d2) * 2)
    assert [p.value for p in s] == [8, 1200.0, None, None, None] and s[3].label == (d3 + (d3 - d2)).strftime("%m.%d")
    ok("series = 모든 달력 날짜(데이터 없는 날도 빈 막대·라벨 MM.DD)")
    end = date(2026, 10, 6)
    assert H.period_range("7d", end) == (date(2026, 9, 30), end)
    assert H.period_range("14d", end) == (date(2026, 9, 23), end)
    assert H.period_range("month", end) == (date(2026, 10, 1), end)
    ok("기간 알약: 최근 7일·최근 14일·이번 달(1일~마지막 데이터일)")
    try:
        H.period_range("1y", end)
    except ValueError:
        ok("알 수 없는 기간 = 오류(조용히 넘기지 않음)")
    else:
        raise AssertionError("알 수 없는 기간이 통과됨")


def check_compare():
    assert H.delta_pct(128, 114) == 12 and H.delta_pct(50, 100) == -50 and H.delta_pct(5, 5) == 0
    assert H.delta_pct(5, None) is None and H.delta_pct(None, 5) is None and H.delta_pct(5, 0) is None
    ok("증감% 반올림 · 비교 불가(없음·이전 0) = None")
    y = date.today().year
    daily = {date(y, 9, 1): H.DayData("09.01", {config.M_SALES: 100}),
             date(y, 9, 8): H.DayData("09.08", {config.M_SALES: 112})}
    assert H.week_compare(daily, "판매현황", date(y, 9, 8)) == (112, 100, 12)
    assert H.week_compare(daily, "판매현황", date(y, 9, 1))[2] is None
    ok("비교 배지 = 지난주 같은 요일 대비(7일 전)")
    assert W.delta_text(12) == ("▲ 12%", "up") and W.delta_text(-5) == ("▼ 5%", "down")
    assert W.delta_text(0) == ("– 0%", "flat") and W.delta_text(None) == ("", "")
    ok("배지 글자·색 구분: ▲ 초록 / ▼ 빨강 / 비교 불가 = 숨김")


def check_bars():
    assert W.bar_ratios([50, None, 100, 0]) == [0.5, 0.0, 1.0, 0.0] and W.bar_ratios([None, 0]) == [0.0, 0.0]
    assert W.label_step(7) == 1 and W.label_step(31) == 4
    assert W.fmt_num(1200.0) == "1,200" and W.fmt_num(None) == "—" and W.fmt_num(2.5) == "2.5"
    ok("막대 높이 = 최댓값 대비 비율(None=빈 트랙) · 날짜 글자 솎기 · 숫자 표기")


def check_ui_smoke():
    from PySide6 import QtWidgets
    from ui import home_page_qt as P
    from ui.theme_qt import QSS
    app = QtWidgets.QApplication.instance() or QtWidgets.QApplication([])
    app.setStyleSheet(QSS)
    missing = P.HomePage(Path("__없는파일__.xlsx"))
    missing.refresh()
    assert "아직 없습니다" in missing.caption.text() and missing.chart.selected_index() == -1
    ok("결과 엑셀 없음 → 카드에 이유 표시(빈 화면 아님)")
    d, d1, d2, _d3 = _daily()
    orig = H.load_master
    H.load_master = lambda _p: d
    try:
        page = P.HomePage(Path(__file__))                # 존재하는 경로(내용은 가짜 집계로 대체)
        page.resize(1000, 520)
        page.refresh()
    finally:
        H.load_master = orig
    assert len(page.chart._pts) == 7 and page.chart.selected_index() == 6
    assert page.kpi.value.text() == "1개" and page.kpi.badge.text() == ""
    ok("홈: 최근 7일 막대 7개·마지막 데이터일 선택·KPI=1페이지 키워드 1개(지난주 값 없어 배지 숨김)")
    page.tabs.set_current("판매현황")
    page.tabs.changed.emit("판매현황")
    page.period.set_current("14d")
    page.period.changed.emit("14d")
    assert len(page.chart._pts) == 14 and page.kpi.value.text() == "1,200건"
    dim = [b for b in page.tabs.findChildren(QtWidgets.QPushButton) if b.objectName() == "pillDim"]
    assert len(dim) == 1 and not dim[0].isEnabled() and "준비 중" in dim[0].text()
    page.grab()                                        # paintEvent 실행(예외 없이 그려지는지)
    ok("탭·기간 전환 반영(14일·판매 1,200건) · 매출·수익 = 회색 '준비 중' 비활성 · 그리기 정상")


def main() -> int:
    for name, fn in [("[1] 엑셀 인덱스 → 날짜별 집계", check_collect), ("[2] 탭별 값", check_values),
                     ("[3] 날짜 줄·기간 알약", check_series_period), ("[4] 비교 배지", check_compare),
                     ("[5] 알약 막대", check_bars), ("[6] 홈 화면 offscreen", check_ui_smoke)]:
        print(name)
        fn()
    print("홈 변동 추이 오프라인 검증 통과")
    return 0


if __name__ == "__main__":
    sys.exit(main())
