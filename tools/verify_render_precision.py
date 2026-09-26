"""정밀 렌더 검증 — run_full 을 목킹으로 end-to-end 실행 후, **각 시트·항목·값**이 새 서식
(항목3 상품명 쿠팡링크·항목4 상품군색·항목5 제목·다계정ID 그룹핑·④ 계정목록 열순서/링크)에 맞게
제대로 채워지는지 렌더된 xlsx + gsheet 요청을 **셀 단위로 대조**한다.

시나리오(현실 재현):
  · 사업자 '로움컨설팅' = 계정ID 2개(loum1·loum2) → 한 시트로 그룹핑(항목5), 상품별 계정ID 태깅
      - loum1: '타프' 로켓그로스 **2옵션**(베이지·그레이) = 같은 상품군(항목4 같은 배경색)
      - loum2: '매트' 판매자배송 단일옵션(다른 상품군 = 다른 배경색)
  · 사업자 '단독상사' = 계정ID solo1, '의자' 로켓그로스 + **체험단(마케팅) 기간**
지표는 항목별로 **서로 다른 값**을 넣어(노출310·판매27·방문88·재고45·판매가19900) 각 지표행에
정확히 매핑되는지 구분 검증한다.

실행: python tools/verify_render_precision.py   (로그인·API·브라우저 없음·결정적)
"""
from __future__ import annotations

import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "tools"))

try:
    sys.stdout.reconfigure(encoding="utf-8")
except AttributeError:
    pass

import openpyxl  # noqa: E402

from coupang_analytics import config  # noqa: E402
from coupang_analytics import pipeline as P  # noqa: E402
from coupang_analytics import gsheet_stats, gsheet_index  # noqa: E402
from coupang_analytics.input_list import Account, InputList, Option, Product  # noqa: E402
from coupang_analytics.report import OptionMetric  # noqa: E402
from coupang_analytics.workbook import OutputWorkbook  # noqa: E402
import simulate_pipeline as SIM  # noqa: E402  (가짜 키워드·순위·브라우저 재사용)

_OK = 0
_FAIL = 0


def _chk(cond: bool, msg: str) -> None:
    global _OK, _FAIL
    mark = "[통과]" if cond else "[실패]"
    print(f"    {mark} {msg}")
    if cond:
        _OK += 1
    else:
        _FAIL += 1


def _n(v) -> str:
    return str(v).strip() if v is not None else ""


# ── 풍부한 수집 대역(항목별 값 구분) ─────────────────────────────
def _rich_discover(a, date_from, date_to, get_password, log, login=True, semi=False):
    metrics: dict[str, OptionMetric] = {}
    inventory: dict[str, int] = {}
    inv_status: dict[str, object] = {}
    vid_meta: dict[str, tuple] = {}
    pid_by_vid: dict[str, str] = {}
    for product in a.products:
        rt = "RFM" if product.kind == config.KIND_CONTRACT else "NORMAL"
        for opt in product.options:
            for vid in opt.vendor_item_ids:
                metrics[vid] = OptionMetric(option_id=vid, product_name=product.name,
                                            option_name=opt.label, item_id=f"item_{vid}",
                                            views=310, sales=27, visitors=88, registration_type=rt,
                                            product_id=f"pid_{vid}")
                pid_by_vid[vid] = f"pid_{vid}"
                if product.kind == config.KIND_CONTRACT:
                    inventory[vid] = 45
                    inv_status[vid] = "판매중"
                    vid_meta[vid] = (19900, "2026-09-01")   # (판매가, 로켓그로스 판매일)
    live = {vid for p in a.products for o in p.options for vid in o.vendor_item_ids}
    log(f"  [{a.business_name}/{a.account_id}] 발견(정밀) 상품 {len(a.products)}개")
    return (Account(a.account_id, a.representative, a.business_name, a.products),
            metrics, inventory, inv_status, set(), live, vid_meta, pid_by_vid)


def _fixture() -> InputList:
    # 로움컨설팅 = 계정ID 2개(다계정ID) → 한 시트
    타프 = Product(name="타프", kind=config.KIND_CONTRACT, options=[
        Option(label="베이지", vendor_item_ids=["vid_beige"], product_ids=[]),
        Option(label="그레이", vendor_item_ids=["vid_gray"], product_ids=[])])
    매트 = Product(name="매트", kind=config.KIND_PERSONAL,
                  options=[Option(label="", vendor_item_ids=["vid_mat"], product_ids=[])])
    의자 = Product(name="의자", kind=config.KIND_CONTRACT,
                  options=[Option(label="", vendor_item_ids=["vid_chair"], product_ids=[])],
                  mkt_start="2026-09-24", mkt_end="2026-09-30", mkt_mon="2026-10-10")   # 체험단중
    accts = [
        Account("loum1", "김대표", "로움컨설팅", [타프]),
        Account("loum2", "김대표", "로움컨설팅", [매트]),
        Account("solo1", "박대표", "단독상사", [의자]),
    ]
    for a in accts:   # 파서 재현: 줄 존재 상품명(⑥ 완전삭제 판정용) — 다계정 reconcile 스코핑 실전 재현
        a.ledger_products = {p.name for p in a.products}
    return InputList(accounts=accts, errors=[],
                     ledger_account_ids={"loum1", "loum2", "solo1"})


def _hdr_row(ws, base: str) -> int:
    for r in range(1, ws.max_row + 1):
        if _n(ws.cell(r, 7).value) == "날짜" and _n(ws.cell(r, 3).value).split(config.NAME_ID_SEP, 1)[0] == base:
            return r
    raise AssertionError(f"헤더행 못 찾음: {base}")


def _metric_val(ws, base: str, metric: str, col: int):
    """상품 base 블록에서 지표행(G=metric)의 col 값."""
    hr = _hdr_row(ws, base)
    for r in range(hr, ws.max_row + 1):
        g = _n(ws.cell(r, 7).value)
        if r > hr and g == "날짜":      # 다음 블록
            break
        if g == metric:
            return ws.cell(r, col).value
    return None


def main() -> int:
    SIM._install_fakes()
    P._login_and_discover = _rich_discover   # 정밀 수집 대역으로 덮어씀
    d = Path(tempfile.mkdtemp())
    final = P.run_full(_fixture(), naver=None, out_dir=str(d), ai_key="sim",
                       date_from="2026-09-25", date_to="2026-09-25", date_label="2026-09-26",
                       resume=False, on_log=lambda m: None)
    wb = openpyxl.load_workbook(final)
    print("=" * 70)
    print("  정밀 렌더 검증 — 시트·항목·값 대조")
    print("=" * 70)

    # ── 시트 구성(항목5 그룹핑) ──
    print("[통계 시트 구성]")
    sheets = set(wb.sheetnames) - {"_상품ID", "계정 목록", "_계정정보", "_수집스탬프", "_마케팅", "_중단"}
    _chk(sheets == {"로움컨설팅", "단독상사"}, f"다계정ID 한 시트 그룹핑 → {sorted(sheets)}")
    ws = wb["로움컨설팅"]

    # ── 항목5: 상단 제목(대표자·사업자·계정ID·다계정 모두) ──
    print("[항목5] 상단 제목")
    title = _n(ws.cell(1, 1).value)
    _chk(title == f"{config.SELDOC_SHEET_TITLE}(김대표, 로움컨설팅, loum1 / loum2)", f"제목='{title}'")

    # ── 상품 블록 헤더(상품명·VID·판매방식·로켓그로스) ──
    print("[상품 블록 헤더]")
    hr = _hdr_row(ws, "타프 (베이지)")
    _chk(_n(ws.cell(hr, 1).value) == "상품명", "pos0 라벨=상품명")
    _chk(_n(ws.cell(hr + 2, 1).value) == "VID" and _n(ws.cell(hr + 2, 3).value) == "vid_beige", "VID 값 렌더")
    _chk(_n(ws.cell(hr + 3, 1).value) == "판매방식" and _n(ws.cell(hr + 3, 3).value) == config.KIND_CONTRACT,
         "판매방식=로켓그로스")

    # ── 항목3: 상품명 → 쿠팡 노출상품 하이퍼링크(새 창) ──
    print("[항목3] 상품명 쿠팡 하이퍼링크")
    lnk = ws.cell(hr, 3).hyperlink
    tgt = _n(getattr(lnk, "target", "")) if lnk else ""
    _chk(tgt.startswith("https://www.coupang.com/"), f"상품명 셀 외부 쿠팡 링크: {tgt[:48]}")

    # ── 지표행 값(항목별 서로 다른 값으로 정확 매핑) ──
    print("[지표행 값 — 항목별 정확 매핑]")
    dcol = max(OutputWorkbook.load(final)._date_col.get("로움컨설팅", {}).values())
    _chk(_metric_val(ws, "타프 (베이지)", config.M_VIEWS, dcol) == 310, "노출량=310")
    _chk(_metric_val(ws, "타프 (베이지)", config.M_SALES, dcol) == 27, "판매량=27")
    _chk(_metric_val(ws, "타프 (베이지)", config.M_VISITORS, dcol) == 88, "방문자=88")
    _chk(_metric_val(ws, "타프 (베이지)", config.M_INVENTORY, dcol) == 45, "재고현황=45")
    _chk(_metric_val(ws, "타프 (베이지)", config.M_SALE_PRICE, dcol) == 19900, "판매가=19900")
    _chk(_n(_metric_val(ws, "타프 (베이지)", config.M_SALE_STATUS, dcol)) == "판매중", "판매상태=판매중")

    # ── 항목2: 상품명 하이퍼링크 = 노출상품 페이지(productId, 판매분석∪재고에서 확보) ──
    wpid = OutputWorkbook.load(final)
    _url = wpid.product_url("로움컨설팅", "타프 (베이지)")
    _chk("/vp/products/pid_" in _url, f"상품 페이지 링크(productId)={_url[:60]}")
    _chk("np/search" not in _url, "검색 링크 폴백 아님(pid 확보됨)")

    # ── 키워드(순위·검색량) ──
    print("[키워드 순위·검색량]")
    wl = OutputWorkbook.load(final)
    kws = wl.product_keywords("로움컨설팅", "타프 (베이지)")
    _chk(kws == ["kw1", "kw2"], f"AI 선정 키워드 적재={kws}")
    _chk(wl.keyword_search("로움컨설팅", "타프 (베이지)", "kw1") == 1000, "검색량(kw1)=1000")
    _chk(P._best is not None, "순위 기록 경로 존재")  # 값 자체는 아래 셀에서
    hrk = _hdr_row(ws, "타프 (베이지)")
    kh = next((r for r in range(hrk, ws.max_row + 1) if _n(ws.cell(r, 1).value) == "키워드"), None)
    _chk(kh is not None and _n(ws.cell(kh + 1, dcol).value) == "3위", "키워드 노출순위=3위")

    # ── 다중옵션: 2차 옵션(그레이) 블록 = 판매정보만(키워드 소헤더 없음) ──
    print("[다중옵션 분리]")
    _chk(wl.has_keyword_section("로움컨설팅", "타프 (베이지)"), "대표 옵션(베이지)=키워드 있음")
    _chk(not wl.has_keyword_section("로움컨설팅", "타프 (그레이)"), "2차 옵션(그레이)=키워드 소헤더 없음")

    # ── 항목5: 상품별 계정ID 태깅(다계정ID) ──
    print("[항목5] 상품별 계정ID 태깅")
    _chk(wl.product_account_id("로움컨설팅", "타프 (베이지)") == "loum1", "타프=loum1")
    _chk(wl.product_account_id("로움컨설팅", "매트") == "loum2", "매트=loum2")

    # ── 항목4: 상품군 배경색(같은 등록명 옵션=같은 색·다른 상품=다른 색) ──
    print("[항목4] 상품군 배경색")
    def _fill(w, r):
        return _n(getattr(w.cell(r, 1).fill.fgColor, "rgb", ""))
    c_beige = _fill(ws, _hdr_row(ws, "타프 (베이지)"))
    c_gray = _fill(ws, _hdr_row(ws, "타프 (그레이)"))
    c_mat = _fill(ws, _hdr_row(ws, "매트"))
    _chk(c_beige == c_gray, "같은 상품군(타프 옵션)=같은 배경색")
    _chk(c_beige != c_mat, "다른 상품(매트)=다른 배경색")

    # ── 계정목록(항목4 열순서·링크·밴드) ──
    print("[계정 목록 시트]")
    idx = wb["계정 목록"]
    heads = [_n(idx.cell(2, c).value) for c in range(1, 10)]
    _chk(heads[:4] == ["대표자", "사업자", "계정ID", "상품명(클릭 이동)"], f"열순서(계정ID 상품 왼쪽)={heads[:4]}")
    _chk(heads[7] == "상태" and heads[8] == "체험단효과", "상태·체험단효과 열")
    # 로움컨설팅 상품 2줄(타프·매트), 각 상품별 계정ID
    rows = {}
    for r in range(3, idx.max_row + 1):
        if _n(idx.cell(r, 2).value) == "로움컨설팅":
            rows[_n(idx.cell(r, 4).value)] = _n(idx.cell(r, 3).value)   # 상품명(대표옵션 블록명) → 계정ID
    # 계정목록 = 상품별 1줄(대표 옵션만): 타프 대표=베이지 블록명 '타프 (베이지)'·매트
    _chk(rows.get("타프 (베이지)") == "loum1" and rows.get("매트") == "loum2", f"계정목록 상품별 계정ID={rows}")
    _chk("타프 (그레이)" not in rows, "2차 옵션(그레이)은 계정목록 제외(상품별 1줄)")
    # 상품명 셀(D열) = 통계 시트 점프 하이퍼링크(항목2)
    drow = next(r for r in range(3, idx.max_row + 1)
                if _n(idx.cell(r, 2).value) == "로움컨설팅" and _n(idx.cell(r, 4).value) == "타프 (베이지)")
    _chk(idx.cell(drow, 4).hyperlink is not None, "계정목록 상품명=통계 점프 링크(항목2)")

    # ── 항목1: gsheet 미러링(통계 셀단위 자동미러 + 상품명 외부링크 =HYPERLINK) ──
    print("[항목1] gsheet 미러링")
    reqs = gsheet_stats.worksheet_to_requests(ws, 42, index_gid=7)
    flat = str(next(r for r in reqs if "updateCells" in r))
    _chk('HYPERLINK("https://www.coupang.com/' in flat, "통계 gsheet: 상품명 쿠팡 외부링크 =HYPERLINK 미러")
    roster = gsheet_index.roster_from_workbook(wl, {"로움컨설팅": 42, "단독상사": 43})
    bands = {r.band for r in roster if r.business == "로움컨설팅"}
    _chk(len(bands) == 1, "계정목록 gsheet: 다계정ID 사업자=한 밴드색")

    print("=" * 70)
    print(f"  결과: 통과 {_OK} · 실패 {_FAIL}   (렌더 파일: {final})")
    print("=" * 70)
    return 1 if _FAIL else 0


if __name__ == "__main__":
    raise SystemExit(main())
