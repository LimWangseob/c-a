"""run_full 파이프라인 시뮬레이션 검증 (브라우저·로그인·네트워크 없이) — 셀독 새 서식.

외부 의존(로그인/브라우저/네이버/순위조회/AI)을 **가짜로 대체**하고 `run_full` 실제 로직만 돌려
새 서식 워크북(시트=사업자, 상품블록 계약/개인, 키워드 노출순위, 일자 누적)·재개·동결·발굴을 검증한다.

시나리오: 1)정상 전체실행 2)크래시→이어서 3)로그인 실패 계정 건너뛰기 4)통계 이어쓰기(동결)+발굴추가
  … 10)새 전체실행 조합(①반자동 판매만→②키워드선정(노출측정 없음)→③반자동 순위 + 백필가드).
실행: python tools/simulate_pipeline.py
"""
from __future__ import annotations

import json
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

try:
    sys.stdout.reconfigure(encoding="utf-8")
except AttributeError:
    pass

import openpyxl  # noqa: E402

from coupang_analytics import config  # noqa: E402
from coupang_analytics import pipeline as P  # noqa: E402
from coupang_analytics.input_list import Account, InputList, Option, Product  # noqa: E402
from coupang_analytics.kw_recommend import TrackKeyword  # noqa: E402
from coupang_analytics.rank import SearchItem  # noqa: E402
from coupang_analytics.report import OptionMetric  # noqa: E402

# 검색결과 매칭 시 반환되는 '정확 노출명'(가짜). run_full 동결 경로에서 계약상품명 갱신을 재현.
_SERP_NAME = "쿠팡실제노출명"

_LOGIN_FAIL_ID = "FAIL"
_STATE = {"select_calls": 0, "crash_at": None, "error_at": None,
          "need_login": set(), "block_login": set()}
# crash_at = 프로세스 강제종료 프록시(KeyboardInterrupt=BaseException, 계정격리 안 됨 → resume 대상)
# error_at = 한 계정 처리 오류 프록시(RuntimeError=일반 예외 → 그 계정만 건너뜀, 전체 완주)
# need_login = 세션 만료(1차 패스에서 NeedLogin→대기열, 2차 로그인 시 정상 수집)
# block_login = Akamai 로그인 차단(2차 패스에서 LoginBlocked → 서킷브레이커 카운트)


# ── 가짜 의존성 ────────────────────────────────────────────────
class _FakeBrowser:
    def __init__(self, *a, **k): pass
    def __enter__(self): return self
    def __exit__(self, *exc): return False


def _fake_login_and_discover(a, date_from, date_to, get_password, log, login=True, semi=False):
    if a.account_id in _STATE["block_login"]:      # Akamai 차단 계정
        if not login:
            raise P.NeedLogin()                    # 1차: 세션 없음 → 대기열
        raise P.LoginBlocked()                     # 2차: 로그인 차단(서킷브레이커 카운트)
    if a.account_id in _STATE["need_login"] and not login:
        raise P.NeedLogin()                        # 1차: 세션 만료 → 대기열(2차 로그인 시 정상)
    if a.account_id == _LOGIN_FAIL_ID:
        log(f"  [{a.business_name}] 로그인 미완료 — 건너뜀(가짜)")
        return None, {}, {}, {}, set(), set(), {}, {}
    metrics: dict[str, OptionMetric] = {}
    inventory: dict[str, int] = {}
    inv_status: dict[str, bool] = {}   # {vid: isSaleSuspended} — 계약 상품은 쿠팡 '판매중'(=False) 가정
    pid_by_vid: dict[str, str] = {}    # {vid: productId} — 상품명 하이퍼링크(항목2)
    for product in a.products:
        rt = "RFM" if product.kind == config.KIND_CONTRACT else "NORMAL"
        for opt in product.options:
            for vid in opt.vendor_item_ids:
                metrics[vid] = OptionMetric(option_id=vid, product_name=product.name,
                                            option_name=opt.label, item_id=f"item_{vid}",
                                            views=100, sales=7, visitors=50, registration_type=rt,
                                            product_id=f"pid_{vid}")
                pid_by_vid[vid] = f"pid_{vid}"
                if product.kind == config.KIND_CONTRACT:   # 계약 상품만 재고현황(가짜)
                    inventory[vid] = 42
                    inv_status[vid] = False                 # 쿠팡 실제 = 판매중
    log(f"  [{a.business_name}] 발견(가짜) 상품 {len(a.products)}개")
    live = {vid for p in a.products for o in p.options for vid in o.vendor_item_ids}
    return (Account(a.account_id, a.representative, a.business_name, a.products),
            metrics, inventory, inv_status, set(), live, {}, pid_by_vid)


def _fake_keywords(title, naver, ai_key=None, n=None, browser=None, log=None,
                   measure_ranks=None, exclude=None):
    """가짜 선정 — TrackKeyword(exposure_best=순위3). exclude면 새 키워드만. crash_at 시점에 크래시."""
    _STATE["select_calls"] += 1
    if _STATE["crash_at"] is not None and _STATE["select_calls"] == _STATE["crash_at"]:
        raise KeyboardInterrupt("시뮬레이션 프로세스 강제종료(계정격리로 안 잡힘 → resume 대상)")
    if _STATE["error_at"] is not None and _STATE["select_calls"] == _STATE["error_at"]:
        raise RuntimeError("시뮬레이션 계정 처리 오류(일반 예외 → 그 계정만 건너뜀)")
    if exclude:
        extra = [("kw3", 1500, "낮음"), ("kw4", 800, "중간")]
        picks = [p for p in extra if p[0] not in exclude][:(n or 1)]
        return [TrackKeyword(k, v, 55.0, comp_idx=c, exposure_best=3) for k, v, c in picks]
    return [TrackKeyword("kw1", 1000, 60.0, comp_idx="중간", exposure_best=3),
            TrackKeyword("kw2", 2000, 55.0, comp_idx="높음", exposure_best=3)]


def _fake_batch(browser, keywords, matchers, max_rank=None, mobile=False, log=None):
    """organic_ranks_batch 대체(동결 시 순위측정) — 상품 매처에 순위 3."""
    return {kw: {lbl: 3 for lbl in matchers} for kw in keywords}


def _fake_organic_ranks(browser, kw, matchers, max_rank=None, mobile=False, log=None, matched_out=None):
    """organic_ranks 대체 — 순위 3. matched_out 주면 매칭 항목(정확 노출명)을 채워 노출명 갱신 재현."""
    if matched_out is not None:
        for lbl in matchers:
            matched_out[lbl] = SearchItem(is_ad=False, product_id="pid1", vendor_item_id="v1",
                                          name=_SERP_NAME)
    return {lbl: 3 for lbl in matchers}


def _fake_track_ranks_semi(wb, path, log, should_stop):
    """③ 반자동 순위 대체 — 실제 브라우저 타이핑 없이 미기입 키워드에 순위 3 기록.

    _track_ranks_semi(wb, path, log, should_stop) 시그니처와 동일. 계정별 최신 일자에 채운다."""
    for biz in wb.account_sheets():
        date = wb.latest_date(biz)
        if not date:
            continue
        for pname in wb.products_of(biz):
            for kw in wb.product_keywords(biz, pname):
                if not wb.is_rank_filled(biz, pname, kw, date):
                    wb.set_keyword_rank(biz, pname, kw, date, 3)
    wb.save(path)
    return path


def _install_fakes():
    P._login_and_discover = _fake_login_and_discover
    P.select_keywords_light = _fake_keywords
    P.recommend_title = lambda *a, **k: "권고 상품명 예시"
    P.organic_ranks_batch = _fake_batch
    P.organic_ranks = _fake_organic_ranks
    P.warmup = lambda browser: None
    P.WingBrowser = _FakeBrowser


# ── 검증 헬퍼 ─────────────────────────────────────────────────
def _accounts(ids) -> InputList:
    """ids 순서대로 계정(상품 1개·옵션 1개). 첫 계정=계약(로켓그로스), 나머지=개인."""
    accts = []
    for idx, aid in enumerate(ids):
        opt = Option(label=f"옵-{aid}", vendor_item_ids=[f"vid-{aid}"], product_ids=[])
        kind = config.KIND_CONTRACT if idx == 0 else config.KIND_PERSONAL
        prod = Product(name=f"상품-{aid}", options=[opt], kind=kind)
        accts.append(Account(aid, f"대표-{aid}", f"비즈-{aid}", [prod]))
    return InputList(accounts=accts, errors=[])


def _account_one_vid(vid: str, name: str = "상품-a1") -> InputList:
    """상품 1개·옵션 1개(vid 지정) 계정 — vid 변경 마이그레이션 검증용."""
    opt = Option(label="", vendor_item_ids=[vid], product_ids=[])
    prod = Product(name=name, options=[opt], kind=config.KIND_CONTRACT)
    return InputList(accounts=[Account("a1", "대표-a1", "비즈-a1", [prod])], errors=[])


def _account_multi_option() -> InputList:
    """2옵션(베이지/그레이) 로켓그로스 상품 1개 계정 — 옵션 분리 검증용."""
    opts = [Option(label="베이지", vendor_item_ids=["vid-beige"], product_ids=[]),
            Option(label="그레이", vendor_item_ids=["vid-gray"], product_ids=[])]
    prod = Product(name="캠핑타프", options=opts, kind=config.KIND_CONTRACT)
    return InputList(accounts=[Account("m1", "대표-m1", "비즈-m1", [prod])], errors=[])


def _sheets(path: Path) -> set[str]:
    # 계정(사업자) 시트만 — 특수 시트(상품ID·목차·계정정보·수집스탬프·중단)는 제외
    return set(openpyxl.load_workbook(path).sheetnames) - {
        "_상품ID", "목차", "계정 목록", "_계정정보", "_수집스탬프", "_마케팅", "_중단"}


def _has_value(path: Path, target) -> bool:
    wb = openpyxl.load_workbook(path)
    return any(c.value == target for ws in wb.worksheets for row in ws.iter_rows() for c in row)


def _keywords_in(path: Path) -> set[str]:
    wb = openpyxl.load_workbook(path)
    out = set()
    for ws in wb.worksheets:
        for r in range(1, ws.max_row + 1):
            # 키워드명 = A열(레이아웃 v4 좌측확장·G=노출순위 행)
            if _n(ws.cell(r, 7).value) == config.M_RANK and ws.cell(r, 1).value:
                out.add(ws.cell(r, 1).value)
    return out


def _date_headers(path: Path) -> set[str]:
    wb = openpyxl.load_workbook(path)
    out = set()
    for ws in wb.worksheets:
        for r in range(1, ws.max_row + 1):
            if _n(ws.cell(r, 7).value) == "날짜":
                for c in range(8, ws.max_column + 1):
                    if ws.cell(r, c).value:
                        out.add(ws.cell(r, c).value)
    return out


def _n(v):
    return str(v).strip() if v is not None else ""


def _check(cond: bool, msg: str) -> None:
    print(f"    {'[통과]' if cond else '[실패]'} {msg}")
    if not cond:
        raise AssertionError(msg)


# ── 시나리오 ──────────────────────────────────────────────────
def scenario_normal():
    print("[시나리오 1] 정상 전체 실행 (3계정: 계약1·개인2)")
    _STATE.update(select_calls=0, crash_at=None, error_at=None)
    d = Path(tempfile.mkdtemp())
    final = P.run_full(_accounts(["a1", "b1", "c1"]), naver=None, out_dir=str(d), ai_key="sim",
                       date_from="2026-09-02", date_to="2026-09-02", resume=False, on_log=lambda m: None)
    _check(final.exists(), f"최종본 생성: {final.name}")
    _check(not P._partial_path(d).exists() and not P._progress_path(d).exists(), "진행 파일 정리됨")
    _check(_sheets(final) == {"비즈-a1", "비즈-b1", "비즈-c1"}, "계정별 시트 3개 존재")
    _check(_has_value(final, "3위"), "키워드 노출순위(3위) 기록됨")
    _check(_has_value(final, 7), "판매량 값 기록됨")
    _check(_has_value(final, 1000), "검색량(1000) 기록됨")
    # 계약(a1)=판매량·방문자·노출량·재고현황, 개인(b1)=판매량·방문자·노출량(재고 없음)
    wb = openpyxl.load_workbook(final)
    g_a1 = {_n(wb["비즈-a1"].cell(r, 7).value) for r in range(1, wb["비즈-a1"].max_row + 1)}
    g_b1 = {_n(wb["비즈-b1"].cell(r, 7).value) for r in range(1, wb["비즈-b1"].max_row + 1)}
    _check(config.M_SALES in g_a1 and config.M_INVENTORY in g_a1, "계약 상품: 판매량·재고현황 지표행")
    _check(config.M_SALES in g_b1 and config.M_VISITORS in g_b1 and config.M_INVENTORY not in g_b1,
           "개인 상품: 판매량·방문자·노출량(재고현황 없음)")


def scenario_crash_resume():
    print("[시나리오 2] 크래시 후 이어서 하기")
    d = Path(tempfile.mkdtemp())
    # 첫 실행: a1 선정(호출1) 완료, b1 선정(호출2)에서 크래시 → done=[a1]
    _STATE.update(select_calls=0, crash_at=2, error_at=None)
    crashed = False
    try:
        P.run_full(_accounts(["a1", "b1", "c1"]), naver=None, out_dir=str(d), ai_key="sim",
                   date_from="2026-09-02", date_to="2026-09-02", resume=False, on_log=lambda m: None)
    except KeyboardInterrupt:   # 프로세스 강제종료 프록시(BaseException) — 계정격리로 안 잡히고 런 중단
        crashed = True
    _check(crashed, "첫 실행이 크래시로 중단됨")
    _check(P._partial_path(d).exists() and P._progress_path(d).exists(), "진행 파일 남음")
    meta = json.loads(P._progress_path(d).read_text(encoding="utf-8"))
    _check(meta["done"] == ["a1"], f"완료 계정=a1 (실제 {meta['done']})")
    # 이어서: a1 건너뛰고 b1·c1 완료
    _STATE.update(select_calls=0, crash_at=None, error_at=None)
    logs: list[str] = []
    final = P.run_full(_accounts(["a1", "b1", "c1"]), naver=None, out_dir=str(d), ai_key="sim",
                       date_from="2026-09-02", date_to="2026-09-02", resume=True, on_log=logs.append)
    _check("이미 완료, 건너뜀" in "\n".join(logs), "완료 계정 a1 건너뜀 로그")
    _check(_sheets(final) == {"비즈-a1", "비즈-b1", "비즈-c1"}, "최종 3계정 시트 존재")
    _check(not P._partial_path(d).exists(), "진행 파일 정리됨")


def scenario_login_fail():
    print("[시나리오 3] 로그인 실패 계정 건너뛰기")
    _STATE.update(select_calls=0, crash_at=None, error_at=None)
    d = Path(tempfile.mkdtemp())
    final = P.run_full(_accounts(["a1", _LOGIN_FAIL_ID, "c1"]), naver=None, out_dir=str(d), ai_key="sim",
                       date_from="2026-09-02", date_to="2026-09-02", resume=False, on_log=lambda m: None)
    _check(_sheets(final) == {"비즈-a1", "비즈-c1"}, "정상 계정만 시트(a1,c1)")
    _check(f"비즈-{_LOGIN_FAIL_ID}" not in _sheets(final), "로그인 실패 계정 제외")


def scenario_carry_forward():
    print("[시나리오 4] 통계 이어쓰기(동결) + 발굴 추가")
    _STATE.update(select_calls=0, crash_at=None, error_at=None)
    d = Path(tempfile.mkdtemp())
    accts = _accounts(["a1"])
    master = P._master_path(d)
    P.run_full(accts, naver=None, out_dir=str(d), ai_key="sim",
               date_from="2026-09-01", date_to="2026-09-01", resume=False, on_log=lambda m: None)
    _check(master.exists() and _keywords_in(master) == {"kw1", "kw2"}, "day1: 마스터+키워드")
    # day2 동결
    logs2: list[str] = []
    P.run_full(accts, naver=None, out_dir=str(d), ai_key="sim", date_from="2026-09-02",
               date_to="2026-09-02", carry_forward=True, on_log=logs2.append)
    _check("(동결)" in "\n".join(logs2), "day2: 키워드 동결 로그")
    _check(_date_headers(master) == {"09.01", "09.02"}, "day2: 2일치 날짜 누적")
    _check(_keywords_in(master) == {"kw1", "kw2"}, "day2: 키워드 변화 없음")
    # day3 발굴 추가
    logs3: list[str] = []
    P.run_full(accts, naver=None, out_dir=str(d), ai_key="sim", date_from="2026-09-03",
               date_to="2026-09-03", carry_forward=True, grow_keywords=True, on_log=logs3.append)
    _check("발굴" in "\n".join(logs3), "day3: 발굴 추가 로그")
    kws3 = _keywords_in(master)
    _check({"kw1", "kw2"} <= kws3 and "kw3" in kws3, f"day3: 기존 유지+발굴 ({sorted(kws3)})")
    _check(len(kws3) <= config.KW_MAX_TRACK, f"day3: 상한 이내 ({len(kws3)})")


def scenario_account_error_isolated():
    print("[시나리오 5] 한 계정 처리 오류 → 격리(건너뜀), 나머지 완주(전체 안 막힘)")
    d = Path(tempfile.mkdtemp())
    _STATE.update(select_calls=0, crash_at=None, error_at=2)   # b1(2번째) 처리 중 일반 예외
    logs: list[str] = []
    final = P.run_full(_accounts(["a1", "b1", "c1"]), naver=None, out_dir=str(d), ai_key="sim",
                       date_from="2026-09-02", date_to="2026-09-02", resume=False, on_log=logs.append)
    joined = "\n".join(logs)
    _check(final is not None, "런이 크래시 없이 완주(최종본 반환)")
    _check("처리 오류" in joined and "건너뜀" in joined, "오류 계정 격리 로그(건너뜀)")
    _check({"비즈-a1", "비즈-c1"} <= _sheets(final), "정상 계정 a1·c1 시트 존재")
    _check(not P._partial_path(d).exists(), "완주 후 진행 파일 정리됨")


def scenario_empty_business_name():
    print("[시나리오 6] 빈 사업자명 → 시트명 label 폴백(대표자명), KeyError 없음")
    d = Path(tempfile.mkdtemp())
    _STATE.update(select_calls=0, crash_at=None, error_at=None)
    opt = Option(label="옵-x", vendor_item_ids=["vid-x"], product_ids=[])
    prod = Product(name="상품-x", options=[opt], kind=config.KIND_PERSONAL)
    ilist = InputList(accounts=[Account("acctX", "대표-X", "", [prod])], errors=[])   # 사업자명 빈값
    final = P.run_full(ilist, naver=None, out_dir=str(d), ai_key="sim", date_from="2026-09-02",
                       date_to="2026-09-02", resume=False, keywords_off=True, on_log=lambda m: None)
    _check(final is not None, "빈 사업자명이어도 크래시 없이 완주")
    _check("대표-X" in _sheets(final), "시트명이 대표자명으로 폴백됨(빈 시트명 KeyError 방지)")


def scenario_session_first():
    print("[시나리오 7] 세션우선 — 세션 만료 계정은 뒤로 미루고 살아있는 계정 먼저 수집")
    d = Path(tempfile.mkdtemp())
    _STATE.update(select_calls=0, crash_at=None, error_at=None,
                  need_login={"b1"}, block_login=set())
    logs: list[str] = []
    final = P.run_full(_accounts(["a1", "b1", "c1"]), naver=None, out_dir=str(d), ai_key="sim",
                       date_from="2026-09-02", date_to="2026-09-02", resume=False, on_log=logs.append)
    joined = "\n".join(logs)
    _check("로그인 대기열" in joined, "세션 만료 b1은 로그인 대기열로 미룸(자동제출 안 함)")
    _check("로그인 필요 계정 1개" in joined, "2차 패스에서 로그인 필요 계정 처리")
    _check(_sheets(final) == {"비즈-a1", "비즈-b1", "비즈-c1"}, "결국 3계정 모두 수집(세션우선+2차 로그인)")


def scenario_circuit_breaker():
    print("[시나리오 8] 로그인 서킷브레이커 — 연속 Akamai 차단 K회 후 이후 로그인 생략")
    d = Path(tempfile.mkdtemp())
    blocked = {"x1", "x2", "x3", "x4"}
    _STATE.update(select_calls=0, crash_at=None, error_at=None,
                  need_login=set(), block_login=blocked)
    logs: list[str] = []
    final = P.run_full(_accounts(["a1", "x1", "x2", "x3", "x4", "c1"]), naver=None, out_dir=str(d),
                       ai_key="sim", date_from="2026-09-02", date_to="2026-09-02",
                       resume=False, on_log=logs.append)
    joined = "\n".join(logs)
    _check(final is not None, "차단 다발에도 크래시 없이 완주")
    _check({"비즈-a1", "비즈-c1"} <= _sheets(final), "세션 있는 a1·c1은 수집됨(세션우선)")
    _check("로그인 생략" in joined, f"연속 {config.LOGIN_BLOCK_CIRCUIT}회 차단 후 이후 로그인 생략(서킷브레이커)")
    _check(all(f"비즈-{x}" not in _sheets(final) for x in blocked), "차단 계정은 미수집(다음에 재시도)")
    _STATE.update(need_login=set(), block_login=set())   # 뒤 시나리오 누수 방지


def _product_names(path: Path, sheet: str) -> list[str]:
    """그 시트의 상품 블록 이름 목록(헤더행=col7 '날짜'의 col3). 중복 블록 감지에 사용."""
    wb = openpyxl.load_workbook(path)
    ws = wb[sheet]
    # 이름칸엔 표시용 vid 꼬리(구분자 뒤)가 붙으므로 순수 상품명(구분자 앞)만 취한다.
    return [_n(ws.cell(r, 3).value).split(config.NAME_ID_SEP, 1)[0]
            for r in range(1, ws.max_row + 1) if _n(ws.cell(r, 7).value) == "날짜"]


def scenario_display_name_rename():
    print("[시나리오 9] 블록 이름=등록상품명 고정(노출명 교체 중단) + vid 앵커 + save/load 왕복")
    _STATE.update(select_calls=0, crash_at=None, error_at=None, need_login=set(), block_login=set())
    d = Path(tempfile.mkdtemp())
    accts = _accounts(["a1"])                      # 계약 상품 1개(단일옵션·vid-a1) → 블록명=등록상품명 '상품-a1'
    master = P._master_path(d)
    # day1: 새 상품 선정 → 블록명 = 등록상품명 '상품-a1'
    P.run_full(accts, naver=None, out_dir=str(d), ai_key="sim",
               date_from="2026-09-01", date_to="2026-09-01", resume=False, on_log=lambda m: None)
    _check("상품-a1" in _product_names(master, "비즈-a1"), "day1: 등록상품명 '상품-a1' 블록")
    # day2: 동결 → 순위측정에서 노출명(_SERP_NAME)을 봐도 **블록명은 등록상품명으로 고정**(교체 중단)
    P.run_full(accts, naver=None, out_dir=str(d), ai_key="sim", date_from="2026-09-02",
               date_to="2026-09-02", carry_forward=True, on_log=lambda m: None)
    names2 = _product_names(master, "비즈-a1")
    _check(names2 == ["상품-a1"], f"day2: 블록명 등록상품명 유지(노출명 교체 안 함) — {names2}")
    _check(_SERP_NAME not in names2, "day2: 노출명으로 안 바뀜(set_display_name 중단)")
    _check(_keywords_in(master) == {"kw1", "kw2"}, "day2: 키워드 동결 유지")
    _check(_date_headers(master) == {"09.01", "09.02"}, "day2: 날짜 2일 누적")
    # day3: 같은 vid·같은 등록상품명 → 같은 블록 재사용(중복 생성 없음)
    P.run_full(accts, naver=None, out_dir=str(d), ai_key="sim", date_from="2026-09-03",
               date_to="2026-09-03", carry_forward=True, on_log=lambda m: None)
    names3 = _product_names(master, "비즈-a1")
    _check(names3 == ["상품-a1"], f"day3: 등록상품명 블록 재사용(중복 없음) — {names3}")
    _check(_date_headers(master) == {"09.01", "09.02", "09.03"}, "day3: 날짜 3일 누적")
    # 워크북 직접 검증 — vid 출처=헤더 이름칸(_block_vids) + save/load 왕복 키 안정
    from coupang_analytics.workbook import OutputWorkbook
    wb2 = OutputWorkbook.load(master)
    _check(wb2.resolve_block_name("비즈-a1", ["vid-a1"]) == "상품-a1",
           "resolve_block_name: vid로 등록상품명 블록 조회")
    _check(wb2.product_vids("비즈-a1", "상품-a1") == ["vid-a1"],
           "product_vids: 헤더 이름칸에서 vid 복원(숨김시트 아님)")
    _check(wb2.product_keywords("비즈-a1", "상품-a1") == ["kw1", "kw2"],
           "load 후에도 (사업자,등록상품명) 키로 키워드 조회됨(키 안정)")
    # day4: **상품조회 상품명이 바뀜(같은 vid)** → 같은 블록 이어받아 이름 갱신·이력(09.01~03) 유지·중복 없음
    #        (결과파일 상품명 = 상품조회 productName 을 매일 반영하되 vid 앵커로 시계열 안 끊김 검증)
    P.run_full(_account_one_vid("vid-a1", "상품-a1-개명"), naver=None, out_dir=str(d), ai_key="sim",
               date_from="2026-09-04", date_to="2026-09-04", carry_forward=True, on_log=lambda m: None)
    names4 = _product_names(master, "비즈-a1")
    _check(names4 == ["상품-a1-개명"], f"day4: 상품조회명 변경 → 블록명 갱신·중복 없음 — {names4}")
    _check(_date_headers(master) == {"09.01", "09.02", "09.03", "09.04"}, "day4: 이력 유지(4일 누적)")
    wb4 = OutputWorkbook.load(master)
    _check(wb4.product_vids("비즈-a1", "상품-a1-개명") == ["vid-a1"], "day4: 같은 vid 승계(정체성 유지)")
    _check(wb4.product_keywords("비즈-a1", "상품-a1-개명") == ["kw1", "kw2"], "day4: 키워드 이력 승계")


def scenario_full_composition():
    """새 전체실행 = do_run_full 전체실행 분기의 3단계 조합(2026-09-15, offscreen 추방).

    ①run_full(keywords_off=True,sales_semi=True,skip_ranks=True)=판매만 → ②select_keywords_stage()=키워드
    (노출측정 없음) → ③track_ranks_stage(semi=True)=반자동 순위. 각 단계가 워크북을 올바르게 진전시키는지,
    그리고 ①(keywords_off)이 offscreen 순위백필을 절대 안 하는지(가드) 검증."""
    print("[시나리오 10] 새 전체실행 조합 — ①반자동(판매만)→②키워드선정→③반자동 순위 + 백필가드")
    _STATE.update(select_calls=0, crash_at=None, error_at=None)
    backfill_calls = {"n": 0}
    orig_backfill, orig_semi = P._backfill_ranks, P._track_ranks_semi
    P._backfill_ranks = lambda *a, **k: backfill_calls.__setitem__("n", backfill_calls["n"] + 1)
    P._track_ranks_semi = _fake_track_ranks_semi
    try:
        d = Path(tempfile.mkdtemp())
        il = _accounts(["a1", "b1"])
        # ① 반자동 판매수집만 — 키워드·순위·백필 없음
        snap = P.run_full(il, naver=None, out_dir=str(d), ai_key="sim",
                          date_from="2026-09-14", date_to="2026-09-14", resume=False,
                          carry_forward=False, grow_keywords=False, skip_ranks=True,
                          redo_today=False, sales_semi=True, keywords_off=True, on_log=lambda m: None)
        _check(snap.exists(), "① 최종본 생성")
        _check(_has_value(snap, 7), "① 판매량(7) 기록됨")
        _check(_keywords_in(snap) == set(), "① 단계엔 키워드 없음(키워드는 ②)")
        _check(not _has_value(snap, "3위"), "① 단계엔 순위 없음(순위는 ③)")
        _check(backfill_calls["n"] == 0, "① keywords_off는 순위백필 안 함")
        # ② 키워드 선정 — 노출측정(measure_ranks) 없이 AI 선정만
        p2 = P.select_keywords_stage(naver=None, ai_key="sim", out_dir=str(d),
                                     grow=False, on_log=lambda m: None)
        _check(_keywords_in(p2) == {"kw1", "kw2"}, "② 키워드 선정됨(kw1·kw2)")
        _check(not _has_value(p2, "3위"), "② 단계엔 순위 없음(순위는 ③)")
        # ③ 반자동 순위
        p3 = P.track_ranks_stage(out_dir=str(d), semi=True, on_log=lambda m: None)
        _check(_has_value(p3, "3위"), "③ 반자동 순위(3위) 기록됨")
        # 가드 직접 검증: keywords_off=True면 skip_ranks=False여도 백필 안 함(옛 잠재버그 차단)
        backfill_calls["n"] = 0
        d2 = Path(tempfile.mkdtemp())
        P.run_full(_accounts(["a1"]), naver=None, out_dir=str(d2), ai_key="sim",
                   date_from="2026-09-14", date_to="2026-09-14", resume=False,
                   skip_ranks=False, keywords_off=True, sales_semi=True, on_log=lambda m: None)
        _check(backfill_calls["n"] == 0, "가드: keywords_off=True는 skip_ranks=False여도 백필 안 함")
    finally:
        P._backfill_ranks, P._track_ranks_semi = orig_backfill, orig_semi


def scenario_option_split():
    print("[시나리오 11] 다중옵션 → 옵션별 블록 분리(대표=키워드/순위·2차=판매정보만)")
    _STATE.update(select_calls=0, crash_at=None, error_at=None, need_login=set(), block_login=set())
    d = Path(tempfile.mkdtemp())
    master = P._master_path(d)
    P.run_full(_account_multi_option(), naver=None, out_dir=str(d), ai_key="sim",
               date_from="2026-09-01", date_to="2026-09-01", resume=False, on_log=lambda m: None)
    names = _product_names(master, "비즈-m1")
    rep, sec = "캠핑타프 (베이지)", "캠핑타프 (그레이)"   # 다중옵션은 대표 포함 모든 옵션에 라벨(등록명+옵션라벨)
    _check(rep in names, f"대표 블록=등록명+첫옵션라벨 '{rep}' — {names}")
    _check(sec in names, f"2차 옵션 블록=등록명+라벨 '{sec}' — {names}")
    _check(len(names) == 2, f"블록 2개(옵션 2개) — {names}")
    from coupang_analytics.workbook import OutputWorkbook
    wb = OutputWorkbook.load(master)
    _check(wb.has_keyword_section("비즈-m1", rep), "대표=키워드 구역 있음")
    _check(not wb.has_keyword_section("비즈-m1", sec), "2차=키워드 구역 없음(판매정보만)")
    _check(wb.product_vids("비즈-m1", rep) == ["vid-beige"], "대표 vid=첫 옵션")
    _check(wb.product_vids("비즈-m1", sec) == ["vid-gray"], "2차 vid=그 옵션")
    _check(sorted(wb.sibling_vids("비즈-m1", rep)) == ["vid-beige", "vid-gray"],
           "sibling_vids=리스팅 전 옵션 합집합(③ 순위 매칭 놓침 방지)")
    _check(wb.product_keywords("비즈-m1", rep) == ["kw1", "kw2"], "대표 키워드 선정됨")
    _check(wb.product_keywords("비즈-m1", sec) == [], "2차 키워드 없음")


def scenario_vid_change_reset():
    print("[시나리오 12] vid 변경 시 이전 데이터 삭제·새로 시작(첫 적용 마이그레이션)")
    _STATE.update(select_calls=0, crash_at=None, error_at=None, need_login=set(), block_login=set())
    d = Path(tempfile.mkdtemp())
    master = P._master_path(d)
    from coupang_analytics.workbook import OutputWorkbook
    # day1: 상품-a1, vid=OLD → 09.01 이력 생성
    P.run_full(_account_one_vid("vidOLD"), naver=None, out_dir=str(d), ai_key="sim",
               date_from="2026-09-01", date_to="2026-09-01", resume=False, on_log=lambda m: None)
    wb1 = OutputWorkbook.load(master)
    _check(wb1.product_vids("비즈-a1", "상품-a1") == ["vidOLD"], "day1: vid=OLD 저장")
    _check(wb1.product_keywords("비즈-a1", "상품-a1") == ["kw1", "kw2"], "day1: 키워드 있음")
    # day2: 같은 상품명, vid=NEW(다름) → 이전 데이터 삭제하고 새로 시작(09.01 이력 소멸)
    P.run_full(_account_one_vid("vidNEW"), naver=None, out_dir=str(d), ai_key="sim",
               date_from="2026-09-02", date_to="2026-09-02", carry_forward=True, on_log=lambda m: None)
    wb2 = OutputWorkbook.load(master)
    _check(wb2.product_vids("비즈-a1", "상품-a1") == ["vidNEW"], "day2: vid=NEW로 교체")
    names = _product_names(master, "비즈-a1")
    _check(names.count("상품-a1") == 1, f"블록 1개(중복 없음) — {names}")
    _check(_date_headers(master) == {"09.02"}, f"이전(09.01) 데이터 삭제·새 날짜만 — {_date_headers(master)}")


def scenario_restore_residue_cleanup():
    print("[시나리오 13] 복원 잔재 정리 — vid 없는 옛 블록(구글시트 복원분) + 옵션분리 = 중복 잔재 제거")
    _STATE.update(select_calls=0, crash_at=None, error_at=None, need_login=set(), block_login=set())
    d = Path(tempfile.mkdtemp())
    master = P._master_path(d)
    from coupang_analytics.workbook import OutputWorkbook
    # 구글시트 복원 흉내: 같은 등록상품명 '캠핑타프' 의 **vid 없는** 옛 블록(개인)만 있는 마스터(숨김 메타 미러 안 됨)
    wb0 = OutputWorkbook.empty()
    wb0.ensure_account("비즈-m1")
    wb0.ensure_product_block("비즈-m1", "캠핑타프", config.KIND_PERSONAL, ["옛키워드"], registered="캠핑타프")
    wb0.save(master)
    _check(wb0.product_vids("비즈-m1", "캠핑타프") == [], "사전조건: 복원 블록은 vid 없음")
    # 옵션분리 ①판매수집: 같은 등록명 '캠핑타프'(옵션 beige/gray, vid 있음)
    P.run_full(_account_multi_option(), naver=None, out_dir=str(d), ai_key="sim",
               date_from="2026-09-02", date_to="2026-09-02", carry_forward=True, on_log=lambda m: None)
    names = _product_names(master, "비즈-m1")
    _check("캠핑타프" not in names, f"vid 없는 옛 블록 '캠핑타프' 삭제됨(잔재 없음) — {names}")
    _check("캠핑타프 (베이지)" in names and "캠핑타프 (그레이)" in names, f"새 옵션 블록 생성 — {names}")
    _check(len(names) == 2, f"블록 2개(중복 잔재 없음) — {names}")


def scenario_gsheet_index_tab_excluded():
    print("[시나리오 14] 구글시트 인덱스 탭 '계정목록'(공백 없음)은 통계 계정으로 오인 안 함")
    from coupang_analytics.workbook import OutputWorkbook
    d = Path(tempfile.mkdtemp())
    master = P._master_path(d)
    wb = OutputWorkbook.empty()
    wb.ensure_account("가게A")
    wb.wb.create_sheet("계정목록")   # 결과 구글시트 미러 인덱스 탭(공백 없음) 이 복원본에 섞인 상황
    wb.save(master)
    wb2 = OutputWorkbook.load(master)
    _check("계정목록" not in wb2.account_sheets(), f"'계정목록'(공백없음) 제외 — {wb2.account_sheets()}")
    _check("가게A" in wb2.account_sheets(), f"실계정 '가게A' 는 포함 — {wb2.account_sheets()}")


def main():
    _install_fakes()
    config.LOGIN_PACE_MIN_SEC = 0   # 시뮬은 로그인 페이싱 sleep 없이(즉시)
    config.LOGIN_PACE_MAX_SEC = 0
    config.RANK_NAV_DELAY_MIN_SEC = 0   # 순위 직렬 네비 간격도 0(즉시)
    config.RANK_NAV_DELAY_MAX_SEC = 0
    import os, tempfile      # 관측 DB는 임시로(실 data/session_state.db 오염 방지)
    config.SESSION_STATE_DB = os.path.join(tempfile.gettempdir(), "sim_session_state.db")
    print("=" * 60)
    print("  run_full 시뮬레이션 검증 (셀독 새 서식)")
    print("=" * 60)
    scenario_normal()
    scenario_crash_resume()
    scenario_login_fail()
    scenario_carry_forward()
    scenario_account_error_isolated()
    scenario_empty_business_name()
    scenario_session_first()
    scenario_circuit_breaker()
    scenario_display_name_rename()
    scenario_full_composition()
    scenario_option_split()
    scenario_vid_change_reset()
    scenario_restore_residue_cleanup()
    scenario_gsheet_index_tab_excluded()
    print("=" * 60)
    print("  [완료] 모든 시나리오 통과")
    print("=" * 60)


if __name__ == "__main__":
    main()
