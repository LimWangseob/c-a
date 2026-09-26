"""워크북 공용 상수·헬퍼·서식 dataclass (leaf — workbook·workbook_render 공유).

순환 import 방지: 이 모듈은 openpyxl·config 등 leaf 만 import 하고 workbook 을 import 하지 않는다.
workbook.py·workbook_render.py 가 `from .workbook_common import *` 로 이어받는다(__all__ 에 밑줄 심볼 포함).
"""
from __future__ import annotations

from dataclasses import dataclass
from datetime import date as _date, datetime as _dt, timedelta as _td
from pathlib import Path
from urllib.parse import quote as _quote

import openpyxl
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink

from . import config


def _unmerge_all(ws) -> None:
    """시트의 모든 병합을 해제. 희소(미실체화) 병합셀은 정규 Cell 로 먼저 채워 unmerge KeyError 를 막는다.

    openpyxl `insert_rows`/재서식이 병합셀과 함께 쓰이면 데이터가 유실되므로, 행 삽입·서식 전에 호출한다.
    """
    for mr in list(ws.merged_cells.ranges):
        for rr in range(mr.min_row, mr.max_row + 1):
            for cc in range(mr.min_col, mr.max_col + 1):
                if (rr, cc) not in ws._cells:
                    ws._cells[(rr, cc)] = Cell(ws, row=rr, column=cc)
        ws.unmerge_cells(str(mr))


@dataclass
class _StyleCtx:
    """apply_style 팔레트(폰트·채움·테두리·정렬) 묶음 — 시트/블록 서식 헬퍼가 공유."""
    font: Font
    bold: Font
    title_font: Font
    f_prod: PatternFill
    f_prod2: PatternFill          # 상품군 교대 배경(같은 등록상품명=한 군, 인접 군을 두 색으로 구분)
    f_prod_lt: PatternFill        # 값칸·키워드 옅은 상품군색(상품군 전체 은은히 통일)
    f_prod2_lt: PatternFill
    f_label: PatternFill
    f_kwhead: PatternFill
    f_kind: PatternFill
    mkt_fill: PatternFill
    box: Border
    center: Alignment
    wrap: Alignment
    thick: Side
    thin: Side
    mid: Side          # 기본↔옵션(같은 상품군 내부) 구분선 — 격자(thin)보다 진하고 그룹 바깥(thick)보다 얇게


def _sty_cell(ws, r, c, sty: _StyleCtx, *, fill=None, fnt=None, align=None, num=False) -> None:
    """셀 서식 적용(폰트·정렬·테두리·선택 채움/숫자서식). 분해 전 apply_style 내부 `cell` 클로저와 동일."""
    x = ws.cell(r, c)
    x.font = fnt or sty.font
    x.alignment = align or sty.center
    x.border = sty.box
    if fill:
        x.fill = fill
    if num and isinstance(x.value, (int, float)):
        x.number_format = "#,##0"


def _sty_merge(ws, r1, c1, r2, c2) -> None:
    if r2 > r1 or c2 > c1:
        ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)


@dataclass
class _IdxStyle:
    """계정 목록(목차) 행 렌더 팔레트 — _build_index / _index_row 공유."""
    font: Font
    gray_font: Font
    link_font: Font
    red_bold: Font
    box: Border
    center: Alignment
    left: Alignment
    mkt_fill: PatternFill


def _rekey_block(mapping: dict, biz: str, old: str, new: str) -> dict:
    """블록 인덱스의 (biz, old[, …]) 키를 (biz, new[, …]) 로 이동한 새 dict — 상품명 리네임용.

    3튜플(_metric_row/_kw_row: (biz, product, metric/kw))·2튜플(_block_vids: (biz, product)) 모두 처리."""
    out = {}
    for k, v in mapping.items():
        out[(biz, new, *k[2:]) if (k[0] == biz and k[1] == old) else k] = v
    return out


def _sty_edge(ws, maxc, row, side, style) -> None:
    """상품 블록 경계(첫 행 상단/마지막 행 하단) 테두리 — 그룹 바깥=굵은선(thick)·변형 사이=얇은선(thin)."""
    for c in range(1, maxc + 1):
        b = ws.cell(row, c).border
        ws.cell(row, c).border = Border(
            left=b.left, right=b.right,
            top=style if side == "top" else b.top,
            bottom=style if side == "bottom" else b.bottom)

_COL_KIND = 1      # A: (레이아웃 v4) 헤더 라벨 칸 / 옛 구분 라벨(마이그레이션 회수용)
_COL_KW = 1        # A: 키워드명(레이아웃 v4 좌측확장 A~E, 병합 앵커=A). 상품 헤더 라벨과 행 종류로 구분(G값)
_COL_NAME = 3      # C: 상품명(헤더행 KEY) / 헤더 값 칸(pos0 C:F)
_COL_SEARCH = 6    # F: 검색량
_COL_METRIC = 7    # G: 지표 라벨
_FIRST_DATE = 8    # H~: 일자
_LABEL_DATE = "날짜"
_NOT_SELLING_STATUSES = ("판매중지", "임시저장", "승인반려", "검토중")   # ③ 순위 조회 제외 대상(판매중/부분판매중만)
_LABEL_KEYWORD = "키워드"
_LABEL_SEARCH = "검색량"
_LABEL_NOTE = "비고"
_ALL_METRICS = frozenset(config.CONTRACT_METRICS + config.PERSONAL_METRICS)
_META_SHEET = "_상품ID"   # 숨김 시트: (사업자,상품)→등록상품명·판매상태·제목캐시·**계정ID(col12, 항목5 상품별)**. ⚠vid(3열)는 폐지—vid 출처=헤더 이름칸(A)
_INDEX_SHEET = "계정 목록"  # 첫 시트: 전 계정(사업자) 목록 + 하이퍼링크 점프 + 요약(계정 100개도 탐색 쉽게)
_ACCT_SHEET = "_계정정보"  # 숨김 시트: (사업자)→계정ID **집합**(항목5: 한 사업자 다계정ID 허용, ' / ' 조인). 목차 표시용(⚠ 비밀번호 절대 저장 안 함)
_STAMP_SHEET = "_수집스탬프"  # 숨김 시트: (계정ID)→판매수집일. **계정 단위**(항목5: 다계정ID 사업자에서 계정마다 따로 수집 판정, 소유자 2026-09-25)
_MKT_SHEET = "_마케팅"     # 숨김 시트: (사업자,상품)→마케팅 시작·종료·모니터링종료. 계정목록 입력을 보존(재생성돼도 유지)
_DISC_SHEET = "_중단"      # 숨김 시트: (사업자,상품) 판매중지/삭제(대장에서 사라짐) 표기. 데이터는 보존, 표시만 구분
_SPECIAL_SHEETS = (_META_SHEET, _INDEX_SHEET, _ACCT_SHEET, _STAMP_SHEET, _MKT_SHEET, _DISC_SHEET)
_MKT_COLS = ("체험단 시작일", "체험단 종료일", "모니터링 종료일")   # 계정목록 편집 열(직원 입력 = 체험단 기간)
_MKT_COLS_LEGACY0 = "마케팅 시작일"   # 옛 라벨('마케팅 시작일') — 기존 마스터 계정목록에서 값 회수 시 인식용


def _norm(v) -> str:
    return str(v).strip() if v is not None else ""


def _key(v) -> str:
    """이름칸 셀 값 → **상품 키**(순수 상품명). 표시용으로 붙은 `⟨SEP⟩<vendorItemId>` 꼬리를 떼어낸다.

    상품 정체성(시계열 키)은 항상 구분자(`config.NAME_ID_SEP`) 앞부분이다. 구분자가 없으면(순수 이름·
    키워드 셀) `_norm` 과 동일하게 동작해 기존 키를 그대로 보존한다. 상품명 내부 개행은 손대지 않는다."""
    return _norm(v).split(config.NAME_ID_SEP, 1)[0]


def _vids_from_cell(v) -> list[str]:
    """이름칸 표시값의 **'VID : a / b' 꼬리**에서 vid 목록을 파싱한다(없으면 빈 리스트).

    vid 의 유일 출처(source of truth) = 헤더 C셀 표시값(소유자 확정 (A)). `_display_name` 이 렌더한
    `"{이름}{SEP}\nVID : a / b"` 를 역파싱한다 — 구분자 뒤 → 콜론 뒤 → '/' 분리. 구분자 없으면(순수
    이름·키워드 셀) vid 없음."""
    s = _norm(v)
    if config.NAME_ID_SEP not in s:
        return []
    tail = s.split(config.NAME_ID_SEP, 1)[1]
    # 꼬리는 여러 줄일 수 있다(VID / 상품판매가 / 로켓그로스 입고일). **'VID :' 줄만** 파싱해야 vid 오염 방지
    # (헤더에 다른 줄을 추가해도 vid 출처가 안 깨짐 — 소유자 2026-09-24 헤더 확장 대비 안전화).
    vid_line = ""
    for line in tail.splitlines():
        if "VID" in line.upper() and ":" in line:
            vid_line = line.split(":", 1)[1]
            break
    if not vid_line:                      # 옛 형식(줄 구분 없이 콜론 하나) 호환
        vid_line = tail.split(":", 1)[1] if ":" in tail else tail
    return [x.strip() for x in vid_line.split("/") if x.strip()]


def _nearest_year(month: int, day: int):
    """년도 없는 '월.일' → **오늘과 가장 가까운 연도**의 date(연말/연초 경계 보정).

    1월에 만난 '12.30'을 올해 12월로 오인하면 정규화가 first~last(1월~12월) 사이 ~363칸을 만들어
    폭발한다. 오늘 기준 작년/올해/내년 중 |날짜-오늘|이 최소인 해를 고른다(2/29 등 그 해에 없는 날은 건너뜀)."""
    today = _date.today()
    best = None
    for y in (today.year - 1, today.year, today.year + 1):
        try:
            cand = _date(y, month, day)
        except ValueError:   # 그 해에 없는 날(예: 평년 2/29)
            continue
        if best is None or abs((cand - today).days) < abs((best - today).days):
            best = cand
    return best


def _parse_date(s):
    """날짜 문자열 → date(못 읽으면 None). YYYY-MM-DD·YY.MM.DD·**MM.DD(년도 없음)**·YYYY/MM/DD·MM/DD 등 허용.

    ⚠ 일자 컬럼 라벨은 2026-09-16부터 **년도 없는 '월.일'(예 09.16)** 로 적는다(사용자 요청·당분간).
    '월.일'은 **오늘과 가장 가까운 연도**로 해석한다(`_nearest_year` — 연말/연초 경계 보정)."""
    s = _norm(s)
    if not s:
        return None
    if isinstance(s, (_dt, _date)):
        return s.date() if isinstance(s, _dt) else s
    no_year = ("%m/%d", "%m-%d", "%m.%d")   # 년도 없는 표기 → 오늘과 가장 가까운 해로 보정
    for fmt in ("%Y-%m-%d", "%Y.%m.%d", "%y.%m.%d", "%Y/%m/%d", "%y/%m/%d", "%m.%d", "%m/%d", "%m-%d"):
        try:
            d = _dt.strptime(s, fmt).date()
        except ValueError:
            continue
        if fmt in no_year:
            ny = _nearest_year(d.month, d.day)
            return ny if ny is not None else d
        return d
    return None


__all__ = [
    'dataclass',
    '_date',
    '_dt',
    '_td',
    'Path',
    '_quote',
    'openpyxl',
    'Cell',
    'Alignment',
    'Border',
    'Font',
    'PatternFill',
    'Side',
    'get_column_letter',
    'Hyperlink',
    'config',
    '_unmerge_all',
    '_StyleCtx',
    '_sty_cell',
    '_sty_merge',
    '_IdxStyle',
    '_rekey_block',
    '_sty_edge',
    '_COL_KIND',
    '_COL_KW',
    '_COL_NAME',
    '_COL_SEARCH',
    '_COL_METRIC',
    '_FIRST_DATE',
    '_LABEL_DATE',
    '_NOT_SELLING_STATUSES',
    '_LABEL_KEYWORD',
    '_LABEL_SEARCH',
    '_LABEL_NOTE',
    '_ALL_METRICS',
    '_META_SHEET',
    '_INDEX_SHEET',
    '_ACCT_SHEET',
    '_STAMP_SHEET',
    '_MKT_SHEET',
    '_DISC_SHEET',
    '_SPECIAL_SHEETS',
    '_MKT_COLS',
    '_MKT_COLS_LEGACY0',
    '_norm',
    '_key',
    '_vids_from_cell',
    '_nearest_year',
    '_parse_date',
]
