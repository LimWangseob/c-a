"""셀독전환자 계정 파일 — 미정산 계약자 중 셀독으로 전환한 계정(①셀독진행·②계정수령_상품미등록)의 정산 로그인 입력
(소유자 2026-10-22·D-111). 정산 계정(원장·D-103)과 출처가 달라 따로 읽는다.

파일 모양(실측 `미정산_셀독전환자_3유형분류_20261022.xlsx`): 시트 '①셀독진행'·'②계정수령_상품미등록'(이름이 ①·②로
시작하는 시트) · 1행 제목 · 2행 머리줄 · 3행부터 한 줄 = 한 계정. 머리줄 이름으로 칸을 찾는다(대표자명·연락처·사업자명·
계정명·비밀번호·상품명 — 열 위치가 바뀌어도 됨). '③계정미수령'·'전체 리스트'는 대상 아님.
- 계정명 빈 줄 = 건너뜀(경고) · 쿠팡 아이디 형식이 아니면 건너뜀(경고) · 같은 계정명이 두 줄이면 오류(어느 비번이 맞는지
  추측하지 않음).
- 비밀번호 앞뒤 공백 = 떼고 씀 + 경고(D-018). 경고·오류 글에는 **비밀번호 값을 넣지 않는다**(행 번호·계정명만).
로그인 비밀번호 = 이 파일 값 먼저 → 이 PC 에 저장된 관리대장 값(credstore) — 파일 칸이 비면 관리대장 값(소유자
2026-10-22 '관리대장 비번 적용'). 두 값이 다르면 최대 2개(D-012 상한·그날 거부된 값은 D-068 이 다시 넣지 않음).
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

DEFAULT_PATH = r"D:\토탈셀러\미정산중 셀독전환자\미정산_셀독전환자_3유형분류_20261022.xlsx"
CONFIG_KEY = "transfer/file"
SHEET_PREFIXES = ("①", "②")         # ①셀독진행 · ②계정수령_상품미등록
HEAD_ROW = 2                          # 1행 제목 · 2행 머리줄
COLS = {"rep": "대표자명", "phone": "연락처", "biz": "사업자명", "aid": "계정명", "pw": "비밀번호", "products": "상품명"}


class TransferFileError(Exception):
    """셀독전환자 파일을 쓸 수 없음(없음·시트·머리줄·중복) — 실행 중단."""


@dataclass
class TransferAccount:
    account_id: str
    label: str                        # 사업자명(없으면 대표자명 → 계정명) — 로그·결과 표시
    representative: str
    business_name: str
    phone: str
    kind: str                         # 시트 이름(①셀독진행 / ②계정수령_상품미등록)
    row: int                          # 시트 행 번호(엑셀 기준)
    products: tuple = ()
    password: str = field(default="", repr=False)   # repr 에 안 나오게(로그 실수 방지)


def _text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v)


def _columns(kind: str, rows: list) -> dict:
    """머리줄(HEAD_ROW 행) 이름 → {칸 키: 열 번호}. 필요한 칸이 없으면 TransferFileError."""
    if len(rows) < HEAD_ROW:
        raise TransferFileError(f"{kind}: 머리줄({HEAD_ROW}행)이 없음")
    head = [_text(h).strip() for h in rows[HEAD_ROW - 1]]
    miss = [name for name in COLS.values() if name not in head]
    if miss:
        raise TransferFileError(f"{kind}: 머리줄({HEAD_ROW}행)에 {', '.join(miss)} 칸이 없음")
    return {k: head.index(name) for k, name in COLS.items()}


def _row_account(kind: str, n: int, cell: dict, warns: list) -> TransferAccount | None:
    """한 줄 → 계정(건너뛸 줄이면 None·경고는 warns 에 — 비밀번호 값은 넣지 않음)."""
    from .registry_model import valid_account_id
    aid = cell["aid"].strip()
    who = cell["biz"].strip() or cell["rep"].strip()
    if not aid:
        warns.append(f"{kind} {n}행({who or '이름 없음'}): 계정명 칸이 비어 건너뜀")
        return None
    if not valid_account_id(aid):
        warns.append(f"{kind} {n}행({who or aid}): 계정명 '{aid}' 이 쿠팡 아이디 형식이 아니라 건너뜀")
        return None
    pw = cell["pw"]
    if pw != pw.strip():
        warns.append(f"{kind} {n}행({aid}): 비밀번호 앞뒤 공백을 떼고 씀(D-018) — 파일 값 확인 권장")
        pw = pw.strip()
    products = tuple(x.strip() for x in cell["products"].splitlines() if x.strip())
    return TransferAccount(aid, who or aid, cell["rep"].strip(), cell["biz"].strip(), cell["phone"].strip(),
                           kind, n, products, pw)


def parse_sheet(kind: str, rows: list) -> tuple[list[TransferAccount], list[str]]:
    """시트 값 격자(1행부터) → (계정 목록, 경고). 머리줄에 필요한 칸이 없으면 TransferFileError."""
    at = _columns(kind, rows)
    out: list[TransferAccount] = []
    warns: list[str] = []
    for n, r in enumerate(rows[HEAD_ROW:], start=HEAD_ROW + 1):
        if not any(_text(v).strip() for v in r):
            continue
        a = _row_account(kind, n, {k: _text(r[i]) if i < len(r) else "" for k, i in at.items()}, warns)
        if a is not None:
            out.append(a)
    return out, warns


def read_transfer_file(path) -> tuple[list[TransferAccount], list[str]]:
    """셀독전환자 파일 → (①·② 계정 목록[시트 순서], 경고). 계정명 중복(시트 사이 포함) = TransferFileError."""
    import openpyxl
    p = Path(path)
    if not p.exists():
        raise TransferFileError(f"셀독전환자 파일이 없음: {p}")
    try:
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    except Exception as exc:   # 열 수 없는 파일(엑셀로 열려 잠김·손상) — 이유를 그대로 올림
        raise TransferFileError(f"{p.name} 을 열 수 없음 — {exc.__class__.__name__}: {str(exc)[:100]}") from exc
    try:
        names = []
        for pre in SHEET_PREFIXES:                     # 접두마다 첫 시트
            found = [s for s in wb.sheetnames if s.startswith(pre)]
            names += found[:1]
        if len(names) != len(SHEET_PREFIXES):
            raise TransferFileError(f"{p.name}: '①…'·'②…' 시트가 모두 있어야 함(지금 시트: {', '.join(wb.sheetnames)})")
        accts: list[TransferAccount] = []
        warns: list[str] = []
        for name in names:
            got, w = parse_sheet(name, [list(r) for r in wb[name].iter_rows(values_only=True)])
            accts += got
            warns += w
    finally:
        wb.close()
    seen: dict[str, TransferAccount] = {}
    for a in accts:
        if a.account_id in seen:
            b = seen[a.account_id]
            raise TransferFileError(f"계정명 {a.account_id} 가 {b.kind} {b.row}행·{a.kind} {a.row}행에 두 번 있음 — "
                                    "어느 줄이 맞는지 정해 주세요")
        seen[a.account_id] = a
    if not accts:
        raise TransferFileError(f"{p.name}: ①·② 시트에 계정이 한 줄도 없음")
    return accts, warns


def login_passwords(a: TransferAccount, store=None) -> tuple[list[str], str]:
    """로그인에 넣을 비밀번호 후보·출처 설명(값은 설명에 안 씀) — 파일 값 먼저 → 관리대장(credstore) 값, 최대 2개."""
    if store is None:
        from .credstore import CredStore
        store = CredStore()
    ledger = store.get_password(a.account_id) or ""
    pws = list(dict.fromkeys(p for p in (a.password, ledger) if p))
    if not pws:
        return [], "비밀번호 없음(파일·관리대장 모두)"
    if not a.password:
        return pws, "파일 칸 비어 관리대장 값"
    if ledger and ledger != a.password:
        return pws, "파일·관리대장 값이 다름 — 파일 먼저(거부되면 관리대장 값 1회)"
    return pws, "파일"
