"""판매분석 파일 파싱·불변식 검증 (POLICY `PARSE_COUPANG_XLSX`). SSOT=designs/SETTLEMENT_MODULE.md §5.

실측(2026-09-29): 판매분석(SELLER_INSIGHTS_VENDOR_ITEM_METRICS) = 시트 'vendor item metrics' 1개·19열·**전 셀 문자열**
(숫자·'7.55%'·아이템위너 '1.10000000000000008881'). 파일 안에 기간 정보가 없다.
- 헤더 이름으로 열을 찾는다(위치 의존 금지). 필수 헤더 누락 = SettlementParseError(누락 목록).
- 숫자 변환 실패 = SettlementParseError(행·열 명시, **조용히 0 으로 두지 않음**). 옵션ID·등록상품ID = 문자열 유지.
- 적재 메타(기간 시작·끝·계정아이디) **필수** — 없으면 적재 거부.
- 불변식: 매출 = 총매출 + 총취소금액 · 판매량 = 총판매수 + 총취소상품수 · 구매전환율 = round(주문/조회×100, 2)
  (방문자 아님·조회>0). 위반 = **경고 + 행 로그**(쿠팡 포맷 변경 감지용·중단 아님). 이전 기간 취소로 음수는 정상.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path

INSIGHTS_SHEET = "vendor item metrics"
_EXCEL_EPOCH = date(1899, 12, 30)

# 필드 → 파일 헤더(실측 19열). 헤더 대조는 공백 무시.
_ID_COLS = {"option_id": "옵션 ID", "product_id": "등록상품ID"}
_TEXT_COLS = {"option_name": "옵션명", "product_name": "상품명", "category": "카테고리", "channel": "판매방식"}
_INT_COLS = {"sales": "매출(원)", "orders": "주문", "units": "판매량", "visitors": "방문자", "views": "조회",
             "cart": "장바구니", "total_sales": "총 매출(원)", "total_units": "총 판매수",
             "cancel_amount": "총 취소 금액(원)", "cancel_units": "총 취소된 상품수",
             "instant_cancel_units": "즉시 취소된 상품수"}
_PCT_COLS = {"cvr_pct": "구매전환율", "item_winner_pct": "아이템위너 비율(%)"}
REQUIRED_HEADERS = tuple({**_ID_COLS, **_TEXT_COLS, **_INT_COLS, **_PCT_COLS}.values())


class SettlementParseError(Exception):
    """파일·값·메타가 규칙에 맞지 않음 — 적재 중단(조용한 보정 없음)."""


@dataclass(frozen=True)
class InsightsMeta:
    period_start: date
    period_end: date
    account_id: str


@dataclass
class InsightRow:
    option_id: str
    product_id: str
    option_name: str
    product_name: str
    category: str
    channel: str
    sales: int
    orders: int
    units: int
    visitors: int
    views: int
    cart: int
    total_sales: int
    total_units: int
    cancel_amount: int
    cancel_units: int
    instant_cancel_units: int
    cvr_pct: Decimal
    item_winner_pct: Decimal


@dataclass
class InsightsFile:
    meta: InsightsMeta
    rows: list = field(default_factory=list)
    warnings: list = field(default_factory=list)


# ── 값 변환 ───────────────────────────────────────────────────────
def _s(v) -> str:
    return "" if v is None else str(v).strip()


def parse_int(v, where: str = "") -> int:
    """'1,234'·'-35000'·'  12 ' → int. 빈칸·글자·소수 = SettlementParseError."""
    s = _s(v).replace(",", "").replace(" ", "")
    try:
        d = Decimal(s)
    except InvalidOperation as exc:
        raise SettlementParseError(f"{where} 정수 변환 실패: {v!r}") from exc
    if d != d.to_integral_value():
        raise SettlementParseError(f"{where} 정수가 아님: {v!r}")
    return int(d)


def parse_pct(v, where: str = "") -> Decimal:
    """'7.55%'·'7.55'·'1.10000000000000008881' → Decimal 소수 2자리(사사오입). 실패 = SettlementParseError."""
    s = _s(v).replace(",", "").rstrip("%").strip()
    try:
        return Decimal(s).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    except InvalidOperation as exc:
        raise SettlementParseError(f"{where} 비율 변환 실패: {v!r}") from exc


def parse_id(v, where: str = "") -> str:
    """옵션ID·등록상품ID — 숫자 문자열 그대로(float 변환 금지). 비었거나 숫자가 아니면 SettlementParseError."""
    s = _s(v)
    if not s.isdigit():
        raise SettlementParseError(f"{where} ID 는 숫자 문자열이어야 함: {v!r}")
    return s


def parse_date(v, where: str = "") -> date:
    """date/datetime·엑셀 날짜 시리얼(1899-12-30 기준)·'YYYY-MM-DD'/'YYYY.MM.DD'/'YYYY/MM/DD' → date."""
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return _EXCEL_EPOCH + timedelta(days=int(v))
    s = _s(v).replace(".", "-").replace("/", "-").rstrip("-")
    try:
        return date.fromisoformat("-".join(p.zfill(2) for p in s.split("-")))
    except ValueError as exc:
        raise SettlementParseError(f"{where} 날짜 변환 실패: {v!r}") from exc


# ── 메타·헤더 ─────────────────────────────────────────────────────
def make_meta(period_start, period_end, account_id) -> InsightsMeta:
    """적재 메타 강제 — 셋 중 하나라도 없거나 기간이 뒤집히면 SettlementParseError(적재 거부)."""
    if not period_start or not period_end or not _s(account_id):
        raise SettlementParseError("판매분석 파일엔 기간 정보가 없어 period_start·period_end·account_id 메타가 "
                                   "필수 — 메타 없는 파일은 적재 거부")
    ps, pe = parse_date(period_start, "period_start"), parse_date(period_end, "period_end")
    if ps > pe:
        raise SettlementParseError(f"기간이 뒤집힘: {ps} > {pe}")
    return InsightsMeta(ps, pe, _s(account_id))


def _header_index(header: list) -> dict:
    nz = ["".join(_s(h).split()) for h in header]
    missing = [h for h in REQUIRED_HEADERS if "".join(h.split()) not in nz]
    if missing:
        raise SettlementParseError(f"판매분석 필수 헤더 누락 {len(missing)}개: {', '.join(missing)}")
    return {h: nz.index("".join(h.split())) for h in REQUIRED_HEADERS}


def _row(raw: list, idx: dict, r: int) -> InsightRow:
    def cell(h):
        i = idx[h]
        return raw[i] if i < len(raw) else None
    vals: dict = {}
    for f, h in _ID_COLS.items():
        vals[f] = parse_id(cell(h), f"{r}행 '{h}'")
    for f, h in _TEXT_COLS.items():
        vals[f] = _s(cell(h))
    for f, h in _INT_COLS.items():
        vals[f] = parse_int(cell(h), f"{r}행 '{h}'")
    for f, h in _PCT_COLS.items():
        vals[f] = parse_pct(cell(h), f"{r}행 '{h}'")
    return InsightRow(**vals)


# ── 불변식 ────────────────────────────────────────────────────────
def check_invariants(row) -> list[str]:
    """행 불변식 위반 목록(빈 목록=정상). row 는 InsightRow 또는 같은 이름의 속성을 가진 객체.
    음수(이전 기간 취소)는 정상 — 위반이 아님."""
    bad = []
    if row.sales != row.total_sales + row.cancel_amount:
        bad.append(f"매출 {row.sales} ≠ 총매출 {row.total_sales} + 총취소금액 {row.cancel_amount}")
    if row.units != row.total_units + row.cancel_units:
        bad.append(f"판매량 {row.units} ≠ 총판매수 {row.total_units} + 총취소상품수 {row.cancel_units}")
    if row.views > 0:
        cvr = (Decimal(row.orders) * 100 / Decimal(row.views)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        if cvr != Decimal(str(row.cvr_pct)).quantize(Decimal("0.01")):
            bad.append(f"구매전환율 {row.cvr_pct} ≠ round(주문 {row.orders} / 조회 {row.views} × 100, 2) = {cvr}")
    return bad


# ── 적재 ──────────────────────────────────────────────────────────
def parse_insights_rows(rows: list, *, period_start, period_end, account_id, on_log=None) -> InsightsFile:
    """판매분석 값 격자(1행=헤더) → InsightsFile. 메타·헤더·값 오류 = SettlementParseError, 불변식 위반 = 경고."""
    log = on_log or (lambda m: None)
    meta = make_meta(period_start, period_end, account_id)
    if not rows:
        raise SettlementParseError("판매분석 파일이 비어 있음(헤더 없음)")
    idx = _header_index(list(rows[0]))
    out = InsightsFile(meta)
    for r, raw in enumerate(rows[1:], start=2):
        if not any(_s(c) for c in raw):
            continue
        row = _row(list(raw), idx, r)
        out.rows.append(row)
        for v in check_invariants(row):
            msg = f"[판매분석 불변식] {meta.account_id} {r}행 옵션 {row.option_id}: {v}"
            out.warnings.append(msg)
            log(f"  ⚠ {msg}")
    log(f"== [판매분석] {meta.account_id} {meta.period_start}~{meta.period_end} — {len(out.rows)}행 적재·"
        f"불변식 경고 {len(out.warnings)}건 ==")
    return out


def read_insights_xlsx(path, *, period_start, period_end, account_id, on_log=None) -> InsightsFile:
    """판매분석 xlsx 적재. 시트가 여러 개면 'vendor item metrics' 를 쓰고, 없으면 SettlementParseError."""
    import openpyxl                                   # 지연 import(순수 계산 모듈과 분리)
    p = Path(path)
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    try:
        if len(wb.sheetnames) == 1:
            ws = wb[wb.sheetnames[0]]
        elif INSIGHTS_SHEET in wb.sheetnames:
            ws = wb[INSIGHTS_SHEET]
        else:
            raise SettlementParseError(f"{p.name}: 시트 '{INSIGHTS_SHEET}' 없음 — 시트 {wb.sheetnames}")
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()
    return parse_insights_rows(rows, period_start=period_start, period_end=period_end, account_id=account_id,
                               on_log=on_log)
