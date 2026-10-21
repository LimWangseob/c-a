"""통합판매현황 상단 대시보드 그리기·차트(엑셀·구글) — SSOT=designs/SALES_SUMMARY.md §5-1(소유자 2026-10-21·D-094·D-098).

①요약 표 → ②월별 막대 + ③주간 선(반씩) → ④일별 막대 → ⑤계정 순위(이번 달 누계 가로 막대 + 5칸 표 + 판매 없음·미수집
한 줄씩). 집계 = sales_dashboard(순수). **모든 도식이 같은 왼쪽 끝·같은 폭 W(A~J열)** — 차트 위치·폭은 W 에 대한 비율로
저장하고 엑셀·구글이 각자 자기 열 픽셀로 계산한다(엑셀 = 열 너비 1 = 6pt 실측, 구글 = 7px+5 — gsheet_stats 와 같은 변환).
차트 자료 = 상세표 오른쪽 **숨긴 열** · 차트 설정(JSON) = 그 첫 칸(1행) → openpyxl 은 파일을 다시 열면 차트를 버리므로
OutputWorkbook.load 가 restore_charts 로 다시 붙이고, 구글 반영도 같은 설정으로 addChart. 열 고정 없음(행 1~2만).
"""
from __future__ import annotations

import json
import math

from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Font
from openpyxl.utils import column_index_from_string, get_column_letter

from .sales_dashboard import short_name
from .workbook_sales import _FILL_HEAD, _FILL_SEC, _FILL_SUB, _FILL_TOTAL, _FN, SALES_SHEET, _cell

SPEC_PREFIX = "차트설정:"
COL_W = {1: 28, 2: 46, 3: 16}   # 열 너비(글자) — A 계정 · B 상품 · C 구분 · 그 뒤 날짜 칸
DATE_W = 12                     # 날짜 칸(상세표)·대시보드 D~J 열 너비 — 모두 같음
DASH_COLS = 10                  # 대시보드 폭 W = A~J 열
XL_PT_PER_CHAR = 6              # 엑셀 열 너비 1 = 6pt(실측 2026-10-21: 28→168pt) · 1pt = 12,700 EMU
XL_ROW_PT = 16.5                # 엑셀 기본 행 높이(실측)
G_ROW_PX = 21                   # 구글 기본 행 높이
CHART_ROWS = 15                 # ②③ · ④ 차트 한 줄이 차지하는 행 수
GAP = 0.012                     # ②와 ③ 사이 간격(W 비율)
DATA_WIDTH = 13                 # 숨긴 차트 자료 열 수(월 3·빈칸·주 2·빈칸·일 3·빈칸·계정 2)
_GREEN, _RED, _GRAY = "1E8E3E", "C00000", "9AA7B6"
_C_DONE, _C_PART, _C_LINE, _C_DAY, _C_MISS, _C_ACC = "4472C4", "A9C4EB", "2F5597", "5B9BD5", "BFBFBF", "4472C4"
_LEFT_WRAP = Alignment(horizontal="left", vertical="center")


def col_width(c: int) -> int:
    return COL_W.get(c, DATE_W)


def write_dashboard(ws, d, r: int, data_col: int, reasons: dict | None = None) -> int:
    """대시보드를 r 행부터 그리고 다음 빈 행 반환. 차트 자료·설정 = data_col 부터 숨긴 열."""
    r = _kpis(ws, r, d) + 2
    r, specs = _charts(ws, r, d, data_col)
    r, rank_spec = _ranks(ws, r + 1, d, data_col + 11, reasons or {})
    specs = specs + rank_spec
    ws.cell(1, data_col, SPEC_PREFIX + json.dumps(specs, ensure_ascii=False))
    for c in range(data_col, data_col + DATA_WIDTH):
        ws.column_dimensions[get_column_letter(c)].hidden = True
    add_xl_charts(ws, specs)
    return r


# ── 표 ───────────────────────────────────────────────────────
def _band(ws, r, title):
    _cell(ws, r, 1, title, font=Font(name=_FN, size=12, bold=True), fill=_FILL_SEC)
    for c in range(2, DASH_COLS + 1):
        _cell(ws, r, c, None, fill=_FILL_SEC)


def _span(ws, r, c1, c2, v, **kw):
    """c1~c2 칸을 한 칸으로(병합) — 열 고정이 없어 어디서나 병합 가능(구글 400 없음)."""
    _cell(ws, r, c1, v, **kw)
    for c in range(c1 + 1, c2 + 1):
        _cell(ws, r, c, None, fill=kw.get("fill"), box=kw.get("box", False))
    if c2 > c1:
        ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)


def paren(names) -> str:
    """미수집 계정 이름 괄호 — '(유라이프, 반달컴퍼니)' · 없으면 ''."""
    return f"({', '.join(names)})" if names else ""


def pct_text(v: float | None) -> tuple[str, str | None]:
    """비율 → ('▲ 12.3%', 초록) · ('▼ 5.0%', 빨강) · 비교 불가 = ('—', 회색)."""
    if v is None:
        return "—", _GRAY
    if round(v * 1000) == 0:
        return "0.0%", None
    return (f"▲ {v * 100:.1f}%", _GREEN) if v > 0 else (f"▼ {-v * 100:.1f}%", _RED)


def _kpis(ws, r, d) -> int:
    _band(ws, r, "① 요약 — 관리 상품 기준(결과 사업자 시트의 상품·판매중지 포함 · 상세표 합계와 같음)")
    h = dict(bold=True, fill=_FILL_HEAD, box=True, center=True)
    for c, t in ((1, "구분"), (2, "판매금액(원)"), (3, "수량")):
        _cell(ws, r + 1, c, t, **h)
    _span(ws, r + 1, 4, 5, "증감", **h)
    _span(ws, r + 1, 6, DASH_COLS, "비교 기준", **h)
    r += 1
    for k in d.kpis:
        r += 1
        t = dict(fill=_FILL_TOTAL, box=True)
        _cell(ws, r, 1, f"{k.title} {k.span}", bold=True, **t)
        _cell(ws, r, 2, k.amount, bold=True, num=True, **t)
        _cell(ws, r, 3, k.qty, num=True, **t)
        text, color = pct_text(k.change) if k.basis else ("", None)
        _span(ws, r, 4, 5, text or None, bold=True, center=True, color=color, **t)
        miss = f"미수집 {k.missing}계정{paren(k.missing_names)} 제외" if k.missing else ""
        note = " · ".join(x for x in (k.basis, miss) if x)
        _span(ws, r, 6, DASH_COLS, note or None, color="595959", **t)
    r += 1
    _cell(ws, r, 1, "미수집 계정(전일)", bold=True, fill=_FILL_SUB, box=True)
    names = f" {paren(d.missing_today_names)}" if d.missing_today_names else ""
    _span(ws, r, 2, DASH_COLS, f"{d.missing_today} / {d.accounts_total} 계정{names}", bold=True, fill=_FILL_SUB,
          box=True, color=_RED if d.missing_today else None)
    return r


def _ranks(ws, r, d, rc, reasons) -> tuple[int, list]:
    """⑤ 이번 달 누계 가로 막대 + 표 5칸(순위·계정·이번 달 누계·지난주·전주 대비) — **전 계정**(소유자 2026-10-21·D-101):
    판매 순 → 판매 0 → 미수집. 순위 번호는 판매 있는 계정만('-' = 판매 0·미수집). 미수집 사유는 표 아래 한 줄."""
    _band(ws, r, "⑤ 계정 순위 — 이번 달 누계 판매금액 순(전 계정 · 판매 0·미수집 포함)")
    acc = d.accounts
    _table(ws, r + 1, rc, ("계정", "이번 달 누계"), [(a.label, a.amount) for a in acc])
    rows = max(6, math.ceil((60 + 26 * len(acc)) / G_ROW_PX) + 1)
    spec = [{"type": "hbar", "title": "⑤ 이번 달 누계 판매금액(전 계정)", "row": r + 1, "rows": rows, "x0": 0.0,
             "x1": 1.0, "cat": [r + 1, r + 1 + len(acc), rc], "series": [[r + 1, r + 1 + len(acc), rc + 1]],
             "colors": [_C_ACC]}] if acc else []
    r = r + 1 + (rows if acc else 0)
    h = dict(bold=True, fill=_FILL_HEAD, box=True, center=True)
    _cell(ws, r, 1, "순위", **h)
    _cell(ws, r, 2, "계정", **h)
    _span(ws, r, 3, 5, "이번 달 누계(원)", **h)
    _span(ws, r, 6, 7, "지난주(원)", **h)
    _span(ws, r, 8, DASH_COLS, "전주 대비", **h)
    for i, a in enumerate(acc, 1):
        r += 1
        f = dict(fill=_FILL_SUB if i % 2 == 0 else None, box=True)
        _cell(ws, r, 1, i if a.amount else "-", center=True, **f)
        _cell(ws, r, 2, a.label, **f)
        for c1, c2, v in ((3, 5, a.amount), (6, 7, a.week)):
            _span(ws, r, c1, c2, v if v is not None else "미수집", num=True,
                  color=None if v is not None else _GRAY, **f)
        text, color = pct_text(a.change)
        _span(ws, r, 8, DASH_COLS, text, center=True, color=color, **f)
    miss = [f"{short_name(a.label)}({reasons.get(a.aid) or '자료 없음'})" for a in acc if a.amount is None]
    if miss:
        r += 1
        _span(ws, r, 1, DASH_COLS, f"미수집 {len(miss)}계정: " + ", ".join(miss), color=_RED)
        ws.cell(r, 1).alignment = _LEFT_WRAP
    return r, spec


# ── 차트 자료·설정 ─────────────────────────────────────────────
def _val(p, n: int):
    """기간 값 — 모든 계정이 그 기간 자료 없음이면 None(빈 칸·막대 없음), 아니면 받은 계정 합."""
    return None if n and p.missing >= n else p.amount


def _charts(ws, r, d, data_col) -> tuple[int, list]:
    _band(ws, r, "② 월별 · ③ 주간(최근 8주·월~일) · ④ 일별(최근 31일) 판매금액(원) — 관리 상품")
    top, n = r + 1, d.accounts_total
    mc, wc, dc = data_col, data_col + 4, data_col + 7

    def split(v, second):   # 두 계열 중 하나에만 값(쌓기 막대 — 다른 쪽 0) · 자료 없음 = 둘 다 빈칸
        return (None, None) if v is None else ((0, v) if second else (v, 0))
    _table(ws, top, mc, ("월", "완료 월", "진행 중(누계)"),
           [(p.label, *split(_val(p, n), p.partial)) for p in d.months])
    _table(ws, top, wc, ("주(월요일)", "판매금액"), [(p.label, _val(p, n)) for p in d.weeks])
    _table(ws, top, dc, ("일", "전 계정 받음", "일부 계정 미수집"),
           [(p.label, *split(_val(p, n), p.missing > 0)) for p in d.days])
    nm, nw, nd = top + len(d.months), top + len(d.weeks), top + len(d.days)
    half = 0.5 - GAP / 2
    specs = [
        {"type": "bar", "title": "② 월별 판매금액", "row": top, "rows": CHART_ROWS, "x0": 0.0, "x1": half,
         "stacked": True, "cat": [top, nm, mc], "series": [[top, nm, mc + 1], [top, nm, mc + 2]],
         "colors": [_C_DONE, _C_PART]},
        {"type": "line", "title": "③ 주간 판매금액(최근 8주)", "row": top, "rows": CHART_ROWS, "x0": 1 - half,
         "x1": 1.0, "cat": [top, nw, wc], "series": [[top, nw, wc + 1]], "colors": [_C_LINE]},
        {"type": "bar", "title": "④ 일별 판매금액(최근 31일 · 빈칸 = 아직 못 받은 날)", "row": top + CHART_ROWS,
         "rows": CHART_ROWS, "x0": 0.0, "x1": 1.0, "stacked": True, "cat": [top, nd, dc],
         "series": [[top, nd, dc + 1], [top, nd, dc + 2]], "colors": [_C_DAY, _C_MISS], "skip": 2},
    ]
    return top + 2 * CHART_ROWS, specs


def _table(ws, r, c, heads, rows) -> None:
    for i, h in enumerate(heads):
        ws.cell(r, c + i, h)
    for j, row in enumerate(rows, 1):
        for i, v in enumerate(row):
            x = ws.cell(r + j, c + i, v)
            if i:
                x.number_format = "#,##0"


def read_specs(ws) -> list:
    """시트 1행에 저장한 차트 설정 — 없으면 빈 목록."""
    for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
        for v in row:
            if isinstance(v, str) and v.startswith(SPEC_PREFIX):
                return json.loads(v[len(SPEC_PREFIX):])
    return []


def place(units: list, x0: float, x1: float) -> tuple[int, float, float]:
    """W(= units 합)의 x0~x1 비율 구간 → (시작 열 0-기준, 그 열 안 오프셋, 폭) — 단위는 units 와 같음."""
    total = sum(units)
    left, c, acc = x0 * total, 0, 0.0
    while c < len(units) - 1 and acc + units[c] <= left + 1e-9:
        acc += units[c]
        c += 1
    return c, left - acc, (x1 - x0) * total


def xl_geometry(sp: dict) -> tuple[int, int, int, int, int]:
    """엑셀 (열, 행, 열 안 오프셋 EMU, 폭 EMU, 높이 EMU) — 열 너비 × 6pt · 1pt = 12,700 EMU."""
    c, off, w = place([col_width(i) * XL_PT_PER_CHAR for i in range(1, DASH_COLS + 1)], sp["x0"], sp["x1"])
    return c, sp["row"] - 1, int(off * 12700), int(w * 12700), int((sp["rows"] - 1) * XL_ROW_PT * 12700)


def g_geometry(sp: dict) -> tuple[int, int, int, int, int]:
    """구글 (열, 행, 오프셋 px, 폭 px, 높이 px) — 열 픽셀 = int(너비×7+5)(gsheet_stats 열 너비 변환과 같음)."""
    c, off, w = place([int(col_width(i) * 7 + 5) for i in range(1, DASH_COLS + 1)], sp["x0"], sp["x1"])
    return c, sp["row"] - 1, int(off), int(w), (sp["rows"] - 1) * G_ROW_PX


def add_xl_charts(ws, specs: list) -> None:
    for sp in specs:
        ch = LineChart() if sp["type"] == "line" else BarChart()
        ch.title = sp["title"]
        ch.title.overlay = False                                   # 제목이 그림 위로 겹치지 않게
        ch.varyColors = False                                      # 점·막대마다 다른 색 안 씀
        ch.visible_cells_only = False                              # 자료 열이 숨김이어도 그림
        for r1, r2, c in sp["series"]:
            ch.add_data(Reference(ws, min_col=c, min_row=r1, max_row=r2), titles_from_data=True)
        r1, r2, c = sp["cat"]
        ch.set_categories(Reference(ws, min_col=c, min_row=r1 + 1, max_row=r2))
        if sp.get("stacked"):
            ch.grouping, ch.overlap = "stacked", 100
        if sp["type"] == "hbar":                                   # 가로 막대 — 1위가 맨 위
            ch.type = "bar"
            ch.x_axis.scaling.orientation = "maxMin"
            ch.y_axis.crosses = "max"
        if len(sp["series"]) == 1:
            ch.legend = None
        else:
            ch.legend.position, ch.legend.overlay = "b", False
        ch.x_axis.delete = ch.y_axis.delete = False
        ch.y_axis.scaling.min = 0                                  # 값 축 = 0부터(자동 최소값은 차이를 부풀림)
        if sp.get("skip"):
            ch.x_axis.tickLblSkip = sp["skip"]                     # 일별 31칸 = 날짜 글자 하나 걸러
        _xl_colors(ch, sp)
        col, row, off, cx, cy = xl_geometry(sp)
        ch.width, ch.height = cx / 360000, cy / 360000             # EMU → cm(정보용 — 위치·크기는 anchor 가 정함)
        ch.anchor = OneCellAnchor(_from=AnchorMarker(col=col, colOff=off, row=row, rowOff=0),
                                  ext=XDRPositiveSize2D(cx, cy))
        ws.add_chart(ch)


def _xl_colors(ch, sp) -> None:
    for s, color in zip(ch.series, sp["colors"]):
        if sp["type"] != "line":
            s.graphicalProperties.solidFill = color
            continue
        s.graphicalProperties.line.solidFill = color
        s.smooth = False
        s.marker.symbol = "circle"
        s.marker.graphicalProperties.solidFill = color
        s.marker.graphicalProperties.line.solidFill = color


def restore_charts(xwb) -> None:
    """파일을 다시 열면(openpyxl) 차트가 빠짐 → 저장해 둔 설정으로 다시 붙임(OutputWorkbook.load)."""
    if SALES_SHEET not in xwb.sheetnames:
        return
    ws = xwb[SALES_SHEET]
    if not ws._charts:
        add_xl_charts(ws, read_specs(ws))


# ── 구글시트 ──────────────────────────────────────────────────
def gsheet_requests(ws, gid: int, old_chart_ids: list) -> tuple[list, list]:
    """(먼저 보낼 요청 = 옛 차트 삭제, 나중 요청 = 숨긴 열 + 차트 추가) — 미러(worksheet_to_requests)는 둘 다 못 옮김."""
    first = [{"deleteEmbeddedObject": {"objectId": i}} for i in old_chart_ids]
    later = _hidden_columns(ws, gid) + [{"addChart": {"chart": _g_chart(sp, gid)}} for sp in read_specs(ws)]
    return first, later


def _hidden_columns(ws, gid: int) -> list:
    def dim(s, e, hidden):
        return {"updateDimensionProperties": {
            "range": {"sheetId": gid, "dimension": "COLUMNS", "startIndex": s, "endIndex": e},
            "properties": {"hiddenByUser": hidden}, "fields": "hiddenByUser"}}
    reqs = [dim(0, ws.max_column, False)]                         # 지난 실행의 숨김 잔재 초기화
    for letter, d in ws.column_dimensions.items():                # 다시 연 파일은 열 묶음(min~max)으로 읽힘
        if d.hidden:
            lo = d.min or column_index_from_string(letter)        # 이번에 그린 시트 = min 없음(저장 때 채워짐)
            reqs.append(dim(lo - 1, min(d.max or lo, ws.max_column), True))
    return reqs


def _g_chart(sp: dict, gid: int) -> dict:
    from .gsheet_stats import _rgb_to_color

    def grid(r1, r2, c):
        return {"sheetId": gid, "startRowIndex": r1 - 1, "endRowIndex": r2, "startColumnIndex": c - 1,
                "endColumnIndex": c}
    value_axis = "BOTTOM_AXIS" if sp["type"] == "hbar" else "LEFT_AXIS"   # 가로 막대 = 값 축이 아래
    series = [{"series": {"sourceRange": {"sources": [grid(*s)]}}, "targetAxis": value_axis,
               "colorStyle": {"rgbColor": _rgb_to_color(color)}} for s, color in zip(sp["series"], sp["colors"])]
    basic = {"chartType": {"bar": "COLUMN", "line": "LINE", "hbar": "BAR"}[sp["type"]],
             "legendPosition": "BOTTOM_LEGEND" if len(series) > 1 else "NO_LEGEND", "headerCount": 1,
             "axis": [{"position": p, **({"viewWindowOptions": {"viewWindowMin": 0}} if p == value_axis else {})}
                      for p in ("BOTTOM_AXIS", "LEFT_AXIS")],                # 값 축 = 0부터
             "domains": [{"domain": {"sourceRange": {"sources": [grid(*sp["cat"])]}}}], "series": series}
    if sp.get("stacked"):
        basic["stackedType"] = "STACKED"
    col, row, off, w, h = g_geometry(sp)
    return {"spec": {"title": sp["title"], "hiddenDimensionStrategy": "SHOW_ALL", "basicChart": basic},
            "position": {"overlayPosition": {
                "anchorCell": {"sheetId": gid, "rowIndex": row, "columnIndex": col},
                "offsetXPixels": off, "widthPixels": w, "heightPixels": h}}}
