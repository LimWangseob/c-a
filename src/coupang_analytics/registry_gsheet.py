"""셀독등록원장 — 구글시트 입출력(시트 5개 생성·서식·읽기·쓰기)·로컬 백업·실행. SSOT=designs/LEDGER_REGISTRY.md.

순서(run_sync): 대장 읽기(실패=전체 중단) → 원장 읽기 → 무결성 점검(어긋나면 중단) → 로컬 백업 → 비교 → 저장.
원장 값은 RAW 로 쓴다(비밀번호 '0012'·'=abc' 가 숫자/수식으로 바뀌지 않게). 로그에는 비밀번호 값을 남기지 않는다.
"""
from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

from .registry import Registry, RegistryIntegrityError, SyncResult, check_integrity, sync
from .registry_core import _replay_status, _row_status
from .registry_model import (ALL_SHEETS, CONFIRM_CHOICES, COUPANG_CHECK_VALUES, GROWTH_FIELDS, HIST_HEADER,
                             HISTORY_SHEETS,
                             MAIN_HEADER, SHEET_MAIN, SHEET_SYNC, ST_ACTIVE, STOCK_FIELD, SYNC_HEADER, HistRow,
                             RegRow, parse_ledger)

_HEADER_BG = {"red": 0xB7 / 255, "green": 0xC9 / 255, "blue": 0xE8 / 255}
_GREY = {"red": 0.6, "green": 0.6, "blue": 0.6}
_HIST_FIELDS = {"번호": "no", "실행번호": "run_id", "일시": "ts", "효력일": "eff", "계정아이디": "account_id",
                "사업자명": "biz", "상품명": "product", "변동유형": "kind", "항목": "item", "이전값": "old",
                "변경값": "new", "출처": "source", "확인상태": "confirm", "확인자": "confirmer",
                "확인일": "confirm_date", "비고": "note"}
_ACCT_COLS = ("대표자명", "사업자명", "비밀번호", "계약금", "체험단주체")


def _headers() -> dict:
    return {SHEET_MAIN: MAIN_HEADER, **HIST_HEADER, SHEET_SYNC: SYNC_HEADER}


# ── 읽기 ─────────────────────────────────────────────────────────
def _read_all(client) -> dict:
    return {s: client.read_values(s) for s in ALL_SHEETS}


def _cellmap(header: list, row: list) -> dict:
    return {h: (row[i] if i < len(row) else "") for i, h in enumerate(header)}


def _reg_row(d: dict) -> RegRow:
    return RegRow(d["계정아이디"], d["상품명"], acct={c: d.get(c, "") for c in _ACCT_COLS},
                  growth={g: d.get(g, "") for g in GROWTH_FIELDS}, stock=d.get(STOCK_FIELD, ""),
                  registered=d.get("등록일", ""), changed=d.get("최종변경일", ""),
                  coupang=d.get("쿠팡확인", ""), coupang_date=d.get("쿠팡확인일", ""))


def _hist_row(sheet: str, d: dict) -> HistRow:
    kw = {attr: d.get(col, "") for col, attr in _HIST_FIELDS.items() if col != "번호"}
    try:
        no = int(str(d.get("번호", "")).strip())
    except ValueError as exc:
        raise RegistryIntegrityError(f"{sheet} 번호 칸이 숫자가 아님: '{d.get('번호')}'") from exc
    return HistRow(sheet=sheet, no=no, **kw)


def registry_from_values(raws: dict) -> Registry:
    """시트 값 격자들 → Registry. 헤더는 이름으로 찾는다(열 순서 변경에도 동작)."""
    reg = Registry()
    main = raws.get(SHEET_MAIN) or []
    for row in main[1:]:
        d = _cellmap(main[0], row)
        if d.get("계정아이디"):
            r = _reg_row(d)
            reg.rows[r.key] = r
    for sheet in HISTORY_SHEETS:
        grid = raws.get(sheet) or []
        for row in grid[1:]:
            if any(str(c).strip() for c in row):
                reg.history.append(_hist_row(sheet, _cellmap(grid[0], row)))
    return reg


def load_registry(client) -> Registry:
    return registry_from_values(_read_all(client))


# ── 쓰기 ─────────────────────────────────────────────────────────
def _main_grid(reg: Registry) -> list[list[str]]:
    st = _replay_status(reg.history)
    out = []
    for r in reg.rows.values():
        status, stop = _row_status(st, r.account_id, r.product)
        out.append([status, stop, r.acct.get("대표자명", ""), r.acct.get("사업자명", ""), r.account_id,
                    r.acct.get("비밀번호", ""), r.acct.get("계약금", ""), r.acct.get("체험단주체", ""), r.product,
                    r.stock, *[r.growth.get(g, "") for g in GROWTH_FIELDS], r.registered, r.changed,
                    r.coupang, r.coupang_date])
    out.sort(key=lambda x: (x[0] != ST_ACTIVE, x[3], x[4], x[8]))   # 관리중 먼저 → 사업자 → 계정 → 상품
    return [list(MAIN_HEADER), *out]


def _hist_line(h: HistRow) -> list[str]:
    vals = {**{col: getattr(h, attr) for col, attr in _HIST_FIELDS.items()}, "번호": str(h.no)}
    return [vals[c] for c in HIST_HEADER[h.sheet]]


def _sync_line(res: SyncResult, run_id: str, ts: str) -> list[str]:
    s = res.summary
    result = "정상" if not res.warnings else f"정상(경고 {len(res.warnings)}건)"
    return [run_id, ts, *[str(s[k]) for k in SYNC_HEADER[2:-1]], result]


def _insert_top(client, lines_by_sheet: dict) -> None:
    """이력·동기화기록 시트 맨 위(헤더 아래)에 새 줄을 끼워 넣는다 — 최신이 위, 기존 줄(사람 확인칸 포함) 보존."""
    reqs = [{"insertDimension": {"range": {"sheetId": client.sheet_id(s), "dimension": "ROWS",
                                           "startIndex": 1, "endIndex": 1 + len(lines)},
                                 "inheritFromBefore": False}}
            for s, lines in lines_by_sheet.items() if lines]
    if not reqs:
        return
    client.batch_update(reqs)
    for s, lines in lines_by_sheet.items():
        if lines:
            client.write_values(s, lines, start="A2", raw=True)


def save_registry(client, reg: Registry, res: SyncResult, now: datetime) -> None:
    client.write_values(SHEET_MAIN, _main_grid(reg), raw=True)
    lines = {s: [_hist_line(h) for h in sorted((e for e in res.events if e.sheet == s),
                                               key=lambda e: e.no, reverse=True)]
             for s in HISTORY_SHEETS}
    run_id = res.events[0].run_id if res.events else now.strftime("R%y%m%d-%H%M%S")
    lines[SHEET_SYNC] = [_sync_line(res, run_id, now.strftime("%Y-%m-%d %H:%M:%S"))]
    _insert_top(client, lines)


# ── 쿠팡확인(§4-1) ───────────────────────────────────────────────
_ISO_DATE = re.compile(r"\d{4}-\d{2}-\d{2}")


def _validate_checks(checks: dict) -> None:
    for key, val in checks.items():
        if not (isinstance(val, tuple) and len(val) == 2):
            raise ValueError(f"쿠팡확인 값은 (값, 확인일) 이어야 함: {key} → {val!r}")
        value, day = val
        if value not in COUPANG_CHECK_VALUES:
            raise ValueError(f"쿠팡확인 값 '{value}'(키 {key}) 은 허용값 아님 — {', '.join(COUPANG_CHECK_VALUES)}")
        if not (isinstance(day, str) and _ISO_DATE.fullmatch(day)):
            raise ValueError(f"쿠팡확인일 '{day}'(키 {key}) 은 YYYY-MM-DD 형식이 아님")


def _check_column(header: list[str]) -> int:
    """쿠팡확인 열 위치(헤더 이름으로 탐지). 필요한 열이 없거나 확인일이 바로 오른쪽이 아니면 무결성 오류."""
    missing = [h for h in ("계정아이디", "상품명", "쿠팡확인", "쿠팡확인일") if h not in header]
    if missing:
        raise RegistryIntegrityError(f"셀독원장 헤더에 열 없음: {missing}")
    i_chk = header.index("쿠팡확인")
    if header.index("쿠팡확인일") != i_chk + 1:
        raise RegistryIntegrityError("셀독원장 '쿠팡확인일' 열이 '쿠팡확인' 바로 오른쪽이 아님 — 헤더 확인 필요")
    return i_chk


def _split_checks(checks: dict) -> dict:
    """키 정규화 — (계정, 상품명) 줄 키는 상품명 공백 정규화, 계정 키는 그대로."""
    return {((k[0], " ".join(str(k[1]).split())) if isinstance(k, tuple) else k): v for k, v in checks.items()}


def _plan_checks(main: list, checks: dict, log) -> tuple[list[list[str]], int, int, int]:
    """셀독원장 값 격자 → (쿠팡확인·확인일 두 열의 새 값 격자, 쿠팡확인 열 위치, 갱신 줄 수, 원장에 없는 키 수).
    줄 키 (계정, 상품명) 이 계정 키보다 우선. 상품명은 공백 정규화로 대조."""
    header = [str(h) for h in main[0]]
    i_chk = _check_column(header)
    keyed = _split_checks(checks)
    grid: list[list[str]] = []
    changed, seen = 0, set[object]()
    for row in main[1:]:
        d = _cellmap(header, row)
        aid, prod = str(d["계정아이디"]), " ".join(str(d["상품명"]).split())
        hit = [k for k in ((aid, prod), aid) if k in keyed]            # 줄 키 먼저 = 우선
        seen.update(hit)
        cur = [str(d["쿠팡확인"]), str(d["쿠팡확인일"])]
        new = list(keyed[hit[0]]) if hit else cur
        changed += new != cur
        grid.append(new)
    unknown = [k for k in keyed if k not in seen]
    for k in unknown:
        log(f"  [원장] ⚠ 쿠팡확인 대상이 원장에 없음 — 건너뜀: {k}")
    return grid, i_chk, changed, len(unknown)


def write_coupang_check(client, checks: dict, *, dry_run: bool = False, on_log=None) -> int:
    """셀독원장의 **쿠팡확인·쿠팡확인일 두 열만** 갱신(§4-1). 이력엔 기록하지 않는다(그로스 재고와 같은 방식).

    checks = {계정아이디: (값, YYYY-MM-DD)} → 그 계정 모든 줄 / {(계정아이디, 상품명): (값, 날짜)} → 그 줄만(우선).
    값 6종·날짜 형식이 아니면 ValueError, 원장 시트 없으면 RegistryIntegrityError, 원장에 없는 키 = 경고+건너뜀.
    반환 = 값이 바뀌는 줄 수(dry_run 이면 쓰지 않고 수만 계산)."""
    log = on_log or (lambda m: None)
    _validate_checks(checks)
    if SHEET_MAIN not in client.sheet_titles():
        raise RegistryIntegrityError(f"원장 시트 '{SHEET_MAIN}' 없음 — 쿠팡확인 기록 불가")
    main = client.read_values(SHEET_MAIN)
    if not main:
        raise RegistryIntegrityError(f"원장 시트 '{SHEET_MAIN}' 가 비어 있음(헤더 없음)")
    grid, i_chk, changed, unknown = _plan_checks(main, checks, log)
    if changed and not dry_run:
        client.write_values(SHEET_MAIN, grid, start=f"{get_column_letter(i_chk + 1)}2", raw=True)
    log(f"== [원장] 쿠팡확인 {'미리보기(저장 안 함) ' if dry_run else ''}— 갱신 {changed}줄"
        + (f" · 원장에 없는 대상 {unknown}건" if unknown else "") + " ==")
    return changed


# ── 시트 생성·서식 ───────────────────────────────────────────────
def _format_requests(ids: dict) -> list[dict]:
    reqs: list[dict] = []
    for sheet, header in _headers().items():
        gid = ids[sheet]
        reqs.append({"updateSheetProperties": {"properties": {"sheetId": gid, "gridProperties": {"frozenRowCount": 1}},
                                               "fields": "gridProperties.frozenRowCount"}})
        reqs.append({"repeatCell": {
            "range": {"sheetId": gid, "startRowIndex": 0, "endRowIndex": 1, "startColumnIndex": 0,
                      "endColumnIndex": len(header)},
            "cell": {"userEnteredFormat": {"textFormat": {"bold": True}, "backgroundColor": _HEADER_BG}},
            "fields": "userEnteredFormat(textFormat,backgroundColor)"}})
        prot: dict = {"range": {"sheetId": gid}, "warningOnly": True,
                      "description": "앱이 관리하는 시트 — 직접 수정하지 마세요(확인상태 칸만 입력)"}
        if sheet in HIST_HEADER:
            c = header.index("확인상태")
            prot["unprotectedRanges"] = [{"sheetId": gid, "startRowIndex": 1, "startColumnIndex": c,
                                          "endColumnIndex": c + 3}]           # 확인상태·확인자·확인일
            reqs.append({"setDataValidation": {
                "range": {"sheetId": gid, "startRowIndex": 1, "startColumnIndex": c, "endColumnIndex": c + 1},
                "rule": {"condition": {"type": "ONE_OF_LIST",
                                       "values": [{"userEnteredValue": v} for v in CONFIRM_CHOICES]},
                         "strict": False, "showCustomUi": True}}})
        reqs.append({"addProtectedRange": {"protectedRange": prot}})
    main = ids[SHEET_MAIN]
    reqs.append({"setBasicFilter": {"filter": {"range": {"sheetId": main, "startRowIndex": 0,
                                                         "startColumnIndex": 0, "endColumnIndex": len(MAIN_HEADER)}}}})
    reqs.append({"addConditionalFormatRule": {"index": 0, "rule": {
        "ranges": [{"sheetId": main, "startRowIndex": 1, "startColumnIndex": 0, "endColumnIndex": len(MAIN_HEADER)}],
        "booleanRule": {"condition": {"type": "CUSTOM_FORMULA", "values": [{"userEnteredValue": '=$A2="관리중단"'}]},
                        "format": {"textFormat": {"foregroundColor": _GREY}}}}}})
    return reqs


def init_sheets(client, log=None) -> None:
    """시트 5개를 만들고(없는 것만) 헤더·서식·보호를 한 번 건다."""
    log = log or (lambda m: None)
    ids = client.ensure_sheets(list(ALL_SHEETS))
    for sheet, header in _headers().items():
        client.write_values(sheet, [list(header)], raw=True)
    client.batch_update(_format_requests(ids))
    log(f"== [원장] 시트 {len(ALL_SHEETS)}개 생성·서식 적용: {', '.join(ALL_SHEETS)} ==")


# ── 백업·실행 ────────────────────────────────────────────────────
def backup_local(raws: dict, backup_dir: str | Path, now: datetime, log) -> Path:
    """원장 파일 전체(시트 5개 값)를 로컬 xlsx 로 저장 — `원장_yymmdd_hhmmss.xlsx`. 비밀번호 포함(로컬 보관)."""
    d = Path(backup_dir)
    d.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet in ALL_SHEETS:
        ws = wb.create_sheet(sheet)
        for row in raws.get(sheet) or []:
            ws.append(list(row))
    path = d / f"원장_{now.strftime('%y%m%d_%H%M%S')}.xlsx"
    wb.save(path)
    log(f"== [원장] 백업 → {path} ==")
    return path


def _describe(h: HistRow) -> str:
    vals = "(값 생략)" if h.item == "비밀번호" else f"'{h.old}' → '{h.new}'" if (h.old or h.new) else ""
    what = f"{h.account_id}" + (f" / {h.product}" if h.product else "")
    return f"  [{h.sheet}] {h.kind} {what} {h.item} {vals}".rstrip()


def _load_existing(client) -> tuple[bool, dict, Registry]:
    """(원장 있음?, 시트 값들, Registry). 원장 시트 일부만 있으면 누가 지운 것 → 무결성 오류."""
    titles = client.sheet_titles()
    if SHEET_MAIN not in titles:
        return False, {}, Registry()
    missing = [s for s in ALL_SHEETS if s not in titles]
    if missing:
        raise RegistryIntegrityError(f"원장 시트 일부 누락: {missing}")
    raws = _read_all(client)
    return True, raws, registry_from_values(raws)


def _log_result(res: SyncResult, log, dry_run: bool) -> None:
    for w in res.warnings:
        log(f"  [원장] ⚠ {w}")
    for h in res.events:
        log(_describe(h))
    log("== [원장] " + ("미리보기(저장 안 함) " if dry_run else "") +
        " · ".join(f"{k} {v}" for k, v in res.summary.items()) + " ==")


def _replay_snapshots(snapshots, log) -> tuple[Registry, list]:
    """과거 대장 사본들을 **날짜순**으로 빈 원장에 재생. 급감 등 예외는 그대로 전파(아무것도 안 씀)."""
    reg, results, seen_warn = Registry(), [], set()
    for ts, label, read in sorted(snapshots, key=lambda s: s[0]):
        title, rows, strike = read()
        res = sync(reg, parse_ledger(rows, strike, title), now=ts)
        results.append((ts, label, res))
        for w in res.warnings:
            if w not in seen_warn:
                seen_warn.add(w)
                log(f"  [소급] ⚠ {w}")
        changes = sum(v for k, v in res.summary.items() if not k.startswith(("대장", "원장")))
        log(f"  [소급] {ts:%m/%d %H:%M} {label} — 변동 {changes}건")
    return reg, results


def run_backfill(client, snapshots: list, *, log=None, dry_run: bool = False,
                 backup_dir: str | Path | None = "output/백업") -> list:
    """과거 대장 사본으로 원장을 **처음부터** 소급 구축. snapshots = [(시각, 이름, read()→(시트명, 값, 취소선))].

    효력일 = 그 변동이 처음 보인 사본의 날짜. 동기화기록엔 사본마다 1줄(출처 사본 이름). 원장에 이미 데이터가 있으면
    거부(중복·이력 순서 꼬임 방지). 재생·무결성이 모두 통과해야 쓴다(중간 실패 = 아무것도 안 씀)."""
    log = log or (lambda m: None)
    exists, _, cur = _load_existing(client)
    if cur.rows or cur.history:
        raise RegistryIntegrityError("원장에 이미 데이터가 있어 소급할 수 없습니다 — 빈 원장 파일에서만 실행하세요")
    reg, results = _replay_snapshots(snapshots, log)
    check_integrity(reg)
    log(f"== [소급] 사본 {len(results)}개 재생 완료 — 원장 {len(reg.rows)}줄·이력 {len(reg.history)}줄"
        + (" (미리보기·저장 안 함)" if dry_run else "") + " ==")
    if dry_run:
        return [r for _, _, r in results]
    if not exists:
        init_sheets(client, log)
    client.write_values(SHEET_MAIN, _main_grid(reg), raw=True)
    lines = {s: [_hist_line(h) for h in sorted((h for h in reg.history if h.sheet == s),
                                               key=lambda h: h.no, reverse=True)]
             for s in HISTORY_SHEETS}
    sync_lines = []
    for ts, label, res in reversed(results):
        line = _sync_line(res, ts.strftime("R%y%m%d-%H%M%S"), ts.strftime("%Y-%m-%d %H:%M:%S"))
        line[-1] = f"{line[-1]} · 소급 출처 {label}"
        sync_lines.append(line)
    lines[SHEET_SYNC] = sync_lines
    _insert_top(client, lines)
    if backup_dir:
        backup_local(_read_all(client), backup_dir, datetime.now(), log)
    return [r for _, _, r in results]


def run_sync(client, read_ledger, *, now: datetime | None = None, log=None, dry_run: bool = False,
             backup_dir: str | Path | None = "output/백업") -> SyncResult:
    """대장 → 원장 동기화 1회. read_ledger() → (시트명, 값격자, 취소선격자). 실패는 예외로 전파(쓰기 없음)."""
    log = log or (lambda m: None)
    now = now or datetime.now()
    title, rows, strike = read_ledger()                       # 실패 = 여기서 중단(원장 미접촉)
    snap = parse_ledger(rows, strike, title)
    exists, raws, reg = _load_existing(client)
    check_integrity(reg)
    if backup_dir and exists and not dry_run:
        backup_local(raws, backup_dir, now, log)
    res = sync(reg, snap, now=now)
    _log_result(res, log, dry_run)
    if dry_run:
        return res
    if not exists:
        init_sheets(client, log)
    save_registry(client, reg, res, now)
    return res
