"""결과 파일 첫 시트 '통합판매현황' 그리기·구글시트 반영 — SSOT=designs/SALES_SUMMARY.md(소유자 2026-10-15).

위 = 대시보드(①요약·②월별·③주간·④일별 차트·⑤계정 순위·⑥상품 상위 20 열지도 — workbook_sales_dash·소유자 2026-10-21) →
아래 = 상세표 [일별 최근 31일·최신이 왼쪽] 전체 합계 → 계정 소계 → 상품 줄(윙/로켓그로스 나눔·판매 있는 상품만).
칸마다 수량·금액. 받지 못한 날 = '미수집'(0 과 다름). 집계 기준 = **관리 상품**(결과 사업자 시트 블록 VID·판매중지 포함).
매번 시트를 지우고 맨 앞(index 0)에 다시 만든다 — 사업자 시트가 아니라 특수 시트(_SPECIAL_SHEETS)라 미러링·삭제 정리·
색인 대상 아님. 원천·집계 = sales_history·sales_summary.
"""
from __future__ import annotations

from datetime import date, datetime, timedelta

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import sales_dashboard, sales_history, sales_summary
from .workbook_common import _SALES_SHEET

SALES_SHEET = _SALES_SHEET
_FN = "맑은 고딕"
_NUM = "#,##0"
_FILL_SEC = PatternFill("solid", fgColor="B7C9E8")
_FILL_HEAD = PatternFill("solid", fgColor="DCE6F1")
_FILL_TOTAL = PatternFill("solid", fgColor="FFF2CC")
_FILL_SUB = PatternFill("solid", fgColor="F2F2F2")
_THIN = Side(style="thin", color="BFBFBF")
_BOX = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_CENTER = Alignment(horizontal="center", vertical="center")
_MISSING = "미수집"


def account_labels(owb) -> dict:
    """{계정ID: '사업자 (계정ID)'} — 결과 파일 사업자 시트의 계정(관리대장 계정). 정산 계정은 3단계에서 더함."""
    out = {}
    for biz in owb.account_sheets():
        for aid in owb.account_ids_of(biz):
            out[aid] = f"{biz} ({aid})"
    return out


def managed_vids(owb) -> dict:
    """{계정ID: 관리 상품 VID 집합} — 결과 사업자 시트 블록(관리대장 상품·판매중지 블록 포함)의 VID."""
    out: dict = {}
    for biz in owb.account_sheets():
        vids = {v for p in owb.products_of(biz) for v in owb.product_vids(biz, p)}
        for aid in owb.account_ids_of(biz):
            out.setdefault(aid, set()).update(vids)
    return out


def refresh_sales_sheet(owb, out_dir="output", log=None, today: date | None = None) -> str:
    """저장된 판매분석 기록으로 첫 시트를 다시 그린다(저장은 호출부). 반환 = 로그 한 줄."""
    log = log or (lambda m: None)
    yesterday = (today or date.today()) - timedelta(days=1)
    try:
        daily, periods = sales_history.load_daily(out_dir), sales_history.load_periods(out_dir)
    except (OSError, ValueError) as exc:   # 기록 파일 손상 등 — 시트는 지난 것 유지(비치명·로그)
        msg = f"[통합판매현황] ⚠ 판매 기록 읽기 실패 — 지난 시트 유지 · {exc.__class__.__name__}: {str(exc)[:80]}"
        log(f"  {msg}")
        return msg
    managed, labels = managed_vids(owb), account_labels(owb)
    daily, periods = sales_summary.scope(daily, managed), sales_summary.scope(periods, managed)
    summary = sales_summary.build(daily, labels, yesterday)
    board = sales_dashboard.build(daily, periods, labels, yesterday)
    write_sales_sheet(owb.wb, summary, _login_reasons(out_dir, log), board)
    day = summary.sections[0]
    key = day.cols[0].key
    q, a = day.total[key]
    month = board.kpis[2]
    msg = (f"[통합판매현황] 갱신(관리 상품 기준) — 전일 {yesterday:%m.%d} {q:,}개·{a:,}원 · 이번 달 누계 "
           f"{month.amount:,}원 · 계정 {len(day.accounts)}개(전일 미수집 {day.missing[key]})")
    log(f"  {msg}")
    return msg


def _login_reasons(out_dir, log) -> dict:
    """{계정ID: 오늘 로그인 실패 사유}(login_today·D-068) — 못 읽으면 로그 후 사유 없이 그림(비치명)."""
    from . import login_today
    try:
        return login_today.reasons_today(out_dir=out_dir)
    except (OSError, ValueError) as exc:
        log(f"  [통합판매현황] ⚠ 로그인 실패 사유 읽기 실패 — 사유 없이 표시 · {exc.__class__.__name__}")
        return {}


def write_sales_sheet(xwb, summary, reasons: dict | None = None, board=None) -> None:
    """openpyxl 워크북 맨 앞에 통합판매현황을 새로 그림(있으면 지우고 다시). reasons = {계정ID: 로그인 실패 사유} —
    미수집 칸이 있는 계정 소계 줄에 '소계 — ⚠ 로그인 비밀번호 오류'처럼 간단히 적음(소유자 2026-10-16).
    board = 대시보드(sales_dashboard) — 있으면 상세표 위에 그림·차트 자료는 상세표 오른쪽 숨긴 열."""
    from . import workbook_sales_dash
    if SALES_SHEET in xwb.sheetnames:
        del xwb[SALES_SHEET]
    ws = xwb.create_sheet(SALES_SHEET, 0)
    ncol = 3 + 2 * max(len(s.cols) for s in summary.sections)
    _cell(ws, 1, 1, "통합 판매현황 — 위탁계정 계정·상품별 판매수량·판매금액(윙/로켓그로스 구분)",
          font=Font(name=_FN, size=14, bold=True))
    _cell(ws, 2, 1, f"기준 전일 {summary.yesterday.isoformat()} · 갱신 {datetime.now():%Y-%m-%d %H:%M} · "
                    "출처 쿠팡 판매분석(주문 기준·정산과 기준 다름) · 집계 = 관리 상품(결과 사업자 시트의 상품·판매중지 포함) · "
                    "'미수집' = 그 기간 자료를 아직 못 받음(0 과 다름)",
          font=Font(name=_FN, size=10, color="595959"))
    r = 4
    if board is not None:
        r = workbook_sales_dash.write_dashboard(ws, board, r, ncol + 3) + 3
    for sec in summary.sections:
        r = _write_section(ws, r, sec, reasons or {}) + 2
    widths = {1: 28, 2: 46, 3: 11}
    for c in range(1, ncol + 1):
        ws.column_dimensions[get_column_letter(c)].width = widths.get(c, 7 if c % 2 == 0 else 13)
    ws.freeze_panes = "D3"


def _write_section(ws, r: int, sec, reasons: dict) -> int:
    """구역 하나(제목·머리 2줄·전체 합계·계정 소계·상품 줄) — 마지막으로 쓴 행 번호 반환."""
    last = 3 + 2 * len(sec.cols)
    _cell(ws, r, 1, sec.title, font=Font(name=_FN, size=12, bold=True), fill=_FILL_SEC)
    # 제목 병합 = 고정 열 A:C 안에서만(freeze D3) — 구글시트는 고정·비고정 열에 걸친 병합을 400 으로 거부하고
    # 그 뒤 계정목록 동기화까지 막음(라이브 2026-10-15 "You can't merge frozen and non-frozen columns").
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    for c in range(4, last + 1):
        _cell(ws, r, c, None, fill=_FILL_SEC)   # 오른쪽은 같은 색만(병합 없이 한 띠로 보이게)
    h1, h2 = r + 1, r + 2
    for c, h in enumerate(("계정", "상품", "구분"), 1):
        _cell(ws, h1, c, h, bold=True, fill=_FILL_HEAD, box=True, center=True)
        _cell(ws, h2, c, None, fill=_FILL_HEAD, box=True)
        ws.merge_cells(start_row=h1, start_column=c, end_row=h2, end_column=c)
    for i, col in enumerate(sec.cols):
        c = 4 + 2 * i
        _cell(ws, h1, c, col.label, bold=True, fill=_FILL_HEAD, box=True, center=True)
        _cell(ws, h1, c + 1, None, fill=_FILL_HEAD, box=True)
        ws.merge_cells(start_row=h1, start_column=c, end_row=h1, end_column=c + 1)
        _cell(ws, h2, c, "수량", bold=True, fill=_FILL_HEAD, box=True, center=True)
        _cell(ws, h2, c + 1, "금액", bold=True, fill=_FILL_HEAD, box=True, center=True)
    r = h2 + 1
    _values_row(ws, r, ("전체 합계", "", ""), sec.cols, sec.total, bold=True, fill=_FILL_TOTAL)
    if any(sec.missing.values()):
        r += 1
        _cell(ws, r, 1, "미수집 계정 수", fill=_FILL_TOTAL, box=True)
        for c in (2, 3):
            _cell(ws, r, c, None, fill=_FILL_TOTAL, box=True)
        for i, col in enumerate(sec.cols):
            _cell(ws, r, 4 + 2 * i, sec.missing[col.key] or None, fill=_FILL_TOTAL, box=True)
            _cell(ws, r, 5 + 2 * i, None, fill=_FILL_TOTAL, box=True)
    for blk in sec.accounts:
        r += 1
        why = reasons.get(blk.aid) if any(v is None for v in blk.total.values()) else None
        _values_row(ws, r, (blk.label, f"소계 — ⚠ {why}" if why else "소계", ""), sec.cols, blk.total,
                    bold=True, fill=_FILL_SUB)
        for row in blk.rows:   # 상품 줄: 그 칸 기록 없음 = 계정이 받은 기간이면 0, 계정 미수집이면 미수집
            r += 1
            vals = {c.key: row.values.get(c.key, (0, 0) if blk.total.get(c.key) is not None else None)
                    for c in sec.cols}
            _values_row(ws, r, (blk.label, row.product, row.kind), sec.cols, vals, gray_first=True)
    return r


def _values_row(ws, r, heads, cols, values: dict, *, bold=False, fill=None, gray_first=False) -> None:
    for c, v in enumerate(heads, 1):
        color = "9AA7B6" if gray_first and c == 1 else None
        _cell(ws, r, c, v or None, bold=bold, fill=fill, box=True, color=color)
    for i, col in enumerate(cols):   # values = {열 키: (수량, 금액) | None(미수집)}
        v = values.get(col.key)
        c = 4 + 2 * i
        if v is None:
            _cell(ws, r, c, _MISSING, fill=fill, box=True, color="9AA7B6", center=True)
            _cell(ws, r, c + 1, None, fill=fill, box=True)
            continue
        _cell(ws, r, c, v[0], bold=bold, fill=fill, box=True, num=True)
        _cell(ws, r, c + 1, v[1], bold=bold, fill=fill, box=True, num=True)


def _cell(ws, r, c, v, *, font=None, bold=False, fill=None, box=False, center=False, color=None, num=False):
    x = ws.cell(r, c, v)
    x.font = font or Font(name=_FN, size=10, bold=bold, color=color)
    if fill:
        x.fill = fill
    if box:
        x.border = _BOX
    if center:
        x.alignment = _CENTER
    if num:
        x.number_format = _NUM
    return x


def push_sales_sheet(client, owb, log=None) -> None:
    """결과 구글시트에 통합판매현황 반영 — 없으면 만들고 값·서식 통째 교체 후 맨 앞(index 0)으로 옮김.
    차트 3개는 미러가 못 옮겨 옛 차트를 지우고 다시 추가(같은 숨긴 열 자료·hiddenDimensionStrategy SHOW_ALL)."""
    from .gsheet_stats import worksheet_to_requests
    if SALES_SHEET not in owb.wb.sheetnames:
        return
    ws = owb.wb[SALES_SHEET]
    from .workbook_sales_dash import gsheet_requests
    gid = client.ensure_sheet(SALES_SHEET)
    first, later = gsheet_requests(ws, gid, client.chart_ids(SALES_SHEET))   # 옛 차트 삭제 · 숨긴 열 · 차트
    reqs = first + worksheet_to_requests(ws, gid) + later
    reqs.append({"updateSheetProperties": {"properties": {"sheetId": gid, "index": 0}, "fields": "index"}})
    client.batch_update(reqs)
    if log:
        log(f"  [구글시트] 통합판매현황 반영({ws.max_row}행×{ws.max_column}열·맨 앞)")
