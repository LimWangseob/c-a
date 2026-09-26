"""판매분석 '상품별 판매 리포트'(VENDOR_ITEM) 파싱.

리포트는 옵션(행) 단위이며 기간 합계다. `parse_by_option` 이 옵션ID(=공개 vendorItemId)
기준으로 {옵션ID: OptionMetric} 을 만든다(합산 없이 옵션 단위).
"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import openpyxl

from . import config


def _norm(value) -> str:
    return str(value).strip() if value is not None else ""


def _to_int(value) -> int:
    if value is None or value == "":
        return 0
    return int(round(float(value)))


def _resolve_header(header: list[str]) -> dict[str, int]:
    required = [
        config.RP_COL_PRODUCT, config.RP_COL_ITEM_ID, config.RP_COL_OPTION_ID,
        config.RP_COL_VIEWS, config.RP_COL_SALES, config.RP_COL_VISITORS,
    ]
    idx: dict[str, int] = {}
    for name in required:
        if name not in header:
            raise ValueError(f"리포트 헤더에 '{name}' 컬럼이 없습니다. 실제 헤더: {header}")
        idx[name] = header.index(name)
    return idx


@dataclass
class OptionMetric:
    option_id: str        # 옵션ID (= 공개 vendorItemId, 확인됨)
    product_name: str     # 리포트 상품명(전체 제목)
    option_name: str      # 옵션명
    item_id: str          # 등록상품ID (상품 그룹 키)
    views: int            # 조회 → 노출건수
    sales: int            # 판매량 → 판매건수
    visitors: int         # 방문자 → 방문자건수
    registration_type: str = ""   # NORMAL(개인/판매자배송) / RFM(계약/로켓그로스) — 상품구분 판별
    product_id: str = ""  # 노출상품ID(productId) — 상품명 하이퍼링크(쿠팡 노출페이지)용. 판매분석 응답서 확보(항목2)


def parse_by_option(path: str | Path) -> dict[str, OptionMetric]:
    """상품별 리포트 → {옵션ID: OptionMetric} (합산 없이 옵션 단위).

    read_only 모드는 .close() 전까지 파일 핸들을 계속 잡고 있어, 안 닫으면 다음 계정에서
    같은 폴더의 리포트 삭제/교체가 PermissionError 로 막힌다(라이브 실측). → 반드시 close.
    """
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        rows = ws.iter_rows(values_only=True)
        header = [_norm(c) for c in next(rows)]
        i_opt = header.index(config.RP_COL_OPTION_ID)
        i_optname = header.index("옵션명") if "옵션명" in header else None
        idx = _resolve_header(header)
        result: dict[str, OptionMetric] = {}
        for row in rows:
            oid = _norm(row[i_opt])
            if not oid:
                continue
            result[oid] = OptionMetric(
                option_id=oid,
                product_name=_norm(row[idx[config.RP_COL_PRODUCT]]),
                option_name=_norm(row[i_optname]) if i_optname is not None else "",
                item_id=_norm(row[idx[config.RP_COL_ITEM_ID]]),
                views=_to_int(row[idx[config.RP_COL_VIEWS]]),
                sales=_to_int(row[idx[config.RP_COL_SALES]]),
                visitors=_to_int(row[idx[config.RP_COL_VISITORS]]),
            )
        return result
    finally:
        wb.close()      # read_only 핸들 해제 — 다음 계정 다운로드 폴더 잠금 방지
