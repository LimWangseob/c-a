"""셀독전환자 정산 1단계 — 로그인 확인(다운로드 없음·소유자 2026-10-22·D-111).

①셀독진행·②계정수령_상품미등록 계정(transfer_accounts)마다 앱 ①과 **같은 로그인 코드**(_login_with_candidates —
D-012 최대 2값·D-068 그날 거부된 값 재제출 금지·2단계 인증 화면 = 건너뜀·프록시 없음)로 로그인만 해 보고 창을 닫는다.
결과 = output/정산_셀독전환/ — `로그인확인_최신.json`(계정별 마지막 결과·비밀번호 없음) + 실행마다
`로그인확인_YYMMDD_HHMMSS.xlsx`(전체 결과 · '연락 필요' = 비밀번호 오류·2단계 인증·비밀번호 없음 등 사람이 연락해 고칠 계정).

안전 규칙(런타임 병렬 금지·사무실 IP 보호):
- 앱 파이프라인(output/_잠금/파이프라인.lock)·정산 자동 다운로드(output/정산/_잠금/_watch.lock)가 돌면 **시작·다음 계정 안 함**.
- until(기본 17:40) 이 지나면 새 계정을 시작하지 않음 — 18:00 무인 실행을 막지 않게.
- 계정 사이 45~75초 쉼(정산 요청 간격과 같음) · **로그인 차단 2계정 연속 = 전체 중단**(IP 차단 의심 — 다른 계정 보호).
"""
from __future__ import annotations

import json
import random
import time
from datetime import datetime
from pathlib import Path

OUT_DIR = Path("output") / "정산_셀독전환"
STATE_FILE = "로그인확인_최신.json"
GAP_SEC = (45, 75)
UNTIL_DEFAULT = "17:40"
BLOCK_STOP = 2
BUSY_LOCKS = (Path("output") / "_잠금" / "파이프라인.lock", Path("output") / "정산" / "_잠금" / "_watch.lock")

OK, FAIL, BLOCK, SKIP = "성공", "실패", "차단", "건너뜀"
CONTACT = "연락 필요"          # 사람이 계약자에게 연락해 고칠 실패(비밀번호·2단계 인증·계정 상태·비밀번호 없음)
HEAD = ("유형", "대표자명", "사업자명", "계정명", "연락처", "결과", "사유", "비밀번호 출처", "시도 일시")
CONTACT_HEAD = ("대표자명", "사업자명", "계정명", "연락처", "유형", "실패 사유", "마지막 시도 일시")


# ── 결과 기록 ─────────────────────────────────────────────────
def load_state(out_dir=OUT_DIR) -> dict:
    p = Path(out_dir) / STATE_FILE
    return json.loads(p.read_text(encoding="utf-8")) if p.exists() else {}


def save_result(out_dir, aid: str, status: str, reason: str, src: str, at: str) -> None:
    """계정 하나 끝날 때마다 바로 저장(중간에 멈춰도 남음). 비밀번호 값·연락처는 넣지 않음."""
    p = Path(out_dir) / STATE_FILE
    p.parent.mkdir(parents=True, exist_ok=True)
    data = load_state(out_dir)
    data[aid] = {"결과": status, "사유": reason, "비밀번호 출처": src, "시도": at}
    tmp = p.with_suffix(".tmp")
    tmp.write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
    tmp.replace(p)


def needs_contact(rec: dict) -> bool:
    return rec.get("결과") == FAIL


def write_report(accts, state: dict, out_dir=OUT_DIR, now: datetime | None = None) -> Path:
    """실행 결과 엑셀 — '로그인 결과'(파일 순서·전 계정) + '연락 필요'(실패만). 아직 안 해 본 계정 = '미확인'."""
    import openpyxl
    from openpyxl.utils import get_column_letter
    now = now or datetime.now()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "로그인 결과"
    ws.append(HEAD)
    need = wb.create_sheet(CONTACT)
    need.append(CONTACT_HEAD)
    for a in accts:
        rec = state.get(a.account_id) or {}
        ws.append((a.kind, a.representative, a.business_name, a.account_id, a.phone, rec.get("결과", "미확인"),
                   rec.get("사유", ""), rec.get("비밀번호 출처", ""), rec.get("시도", "")))
        if needs_contact(rec):
            need.append((a.representative, a.business_name, a.account_id, a.phone, a.kind, rec.get("사유", ""),
                         rec.get("시도", "")))
    for sh, widths in ((ws, (22, 10, 22, 16, 15, 8, 46, 30, 19)), (need, (10, 22, 16, 15, 22, 46, 19))):
        for i, w in enumerate(widths):
            sh.column_dimensions[get_column_letter(i + 1)].width = w
        sh.freeze_panes = "A2"
    p = Path(out_dir) / f"로그인확인_{now:%y%m%d_%H%M%S}.xlsx"
    p.parent.mkdir(parents=True, exist_ok=True)
    wb.save(p)
    return p


def counts(accts, state: dict) -> dict:
    out = {OK: 0, FAIL: 0, BLOCK: 0, SKIP: 0, "미확인": 0}
    for a in accts:
        out[(state.get(a.account_id) or {}).get("결과", "미확인")] += 1
    return out


# ── 실행 조건 ─────────────────────────────────────────────────
def busy_reason(locks=BUSY_LOCKS) -> str:
    """앱 파이프라인·정산 자동 다운로드가 도는 중이면 그 이유(아니면 '') — 잠금을 잠깐 잡아 보고 바로 놓음."""
    from .registry_lock import RegistryLockError, registry_lock
    for p in locks:
        try:
            with registry_lock(p, wait_sec=0):
                pass
        except RegistryLockError:
            return f"{'앱 실행(판매수집·순위)' if p.name == '파이프라인.lock' else '정산 자동 다운로드'}이 도는 중({p})"
    return ""


def past_until(until: str, now: datetime | None = None) -> bool:
    if not until:
        return False
    h, m = (int(x) for x in until.split(":"))
    now = now or datetime.now()
    return (now.hour, now.minute) >= (h, m)


def pick(accts, state: dict, retry: bool) -> list:
    """retry = 마지막 결과가 '성공'이 아닌 계정만(실패·차단·건너뜀·미확인)."""
    return [a for a in accts if not retry or (state.get(a.account_id) or {}).get("결과") != OK]


# ── 계정 하나 ─────────────────────────────────────────────────
def check_one(a, pws: list, hidden: bool, log) -> tuple[str, str]:
    """실제 Chrome 으로 로그인만 해 보고 닫음 → (결과, 사유). 로그인 코드·판정은 앱 ①과 같음."""
    from . import login_today
    from .browser import WingBrowser
    from .pipeline_sales import LoginBlocked, LoginCredentialError, _login_with_candidates, account_profile
    with WingBrowser(profile_dir=account_profile(a.account_id), offscreen=hidden) as b:
        try:
            ok = _login_with_candidates(b, a, pws, log, login=True, semi=not hidden)
        except LoginBlocked:
            return BLOCK, "로그인 차단(Akamai) — 사무실 IP 차단 의심·나중에 다시"
        except LoginCredentialError:
            code, detail = b.classify_login()
            why = login_today.reasons_today().get(a.account_id) or login_today.PW_WRONG
            return FAIL, f"{why}" + (f" — 쿠팡 안내: {detail[:80]}" if code == "error" and detail else "")
        code, detail = b.classify_login()
        if ok:
            return OK, "로그인 됨"
        if code == "otp":
            return FAIL, f"{login_today.OTP} — 계약자 휴대폰 인증번호 필요"
        if code in ("blocked", "akamai"):
            return BLOCK, f"로그인 차단({code}) — 나중에 다시"
        return SKIP, f"{login_today.UNFINISHED}({code}: {detail[:60]}) — 다음 실행에서 다시"


# ── 전체 ─────────────────────────────────────────────────────
def run(accts, log, *, hidden: bool = False, until: str = UNTIL_DEFAULT, retry: bool = False, out_dir=OUT_DIR,
        check=check_one, passwords=None, busy=busy_reason, sleep=time.sleep, now=datetime.now) -> int:
    """계정마다 로그인 확인 → 결과 기록 → 끝에 엑셀. 반환 = 0(끝까지·멈춤 시각) | 1(차단 연속 중단) | 2(다른 실행 중)."""
    from . import transfer_accounts as TA
    passwords = passwords or TA.login_passwords
    targets = pick(accts, load_state(out_dir), retry)
    log(f"[셀독전환 로그인 확인] 대상 {len(targets)}계정 / 파일 {len(accts)}계정"
        + (" (성공 안 된 계정만)" if retry else "") + f" · 멈춤 시각 {until or '없음'}")
    rc, blocked, tried = 0, 0, 0
    for a in targets:
        why = busy()
        if why:
            log(f"  ⛔ {why} — 로그인 확인 멈춤(남은 계정은 다음에 '성공 안 된 계정만'으로 이어서)")
            rc = 2
            break
        if past_until(until, now()):
            log(f"  ⏸ 멈춤 시각 {until} 지남 — 남은 계정은 다음에 이어서(18:00 앱 실행 보호)")
            break
        pws, src = passwords(a)
        at = now().isoformat(sep=" ", timespec="seconds")
        if not pws:
            save_result(out_dir, a.account_id, FAIL, "비밀번호 없음(파일·관리대장 모두)", src, at)
            log(f"  [{a.label}] 비밀번호 없음 — 연락 필요로 기록(로그인 안 함)")
            continue
        if tried:
            sleep(random.uniform(*GAP_SEC))
        tried += 1
        log(f"== [{a.kind}] {a.label} ({a.account_id}) · 비밀번호 {src} ==")
        try:
            status, reason = check(a, pws, hidden, log)
        except Exception as exc:   # 브라우저 오류 등 — 이 계정만 '건너뜀'(무음 아님·다음 계정 진행)
            status, reason = SKIP, f"확인 중 오류 — {exc.__class__.__name__}: {str(exc)[:80]}"
        save_result(out_dir, a.account_id, status, reason, src, at)
        log(f"  → {status}: {reason}")
        blocked = blocked + 1 if status == BLOCK else 0
        if blocked >= BLOCK_STOP:
            log(f"  ⛔ 로그인 차단 {BLOCK_STOP}계정 연속 — 전체 중단(사무실 IP 차단 의심·다른 계정 보호). 몇 시간 뒤 다시")
            rc = 1
            break
    state = load_state(out_dir)
    report = write_report(accts, state, out_dir, now())
    c = counts(accts, state)
    log(f"[셀독전환 로그인 확인] 끝 — 성공 {c[OK]} · 실패(연락 필요) {c[FAIL]} · 차단 {c[BLOCK]} · 건너뜀 {c[SKIP]} · "
        f"미확인 {c['미확인']} · 결과 {report}")
    return rc


# ── 명령줄(정산다운로드.exe transfer-login — 앱 '셀독전환 정산' 탭이 띄움) ────────────────────────
def main_cli(args) -> int:
    """args: file·retry·until·hidden·accounts. 로그 = 화면 + output/정산_셀독전환/로그/로그인확인_YYMMDD_HHMMSS.log."""
    import sys
    from . import appconfig
    from . import transfer_accounts as TA
    log_path = OUT_DIR / "로그" / f"로그인확인_{datetime.now():%y%m%d_%H%M%S}.log"
    log_path.parent.mkdir(parents=True, exist_ok=True)

    def log(m: str) -> None:
        line = f"{datetime.now():%H:%M:%S} {m}"
        print(line, flush=True)
        with log_path.open("a", encoding="utf-8") as f:
            f.write(line + "\n")
    path = args.file or appconfig.get(TA.CONFIG_KEY, "") or TA.DEFAULT_PATH
    try:
        accts, warns = TA.read_transfer_file(path)
    except TA.TransferFileError as exc:
        log(f"[셀독전환] ⛔ 파일 오류 — {exc}")
        return 2
    for w in warns:
        log(f"  [파일] ⚠ {w}")
    only = {x.strip() for x in (args.accounts or "").split(",") if x.strip()}
    if only:
        missing = only - {a.account_id for a in accts}
        if missing:
            log(f"[셀독전환] ⛔ 파일에 없는 계정: {sorted(missing)}")
            return 2
        accts = [a for a in accts if a.account_id in only]
    log(f"[셀독전환] 파일 {path} · ①② {len(accts)}계정")
    why = busy_reason()
    if why:
        log(f"[셀독전환] ⛔ {why} — 끝난 뒤 다시 실행하세요")
        return 2
    rc = run(accts, log, hidden=args.hidden, until=args.until if args.until is not None else UNTIL_DEFAULT,
             retry=args.retry)
    sys.stdout.flush()
    return rc
